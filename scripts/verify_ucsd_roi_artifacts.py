"""Verify actual six-case arrays, processing tables, notebook images, Excel and prior work."""
from pathlib import Path
import base64,hashlib,json,subprocess
import nbformat
import nibabel as nib
import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1];BASE=ROOT/'reports/ucsd_roi_pilot';OUT=ROOT/'outputs/ucsd-roi-pilot-20261029'
REFERENCE='9cba400e0afd18f012db9d18b085e42772caa6da'
PRESERVE=[
 'data/raw/reference/ucsd_ptgbm_clinical_v3.xlsx','data/raw/reference/ucsd_ptgbm_negative_cases.xlsx',
 'data/processed/ucsd_clinical_revalidated_cohort.csv','data/processed/ucsd_clinical_revalidated_oof.csv',
 'reports/ucsd_clinical_revalidation/metrics.json','outputs/ucsd-clinical-revalidation-20261029/UCSD_Clinical_Revalidation.xlsx',
 'src/rn_public/ucsd_clinical_revalidation.py','src/rn_public/ucsd_mri_intake.py','src/rn_public/ucsd_mri_mirror.py',
 'notebooks/13_ucsd_clinical_revalidation.ipynb','notebooks/14_ucsd_mri_intake_and_qc.ipynb','notebooks/15_ucsd_mri_mirror_samples.ipynb',
 'reports/ucsd_mri_mirror/sample_qc_receipt.json','data/metadata/ucsd_mirror_asset_inventory.csv']

def digest(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()

def main():
 r=json.loads((BASE/'pilot_receipt.json').read_text());s=r['summary']
 assert s['protocol_sha256']==digest(ROOT/'reports/ucsd_roi_pilot_protocol.md')
 assert s['code_sha256']==digest(ROOT/'src/rn_public/ucsd_roi_pilot.py')
 c=json.loads((BASE/'normalization_choice.json').read_text())
 assert c['amendment_sha256']==digest(ROOT/'reports/ucsd_roi_pilot_intensity_amendment.md')
 assert c['code_sha256']==digest(ROOT/'src/rn_public/ucsd_intensity_followup.py')
 assert len(r['files'])==s['source_files_available']==s['source_files_decoded']==36
 hashes={}
 for x in r['files']:
  p=ROOT/x['path'];hashes[x['path']]=digest(p)
  assert hashes[x['path']]==x['sha256'] and p.stat().st_size==x['bytes']
  assert x['gzip_crc_verified'] and x['release_digest_matches'] and x['nonfinite_voxels']==0
  assert x['relative_grid']['status']=='native_matrix_match' and not x['technical_pass']
 assert len(set(hashes.values()))==36
 masks={}
 for x in r['cases']:
  assert x['success'];p=ROOT/x['derived_mask'];assert digest(p)==x['derived_mask_sha256']
  original=next(y for y in r['files'] if y['ID']==x['ID'] and y['sequence']=='BraTS_tumor_seg')
  source=nib.load(ROOT/original['path']);mask=nib.load(p)
  np.testing.assert_array_equal(np.asarray(mask.dataobj),np.asarray(source.dataobj)>0)
  np.testing.assert_array_equal(mask.affine,source.affine)
  assert mask.header.get_xyzt_units()[0]=='unknown'
  masks[x['derived_mask']]=digest(p)
 cases=pd.read_csv(BASE/'case_ROI_diagnostics.csv');assert len(cases)==cases.patient_id.nunique()==6
 assert cases.tumor_label.value_counts().to_dict()=={0:3,1:3}
 for file in ['fixed_32_bin_histograms.csv','unclipped_fixed_32_bin_histograms.csv']:
  h=pd.read_csv(BASE/file);assert len(h)==768
  for case,sub in h.groupby('ID'):
   row=cases[cases.ID.eq(case)].iloc[0]
   for seq,d in sub.groupby('sequence'):
    assert set(d.bin)==set(range(1,33)) and d.voxel_count.sum()==row.roi_voxels
    assert d.eroded_voxel_count.sum()==row.erosion_voxels
    assert np.isclose(d.probability.sum(),1)
 data=json.loads((ROOT/'reports/ucsd_roi_workbook_input.json').read_text())
 path=OUT/'UCSD_ROI_Intensity_Pilot.xlsx';wb=load_workbook(path,data_only=False)
 assert wb.sheetnames==['Summary','ROI','Processing','Files','Histograms']
 for tab,key in [('ROI','cases'),('Processing','processing'),('Files','files'),('Histograms','histograms')]:
  sh=wb[tab];headers=[c.value for c in sh[1]];source=data[key]
  assert set(headers)==set(source['headers']) and sh.max_row==len(source['rows'])+1
  for i,row in enumerate(source['rows'],2):
   expected=dict(zip(source['headers'],row))
   for j,k in enumerate(headers,1):
    value=expected[k];actual=sh.cell(i,j).value
    if isinstance(value,(int,float)) and not isinstance(value,bool):assert np.isclose(actual,value,rtol=1e-12,atol=1e-12),(tab,i,k)
    else:assert actual==(None if value=='' else value),(tab,i,k,actual,value)
 cached=load_workbook(path,data_only=True)['Summary']
 for cell,value in {'C5':6,'C6':36,'C7':43455955,'C8':36,'C9':36,'C10':0,'C11':6,'C12':24}.items():assert cached[cell].value==value,(cell,cached[cell].value,value)
 nb=nbformat.read(ROOT/'notebooks/16_ucsd_roi_and_intensity_pilot.ipynb',as_version=4);nbformat.validate(nb)
 cells=[c for c in nb.cells if c.cell_type=='code'];assert [c.execution_count for c in cells]==[1,2,3,4,5]
 codehash=hashlib.sha256('\n'.join(c.source for c in cells).encode()).hexdigest()
 kernel=json.loads((ROOT/'reports/kernel_execution_16.json').read_text());assert kernel['success']
 assert kernel['notebooks'][0]['code_sha256']==codehash and kernel['notebooks'][0]['errors']==0
 images=[o.data['image/png'] for c in cells for o in c.outputs if o.output_type in ['display_data','execute_result'] and 'image/png' in o.data]
 names=['six_case_ROI_review.png','six_case_intensity_histograms.png','six_case_unclipped_histograms.png'];assert len(images)==3
 for encoded,name in zip(images,names):assert base64.b64decode(encoded)==(OUT/name).read_bytes()
 tests=json.loads((BASE/'software_test_receipt.json').read_text());assert tests['success'] and tests['tests_run']==67
 preserved={}
 for p in PRESERVE:
  prior=subprocess.run(['git','show',f'{REFERENCE}:{p}'],cwd=ROOT,check=True,capture_output=True).stdout
  assert prior==(ROOT/p).read_bytes(),p;preserved[p]=hashlib.sha256(prior).hexdigest()
 result=dict(success=True,actual_source_files_verified=36,source_bytes=43455955,
             derived_ROIs_verified=6,processing_records_verified=24,histogram_comparison_rows_verified=768,
             notebook_code_sha256=codehash,actual_notebook_images_verified=3,
             workbook_sha256=digest(path),MRI_model_fitted=False,physical_scale_confirmed=False,
             original_clipping_results_preserved=True,source_sample_sha256=hashes,
             derived_mask_sha256=masks,preserved_reference_commit=REFERENCE,preserved_artifacts_sha256=preserved)
 (BASE/'artifact_verification.json').write_text(json.dumps(result,indent=2)+'\n')
 print(json.dumps({k:v for k,v in result.items() if not isinstance(v,dict)},indent=2))

if __name__=='__main__':main()
