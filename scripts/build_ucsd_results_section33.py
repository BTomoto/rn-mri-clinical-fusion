"""Draft R3 with fixed OOF metrics, editable Table 2 and unchanged Figure 2."""
from pathlib import Path
import argparse,hashlib,json,re,sys,zipfile
import numpy as np
import pandas as pd
from lxml import etree
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION,WD_ORIENT

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'outputs/ucsd-manuscript-step35-20261029'
REPORT=ROOT/'reports/ucsd_manuscript_step35'
PARAGRAPHS=[
 {'id':'R3','citation_ids':[],
  'english': 'Each of the four fixed models provided an OOF probability for all 182 examinations under the same patient-grouped five-fold partition. The pooled ROC AUC was 0.526 for M0 (95% patient-bootstrap percentile interval, 0.404 to 0.644), 0.672 for M1 (0.553 to 0.778), 0.762 for MRI (0.683 to 0.838) and 0.782 for Fusion (0.698 to 0.859; Table 2 and Figure 2A). Treatment-related-change AP was 0.261, 0.354, 0.444 and 0.505 for M0, M1, MRI and Fusion, respectively, against a treatment-related-change prevalence of 0.225 in this cohort (Figure 2B). The corresponding Brier scores were 0.177, 0.163, 0.151 and 0.144. Fusion had the largest AUC and treatment-related-change AP point estimates and the lowest Brier score among the four models in this partition. The intervals summarized patient-sampling uncertainty conditional on the saved OOF probabilities; they did not include the full variability of model retraining or alternative fold assignments. These results represent internal validation in the same imaging cohort. The prespecified paired Fusion-minus-M1 comparison is reported separately, because the ordering of model-specific point estimates alone does not quantify the uncertainty in their difference.',
  'chinese': '四个固定模型均在相同的患者分组五折划分下，为全部182次检查提供OOF概率。汇总ROC AUC分别为M0的0.526（95%患者bootstrap百分位区间0.404至0.644）、M1的0.672（0.553至0.778）、MRI的0.762（0.683至0.838）及Fusion的0.782（0.698至0.859；Table 2和Figure 2A）。M0、M1、MRI与Fusion的治疗相关改变AP分别为0.261、0.354、0.444和0.505，该队列治疗相关改变的比例为0.225（Figure 2B）。相应Brier评分分别为0.177、0.163、0.151和0.144。在本次划分中，Fusion在四模型中具有最大的AUC与治疗相关改变AP点估计，以及最低的Brier评分。这些区间概括条件于已保存OOF概率的患者抽样不确定性，不涵盖模型重新训练或其他分折的全部变异。上述结果属于同一影像队列中的内部验证。预设的Fusion减去M1配对比较另节报告，因为仅凭各模型点估计排序，无法量化模型差值的不确定性。'}
]
MODELS=['M0','M1','MRI','Fusion']
FIGURE=ROOT/'outputs/ucsd-paper-materials-20261029/figures/Figure2_Main_Comparison.png'
FIGURE_CAPTION='Figure 2. Internal validation of the four fixed models in 182 examinations from 135 patients. A, ROC curves based on pooled patient-grouped OOF probabilities. B, Precision-recall curves for treatment-related change; the dashed horizontal line marks its examination-level prevalence. AP denotes average precision. C, Descriptive calibration using predicted-probability quantile groups, with mean predicted tumor probability plotted against observed tumor proportion. The diagonal line in A is a reference for nondiscrimination and in C for agreement. Calibration summaries were not used to recalibrate the evaluated probabilities.'
TABLE_SPEC={'key':'Table2','title':'Table 2. Performance of the four fixed models',
  'headers':['Model','ROC AUC (95% interval)','Brier score','AP for treatment change','Treatment change correct, n/N'],
  'notes':['M0: age and sex; M1: expanded clinical inputs; MRI: eight MRI features; Fusion: clinical and MRI inputs. All models used the same 182 examinations and patient-grouped five-fold partition.',
  'AUC intervals are the 2.5th and 97.5th percentiles of 2000 paired patient-cluster bootstrap attempts using fixed OOF probabilities without refitting. They do not capture full retraining or partition variability.',
  'AP is average precision for treatment-related change, scored by one minus the tumor probability; its prevalence is 41/182 (0.225). Lower Brier scores indicate smaller probability error. Correct treatment-change counts use tumor probability below 0.5. Displayed metrics are rounded to three decimals; the accompanying CSV retains the original precision.']}

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def wc(s):return len(re.findall(r"[A-Za-z0-9]+(?:['/-][A-Za-z0-9]+)*",s))
def close(a,b):assert np.isclose(a,b,rtol=1e-12,atol=1e-12),(a,b)

def audit():
    old=json.loads((ROOT/'reports/ucsd_manuscript_step34/results_section32.json').read_text());frozen=old['source_audit']['frozen_file_sha256'].copy()
    frozen.update(json.loads((ROOT/'reports/ucsd_manuscript_step34/artifact_verification.json').read_text())['sha256'])
    for p,h in frozen.items():assert sha(ROOT/p)==h,p
    extra=['reports/ucsd_fusion/model_metrics.csv','outputs/ucsd-paper-materials-20261029/tables/Table2.csv',str(FIGURE.relative_to(ROOT)),
      'outputs/ucsd-paper-materials-20261029/figures/Figure2_Main_Comparison.svg','reports/ucsd_paper_materials/generation_receipt.json']
    frozen.update({p:sha(ROOT/p) for p in extra})
    pred=pd.read_csv(ROOT/'data/processed/ucsd_fusion_oof.csv');y=pred.tumor_label.to_numpy()
    assert len(pred)==182 and pred.validation_group.nunique()==135 and y.sum()==141 and (y==0).sum()==41
    saved=json.loads((ROOT/'reports/ucsd_fusion/metrics.json').read_text());table=pd.read_csv(ROOT/'outputs/ucsd-paper-materials-20261029/tables/Table2.csv')
    modelcsv=pd.read_csv(ROOT/'reports/ucsd_fusion/model_metrics.csv').set_index('model');boot=pd.read_csv(ROOT/'reports/ucsd_fusion/paired_patient_bootstrap.csv')
    assert table.shape==(4,8) and saved['same_cases_and_folds'] and saved['n_splits']==5
    assert saved['bootstrap_attempts']==saved['bootstrap_valid']==2000 and saved['bootstrap_models_refitted'] is False
    def manual_ap(labels,scores):
        order=np.argsort(-scores,kind='stable');labels=labels[order];scores=scores[order]
        ends=np.flatnonzero(np.r_[scores[:-1]!=scores[1:],True]);tp=np.cumsum(labels)[ends]
        return float(np.sum(np.diff(np.r_[0,tp])/labels.sum()*tp/(ends+1)))
    records=[];wordrows=[]
    for i,name in enumerate(MODELS):
        p=pred[name+'_oof_probability'].to_numpy();assert np.isfinite(p).all() and np.all((p>0)&(p<1))
        diff=p[y==1][:,None]-p[y==0][None,:];auc=float(np.mean((diff>0)+.5*(diff==0)))
        ap=manual_ap(1-y,1-p);brier=float(np.mean((p-y)**2));correct=int(((p<.5)&(y==0)).sum())
        for key,value in [('roc_auc',auc),('ap_treatment_change',ap),('brier',brier),('tn',correct)]:close(value,saved['metrics'][name][key]);close(value,modelcsv.loc[name,key])
        lo=float(boot[name+'__roc_auc'].quantile(.025));hi=float(boot[name+'__roc_auc'].quantile(.975))
        for key,value in [('lower',lo),('upper',hi)]:close(value,saved['patient_cluster_bootstrap_95ci'][name+'__roc_auc'][key])
        for key,value in [('roc_auc_ci_lower',lo),('roc_auc_ci_upper',hi)]:close(value,modelcsv.loc[name,key])
        for key,value in [('AUC',auc),('CI_lower',lo),('CI_upper',hi),('Brier',brier),('AP_TC',ap),('TC_correct',correct),('TC_total',41)]:close(value,table.iloc[i][key])
        records.append({'Model':name,'ROC_AUC':float(table.iloc[i].AUC),'AUC_interval_lower':float(table.iloc[i].CI_lower),
          'AUC_interval_upper':float(table.iloc[i].CI_upper),'Brier_score':float(table.iloc[i].Brier),
          'AP_treatment_change':float(table.iloc[i].AP_TC),'Treatment_change_correct':int(table.iloc[i].TC_correct),'Treatment_change_total':41})
        wordrows.append([name,f'{auc:.3f} [{lo:.3f}, {hi:.3f}]',f'{brier:.3f}',f'{ap:.3f}',f'{correct}/41'])
    svg=etree.parse(str(FIGURE.with_suffix('.svg')));texts=svg.xpath('//s:text/text()',namespaces={'s':'http://www.w3.org/2000/svg'})
    for r in records:
        assert f"{r['Model']}: {r['ROC_AUC']:.3f}" in texts and f"{r['Model']}: AP={r['AP_treatment_change']:.3f}" in texts
    assert not re.search(r'[\u3400-\u9fff]',' '.join(texts))
    receipt=json.loads((ROOT/'reports/ucsd_paper_materials/generation_receipt.json').read_text());assert receipt['output_sha256'][str(FIGURE.relative_to(ROOT))]==sha(FIGURE)
    spec=dict(**TABLE_SPEC,rows=wordrows)
    assert not re.search(r'[\u3400-\u9fff]',' '.join(spec['headers']+sum(wordrows,[])+spec['notes']+[FIGURE_CAPTION]))
    return {'frozen_file_sha256':frozen,'translated_table':spec,'full_precision_table_records':records,
      'original_table_numeric_values_checked':28,'point_or_count_values_reconstructed':20,'bootstrap_auc_interval_endpoints_checked':8,
      'manual_ap_method':'sort scores descending, aggregate ties, weight precision by positive-recall increments',
      'figure_legend_values_checked':8,'figure_unchanged_from_saved_paper_materials':True,
      'treatment_change_prevalence':41/182,'numeric_values_changed':0,'models_refitted':0,'features_reextracted':0,'new_hypothesis_tests':0,
      'verification_scope':'Reconstruct AUC, treatment-change AP, Brier and class counts from fixed OOF probabilities; verify saved percentile endpoints and Table 2; compare figure legends with saved values without regenerating the figure'}

def add_table(doc,spec):
    title=doc.add_paragraph(spec['title'],style='Heading 1');title.paragraph_format.page_break_before=True
    title.paragraph_format.keep_with_next=True
    table=doc.add_table(rows=1,cols=5);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    widths=[.64,2.12,.86,1.44,1.86]
    for col,width in zip(table.columns,widths):col.width=Inches(width)
    for cell,width in zip(table.rows[0].cells,widths):cell.width=Inches(width)
    for cell,label in zip(table.rows[0].cells,spec['headers']):cell.text=label
    for values in spec['rows']:
        for cell,value,width in zip(table.add_row().cells,values,widths):cell.text=value;cell.width=Inches(width)
    props=table._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side)
        for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:e.set(qn('w:'+k),v)
        borders.append(e)
    props.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for side in ['top','left','bottom','right']:
        e=OxmlElement('w:'+side);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
    props.append(margins)
    for ri,row in enumerate(table.rows):
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if ri==0:trpr.append(OxmlElement('w:tblHeader'))
        for ci,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7EBEF');cell._tc.get_or_add_tcPr().append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:run.font.size=Pt(10.5);run.font.bold=ri==0
    for note in spec['notes']:
        p=doc.add_paragraph(note);p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(4)
        p.paragraph_format.keep_together=True
        for r in p.runs:r.font.size=Pt(10)

def build():
    OUT.mkdir(parents=True,exist_ok=True);REPORT.mkdir(parents=True,exist_ok=True);a=audit()
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(.79)
    for name in ['Normal','Title','Heading 1']:
        st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        rf=st.element.get_or_add_rPr().rFonts;rf.set(qn('w:eastAsia'),'Noto Serif CJK SC')
        for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:rf.attrib.pop(qn('w:'+key),None)
    normal=doc.styles['Normal'];normal.font.size=Pt(12);normal.paragraph_format.line_spacing=1.;normal.paragraph_format.space_after=Pt(9)
    doc.styles['Title'].font.size=Pt(16);doc.styles['Title'].paragraph_format.space_after=Pt(12)
    doc.styles['Heading 1'].font.size=Pt(13);doc.styles['Heading 1'].paragraph_format.space_before=Pt(8);doc.styles['Heading 1'].paragraph_format.space_after=Pt(9)
    doc.core_properties.title='Results 3.3 Four-model discrimination';doc.core_properties.author='Tianyu Feng'
    doc.add_paragraph('Results',style='Title');doc.add_paragraph('3.3 Four-model discrimination',style='Heading 1')
    for item in PARAGRAPHS:
        en=doc.add_paragraph(item['english']);en.paragraph_format.page_break_before=False
        en.paragraph_format.keep_together=True;en.paragraph_format.keep_with_next=True
        zh=doc.add_paragraph(item['chinese']);zh.paragraph_format.keep_together=True
        for r in zh.runs:r.font.size=Pt(11)
    add_table(doc,a['translated_table'])
    pd.DataFrame(a['full_precision_table_records']).to_csv(OUT/'Table2_EN.csv',index=False)
    landscape=doc.add_section(WD_SECTION.NEW_PAGE);landscape.orientation=WD_ORIENT.LANDSCAPE
    landscape.page_width=Inches(11);landscape.page_height=Inches(8.5)
    landscape.top_margin=landscape.bottom_margin=landscape.left_margin=landscape.right_margin=Inches(.79)
    doc.add_paragraph('Figure 2',style='Heading 1')
    doc.add_picture(str(FIGURE),width=Inches(9.3))
    p=doc.add_paragraph(FIGURE_CAPTION);p.paragraph_format.keep_together=True
    for r in p.runs:r.font.size=Pt(10)
    for e in [*doc.styles.element.xpath('.//w:pBdr'),*doc.element.xpath('.//w:pBdr')]:e.getparent().remove(e)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    path=OUT/'UCSD_Manuscript_Results_3_3_EN_ZH.docx';doc.save(path)
    (OUT/'UCSD_Manuscript_Results_3_3_EN.txt').write_text('\n\n'.join(p['english'] for p in PARAGRAPHS)+'\n')
    d={'stage':35,'date':'2026-10-29','subsection':'3.3 Four-model discrimination','v2_paragraph_ids':['R3'],'paragraphs':PARAGRAPHS,
      'english_word_count':sum(wc(p['english']) for p in PARAGRAPHS),'paragraph_word_counts':{p['id']:wc(p['english']) for p in PARAGRAPHS},
      'word_count_scope':'English narrative only; excludes headings, tables/notes, figure/caption and Chinese translation','v2_approximate_target_words':200,
      'equations':{'count':0,'numbers':[],'cumulative_previous_numbered_groups':10},
      'citation_mapping':{},'new_manuscript_references':0,'cumulative_unique_cited_sources':5,'editable_english_tables':['Table 2'],'table_numeric_values_changed':0,'included_figure':str(FIGURE.relative_to(ROOT)),'figure_sha256':sha(FIGURE),'figure_caption':FIGURE_CAPTION,'figure_reused_without_regeneration':True,'source_audit':a,'models_refitted':0,'new_hypothesis_tests':0,
      'next_section':'Results 3.4 Paired incremental comparison (R4)','docx':str(path.relative_to(ROOT)),'docx_sha256':sha(path)}
    (REPORT/'results_section33.json').write_text(json.dumps(d,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'docx':d['docx'],'english_word_count':d['english_word_count'],'paragraph_word_counts':d['paragraph_word_counts'],'prior_files':len(a['frozen_file_sha256'])},indent=2))

def verify():
    from pypdf import PdfReader
    d=json.loads((REPORT/'results_section33.json').read_text())
    for p,h in d['source_audit']['frozen_file_sha256'].items():assert sha(ROOT/p)==h,p
    path=ROOT/d['docx'];assert sha(path)==d['docx_sha256'];texts=[p.text for p in Document(path).paragraphs]
    for p in PARAGRAPHS:
        assert p['english'] in texts and p['chinese'] in texts and texts.index(p['chinese'])==texts.index(p['english'])+1
        assert not re.search(r'[\u3400-\u9fff，。；：（）【】、]',p['english'])
    assert sum(wc(p['english']) for p in PARAGRAPHS)==d['english_word_count']
    doc=Document(path);assert len(doc.tables)==1
    spec=d['source_audit']['translated_table'];expected=[spec['headers'],*spec['rows']]
    assert [[c.text for c in row.cells] for row in doc.tables[0].rows]==expected
    assert all(not re.search(r'[\u3400-\u9fff]',c.text) for row in doc.tables[0].rows for c in row.cells)
    for note in spec['notes']:assert note in texts
    full=pd.read_csv(OUT/'Table2_EN.csv');assert len(full)==4
    for (_,row),r in zip(full.iterrows(),d['source_audit']['full_precision_table_records']):
        assert row.Model==r['Model']
        for key,value in r.items():
            if key!='Model':close(row[key],value)
    assert FIGURE_CAPTION in texts and len(doc.sections)==2 and doc.sections[1].orientation==WD_ORIENT.LANDSCAPE
    ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
    with zipfile.ZipFile(path) as z:
        t=etree.fromstring(z.read('word/document.xml'));assert len(t.xpath('//m:oMath',namespaces=ns))==0
        media=[n for n in z.namelist() if n.startswith('word/media/')];assert len(media)==1
        assert hashlib.sha256(z.read(media[0])).hexdigest()==sha(FIGURE)
    pages=len(PdfReader(OUT/'docx_render/UCSD_Manuscript_Results_3_3_EN_ZH.pdf').pages);assert pages==3
    v={'success':True,'stage':35,'date':'2026-10-29','subsection':d['subsection'],'paragraph_ids':['R3'],
      'english_word_count':d['english_word_count'],'paragraph_word_counts':d['paragraph_word_counts'],'matched_chinese_paragraphs':1,
      'word_pages':pages,'visual_review_pages':[1,2,3],'native_editable_equations':0,'equation_numbers':[],'equation_images':0,
      'prior_files_unchanged':len(d['source_audit']['frozen_file_sha256']),'original_table_numeric_values_checked':28,'point_or_count_values_reconstructed':20,'bootstrap_auc_interval_endpoints_checked':8,'editable_english_tables':1,'table_numeric_values_changed':0,'figure_legend_values_checked':8,'unchanged_figure_embedded':True,'figure_page_landscape':True,
      'models_refitted':0,'new_hypothesis_tests':0,'new_manuscript_references':0,'cumulative_unique_cited_sources':5,'next_section':d['next_section'],
      'sha256':{str(p.relative_to(ROOT)):sha(p) for p in OUT.iterdir() if p.is_file() and p.suffix in {'.docx','.txt','.csv'}}}
    (REPORT/'artifact_verification.json').write_text(json.dumps(v,ensure_ascii=False,indent=2)+'\n');print(json.dumps(v,indent=2))

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--verify',action='store_true');args=parser.parse_args()
    verify() if args.verify else build()
