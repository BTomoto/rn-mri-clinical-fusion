"""Draft D11 from frozen performance records and primary metrics and utility guidance."""
from pathlib import Path
import argparse, hashlib, json, re, zipfile
import pandas as pd
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'outputs/ucsd-manuscript-step51-20261030'
REPORT = ROOT/'reports/ucsd_manuscript_step51'
EVIDENCE = REPORT/'literature_evidence.json'
MAPPING = {'15': 'S25', '17': 'S37'}
PARAGRAPHS = [{'id': 'D11', 'citation_ids': ['S37', 'S25'], 'english': "Fusion's AUC of 0.782 indicates discrimination within this dataset, but does not establish reliable individual probabilities or a decision rule. Metrics Reloaded distinguishes classification at a threshold, discrimination across thresholds and calibration as complementary dimensions of evaluation (17). Here, the descriptive calibration slope of 0.667 raises concern about the spread of predicted probabilities, while the paired Brier-score and logarithmic-loss difference intervals included zero. At the unchanged threshold of 0.5, Fusion identified 130/141 tumor examinations but only 12/41 treatment-related-change examinations. For a task intended to distinguish treatment effects from tumor, the 29 treatment-change examinations classified as tumor remain a relevant limitation. These counts do not quantify the consequences of using the model, because no treatment decisions or downstream outcomes were evaluated. Nor does the small accuracy advantage over assigning all examinations to tumor establish clinical value. A clinically defensible threshold would require a defined action and consideration of the consequences of each type of error. Guidance recommends assessing net benefit across relevant, prespecified thresholds when predictions are intended to direct care (15). We did not perform that evaluation, recalibrate on independent data or compare the model with clinician decisions. The present evidence supports further evaluation of an incremental imaging signal, without establishing that model-guided care improves outcomes.", 'chinese': 'Fusion的AUC为0.782，表明其在本数据集内具有鉴别能力，但不能据此确认个体概率可靠或确立决策规则。Metrics Reloaded将固定阈值分类、跨阈值鉴别及校准区分为相互补充的评价维度（17）。本研究描述性的校准斜率为0.667，提示预测概率的离散程度值得关注；同时，配对Brier评分及对数损失差值的区间均包含零。在保持不变的0.5阈值下，Fusion识别出130/141次肿瘤检查，但仅识别出12/41次治疗相关改变检查。对于旨在区分治疗效应与肿瘤的任务，29次治疗相关改变检查被判为肿瘤仍是重要局限。由于未评价任何治疗决策或后续结局，这些计数不能量化使用模型的后果。相较将所有检查均判为肿瘤的规则，准确率的小幅优势也不能确立临床价值。具有临床依据的阈值需要明确相应处置，并考虑每类错误的后果。指南建议，在预测用于指导诊疗时，应在相关且预先指定的阈值范围内评价净获益（15）。本研究未进行该项评价，未在独立数据上重新校准，也未与临床医师决策比较。现有证据支持进一步评价影像的增量信号，但不能证明模型指导诊疗能够改善结局。'}]

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def wc(text):
    return len(re.findall(r"[A-Za-z0-9]+(?:['/-][A-Za-z0-9]+)*", text))

def audit():
    previous=json.loads((ROOT/'reports/ucsd_manuscript_step50/discussion_section50.json').read_text())
    frozen=previous['source_audit']['frozen_file_sha256'].copy()
    frozen.update(json.loads((ROOT/'reports/ucsd_manuscript_step50/artifact_verification.json').read_text())['sha256'])
    for rel in ['reports/ucsd_manuscript_step36/results_section34.json','reports/ucsd_manuscript_step37/results_section35.json','reports/ucsd_manuscript_step30/methods_section28.json']:
        if rel not in frozen:frozen[rel]=sha(ROOT/rel)
    frozen[str(EVIDENCE.relative_to(ROOT))]=sha(EVIDENCE)
    for p,h in frozen.items():assert sha(ROOT/p)==h,p
    evidence=json.loads(EVIDENCE.read_text());records={s['reference_id']:s for s in evidence['studies']}
    refs={s['id']:s for s in json.loads((ROOT/'reports/ucsd_manuscript_step23/reference_library.json').read_text())['references']}
    selected=[]
    for number,rid in MAPPING.items():
        ref=refs[rid];record=records[rid];assert ref['doi']==record['doi'] and int(number)==record['provisional_citation_number']
        citation=ref['citation']
        if rid=='S37':
            assert '2024;21(2):195.' in citation and record['facts']['pages']=='195-212'
            citation=citation.replace('2024;21(2):195.','2024;21(2):195-212.')
        selected.append({'number':int(number),'reference_id':rid,'citation':citation,'doi':ref['doi'],'read_level_this_stage':record['read_level_this_stage']})
    metrics=json.loads((ROOT/'reports/ucsd_fusion/metrics.json').read_text());m=metrics['metrics']['Fusion'];cal=metrics['calibration']['Fusion']
    assert f"{m['roc_auc']:.3f}"=='0.782' and f"{cal['calibration_slope']:.3f}"=='0.667' and cal['error'] is None
    paired={x['metric']:x for x in metrics['paired_comparisons'] if x['comparison']=='Fusion_minus_M1'}
    for key in ['brier','log_loss']:assert paired[key]['lower']<0<paired[key]['upper']
    pred=pd.read_csv(ROOT/'data/processed/ucsd_fusion_oof.csv');assert len(pred)==182 and metrics['threshold']==.5
    y=pred.tumor_label;positive=pred.Fusion_oof_probability.ge(.5)
    counts={'tp':int((positive & y.eq(1)).sum()),'tn':int((~positive & y.eq(0)).sum()),'fp':int((positive & y.eq(0)).sum()),'fn':int((~positive & y.eq(1)).sum())}
    assert counts=={'tp':130,'tn':12,'fp':29,'fn':11} and all(m[k]==v for k,v in counts.items())
    assert int(y.sum())==141 and int(y.eq(0).sum())==41 and sum(counts[k] for k in ['tp','tn'])==142
    assert not metrics['independent_external_validation']
    prior_ids=set()
    for pattern in ['methods_section*.json','results_section*.json','discussion_section*.json']:
        for p in ROOT.glob('reports/ucsd_manuscript_step*/'+pattern):
            if int(p.parent.name.rsplit('step',1)[1])<51:prior_ids.update(r for item in json.loads(p.read_text())['paragraphs'] for r in item['citation_ids'])
    assert len(prior_ids)==16 and 'S25' in prior_ids and 'S37' not in prior_ids
    return {'frozen_file_sha256':frozen,'reference_records':selected,'source_evidence_record':str(EVIDENCE.relative_to(ROOT)),
            'new_full_texts_checked_for_current_claims':['S37'],'previous_full_text_revisited':['S25'],
            'fusion_AUC':m['roc_auc'],'descriptive_calibration_slope':cal['calibration_slope'],
            'fixed_threshold':.5,'classification_counts_checked':counts,'tumor_examinations':141,'treatment_change_examinations':41,
            'correct_examinations':142,'always_tumor_correct_examinations':141,'cross_zero_probability_error_intervals_checked':{k:paired[k] for k in ['brier','log_loss']},
            'clinical_actions_or_downstream_outcomes_evaluated':False,'clinical_net_benefit_calculated':False,'independent_probability_recalibration_performed':False,
            'clinician_comparison_performed':False,'threshold_optimized':False,'new_threshold_recommended':False,'clinical_benefit_claimed_established':False,
            'complete_S37_pagination_used_without_modifying_candidate_library':True,'previous_unique_cited_sources':sorted(prior_ids),
            'models_refitted':0,'new_hypothesis_tests':0,'verification_scope':'Read-only checks of saved probabilities, classification counts, descriptive calibration and paired probability-error intervals; primary full-text review, no recalibration, threshold search or net-benefit calculation'}

def build():
    a = audit(); OUT.mkdir(parents=True, exist_ok=True); REPORT.mkdir(parents=True, exist_ok=True)
    doc = Document(); sec = doc.sections[0]; sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(.79)
    for name in ['Normal', 'Title', 'Heading 1']:
        st = doc.styles[name]; st.font.name = 'Times New Roman'; st.font.color.rgb = RGBColor(0, 0, 0)
        rf = st.element.get_or_add_rPr().rFonts; rf.set(qn('w:eastAsia'), 'Noto Serif CJK SC')
        for key in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']: rf.attrib.pop(qn('w:'+key), None)
    normal = doc.styles['Normal']; normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.; normal.paragraph_format.space_after = Pt(9)
    doc.styles['Title'].font.size = Pt(16); doc.styles['Title'].paragraph_format.space_after = Pt(12)
    heading = doc.styles['Heading 1']; heading.font.size = Pt(13)
    heading.paragraph_format.space_before = Pt(8); heading.paragraph_format.space_after = Pt(9)
    doc.core_properties.title = 'Discussion 4.11 Clinical interpretation'; doc.core_properties.author = 'Tianyu Feng'
    doc.add_paragraph('Discussion', style='Title'); doc.add_paragraph('4.11 Clinical interpretation', style='Heading 1')
    for item in PARAGRAPHS:
        en = doc.add_paragraph(item['english']); en.paragraph_format.keep_together = True; en.paragraph_format.keep_with_next = True
        zh = doc.add_paragraph(item['chinese']); zh.paragraph_format.keep_together = True
        for run in zh.runs: run.font.size = Pt(11)
    if a['reference_records']:
        title = doc.add_paragraph('References', style='Heading 1'); title.paragraph_format.page_break_before = True
    for ref in a['reference_records']:
        p = doc.add_paragraph(str(ref['number'])+'. '+ref['citation']); p.paragraph_format.keep_together = True
    for e in [*doc.styles.element.xpath('.//w:pBdr'), *doc.element.xpath('.//w:pBdr')]: e.getparent().remove(e)
    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    path = OUT/'UCSD_Manuscript_Discussion_4_11_EN_ZH.docx'; doc.save(path)
    (OUT/'UCSD_Manuscript_Discussion_4_11_EN.txt').write_text('\n\n'.join(p['english'] for p in PARAGRAPHS)+'\n')
    d = {'stage': 51, 'date': '2026-10-30', 'subsection': '4.11 Clinical interpretation', 'v2_paragraph_ids': ['D11'],
         'paragraphs': PARAGRAPHS, 'english_word_count': sum(wc(p['english']) for p in PARAGRAPHS),
         'paragraph_word_counts': {p['id']: wc(p['english']) for p in PARAGRAPHS},
         'word_count_scope': 'English narrative including in-text citation numbers; excludes headings, reference list and Chinese translation',
         'v2_approximate_target_words': 200, 'equations': {'count': 0, 'numbers': [], 'cumulative_previous_numbered_groups': 10},
         'citation_mapping': MAPPING, 'citation_numbering': 'Existing number 15 preserved and new first-appearance number 17 added; renumber globally when Introduction is assembled',
         'new_manuscript_references': 1, 'cumulative_unique_cited_sources': 17,
         'source_audit': a, 'models_refitted': 0, 'new_hypothesis_tests': 0,
         'next_section': 'Discussion 4.12 Priorities for further evaluation (D12)',
         'docx': str(path.relative_to(ROOT)), 'docx_sha256': sha(path)}
    (REPORT/'discussion_section51.json').write_text(json.dumps(d, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps({'docx': d['docx'], 'english_word_count': d['english_word_count'], 'frozen_files': len(a['frozen_file_sha256'])}))

def verify():
    from pypdf import PdfReader
    d = json.loads((REPORT/'discussion_section51.json').read_text()); a = audit(); assert a == d['source_audit']
    path = ROOT/d['docx']; assert sha(path) == d['docx_sha256']; doc = Document(path); texts = [p.text for p in doc.paragraphs]
    for p in PARAGRAPHS:
        assert p['english'] in texts and p['chinese'] in texts and texts.index(p['chinese']) == texts.index(p['english'])+1
        assert not re.search(r'[\u3400-\u9fff，。；：（）【】、]', p['english'])
    assert sum(wc(p['english']) for p in PARAGRAPHS) == d['english_word_count']
    assert re.findall(r'\((\d+)\)', PARAGRAPHS[0]['english']) == ['17','15']
    for ref in a['reference_records']: assert str(ref['number'])+'. '+ref['citation'] in texts
    assert not doc.tables and len(doc.sections) == 1
    with zipfile.ZipFile(path) as z:
        xml = etree.fromstring(z.read('word/document.xml'))
        assert not xml.xpath('//m:oMath', namespaces={'m': 'http://schemas.openxmlformats.org/officeDocument/2006/math'})
        assert not [n for n in z.namelist() if n.startswith('word/media/')]
    pdf = PdfReader(OUT/'docx_render/UCSD_Manuscript_Discussion_4_11_EN_ZH.pdf'); pages = len(pdf.pages); assert pages == 2
    assert '4.11 Clinical' in pdf.pages[0].extract_text() and 'References' in pdf.pages[1].extract_text()
    v = {'success': True, 'stage': 51, 'date': '2026-10-30', 'subsection': d['subsection'], 'paragraph_ids': ['D11'],
         'english_word_count': d['english_word_count'], 'matched_chinese_paragraphs': 1, 'word_pages': pages,
         'visual_review_pages': [1,2], 'native_editable_equations': 0, 'equation_images': 0,
         'frozen_files_unchanged': len(a['frozen_file_sha256']), 'prior_files_unchanged': len(a['frozen_file_sha256'])-1,
         'new_full_texts_checked_for_current_claims': ['S37'], 'classification_counts_checked': a['classification_counts_checked'],
         'saved_AUC_and_descriptive_calibration_slope_checked': True, 'paired_probability_error_intervals_include_zero': True,
         'clinical_actions_or_downstream_outcomes_evaluated': False, 'clinical_net_benefit_calculated': False,
         'independent_probability_recalibration_performed': False, 'clinician_comparison_performed': False,
         'threshold_optimized': False, 'new_threshold_recommended': False, 'clinical_benefit_claimed_established': False,
         'complete_S37_pagination_used_without_modifying_candidate_library': True,
         'new_manuscript_references': 1, 'provisional_citation_mapping': MAPPING, 'cumulative_unique_cited_sources': 17,
         'models_refitted': 0, 'new_hypothesis_tests': 0, 'next_section': d['next_section'],
         'sha256': {str(p.relative_to(ROOT)): sha(p) for p in OUT.iterdir() if p.is_file() and p.suffix in {'.docx', '.txt'}}}
    (REPORT/'artifact_verification.json').write_text(json.dumps(v, ensure_ascii=False, indent=2)+'\n'); print(json.dumps(v, indent=2))

if __name__ == '__main__':
    parser = argparse.ArgumentParser(); parser.add_argument('--verify', action='store_true'); args = parser.parse_args()
    verify() if args.verify else build()
