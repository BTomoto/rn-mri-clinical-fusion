"""Build the literature learning report and paragraph-level writing outline.

No manuscript body is drafted. Frozen analysis tables supply result anchors.
"""
import hashlib
import json
from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'outputs/ucsd-writing-preparation-20261029'
REPORT = ROOT / 'reports/ucsd_writing_preparation'
TITLE = 'Clinical-MRI Feature Fusion for Distinguishing Residual or Recurrent Tumor from Treatment-Related Changes after Glioma Treatment: A Public-Data Internal Validation Study'

# Paragraph IDs remain stable when a later user-supplied template changes layout.
INTRO = [
 ('I1',130,'Clinical diagnostic problem','交代胶质瘤治疗后残余或复发肿瘤与治疗相关改变在MRI上表现重叠，以及误判可能带来的处理困难。尽快进入鉴别问题，避免长篇流行病学或生存背景。此处使用广义治疗相关改变；假性进展只作为其中一种情况。','[L1-L4, L6]；承接：为什么需要额外的定量信息。'),
 ('I2',160,'Quantitative MRI and information fusion','介绍常规MRI的强度和纹理计算如何补充视觉评估，以及影像信息与年龄、性别、治疗及时间信息可能互补。指出已有方法从少量参数到高维radiomics均有探索，避免暗示特征越多必然越好。','[L1-L5]；承接：已有研究为什么仍不足以回答我们的增量问题。'),
 ('I3',170,'Evidence gap and computational motivation','归纳既有结果不一致、终点定义不同、样本有限及验证不足。将本研究缺口限定为：在这一公开派生影像队列中，采用固定低维特征、同病例同折比较和患者隔离，检查影像相对临床模型的增量及稳定性。不能声称此前从无人融合MRI与临床信息。','[L1, L2, L4, L5]；承接：从文献缺口导向可检验目标。'),
 ('I4',140,'Study objective and prespecified comparisons','明确目标为扫描当时的二分类鉴别；比较M0、M1、MRI和Fusion，主要比较为Fusion减M1的折外AUC。说明患者分组验证、边界敏感性、重复分组及每患者单扫描分析的作用。贡献定位为透明计算流程与应用验证。','本项目第18-19步固定方案；收束引言并引出方法，不预告最佳结果。'),
]
METHODS = [
 ('M1',110,'Study design and public data sources','写明公开去标识数据的回顾性二次计算分析、UCSD-PTGBM及实际使用的公开镜像。区分源数据采集机构与本研究分析者；本研究没有新增招募或实验。源数据伦理信息引用描述论文，本研究机构适用说明以作者实际信息为准。','[L6]；数据URL、版本和哈希见项目清单。不得编造本研究伦理批准号。'),
 ('M2',130,'Cohort assembly and unit of analysis','说明源243扫描/178患者，经双ID冲突隔离5扫描/2患者后临床238扫描/176患者；可分析影像182扫描/135患者，另56扫描未覆盖。区分扫描数和患者数，解释不合并其他历史数据库，样本量由可用数据决定。','Figure 1、Table 1、Table S1；不得把182写成182名患者。'),
 ('M3',140,'Outcome labels and reference standard','定义1为残余或复发肿瘤、0为治疗相关改变。阴性病例来自公开名单，肿瘤来自冻结总体分组的补集；沿用源数据库临床/病理混合参考。说明治疗相关改变包含多种状态，并非仅假性进展或仅放射性坏死。','[L6]及标签审计；不得声称182例均经病理证实或重新独立判读。'),
 ('M4',150,'Clinical predictors and missing data','列明M1五项：年龄、性别、log1p初始事件间隔、扫描前其他治疗、间隔缺失标记。解释源时间字段包括初次手术/治疗/诊断，负偏移映射为正间隔，0、正偏移或Unknown按既定规则缺失；36条间隔缺失，仅在训练折中位数填补。','Table S2；MGMT/IDH及未来随访等未纳入，理由为时点或字段核查不足。不能把初始事件间隔改称距放疗时间。'),
 ('M5',130,'Derived MRI inputs and source annotation region','实际输入为FLAIR、增强T1及现成普通BraTS标注；派生数组256立方、uint8，共546个影像/标注文件。ROI取全部正标签并集，保留全部连通分量，称source annotation region。说明数组/仿射一致性核查与空间单位unknown。','第15-17步质控；不写本研究训练了分割网络、不称纯肿瘤ROI、不推断毫米或毫升。'),
 ('M6',190,'Intensity normalization and discretization','按单幅图像有限且大于0的支持区域计算p1/p99，线性变换但不截断，ROI内原始0保留。说明该支持区域不是经验证脑掩膜；32级ROI范围离散化，最大值归入最后一级。完整公式及边界规则放补充方法。','第16步强度修订及第17步固定代码；如实交代p1/p99不截断是观察到尾部问题后的工程修订，不能称全流程预注册。'),
 ('M7',170,'Fixed quantitative MRI features','列出FLAIR和T1post各自的均值、标准差、熵和三维GLCM对比度，共8项。熵基于32级直方图；GLCM在13个无向原生单体素偏移下汇总对称计数后归一化，再求对比度。明确未做高维筛选、形状或定量ADC计算。','Table S2及源码函数；不套用别人的PyRadiomics默认配置，不未经核查宣称IBSI一致。'),
 ('M8',130,'Models and feature-level fusion','四模型输入数分别2、5、8、13；Fusion为特征级拼接。全部L2 Logistic，C=1、lbfgs，连续变量折内标准化，二元变量不缩放。说明固定参数、无特征选择、无类别加权或超参数搜索；完整软件/参数放补充。','第18步analysis_lock与模型工厂；不能称新型融合算法、深度学习或自动特征学习。'),
 ('M9',140,'Patient-grouped internal validation','固定StratifiedGroupKFold五折及随机种子20261029；所有模型用相同病例及折，源患者、BraTS患者和保守分组均隔离。各扫描获得一次折外概率，填补/标准化仅从训练折学习。解释评价是扫描级、验证按患者隔离。','折分配清单和fold_audit；称internal cross-validation，不称独立外部测试。'),
 ('M10',190,'Performance and uncertainty','主要终点为Fusion-M1的AUC差；同时给ROC AUC、治疗改变AP、Brier、log loss、0.5阈值混淆矩阵与校准诊断。患者聚类配对bootstrap 2000次给百分位区间；固定OOF概率，不重拟合，区间不包含全部训练/分折变异。','Table 2、Table 3、Table S3；校准斜率及截距定义遵循原方案，未重校准或优化阈值。'),
 ('M11',160,'Sensitivity and split-stability analyses','依序描述一次原生单体素六邻域腐蚀、20套固定种子的重复患者分组五折、最小可用会话编号的每患者单扫描重建验证。重复分布只作描述；单扫描为135人/135扫描的新分析样本，单独bootstrap。','Tables S4-S7及Figure 3；不把20套重复当独立样本，不选最大AUC，不称最小会话为真实首次就诊。'),
 ('M12',110,'Reproducibility and reporting','说明公开仓库、Jupyter Notebook、逐例OOF、冻结清单、特征定义、校验和和源文件恢复路径。大影像以实际可访问的独立归档及清单说明，不能宣称全部原始MRI都在普通Git文件中。按TRIPOD+AI并参考CLAIM核对报告。','[G1, G2]；报告规范是透明性工具，不能据此宣布研究低偏倚。'),
]
RESULTS = [
 ('R1',120,'Cohort characteristics and imaging coverage','报告182扫描/135患者、141肿瘤/41治疗改变，主要临床分布和36条间隔缺失；说明56扫描未进入影像分析。年龄等按扫描汇总。不同类别的去重患者有重叠，不能把114与29直接相加。','Figure 1、Table 1、Table S1；先交代分析样本，再呈现模型。'),
 ('R2',120,'Primary model discrimination','依次报告M0、M1、MRI、Fusion的主折外AUC及区间，重点写M1 0.672、MRI 0.762、Fusion 0.782。只使用原固定五折结果，不用较高的腐蚀或重复结果替代。表中给全指标，正文避免逐格复述。','Table 2、Figure 2；承接主要增量比较。'),
 ('R3',100,'Incremental value over the clinical model','报告Fusion-M1的AUC差0.110，患者聚类95%区间0.021至0.214；同时交代Brier差-0.019，其区间跨0。区分排序区分度与概率误差，不用AUC区间不跨0推出各项性能全面改善。','Table 3；不增加事后显著性检验或未经计算的p值。'),
 ('R4',100,'Calibration and class-specific errors','报告主Fusion校准斜率0.667；0.5阈值下肿瘤正确130/141、治疗改变正确12/41。说明少数类识别不足，用两类分母呈现。整体准确率78.0%接近全部判肿瘤77.5%的参考，不能以准确率宣称临床可用。','Figure 2、Table S3；解释留到讨论，结果段仅陈述。'),
 ('R5',130,'Boundary sensitivity and repeated splits','腐蚀Fusion AUC 0.788，相对原版本差的区间跨0；20套重复Fusion AUC中位数0.784，范围0.759至0.809，Fusion-M1差全部为正。交代所有重复成功，不把范围称95%区间或当成20次独立验证。','Table S4、Tables S5A/S5B、Figure 3、Figure S1；稳定性不能解决标签或外部推广问题。'),
 ('R6',130,'One-scan-per-patient sensitivity','报告135扫描/135患者、113肿瘤/22治疗改变。Fusion AUC 0.730，Fusion-M1差0.120，其区间-0.015至0.252跨0；治疗改变仅正确2/22。明确不同人群/类别组成/训练集，不能将0.782与0.730的差归因于重复扫描造成偏倚。','Tables S6/S7、Figure S2；完整呈现较弱结果，与讨论呼应。'),
]
DISCUSSION = [
 ('D1',120,'Principal findings','开头直接回答目标：本公开队列主分析存在MRI对临床模型的区分度增量信号，20套分折方向一致；单扫描增量仍不确定。随即指出校准和治疗改变识别不足，避免先作过强临床结论再在末尾补限制。','对应R2-R6；使用showed、suggests、remains uncertain等匹配证据强度的表达。'),
 ('D2',180,'Relation to prior evidence','把融合获益研究[L1,L5]、增量不稳定研究[L2]及简洁参数验证[L4]放在同一论证中。比较信息类型、ROI、样本、参考标准及验证设计，解释为何现有证据并不一致。不能仅比较AUC大小为研究排名。','文献间是情境比较，不能写我们的模型优于这些论文；源结局较广也不能等同纯PsP。'),
 ('D3',150,'Computational contribution and plausible interpretation','解释8项固定MRI强度/纹理特征与临床特征级融合的应用价值，以及少量输入、可审计实现和患者隔离验证的意义。纹理反映灰度分布及局部空间变化；只能提出与组织异质性相容的解释，不能证明血管、细胞或病理机制。','对应M5-M9；计算机相关贡献为影像计算与融合验证，不是Python本身或算法原创。'),
 ('D4',140,'What robustness analyses establish','区分三种敏感性所检验的问题：边界依赖、分折变化、现有影像内单扫描选择。说明分折稳定性是有利信号，但单扫描少数类只有22条、增量区间宽，证据尚不稳固。没有独立数据，不能把重复CV改称多中心验证。','对应R5-R6；不在讨论中引入未执行的实验。'),
 ('D5',140,'Outcome and annotation limitations','优先讨论标签补集、非统一病理参考、混合治疗相关改变；源标注盲法未确认，ROI可能包含诊断信息。这些问题影响可解释性和有效性，患者分组也不能消除。不能将未知写成已证实泄漏，但必须清楚说明可能性。','[L6]及本项目审计；本研究没有新增病理、阅片或观察者一致性实验。'),
 ('D6',140,'Derived images and preprocessing limitations','说明镜像为8位派生强度、空间单位unknown、ROI为全部正标签并集，不能恢复原生定量强度或毫米体积。说明现成标注依赖、缺少独立观察者重画和完整人工解剖配准验证；特征对处理流程具有依赖性。','不能把文件完整性/数组一致性等同于临床影像质量全面合格。'),
 ('D7',110,'Sampling and validation limitations','讨论影像仅覆盖238临床扫描中的182条、单中心来源、135患者和少数类有限、13输入、重复扫描及扫描权重；36条间隔缺失和时间字段非统一放疗终点。说明bootstrap条件区间以及未有外部验证。','不补造事前样本量计算、统计功效、事件数充分性或缺失机制结论。'),
 ('D8',80,'Next research needed','按优先顺序提出独立队列、与标签独立的标注/参考标准、标准化影像及少数类与校准评估。将其写成后续研究要求，不承诺当前模型能指导治疗，不为了补充讨论自动启动新数据或新模型。','回应局限和临床定位，过渡到简短结论。'),
]
ABSTRACT = [
 ('A1',35,'Background','一句临床鉴别困难，一句研究目的：评估固定MRI特征相对临床信息的增量。'),
 ('A2',80,'Methods','公开二次分析，182扫描/135患者，5临床加8MRI，固定L2 Logistic及患者分组五折，主要配对AUC差；交代重复及单扫描敏感性。'),
 ('A3',105,'Results','列M1与Fusion主AUC、Fusion区间及增量区间；补一句单扫描增量不确定及治疗改变识别不足。若字数受限保留这些关键信息，不用最高重复AUC。'),
 ('A4',30,'Conclusions','限定为内部验证下的探索性增量信号，外部验证和独立标注仍必要。避免ready for clinical use之类过强结论。'),
]
CONCLUSION = ('C1',70,'Conclusion','简洁回答MRI特征在主内部验证中是否增加临床模型区分度，并同时承认单扫描不确定及少数类/校准限制。提出需要独立验证。结论只用已有证据，不加入生物机制、临床获益或投稿档次判断。','对应I4目标及D1主发现。')


def p(doc, text, boldlead=False):
    para = doc.add_paragraph()
    if boldlead and '：' in text:
        a,b = text.split('：',1);para.add_run(a+'：').bold=True;para.add_run(b)
    else:
        para.add_run(text)
    return para


def h(doc, text, level=1):
    doc.add_paragraph(text, style=f'Heading {level}')


def table(doc, headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for j,x in enumerate(headers):t.rows[0].cells[j].text=str(x)
    for row in rows:
        cells=t.add_row().cells
        for j,x in enumerate(row):cells[j].text=str(x)
    for i,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if not i:pr.append(OxmlElement('w:tblHeader'))
        for j,c in enumerate(row.cells):
            c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            mar=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                x=OxmlElement('w:'+side);x.set(qn('w:w'),'80');x.set(qn('w:type'),'dxa');mar.append(x)
            c._tc.get_or_add_tcPr().append(mar)
            if not i:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EEF1F3');c._tc.get_or_add_tcPr().append(sh)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=1.05
                para.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
                for r in para.runs:r.font.size=Pt(9.5);r.bold=i==0
    return t


def items(doc, rows):
    for row in rows:
        ident,words,title,task,source=row
        h(doc,f'{ident}  {title}  |  {words} English words',2)
        p(doc,'写作任务：'+task,True)
        p(doc,'依据与衔接：'+source,True)


def newpage(doc, title):
    doc.add_page_break();h(doc,title)


def build():
    OUT.mkdir(parents=True,exist_ok=True);REPORT.mkdir(parents=True,exist_ok=True)
    countdata=json.loads((REPORT/'literature_structure_counts.json').read_text())
    refs=[
      ('L1','McKenney AS, Weg E, Bale TA, et al. Radiomic Analysis to Predict Histopathologically Confirmed Pseudoprogression in Glioblastoma Patients. Advances in Radiation Oncology. 2023;8:100916. Online publication: 2022.','10.1016/j.adro.2022.100916'),
      ('L2','Leone R, Meredig H, Foltyn-Dumitru M, et al. Assessing the added value of apparent diffusion coefficient, cerebral blood volume, and radiomic magnetic resonance features for differentiation of pseudoprogression versus true tumor progression in patients with glioblastoma. Neuro-Oncology Advances. 2023;5:vdad016.','10.1093/noajnl/vdad016'),
      ('L3','Sun YZ, Yan LF, Han Y, et al. Differentiation of Pseudoprogression from True Progression in Glioblastoma Patients after Standard Treatment: A Machine Learning Strategy Combined with Radiomics Features from T1-weighted Contrast-enhanced Imaging. BMC Medical Imaging. 2021;21:17.','10.1186/s12880-020-00545-5'),
      ('L4','de Godoy LL, Mohan S, Wang S, et al. Validation of multiparametric MRI based prediction model in identification of pseudoprogression in glioblastomas. Journal of Translational Medicine. 2023;21:287.','10.1186/s12967-023-03941-x'),
      ('L5','Zhou J, et al. Radiomics and clinical data predict pseudoprogression after radiotherapy in high-grade glioma. Frontiers in Oncology. 2026;16:1892760.','10.3389/fonc.2026.1892760'),
      ('L6','Gagnon L, Gupta D, Nguyen U, et al. The University of California San Diego Post-Treatment Glioblastoma (UCSD-PTGBM) annotated multimodal MRI Dataset. Scientific Data. Published 6 January 2026.','10.1038/s41597-025-06499-z'),
      ('G1','Collins GS, Moons KGM, Dhiman P, et al. TRIPOD+AI statement: updated guidance for reporting clinical prediction models that use regression or machine learning methods. BMJ. 2024;385:e078378.','10.1136/bmj-2023-078378'),
      ('G2','Tejani AS, Klontzas ME, Gatti AA, et al. Checklist for Artificial Intelligence in Medical Imaging (CLAIM): 2024 Update. Radiology: Artificial Intelligence. 2024;6:e240300.','10.1148/ryai.240300'),
    ]
    # Detailed tasks are original to this project, not paraphrases of a single paper.
    plan={'date':'2026-10-29','stage':21,'provisional_title':TITLE,'english_main_word_target':4180,'english_abstract_word_target':250,
          'english_main_word_range':[3800,4500], 'manuscript_body_drafted':False,'awaiting_user_template':True,
          'paragraphs':{'abstract':ABSTRACT,'introduction':INTRO,'methods':METHODS,'results':RESULTS,'discussion':DISCUSSION,'conclusion':[CONCLUSION]},
          'references':[{'id':a,'citation':b,'doi':c,'url':'https://doi.org/'+c} for a,b,c in refs]}
    assert sum(x[1] for rows in [INTRO,METHODS,RESULTS,DISCUSSION,[CONCLUSION]] for x in rows)==4180
    assert sum(x[1] for x in ABSTRACT)==250
    (REPORT/'writing_outline.json').write_text(json.dumps(plan,ensure_ascii=False,indent=2)+'\n')
    doc=Document();s=doc.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11)
    s.top_margin=Inches(.65);s.bottom_margin=Inches(.65);s.left_margin=Inches(.8);s.right_margin=Inches(.8)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
        st=doc.styles[name];st.font.name='Nimbus Roman';st.font.color.rgb=RGBColor(0,0,0)
        st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC')
    st=doc.styles['Normal'];st.font.size=Pt(11);st.paragraph_format.line_spacing=1.16;st.paragraph_format.space_after=Pt(6)
    doc.styles['Title'].font.size=Pt(21);doc.styles['Title'].paragraph_format.space_after=Pt(10)
    doc.styles['Heading 1'].font.size=Pt(15);doc.styles['Heading 1'].paragraph_format.space_after=Pt(9)
    doc.styles['Heading 2'].font.size=Pt(11.5);doc.styles['Heading 2'].paragraph_format.space_before=Pt(8);doc.styles['Heading 2'].paragraph_format.space_after=Pt(4)
    doc.core_properties.title='UCSD public-data study: literature learning and writing outline'
    doc.core_properties.subject='Literature structure analysis and paragraph-level bilingual manuscript plan'
    doc.core_properties.author='Tianyu Feng'
    doc.add_paragraph('文献学习总结与详细写作大纲',style='Title')
    p(doc,'胶质瘤治疗后鉴别研究  |  2026年10月29日')
    h(doc,'研究定位')
    p(doc,'研究问题是：在现有公开队列中，将FLAIR及增强T1的8项固定计算特征与5项临床输入融合，是否比仅临床信息更能区分扫描当时的残余或复发肿瘤与治疗相关改变？重点是增量价值、患者隔离的内部验证和结果的稳定程度。')
    p(doc,'论文属于医学影像计算与信息融合的应用验证研究。计算贡献落实在三维图像处理、定量特征定义、特征级融合和可复现验证，不以成熟Logistic算法或Python软件本身作为原创方法。')
    h(doc,'建议暂定题目')
    p(doc,TITLE)
    p(doc,'中文释义：临床与MRI特征融合鉴别胶质瘤治疗后的残余或复发肿瘤与治疗相关改变：一项公开数据内部验证研究。')
    p(doc,'题目使用glioma以准确覆盖实际来源中并非全部为现代分子分型胶质母细胞瘤的记录；这不改变既有队列、标签或分析。最终题目在收到写作模板后再定。')
    h(doc,'本轮交付与后续方式')
    p(doc,'本轮交付文献学习、实际词数统计和逐段大纲。论文正文将在收到作者写作要求模板后开始：每个英文段落后紧跟对应中文译文；图、表、图题、表题及图表注释只使用英文。英文正文建议约3,800-4,500词，中文译文不计入英文篇幅。')

    newpage(doc,'1  文献学习：融合获益与相反证据')
    h(doc,'L1  McKenney等：以病理参考与融合对照组织论文',2)
    p(doc,'74名患者、手术病理二分类；影像特征与MGMT结合的交叉验证AUC由0.81增至0.89。结构先讲鉴别困难，再界定参考标准和计算流程，结果从人群到模型及对照，讨论再回到补充信息与验证不足。写法以方法细节支撑发现。[L1]')
    p(doc,'对我们的启发：先说明标签，再说明特征，最后比较增量；但其病理标准、灌注输入及MGMT均不能移植成我们的已完成工作。')
    h(doc,'L2  Leone等：围绕“额外信息是否有用”组织负向结果',2)
    p(doc,'105名患者；临床加MGMT模型AUC约0.69，增加影像信息并未稳定改善，高维特征模型反而较弱。引言逐步提出增量问题，方法按信息组合列模型，结果与讨论均围绕同一临床参考展开，结论保留阴性发现。[L2]')
    p(doc,'对我们的启发：主要对照应固定为M1，不能只展示Fusion的单个AUC。单扫描增量不确定与Brier区间跨0，应与有利结果同样清楚地写出。')
    h(doc,'两篇论文共同支持的写作顺序',2)
    p(doc,'临床问题 → 可用信息及缺口 → 参考标准与比较设计 → 同条件性能比较 → 限定适用范围。我们采用这一论证顺序，所有数值和方法仍以项目已冻结资料为准。')

    newpage(doc,'2  文献学习：常规序列、简洁模型与近期研究')
    h(doc,'L3  Sun等：常规增强T1与计算流程',2)
    p(doc,'77名患者，增强T1高维纹理和随机森林；总体准确率约72.8%。背景从常规MRI的可及性进入，方法依次讲分割、特征、筛选和分类，结果再作读者比较。讨论将纹理解释与局限分开，但部分措辞偏强。[L3]')
    p(doc,'对我们的启发：常规序列的计算分析可以构成文章主线；借鉴清楚的流程顺序，避免照搬长句、重复背景或过强结论。我们没有读者比较，不能声称优于放射科医师。')
    h(doc,'L4  de Godoy等：少量MRI参数与新队列验证',2)
    p(doc,'56名患者，用三个MRI参数构成的既有Logistic模型在新队列中验证，AUC约0.757。文章先解释为何要验证，再给诊断参考、参数和模型，结果重点报告类别识别；讨论解释性能变化及分组标准差异。[L4]')
    p(doc,'对我们的启发：模型简单不妨碍提出明确研究问题；证据质量、对照和验证比输入数量更关键。其新队列验证性质不能套用为我们的外部验证。')
    h(doc,'L5  2026年近期融合研究',2)
    p(doc,'Zhou等在222名高级别胶质瘤患者中整合影像、分子及其他临床指标，内部留出验证AUC约0.811。结构明确区分队列、终点、预处理、建模和验证；讨论同时呈现增量、可能解释和外部验证需求。[L5]')
    p(doc,'对我们的启发：可将影像与临床模型并列展示，并区分内部验证与外部推广。本文学习了其结构，不复制其17项特征、额外生物标志物、列线图或决策曲线。L5只作近期结构补充，未纳入下页统一XML词数统计。')

    newpage(doc,'3  实际英文词数与本研究篇幅安排')
    p(doc,'下表为对4篇全文采用同一脚本计算的近似英文词数。Main为引言、方法、结果、讨论及独立结论之和；不包括参考文献、图表内容、标题、声明和补充材料。Total仅为Abstract + Main，不代表投稿系统的全文字数。')
    rows=[]
    for r in countdata['papers']:
        c=r['section_words'];rows.append([r['key'],r['abstract_words'],c['Introduction'],c['Methods'],c['Results'],c['Discussion'],c['Conclusion'] or 'in D',r['main_prose_words'],r['abstract_plus_main_words']])
    table(doc,['Study','Abs.','Intro.','Meth.','Res.','Disc.','Conc.','Main','Total'],rows,[.55,.55,.6,.65,.55,.65,.55,.7,.7])
    p(doc,'计数说明：读取全文JATS的叙述段落，去除引用标记、公式、图注和表格等；字母/数字及词内连字符、斜杠或撇号按固定规则识别。L2无标题的开头4段计入Introduction；L3的Methods和Radiomics Strategy合并；L2/L4的结论仍计在Discussion。原文个别连写、公式及各软件规则会造成与Word不同的结果，因此这些不是出版社公布的字数。')
    p(doc,'四篇Main约2,831-4,814词。我们的方法、聚类验证与敏感性内容较多，建议正文目标4,180词，而不是为追随某篇论文强行压缩方法。以下是规划值，收到模板或选定期刊后调整，不是期刊字数限制。')
    table(doc,['Section','Target words','Paragraphs'],[['Abstract',250,4],['Introduction',600,4],['Methods',1750,12],['Results',700,6],['Discussion',1060,8],['Conclusion',70,1],['Main text',4180,31],['Abstract + main',4430,35]],[2.9,1.4,1.2])
    p(doc,'英文图表注释、参考文献和声明另计；中文译文不计入上述目标。独立Conclusion可按模板合入Discussion。')

    newpage(doc,'4  论文各部分的逻辑关系与写作风格')
    h(doc,'各部分只承担一个主要任务',2)
    for t in [
      '引言：从临床鉴别困难缩小到可检验的增量问题，最后明确研究目标。不要提前罗列结果。',
      '方法：按数据 → 标签 → 输入 → 图像计算 → 模型 → 验证 → 评价 → 敏感性的顺序，使读者能复现结果。',
      '结果：顺序与方法对应，先样本，再主比较，再类别错误及校准，最后敏感性。仅陈述观察，不讨论未经验证的机制。',
      '讨论：先回答目标，再比较相关证据和解释计算贡献，随后说明敏感性含义、局限及后续验证要求。',
      '结论：只保留研究问题的答案及适用边界，与摘要结论一致，不能比讨论更强。']:
        p(doc,t,True)
    h(doc,'英文规范与中译方式',2)
    p(doc,'每段采用“主题句 → 必要证据或细节 → 解释/衔接”的组织方式。已完成的方法和结果以过去时为主；公认背景通常用现在时；推断用may、suggest或is consistent with，不能写成因果或机制证明。优先具体、简短句，减少大而空的临床价值表述。')
    p(doc,'首次出现给出全称与缩写，后续统一使用。肿瘤为正类，所有敏感度、特异度、AP必须标明针对哪一类；不同论文正类不同，不能直接把同名指标作排名。数值以同一精度报告，AUC及差值建议三位小数，比例一位小数并附分子/分母。')
    p(doc,'中文是紧跟英文段落的忠实翻译，不另加解释或新结果。先组织符合英文论文逻辑的原文，再翻译；引用号、数值、否定和不确定性必须对应。图表及其注释仅英文，补充方法的英文段落仍可附中译。')
    h(doc,'报告规范的具体用途',2)
    p(doc,'TRIPOD+AI适用于回归及机器学习预测模型，要求交代数据、模型、验证、性能、局限及开放科学信息；本研究按这些方面逐项核对。[G1] CLAIM用于补充核对影像来源、参考标注、预处理、分组隔离和错误分析。[G2] 采用规范不表示已消除偏倚或已可临床应用。')

    newpage(doc,'5  逐段大纲：摘要')
    p(doc,'以下为段落任务与证据安排，尚不是论文正文。每个ID在后续逐段写作和模板核对时保留。字数指英文计划值。')
    h(doc,'Abstract  |  250 words',2)
    for ident,words,title,task in ABSTRACT:
        p(doc,f'{ident}  {title}（{words}词）：{task}',True)
    newpage(doc,'6  逐段大纲：引言')
    p(doc,'Introduction  |  600 words')
    items(doc,INTRO)

    newpage(doc,'7  逐段大纲：研究设计、队列、标签与临床输入')
    items(doc,METHODS[:4])
    p(doc,'逻辑检查：这4段先确定“研究谁、判断什么、输入在什么时候可获得”。没有先界定标签和时间变量，就不能解释后面的模型增量。')

    newpage(doc,'8  逐段大纲：影像计算与融合模型')
    items(doc,METHODS[4:8])
    p(doc,'逻辑检查：这4段使计算机相关内容具体可查。正文写清核心规则，补充方法给公式、离散化边界、GLCM计数合并和完整参数，不用堆砌技术名词增加复杂感。')

    newpage(doc,'9  逐段大纲：验证、评价、敏感性与复现')
    items(doc,METHODS[8:])
    p(doc,'逻辑检查：同病例同折保证主要增量比较的条件一致；患者分组防止同一患者跨折，但不会自动消除源标注、标签或可用影像选择问题。所有限制应在讨论中对应出现。')

    newpage(doc,'10  逐段大纲：结果')
    items(doc,RESULTS[:3])
    newpage(doc,'10  逐段大纲：结果（续）')
    items(doc,RESULTS[3:])

    newpage(doc,'11  逐段大纲：讨论的发现、文献与计算意义')
    items(doc,DISCUSSION[:4])
    p(doc,'文献比较采用“本研究观察 → 与哪类证据相近/不同 → 哪些设计差异可能相关 → 当前能支持的解释”的顺序。文献不能替本研究补做未完成的实验或证明未验证的病理机制。')

    newpage(doc,'12  逐段大纲：局限、后续方向与结论')
    items(doc,DISCUSSION[4:])
    items(doc,[CONCLUSION])

    newpage(doc,'13  图表安排与逐段写作流程')
    p(doc,'以下沿用第20步已完成材料，不增加新模型。主文拟用3表3图，补充使用其余表图；最终可随目标期刊调整。图表内容和注释全部英文，现有中文表头在正式组稿时转换为英文。')
    table(doc,['Item','Purpose','Paragraphs'],[
      ['Table 1','Main-cohort characteristics','M2, R1'],['Table 2','Four-model OOF performance','M10, R2'],['Table 3','Fusion minus M1 paired differences','M10, R3'],
      ['Figure 1','Cohort flow','M2, R1'],['Figure 2','ROC, treatment-change PR, calibration','R2-R4'],['Figure 3','Repeated grouped-CV stability','M11, R5'],
      ['Tables S1-S2','Coverage and input definitions','M2, M4-M8'],['Tables S3-S4','Threshold errors and boundary sensitivity','R4-R5'],['Tables S5A-S5B','Repeated-CV distributions','R5'],['Tables S6-S7','One-scan analysis and paired differences','R6'],['Figures S1-S2','Boundary and one-scan analyses','R5-R6']],[1.25,3.7,1.1])
    h(doc,'收到写作要求模板后的执行顺序',2)
    p(doc,'第一步，逐条读取模板，建立“要求 → 适用段落/图表 → 完成状态”的核对清单，调整大纲、篇幅与参考文献格式。涉及作者、单位、资助或声明等事实，使用作者实际提供的信息，不填造内容。')
    p(doc,'第二步，按M1-M12、R1-R6、I1-I4、D1-D8与C1、A1-A4的顺序逐段写作；最终Word仍按Title、Abstract、Introduction、Methods、Results、Discussion、Conclusion、Declarations、References及图表的论文顺序排列。每段英文后接对应中文。')
    p(doc,'第三步，每完成一个阶段核对源数值、引文是否支持陈述、英文与中文是否一致，并总结完成内容、问题及下一步。未完成的事实留在作者核对清单，不以虚构信息补齐。')
    p(doc,'第四步，整篇完成后核对模板、报告规范及图表引用，统一语言和数值，再生成可编辑Word；按需要另导出英文投稿版。当前仅完成写作准备，下一阶段等待作者模板及具体要求。')

    newpage(doc,'14  检索记录与文献来源')
    p(doc,'检索与核对日期：2026年10月29日。以公开网页检索定位出版社、PubMed/PMC及规范原文，使用Europe PMC全文XML读取L1-L4，并计算结构与词数；L5读取出版社全文，L6读取数据描述全文，G1读取全文，G2核对官方完整清单。此为针对性写作准备，不是系统综述，未作穷尽检索或文献质量评分。')
    p(doc,'代表性检索概念：glioblastoma/pseudoprogression、radiomics/clinical combined model、added value、multiparametric MRI validation、UCSD-PTGBM、TRIPOD+AI与CLAIM 2024。选择理由为研究任务相近、全文可读，并同时覆盖融合获益、无获益及简洁模型验证。')
    for ident,citation,doi in refs:
        para=p(doc,f'[{ident}] {citation}')
        para.paragraph_format.space_after=Pt(2)
        para=p(doc,'https://doi.org/'+doi);para.paragraph_format.space_after=Pt(7)
        for r in para.runs:r.font.size=Pt(9.5)


    footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    for e in [*doc.styles.element.xpath('.//w:pBdr'),*doc.element.xpath('.//w:pBdr')]:e.getparent().remove(e)
    path=OUT/'UCSD_Literature_Learning_and_Writing_Outline.docx';doc.save(path)
    receipt={'stage':21,'date':'2026-10-29','deliverable':str(path.relative_to(ROOT)),
        'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'paragraph_ids':35,
        'main_word_target':4180,'abstract_word_target':250,'core_fulltext_word_counts':4,
        'empirical_papers_reviewed':5,'dataset_papers_reviewed':1,'reporting_guidelines':2,
        'manuscript_body_drafted':False,'awaiting_user_template':True,'models_refitted':0,
        'prior_results_replaced':False}
    (REPORT/'generation_receipt.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(receipt,ensure_ascii=False,indent=2))


if __name__=='__main__':build()
