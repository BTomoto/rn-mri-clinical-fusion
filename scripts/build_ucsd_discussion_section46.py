"""Draft D6 from saved sensitivity results and checked primary-source context."""
from pathlib import Path
from copy import deepcopy
import argparse, hashlib, json, re, zipfile
import pandas as pd
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'outputs/ucsd-manuscript-step46-20261030'
REPORT = ROOT/'reports/ucsd_manuscript_step46'
EVIDENCE = REPORT/'literature_evidence.json'
MAPPING = {'13': 'S30', '14': 'S27'}
PARAGRAPHS = [{
    'id': 'D6', 'citation_ids': ['S30', 'S27'],
    'english': 'Boundary perturbation and repeated partitioning probe different sources of sensitivity. For Fusion, the eroded-minus-original AUC difference was 0.0062, with a conditional paired interval spanning zero. Preserving the original normalization and binning references while eroding the ROI suggests limited sensitivity to that specific boundary change under the fixed workflow. The result does not demonstrate equivalence or segmentation accuracy. Martin et al. examined scan-rescan repeatability of diffusion and perfusion radiomics (13), a measurement property that our erosion analysis does not test. Repeated patient-grouped validation instead examined sensitivity to the allocation of patients across training folds. The Fusion-minus-M1 AUC difference remained positive in all 20 partitions, supporting directional consistency within the available cohort. Because the same patients and labels were reused, the repetitions were neither independent replications nor external validation. Patient grouping addresses the dependence risk illustrated by Rosenblatt et al. using duplicated scans in connectome prediction (14), although that experiment does not quantify bias in our diagnostic task. Together, the analyses strengthen the internal sensitivity assessment while leaving uncertainty about source annotations, reference diagnoses and performance at other institutions unresolved. The original ROI and prespecified primary partition remain the basis of the main comparison.',
    'chinese': '边界扰动和重复分折考察的是不同的敏感性来源。对于Fusion，腐蚀后减去原始ROI的AUC差值为0.0062，其条件性配对区间跨零。在保留原归一化及分箱参考的同时腐蚀ROI，提示在固定流程下，结果对这一特定边界变化的敏感性有限；该结果不能证明等价性或分割准确性。Martin等考察了扩散和灌注影像组学的扫描—重扫描可重复性（13），这是我们的腐蚀分析未检验的一种测量属性。重复患者分组验证则考察患者在训练折间分配所带来的敏感性。全部20套分折的Fusion减去M1的AUC差值均为正，支持现有队列内增量方向的一致性。由于重复使用了相同患者和标签，这些重复并非独立重复研究或外部验证。患者分组有助于应对Rosenblatt等通过连接组预测中的扫描复制实验所展示的依赖性风险（14），但该实验不能量化我们诊断任务中的偏倚。综合来看，这些分析加强了内部敏感性评价，但源标注、参考诊断以及其他机构中的表现仍存在未解决的不确定性。主要比较仍以原始ROI和预先指定的主分折为依据。'
}]

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def wc(text):
    return len(re.findall(r"[A-Za-z0-9]+(?:['/-][A-Za-z0-9]+)*", text))

def audit():
    previous = json.loads((ROOT/'reports/ucsd_manuscript_step45/discussion_section45.json').read_text())
    frozen = previous['source_audit']['frozen_file_sha256'].copy()
    frozen.update(json.loads((ROOT/'reports/ucsd_manuscript_step45/artifact_verification.json').read_text())['sha256'])
    extra = ['reports/ucsd_manuscript_step31/methods_section29.json', 'reports/ucsd_manuscript_step38/results_section36.json',
             'reports/ucsd_manuscript_step39/results_section37.json']
    frozen.update({p: sha(ROOT/p) for p in extra if p not in frozen})
    frozen[str(EVIDENCE.relative_to(ROOT))] = sha(EVIDENCE)
    for p, h in frozen.items(): assert sha(ROOT/p) == h, p
    evidence = json.loads(EVIDENCE.read_text()); records = {r['reference_id']: r for r in evidence['studies']}
    refs = {r['id']: r for r in json.loads((ROOT/'reports/ucsd_manuscript_step23/reference_library.json').read_text())['references']}
    selected = []
    for number, rid in MAPPING.items():
        ref = deepcopy(refs[rid]); r = records[rid]
        assert ref['doi'] == r['doi'] and r['provisional_citation_number'] == int(number)
        assert ref['english_flag'] == ref['recent_5yr'] == ref['recent_3yr'] == 1
        if rid == 'S30':
            ref['pages_or_article'] = r['facts']['full_pagination']; ref['citation'] = ref['citation'].replace(':1691.', ':1691-1702.')
            assert ref['pages_or_article'] == '1691-1702' and ':1691-1702.' in ref['citation']
        selected.append({'number': int(number), 'reference_id': rid, 'citation': ref['citation'], 'doi': ref['doi'],
                         'eligibility_year': ref['eligibility_year'], 'read_level_this_stage': r['read_level_this_stage']})
    main = json.loads((ROOT/'reports/ucsd_fusion/metrics.json').read_text())
    boundary = next(r for r in main['paired_comparisons'] if r['comparison'] == 'Fusion_eroded_minus_Fusion' and r['metric'] == 'roc_auc')
    assert f"{boundary['difference']:.4f}" == '0.0062' and boundary['lower'] < 0 < boundary['upper']
    assert main['bootstrap_models_refitted'] == 0 and not main['independent_external_validation']
    code = (ROOT/'src/rn_public/ucsd_batch_features.py').read_text()
    assert 'structure=ndimage.generate_binary_structure(3,1),iterations=1' in code
    assert 'calculate(eroded,affine_values(np.asarray(image)[eroded],params))' in code
    assert code.index('quantized,lo,hi=fixed_bins(v)') < code.index('def calculate(mask,values):')
    assert 'bins[roi]=quantized' in code and 'counts=np.bincount(bins[mask]' in code
    stable = json.loads((ROOT/'reports/ucsd_stability/metrics.json').read_text())
    differences = pd.read_csv(ROOT/'reports/ucsd_stability/repeated_paired_differences.csv')
    auc = differences.loc[differences.metric.eq('roc_auc') & differences.comparison.eq('Fusion_minus_M1')]
    assert len(auc) == auc['repeat'].nunique() == stable['repeats_complete'] == 20 and auc.difference.gt(0).all()
    summary = next(r for r in stable['repeated_primary_difference_distribution'] if r['metric'] == 'roc_auc')
    assert summary['n_complete'] == summary['positive_count'] == 20 and summary['minimum'] > 0
    assert not stable['repeats_are_independent_samples'] and not stable['repeated_ranges_are_confidence_intervals']
    assert not stable['prior_primary_analysis_replaced'] and not stable['external_validation']
    prior_boundary = json.loads((ROOT/extra[1]).read_text())['source_audit']
    prior_repeat = json.loads((ROOT/extra[2]).read_text())['source_audit']
    assert prior_boundary['eroded_complete_cases_checked'] == 182 and prior_boundary['paired_intervals_include_zero']
    assert prior_repeat['folds_checked'] == 100 and prior_repeat['saved_successful_fold_model_records_checked'] == 400
    assert prior_repeat['repetitions_checked'] == 20 and prior_repeat['descriptive_distributions_not_intervals']
    assert records['S30']['read_level_this_stage'].startswith('Primary publisher abstract')
    assert records['S27']['facts']['subject_leakage_experiment'].startswith('Duplicated the exact same scans')
    prior_ids = set()
    for pattern in ['methods_section*.json', 'results_section*.json', 'discussion_section*.json']:
        for p in ROOT.glob('reports/ucsd_manuscript_step*/'+pattern):
            if int(p.parent.name.rsplit('step', 1)[1]) < 46:
                prior_ids.update(r for item in json.loads(p.read_text())['paragraphs'] for r in item['citation_ids'])
    assert len(prior_ids) == 12 and not prior_ids.intersection(MAPPING.values())
    return {'frozen_file_sha256': frozen, 'reference_records': selected, 'source_evidence_record': str(EVIDENCE.relative_to(ROOT)),
            'new_full_texts_checked_for_current_claims': ['S27'], 'limited_abstract_claim_sources': ['S30'],
            'paired_boundary_AUC_record_checked': boundary, 'within_partition_AUC_differences_checked': 20,
            'historical_grouped_fold_checks_confirmed': 100, 'historical_successful_fold_model_records_confirmed': 400,
            'original_normalization_and_discretization_references_retained': True,
            'equivalence_or_segmentation_accuracy_claimed': False, 'acquisition_repeatability_established': False,
            'repeated_partitions_claimed_independent': False, 'source_label_validity_or_external_performance_established': False,
            'S30_complete_pagination_used': '1691-1702', 'previous_unique_cited_sources': sorted(prior_ids),
            'models_refitted': 0, 'features_reextracted': 0, 'new_hypothesis_tests': 0,
            'verification_scope': 'Inspect frozen erosion rules, paired AUC record, all saved repeated AUC differences and prior fold audits, plus manually reviewed source claims; no retraining, resampling or re-extraction'}

def build():
    a = audit(); OUT.mkdir(parents=True, exist_ok=True); REPORT.mkdir(parents=True, exist_ok=True)
    doc = Document(); sec = doc.sections[0]; sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(.79)
    for name in ['Normal', 'Title', 'Heading 1']:
        st = doc.styles[name]; st.font.name = 'Times New Roman'; st.font.color.rgb = RGBColor(0, 0, 0)
        rf = st.element.get_or_add_rPr().rFonts; rf.set(qn('w:eastAsia'), 'Noto Serif CJK SC')
        for key in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']: rf.attrib.pop(qn('w:'+key), None)
    normal = doc.styles['Normal']; normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.; normal.paragraph_format.space_after = Pt(9)
    doc.styles['Title'].font.size = Pt(16); doc.styles['Title'].paragraph_format.space_after = Pt(12)
    heading = doc.styles['Heading 1']; heading.font.size = Pt(13)
    heading.paragraph_format.space_before = Pt(8); heading.paragraph_format.space_after = Pt(9)
    doc.core_properties.title = 'Discussion 4.6 Boundary and split robustness'; doc.core_properties.author = 'Tianyu Feng'
    doc.add_paragraph('Discussion', style='Title'); doc.add_paragraph('4.6 Boundary and split robustness', style='Heading 1')
    for item in PARAGRAPHS:
        en = doc.add_paragraph(item['english']); en.paragraph_format.keep_together = True; en.paragraph_format.keep_with_next = True
        zh = doc.add_paragraph(item['chinese']); zh.paragraph_format.keep_together = True
        for run in zh.runs: run.font.size = Pt(11)
    title = doc.add_paragraph('References', style='Heading 1'); title.paragraph_format.page_break_before = True
    for ref in a['reference_records']:
        p = doc.add_paragraph(str(ref['number'])+'. '+ref['citation']); p.paragraph_format.keep_together = True
    for e in [*doc.styles.element.xpath('.//w:pBdr'), *doc.element.xpath('.//w:pBdr')]: e.getparent().remove(e)
    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    path = OUT/'UCSD_Manuscript_Discussion_4_6_EN_ZH.docx'; doc.save(path)
    (OUT/'UCSD_Manuscript_Discussion_4_6_EN.txt').write_text('\n\n'.join(p['english'] for p in PARAGRAPHS)+'\n')
    d = {'stage': 46, 'date': '2026-10-30', 'subsection': '4.6 Boundary and split robustness', 'v2_paragraph_ids': ['D6'],
         'paragraphs': PARAGRAPHS, 'english_word_count': sum(wc(p['english']) for p in PARAGRAPHS),
         'paragraph_word_counts': {p['id']: wc(p['english']) for p in PARAGRAPHS},
         'word_count_scope': 'English narrative including in-text citation numbers; excludes headings, reference list and Chinese translation',
         'v2_approximate_target_words': 200, 'equations': {'count': 0, 'numbers': [], 'cumulative_previous_numbered_groups': 10},
         'citation_mapping': MAPPING, 'citation_numbering': 'First-appearance numbers 13-14; renumber globally when Introduction is assembled',
         'new_manuscript_references': 2, 'cumulative_unique_cited_sources': 14,
         'source_audit': a, 'models_refitted': 0, 'features_reextracted': 0, 'new_hypothesis_tests': 0,
         'next_section': 'Discussion 4.7 Meaning of the single-scan result (D7)',
         'docx': str(path.relative_to(ROOT)), 'docx_sha256': sha(path)}
    (REPORT/'discussion_section46.json').write_text(json.dumps(d, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps({'docx': d['docx'], 'english_word_count': d['english_word_count'], 'frozen_files': len(a['frozen_file_sha256'])}))

def verify():
    from pypdf import PdfReader
    d = json.loads((REPORT/'discussion_section46.json').read_text()); a = audit(); assert a == d['source_audit']
    path = ROOT/d['docx']; assert sha(path) == d['docx_sha256']; doc = Document(path); texts = [p.text for p in doc.paragraphs]
    for p in PARAGRAPHS:
        assert p['english'] in texts and p['chinese'] in texts and texts.index(p['chinese']) == texts.index(p['english'])+1
        assert not re.search(r'[\u3400-\u9fff，。；：（）【】、]', p['english'])
    assert sum(wc(p['english']) for p in PARAGRAPHS) == d['english_word_count']
    assert re.findall(r'\((\d+)\)', PARAGRAPHS[0]['english']) == list(MAPPING)
    for ref in a['reference_records']: assert str(ref['number'])+'. '+ref['citation'] in texts
    assert not doc.tables and len(doc.sections) == 1
    with zipfile.ZipFile(path) as z:
        xml = etree.fromstring(z.read('word/document.xml'))
        assert not xml.xpath('//m:oMath', namespaces={'m': 'http://schemas.openxmlformats.org/officeDocument/2006/math'})
        assert not [n for n in z.namelist() if n.startswith('word/media/')]
    pdf = PdfReader(OUT/'docx_render/UCSD_Manuscript_Discussion_4_6_EN_ZH.pdf'); pages = len(pdf.pages); assert pages == 2
    assert '4.6 Boundary' in pdf.pages[0].extract_text() and 'References' in pdf.pages[1].extract_text()
    v = {'success': True, 'stage': 46, 'date': '2026-10-30', 'subsection': d['subsection'], 'paragraph_ids': ['D6'],
         'english_word_count': d['english_word_count'], 'matched_chinese_paragraphs': 1, 'word_pages': pages,
         'visual_review_pages': [1, 2], 'native_editable_equations': 0, 'equation_images': 0,
         'frozen_files_unchanged': len(a['frozen_file_sha256']), 'prior_files_unchanged': len(a['frozen_file_sha256'])-1,
         'new_full_texts_checked_for_current_claims': ['S27'], 'limited_abstract_claim_sources': ['S30'],
         'paired_boundary_AUC_record_checked': True, 'within_partition_AUC_differences_checked': 20,
         'historical_grouped_fold_checks_confirmed': 100, 'new_manuscript_references': 2,
         'provisional_citation_mapping': MAPPING, 'cumulative_unique_cited_sources': 14,
         'equivalence_or_segmentation_accuracy_claimed': False, 'repeated_partitions_claimed_independent': False,
         'models_refitted': 0, 'features_reextracted': 0, 'new_hypothesis_tests': 0, 'next_section': d['next_section'],
         'sha256': {str(p.relative_to(ROOT)): sha(p) for p in OUT.iterdir() if p.is_file() and p.suffix in {'.docx', '.txt'}}}
    (REPORT/'artifact_verification.json').write_text(json.dumps(v, ensure_ascii=False, indent=2)+'\n'); print(json.dumps(v, indent=2))

if __name__ == '__main__':
    parser = argparse.ArgumentParser(); parser.add_argument('--verify', action='store_true'); args = parser.parse_args()
    verify() if args.verify else build()
