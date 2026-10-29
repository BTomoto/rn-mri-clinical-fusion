"""Draft R4 with paired fixed-OOF differences and an editable English Table 3."""
from pathlib import Path
import argparse,hashlib,json,re,sys,zipfile
import numpy as np
import pandas as pd
from lxml import etree
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'outputs/ucsd-manuscript-step36-20261029'
REPORT=ROOT/'reports/ucsd_manuscript_step36'
PARAGRAPHS=[
 {'id':'R4','citation_ids':[],
  'english': 'For the prespecified primary comparison, the Fusion-minus-M1 ROC AUC difference was 0.110 (95% paired patient-bootstrap percentile interval, 0.021 to 0.214; Table 3). The treatment-related-change AP difference was 0.152 (0.026 to 0.268). Both intervals lay above zero under the conditional bootstrap analysis. Probability-error differences also favored Fusion in their point estimates: the Brier-score difference was -0.019 (-0.040 to 0.002), and the logarithmic-loss difference was -0.037 (-0.107 to 0.039). However, both probability-error intervals included zero, leaving these improvements uncertain. ROC AUC was the sole primary incremental metric; AP, Brier score and logarithmic loss were auxiliary summaries. Differences were calculated from the unrounded model metrics, rather than from the displayed three-decimal values. All contrasts used the same 182 examinations, and each bootstrap attempt applied the same resampled patient groups to both models. The resulting intervals described uncertainty in the paired differences conditional on the saved OOF predictions. They were not obtained by subtracting separate model interval bounds and did not include full retraining or partition variability. The primary result supported an AUC increment from adding the eight MRI features to the fixed clinical model within this cohort, while the probability-error contrasts remained uncertain.',
  'chinese': '在预设主要比较中，Fusion减去M1的ROC AUC差值为0.110（95%配对患者bootstrap百分位区间0.021至0.214；Table 3）。治疗相关改变AP差值为0.152（0.026至0.268）。在该条件bootstrap分析中，两项区间均位于零以上。概率误差差值的点估计也有利于Fusion：Brier评分差值为−0.019（−0.040至0.002），对数损失差值为−0.037（−0.107至0.039）。但两项概率误差区间均包含零，因此这些改善仍存在不确定性。ROC AUC是唯一主要增量指标，AP、Brier评分和对数损失为辅助汇总。差值由未四舍五入的模型指标计算，而非由显示为三位小数的数值相减。全部比较使用相同的182次检查，每次bootstrap对两个模型采用相同的重采样患者组。所得区间表示条件于已保存OOF预测的配对差值不确定性，不是通过两个模型各自的区间端点相减得到，也不包含完整重训或分折变异。主要结果支持在本队列内向固定临床模型加入八项MRI特征后获得AUC增量，而概率误差方面的差异仍不确定。'}
]
METRICS=['roc_auc','ap_treatment_change','brier','log_loss']
METRIC_LABELS=['ROC AUC','Treatment-related-change AP','Brier score','Logarithmic loss']
TABLE_SPEC={'key':'Table3','title':'Table 3. Paired incremental performance of Fusion versus M1',
  'headers':['Metric','Fusion minus M1','95% paired interval','Role'],
  'notes':['All differences are Fusion minus M1 on the same 182 examinations from 135 patients. Positive AUC or AP differences favor Fusion; negative Brier-score or logarithmic-loss differences favor Fusion.',
  'Intervals are the 2.5th and 97.5th percentiles of 2000 paired patient-cluster bootstrap attempts. Both models use identical resampled groups and fixed OOF probabilities without refitting. The intervals do not include full retraining or partition variability.',
  'ROC AUC is the sole primary incremental metric; other metrics are auxiliary summaries. AP denotes average precision for treatment-related change. Differences are calculated from unrounded metrics and displayed to three decimals.']}

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def wc(s):return len(re.findall(r"[A-Za-z0-9]+(?:['/-][A-Za-z0-9]+)*",s))
def close(a,b):assert np.isclose(a,b,rtol=1e-12,atol=1e-12),(a,b)

def audit():
    old=json.loads((ROOT/'reports/ucsd_manuscript_step35/results_section33.json').read_text());frozen=old['source_audit']['frozen_file_sha256'].copy()
    frozen.update(json.loads((ROOT/'reports/ucsd_manuscript_step35/artifact_verification.json').read_text())['sha256'])
    for p,h in frozen.items():assert sha(ROOT/p)==h,p
    extra=['outputs/ucsd-paper-materials-20261029/tables/Table3.csv','reports/ucsd_fusion/paired_comparisons.csv']
    frozen.update({p:sha(ROOT/p) for p in extra})
    pred=pd.read_csv(ROOT/'data/processed/ucsd_fusion_oof.csv');y=pred.tumor_label.to_numpy()
    saved=json.loads((ROOT/'reports/ucsd_fusion/metrics.json').read_text());boot=pd.read_csv(ROOT/'reports/ucsd_fusion/paired_patient_bootstrap.csv')
    paircsv=pd.read_csv(ROOT/'reports/ucsd_fusion/paired_comparisons.csv');paircsv=paircsv[paircsv.comparison.eq('Fusion_minus_M1')].set_index('metric')
    table=pd.read_csv(ROOT/'outputs/ucsd-paper-materials-20261029/tables/Table3.csv')
    assert len(pred)==182 and pred.validation_group.nunique()==135 and len(paircsv)==4 and table.shape==(4,4)
    assert saved['same_cases_and_folds'] and saved['bootstrap_attempts']==saved['bootstrap_valid']==2000 and saved['bootstrap_models_refitted'] is False
    assert boot.shape==(2000,46)
    primary=[r for r in saved['paired_comparisons'] if r['primary_comparison']];assert len(primary)==1 and primary[0]['comparison']=='Fusion_minus_M1' and primary[0]['metric']=='roc_auc'
    def manual_metrics(p):
        diff=p[y==1][:,None]-p[y==0][None,:];auc=float(np.mean((diff>0)+.5*(diff==0)))
        labels=1-y;score=1-p;order=np.argsort(-score,kind='stable');labels=labels[order];score=score[order]
        ends=np.flatnonzero(np.r_[score[:-1]!=score[1:],True]);tp=np.cumsum(labels)[ends]
        ap=float(np.sum(np.diff(np.r_[0,tp])/labels.sum()*tp/(ends+1)))
        return dict(roc_auc=auc,ap_treatment_change=ap,brier=float(np.mean((p-y)**2)),log_loss=float(-np.mean(y*np.log(p)+(1-y)*np.log1p(-p))))
    model_values={name:manual_metrics(pred[name+'_oof_probability'].to_numpy()) for name in ['M1','Fusion']}
    records=[];rows=[];numeric_checks=0
    for i,metric in enumerate(METRICS):
        for name in ['M1','Fusion']:close(model_values[name][metric],saved['metrics'][name][metric])
        difference=model_values['Fusion'][metric]-model_values['M1'][metric]
        key='Fusion_minus_M1__'+metric
        np.testing.assert_allclose(boot[key],boot['Fusion__'+metric]-boot['M1__'+metric],rtol=1e-12,atol=1e-12)
        lower=float(boot[key].quantile(.025));upper=float(boot[key].quantile(.975))
        summary=next(r for r in saved['paired_comparisons'] if r['comparison']=='Fusion_minus_M1' and r['metric']==metric)
        for k,v in [('difference',difference),('lower',lower),('upper',upper)]:close(v,summary[k]);close(v,paircsv.loc[metric,k])
        for k,v in [('lower',lower),('upper',upper)]:close(v,saved['patient_cluster_bootstrap_95ci'][key][k])
        for k,v in [('差值',difference),('CI_lower',lower),('CI_upper',upper)]:close(v,table.iloc[i][k]);numeric_checks+=1
        assert bool(paircsv.loc[metric,'primary_comparison'])==(metric=='roc_auc')
        original_difference=float(table.iloc[i]['差值']);original_lower=float(table.iloc[i].CI_lower);original_upper=float(table.iloc[i].CI_upper)
        records.append({'Metric':METRIC_LABELS[i],'Difference_Fusion_minus_M1':original_difference,'Interval_lower':original_lower,'Interval_upper':original_upper,'Primary_metric':metric=='roc_auc'})
        rows.append([METRIC_LABELS[i],f'{difference:+.3f}',f'[{lower:.3f}, {upper:.3f}]','Primary' if metric=='roc_auc' else 'Auxiliary'])
        assert lower>0 if metric in ['roc_auc','ap_treatment_change'] else lower<0<upper
    spec=dict(**TABLE_SPEC,rows=rows)
    assert not re.search(r'[\u3400-\u9fff]',' '.join(spec['headers']+sum(rows,[])+spec['notes']))
    return {'frozen_file_sha256':frozen,'translated_table':spec,'full_precision_table_records':records,
      'individual_model_metric_values_reconstructed':8,'paired_point_differences_checked':4,'bootstrap_paired_difference_values_checked':8000,
      'paired_interval_endpoints_checked':8,'original_table_numeric_values_checked':numeric_checks,
      'only_primary_incremental_metric':'roc_auc','intervals_above_zero':['roc_auc','ap_treatment_change'],'intervals_include_zero':['brier','log_loss'],
      'numeric_values_changed':0,'models_refitted':0,'features_reextracted':0,'new_hypothesis_tests':0,
      'verification_scope':'Reconstruct the two models’ four metrics from fixed OOF probabilities, verify all 8000 saved paired differences and eight percentile endpoints, and compare the original Table 3 without new resampling or model fits'}

def add_table(doc,spec):
    title=doc.add_paragraph(spec['title'],style='Heading 1');title.paragraph_format.page_break_before=True
    title.paragraph_format.keep_with_next=True
    table=doc.add_table(rows=1,cols=4);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    widths=[2.14,1.30,2.50,.98]
    for col,width in zip(table.columns,widths):col.width=Inches(width)
    for cell,width in zip(table.rows[0].cells,widths):cell.width=Inches(width)
    for cell,label in zip(table.rows[0].cells,spec['headers']):cell.text=label
    for values in spec['rows']:
        for cell,value,width in zip(table.add_row().cells,values,widths):cell.text=value;cell.width=Inches(width)
    props=table._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side)
        for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:e.set(qn('w:'+k),v)
        borders.append(e)
    props.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for side in ['top','left','bottom','right']:
        e=OxmlElement('w:'+side);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
    props.append(margins)
    for ri,row in enumerate(table.rows):
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if ri==0:trpr.append(OxmlElement('w:tblHeader'))
        for ci,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7EBEF');cell._tc.get_or_add_tcPr().append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:run.font.size=Pt(10.5);run.font.bold=ri==0
    for note in spec['notes']:
        p=doc.add_paragraph(note);p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(4)
        p.paragraph_format.keep_together=True
        for r in p.runs:r.font.size=Pt(10)

def build():
    OUT.mkdir(parents=True,exist_ok=True);REPORT.mkdir(parents=True,exist_ok=True);a=audit()
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(.79)
    for name in ['Normal','Title','Heading 1']:
        st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        rf=st.element.get_or_add_rPr().rFonts;rf.set(qn('w:eastAsia'),'Noto Serif CJK SC')
        for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:rf.attrib.pop(qn('w:'+key),None)
    normal=doc.styles['Normal'];normal.font.size=Pt(12);normal.paragraph_format.line_spacing=1.;normal.paragraph_format.space_after=Pt(9)
    doc.styles['Title'].font.size=Pt(16);doc.styles['Title'].paragraph_format.space_after=Pt(12)
    doc.styles['Heading 1'].font.size=Pt(13);doc.styles['Heading 1'].paragraph_format.space_before=Pt(8);doc.styles['Heading 1'].paragraph_format.space_after=Pt(9)
    doc.core_properties.title='Results 3.4 Paired incremental comparison';doc.core_properties.author='Tianyu Feng'
    doc.add_paragraph('Results',style='Title');doc.add_paragraph('3.4 Paired incremental comparison',style='Heading 1')
    for item in PARAGRAPHS:
        en=doc.add_paragraph(item['english']);en.paragraph_format.page_break_before=False
        en.paragraph_format.keep_together=True;en.paragraph_format.keep_with_next=True
        zh=doc.add_paragraph(item['chinese']);zh.paragraph_format.keep_together=True
        for r in zh.runs:r.font.size=Pt(11)
    add_table(doc,a['translated_table'])
    pd.DataFrame(a['full_precision_table_records']).to_csv(OUT/'Table3_EN.csv',index=False)
    for e in [*doc.styles.element.xpath('.//w:pBdr'),*doc.element.xpath('.//w:pBdr')]:e.getparent().remove(e)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    path=OUT/'UCSD_Manuscript_Results_3_4_EN_ZH.docx';doc.save(path)
    (OUT/'UCSD_Manuscript_Results_3_4_EN.txt').write_text('\n\n'.join(p['english'] for p in PARAGRAPHS)+'\n')
    d={'stage':36,'date':'2026-10-29','subsection':'3.4 Paired incremental comparison','v2_paragraph_ids':['R4'],'paragraphs':PARAGRAPHS,
      'english_word_count':sum(wc(p['english']) for p in PARAGRAPHS),'paragraph_word_counts':{p['id']:wc(p['english']) for p in PARAGRAPHS},
      'word_count_scope':'English narrative only; excludes headings, tables/notes and Chinese translation','v2_approximate_target_words':200,
      'equations':{'count':0,'numbers':[],'cumulative_previous_numbered_groups':10},
      'citation_mapping':{},'new_manuscript_references':0,'cumulative_unique_cited_sources':5,'editable_english_tables':['Table 3'],'table_numeric_values_changed':0,'source_audit':a,'models_refitted':0,'new_hypothesis_tests':0,
      'next_section':'Results 3.5 Calibration and errors (R5)','docx':str(path.relative_to(ROOT)),'docx_sha256':sha(path)}
    (REPORT/'results_section34.json').write_text(json.dumps(d,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'docx':d['docx'],'english_word_count':d['english_word_count'],'paragraph_word_counts':d['paragraph_word_counts'],'prior_files':len(a['frozen_file_sha256'])},indent=2))

def verify():
    from pypdf import PdfReader
    d=json.loads((REPORT/'results_section34.json').read_text())
    for p,h in d['source_audit']['frozen_file_sha256'].items():assert sha(ROOT/p)==h,p
    path=ROOT/d['docx'];assert sha(path)==d['docx_sha256'];texts=[p.text for p in Document(path).paragraphs]
    for p in PARAGRAPHS:
        assert p['english'] in texts and p['chinese'] in texts and texts.index(p['chinese'])==texts.index(p['english'])+1
        assert not re.search(r'[\u3400-\u9fff，。；：（）【】、]',p['english'])
    assert sum(wc(p['english']) for p in PARAGRAPHS)==d['english_word_count']
    doc=Document(path);assert len(doc.tables)==1
    spec=d['source_audit']['translated_table'];expected=[spec['headers'],*spec['rows']]
    assert [[c.text for c in row.cells] for row in doc.tables[0].rows]==expected
    assert all(not re.search(r'[\u3400-\u9fff]',c.text) for row in doc.tables[0].rows for c in row.cells)
    for note in spec['notes']:assert note in texts
    full=pd.read_csv(OUT/'Table3_EN.csv');assert len(full)==4
    for (_,row),r in zip(full.iterrows(),d['source_audit']['full_precision_table_records']):
        assert row.Metric==r['Metric']
        for key,value in r.items():
            if key!='Metric':close(row[key],value)
    assert len(doc.sections)==1
    ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
    with zipfile.ZipFile(path) as z:
        t=etree.fromstring(z.read('word/document.xml'));assert len(t.xpath('//m:oMath',namespaces=ns))==0
        assert not any(n.startswith('word/media/') for n in z.namelist())
    pages=len(PdfReader(OUT/'docx_render/UCSD_Manuscript_Results_3_4_EN_ZH.pdf').pages);assert pages==2
    v={'success':True,'stage':36,'date':'2026-10-29','subsection':d['subsection'],'paragraph_ids':['R4'],
      'english_word_count':d['english_word_count'],'paragraph_word_counts':d['paragraph_word_counts'],'matched_chinese_paragraphs':1,
      'word_pages':pages,'visual_review_pages':[1,2],'native_editable_equations':0,'equation_numbers':[],'equation_images':0,
      'prior_files_unchanged':len(d['source_audit']['frozen_file_sha256']),'original_table_numeric_values_checked':12,'individual_model_metric_values_reconstructed':8,'paired_point_differences_checked':4,'bootstrap_paired_difference_values_checked':8000,'paired_interval_endpoints_checked':8,'editable_english_tables':1,'table_numeric_values_changed':0,'probability_error_intervals_include_zero_disclosed':True,
      'models_refitted':0,'new_hypothesis_tests':0,'new_manuscript_references':0,'cumulative_unique_cited_sources':5,'next_section':d['next_section'],
      'sha256':{str(p.relative_to(ROOT)):sha(p) for p in OUT.iterdir() if p.is_file() and p.suffix in {'.docx','.txt','.csv'}}}
    (REPORT/'artifact_verification.json').write_text(json.dumps(v,ensure_ascii=False,indent=2)+'\n');print(json.dumps(v,indent=2))

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--verify',action='store_true');args=parser.parse_args()
    verify() if args.verify else build()
