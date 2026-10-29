"""Draft M17-M19 from frozen sensitivity protocols and read-only result audits."""
from pathlib import Path
import argparse,hashlib,json,re,sys,zipfile
import numpy as np
import pandas as pd
from lxml import etree
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'outputs/ucsd-manuscript-step31-20261029'
REPORT=ROOT/'reports/ucsd_manuscript_step31'
PARAGRAPHS=[
 {'id':'M17','citation_ids':[],
  'english': 'Sensitivity to the annotation boundary was examined by eroding the full ROI once using a three-dimensional six-neighbor structuring element. Erosion operated on the supplied voxel grid and did not represent a fixed distance in millimeters. We recalculated the same eight MRI features within the eroded region, retaining the original image-specific percentile normalization parameters and the discretization boundaries derived from the full ROI. This kept the intensity and binning reference consistent when changing the region. Eroded MRI and Fusion models used the original examination set, patient groups and five-fold assignments, with their preprocessing and regression fitted within each training fold under the same fixed settings. The clinical inputs were unchanged. Comparisons included eroded Fusion versus M1 and eroded versus original Fusion, using paired patient-cluster bootstrap intervals conditional on the saved OOF predictions. This perturbation assessed dependence on the supplied boundary; it did not establish segmentation accuracy. The original ROI remained the primary analysis.',
  'chinese': '采用三维六邻域结构元素，对完整ROI执行一次腐蚀，以检查对标注边界的敏感性。腐蚀在所提供的体素网格上进行，不代表以毫米表示的固定距离。在腐蚀后区域内重新计算相同的八项MRI特征，保留原图像特异的百分位归一化参数，以及由完整ROI确定的离散化边界。这样，在改变区域时保持强度与分箱的参考一致。腐蚀后MRI与Fusion模型使用原有检查集合、患者分组及五折划分，并在相同固定设置下于各训练折内拟合预处理和回归。临床输入保持不变。比较包括腐蚀后Fusion与M1，以及腐蚀后Fusion与原始Fusion；采用条件于已保存OOF预测的配对患者聚类bootstrap区间。该扰动用于评估结果对所提供边界的依赖，不用于证明分割准确性。完整ROI分析仍为主分析。'},
 {'id':'M18','citation_ids':[],
  'english': 'To examine sensitivity to fold assignment, we repeated stratified five-fold validation by patient group on the full imaging cohort using 20 seeds fixed as the consecutive integers 20261029 through 20261048. The first partition reproduced the primary analysis. Within each repetition, M0, M1, MRI and Fusion shared the same examinations and folds, and preprocessing and regression were fitted separately within each training fold. Each examination received one OOF prediction per model per repetition. We summarized performance and the Fusion-minus-M1 differences using the median, quartiles and observed minimum and maximum, and counted repetitions with positive differences. The protocol required recording any failed repetitions without replacing their seeds. Because all repetitions reused the same patients, their distributions described sensitivity to partitioning and were not confidence intervals or independent replications. We did not pool repeated predictions as independent observations, average them into an ensemble or choose the highest-performing repetition to replace the primary estimate. Model settings and the classification threshold remained fixed.',
  'chinese': '为检查分折敏感性，在完整影像队列上按患者组重复进行分层五折验证，采用预先固定的20个连续整数种子20261029至20261048。第一套划分复现主分析。每次重复中，M0、M1、MRI与Fusion共用相同检查及分折，各训练折内分别拟合预处理和回归。每次检查在每个模型、每次重复中获得一次OOF预测。使用中位数、四分位数、实际最小值和最大值概括性能及Fusion减去M1的差值，并统计差值为正的重复次数。方案要求在重复失败时保留记录，不更换种子。由于全部重复使用同一批患者，其分布描述对划分的敏感性，不是置信区间或独立重复研究。不将重复预测合并为独立观测，不取平均构建集成模型，也不选择表现最好的重复替代主估计。模型设置与分类阈值保持固定。'},
 {'id':'M19','citation_ids':[],
  'english': 'A separate sensitivity analysis retained one examination per patient, selecting the smallest numeric session suffix available in the imaging subset, with the complete examination identifier resolving ties. Selection used neither diagnostic labels nor model predictions. This rule identified the earliest numbered available session, without establishing the first clinical visit or the earliest MRI in the source database. We rebuilt patient-grouped five-fold partitions with seed 20261029 and refitted the four fixed models within these training folds. Patient bootstrap sampling used 2000 attempts with seed 20261101 and fixed OOF probabilities. Percentile intervals for ROC AUC, treatment-related-change AP, Brier score, logarithmic loss and their Fusion-minus-M1 differences followed the conditional interpretation above. Draws containing only one diagnostic class were ineligible and would not be replaced. The selected cohort differed from the primary cohort in examination composition, class distribution and training data. Accordingly, differences between their performance estimates could not be attributed solely to removing repeated scans, and the single-examination analysis did not replace the primary comparison.',
  'chinese': '另开展每患者仅保留一次检查的敏感性分析：选择影像子集中可用的最小数字会话后缀；如有并列，按完整检查标识符排序决定。选择不使用诊断标签或模型预测。该规则确定的是可用会话中编号最早的一次，不能据此认定为首次临床就诊或源数据库中的最早MRI。采用种子20261029重建患者分组五折，并在这些训练折内重新拟合四个固定模型。患者bootstrap使用种子20261101进行2000次抽样尝试，OOF概率保持固定。ROC AUC、治疗相关改变AP、Brier评分、对数损失及其Fusion减去M1差值的百分位区间，沿用前述条件性解释。仅含一种诊断类别的抽样不符合条件，且不追加抽样替代。所选队列与主队列在检查构成、类别分布及训练数据方面不同。因此，两者性能估计的差异不能仅归因于去除重复扫描，单次检查分析也不替代主要比较。'}
]

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def wc(s):return len(re.findall(r"[A-Za-z0-9]+(?:['/-][A-Za-z0-9]+)*",s))
def close(a,b):assert np.isclose(a,b,rtol=1e-12,atol=1e-12),(a,b)

def audit():
    old=json.loads((ROOT/'reports/ucsd_manuscript_step30/methods_section28.json').read_text());frozen=old['source_audit']['frozen_file_sha256'].copy()
    frozen.update(json.loads((ROOT/'reports/ucsd_manuscript_step30/artifact_verification.json').read_text())['sha256'])
    for p,h in frozen.items():assert sha(ROOT/p)==h,p
    directory=ROOT/'reports/ucsd_stability'
    extra=['src/rn_public/ucsd_stability.py','reports/ucsd_stability_protocol.md','data/processed/ucsd_stability_repeated_oof.csv','data/processed/ucsd_stability_single_scan_oof.csv']
    extra += [str(p.relative_to(ROOT)) for p in directory.rglob('*') if p.is_file() and p.suffix in {'.csv','.json'}]
    frozen.update({p:sha(ROOT/p) for p in extra})
    lock=json.loads((directory/'analysis_lock.json').read_text());summary=json.loads((directory/'metrics.json').read_text())
    assert lock['seeds']==list(range(20261029,20261049)) and lock['single_scan_seed']==20261029 and lock['single_scan_bootstrap_seed']==20261101
    assert sha(ROOT/'src/rn_public/ucsd_stability.py')==lock['analysis_code_sha256']
    assert sha(ROOT/'reports/ucsd_stability_protocol.md')==lock['protocol_sha256']
    for p,h in lock['manifest_sha256'].items():assert sha(directory/p)==h,p
    for p,h in lock['input_sha256'].items():assert sha(ROOT/p)==h,p
    assert summary['repeats_complete']==20 and summary['repeats_failed']==summary['model_fold_failure_records']==0
    assert not any(summary[k] for k in ['repeats_are_independent_samples','repeated_ranges_are_confidence_intervals','probability_ensemble_created','bootstrap_models_refitted','prior_primary_analysis_replaced'])
    assert json.loads((directory/'failures.json').read_text())['records']==[]
    models=['M0','M1','MRI','Fusion'];main=pd.read_csv(ROOT/'data/processed/ucsd_fusion_oof.csv')
    cohort=pd.read_csv(ROOT/'data/processed/ucsd_clinical_image_features.csv').sort_values('ID').reset_index(drop=True)
    repeated=pd.read_csv(ROOT/'data/processed/ucsd_stability_repeated_oof.csv');manifest=pd.read_csv(directory/'repeated_fold_manifest.csv')
    assert len(manifest)==len(repeated)==3640
    # Check every saved split, its class availability and patient separation; do not regenerate or fit.
    folds_checked=0;fold_models_checked=0
    def check_plan(record,part):
        nonlocal folds_checked,fold_models_checked
        ids=set(part.ID);all_test=[]
        for f in record['folds']:
            tr=set(f['train_IDs']);te=set(f['test_IDs']);assert not tr&te and tr|te==ids
            assert set(part.loc[part.validation_fold.eq(f['fold']),'ID'])==te
            train=part[part.ID.isin(tr)];test=part[part.ID.isin(te)]
            for key in ['patient_id','brats_patient_id','validation_group']:assert not set(train[key])&set(test[key])
            assert train.tumor_label.nunique()==test.tumor_label.nunique()==2 and f['valid'] and f['errors']==[]
            assert set(f['models'])==set(models) and all(v['success'] for v in f['models'].values())
            folds_checked+=1;fold_models_checked+=len(f['models']);all_test.extend(te)
        assert len(all_test)==len(ids) and set(all_test)==ids
    metric_table=pd.read_csv(directory/'repeated_metrics.csv');delta=pd.read_csv(directory/'repeated_paired_differences.csv')
    auc_checks=0
    for rep,seed in enumerate(lock['seeds'],1):
        part=manifest[manifest['repeat'].eq(rep)];pred=repeated[repeated['repeat'].eq(rep)]
        assert len(part)==len(pred)==182 and part.seed.eq(seed).all() and pred.seed.eq(seed).all()
        assert not pred.ID.duplicated().any() and set(pred.ID)==set(main.ID)
        pd.testing.assert_frame_equal(part.reset_index(drop=True),pred[part.columns].reset_index(drop=True))
        a=json.loads((directory/f'repeated_fold_audits/repeat_{rep:02d}.json').read_text())
        assert a['complete'] and a['errors']==[] and a['repeat']==rep and a['seed']==seed;check_plan(a,part)
        y=pred.tumor_label.to_numpy()
        for name in models:
            p=pred[name+'_oof_probability'].to_numpy();assert np.isfinite(p).all()
            diff=p[y==1][:,None]-p[y==0][None,:];auc=float(np.mean((diff>0)+.5*(diff==0)))
            row=metric_table[metric_table['repeat'].eq(rep)&metric_table.model.eq(name)].iloc[0];close(auc,row.roc_auc);auc_checks+=1
        if rep==1:
            assert pred.validation_fold.tolist()==main.set_index('ID').loc[pred.ID,'validation_fold'].tolist()
            for name in models:np.testing.assert_allclose(pred[name+'_oof_probability'],main.set_index('ID').loc[pred.ID,name+'_oof_probability'],rtol=1e-12,atol=1e-12)
        by=metric_table[metric_table['repeat'].eq(rep)].set_index('model')
        for _,row in delta[delta['repeat'].eq(rep)].iterrows():close(row.difference,by.loc['Fusion',row.metric]-by.loc['M1',row.metric])
    # Verify saved descriptive summaries, including unfavorable signs; no seed selection.
    summary_rows=0
    for filename,source in [('repeated_metric_distribution.csv',metric_table),('repeated_delta_distribution.csv',delta)]:
        for _,row in pd.read_csv(directory/filename).iterrows():
            if filename.startswith('repeated_metric'):values=source.loc[source.model.eq(row.model),row.metric]
            else:values=source.loc[source.metric.eq(row.metric),'difference']
            assert row.n_complete==len(values)==20
            for key,value in [('median',values.median()),('q25',values.quantile(.25)),('q75',values.quantile(.75)),('minimum',values.min()),('maximum',values.max())]:close(row[key],value)
            assert row.positive_count==values.gt(0).sum() and row.zero_count==values.eq(0).sum() and row.negative_count==values.lt(0).sum()
            assert row.interpretation=='descriptive_split_distribution_not_confidence_interval';summary_rows+=1
    # Independently reconstruct the selection ledger from identifiers alone.
    identifiers=cohort[['ID','patient_id']].copy();identifiers['session']=identifiers.ID.str.rsplit('_',n=1).str[-1].astype(int)
    chosen=set(identifiers.sort_values(['patient_id','session','ID']).drop_duplicates('patient_id').ID)
    ledger=pd.read_csv(directory/'single_scan_selection.csv');assert len(ledger)==182 and len(chosen)==135
    assert set(ledger.loc[ledger.selected,'ID'])==chosen and set(ledger.ID)==set(cohort.ID)
    single=pd.read_csv(ROOT/'data/processed/ucsd_stability_single_scan_oof.csv');sm=pd.read_csv(directory/'single_scan_fold_manifest.csv')
    assert len(single)==135 and set(single.ID)==chosen and single.patient_id.nunique()==single.brats_patient_id.nunique()==135
    assert int(single.tumor_label.sum())==113 and int(single.tumor_label.eq(0).sum())==22
    pd.testing.assert_frame_equal(sm,single[sm.columns])
    sa=json.loads((directory/'single_scan_fold_audit.json').read_text());assert sa['complete'] and sa['errors']==[];check_plan(sa,sm)
    boot=pd.read_csv(directory/'single_scan_patient_bootstrap.csv');info=summary['single_scan'];assert boot.shape==(2000,21) and info['bootstrap_valid']==2000 and info['bootstrap_skipped']==0
    for k in ['roc_auc','ap_treatment_change','brier','log_loss']:
        np.testing.assert_allclose(boot['Fusion_minus_M1__'+k],boot['Fusion__'+k]-boot['M1__'+k],rtol=1e-12,atol=1e-12)
    for key in boot.columns[1:]:
        close(boot[key].quantile(.025),info['bootstrap_95ci'][key]['lower']);close(boot[key].quantile(.975),info['bootstrap_95ci'][key]['upper'])
    batch=json.loads((ROOT/'reports/ucsd_batch_features/batch_receipt.json').read_text())
    assert len(batch['cases'])==182
    for case in batch['cases']:
        assert case['status']=='complete' and 0<case['eroded_voxels']<case['roi_voxels'] and len(case['eroded'])==8
        assert all(d['erosion_error'] is None for d in case['diagnostics'])
    erosion=pd.read_csv(ROOT/'data/processed/ucsd_image_features_eroded.csv');assert len(erosion)==182 and set(erosion.ID)==set(main.ID)
    assert np.isfinite(erosion.drop(columns='ID').to_numpy()).all()
    return {'frozen_file_sha256':frozen,'boundary_cases_checked':182,'boundary_grid_rule':'one iteration of 3D six-neighbor binary erosion; original per-image normalization and full-ROI bins retained',
      'repeated_manifest_rows_checked':3640,'folds_checked':folds_checked,'saved_successful_fold_model_records_checked':fold_models_checked,
      'repeated_auc_values_checked':auc_checks,'repeated_paired_differences_checked':80,'descriptive_summary_rows_checked':summary_rows,
      'single_scan_ledger_rows_checked':182,'single_scan_selected':135,'single_scan_tumor':113,'single_scan_treatment_change':22,
      'single_bootstrap_paired_values_checked':8000,'single_bootstrap_interval_endpoints_checked':40,
      'models_refitted':0,'features_reextracted':0,'new_hypothesis_tests':0,
      'verification_scope':'Read fixed protocols and implementation; verify saved fold separation, OOF AUC, descriptive summaries, identifier-only selection, existing bootstrap differences/endpoints and erosion records; no study rerun'}

def build():
    OUT.mkdir(parents=True,exist_ok=True);REPORT.mkdir(parents=True,exist_ok=True);a=audit()
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(.79)
    for name in ['Normal','Title','Heading 1']:
        st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        rf=st.element.get_or_add_rPr().rFonts;rf.set(qn('w:eastAsia'),'Noto Serif CJK SC')
        for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:rf.attrib.pop(qn('w:'+key),None)
    normal=doc.styles['Normal'];normal.font.size=Pt(12);normal.paragraph_format.line_spacing=1.;normal.paragraph_format.space_after=Pt(9)
    doc.styles['Title'].font.size=Pt(16);doc.styles['Title'].paragraph_format.space_after=Pt(12)
    doc.styles['Heading 1'].font.size=Pt(13);doc.styles['Heading 1'].paragraph_format.space_before=Pt(8);doc.styles['Heading 1'].paragraph_format.space_after=Pt(9)
    doc.core_properties.title='Materials and methods 2.9 Sensitivity analyses';doc.core_properties.author='Tianyu Feng'
    doc.add_paragraph('Materials and methods',style='Title');doc.add_paragraph('2.9 Sensitivity analyses',style='Heading 1')
    for item in PARAGRAPHS:
        en=doc.add_paragraph(item['english']);en.paragraph_format.page_break_before=item['id']!='M17'
        en.paragraph_format.keep_together=True;en.paragraph_format.keep_with_next=True
        zh=doc.add_paragraph(item['chinese']);zh.paragraph_format.keep_together=True
        for r in zh.runs:r.font.size=Pt(11)
    for e in [*doc.styles.element.xpath('.//w:pBdr'),*doc.element.xpath('.//w:pBdr')]:e.getparent().remove(e)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    path=OUT/'UCSD_Manuscript_Methods_2_9_EN_ZH.docx';doc.save(path)
    (OUT/'UCSD_Manuscript_Methods_2_9_EN.txt').write_text('\n\n'.join(p['english'] for p in PARAGRAPHS)+'\n')
    d={'stage':31,'date':'2026-10-29','subsection':'2.9 Sensitivity analyses','v2_paragraph_ids':['M17','M18','M19'],'paragraphs':PARAGRAPHS,
      'english_word_count':sum(wc(p['english']) for p in PARAGRAPHS),'paragraph_word_counts':{p['id']:wc(p['english']) for p in PARAGRAPHS},
      'word_count_scope':'English narrative; excludes equation, headings and Chinese translation','v2_approximate_target_words':450,
      'equations':{'count':0,'numbers':[],'cumulative_previous_numbered_groups':10},
      'citation_mapping':{},'new_manuscript_references':0,'cumulative_unique_cited_sources':4,'source_audit':a,'models_refitted':0,'new_hypothesis_tests':0,
      'next_section':'Reproducibility and software (M20)','docx':str(path.relative_to(ROOT)),'docx_sha256':sha(path)}
    (REPORT/'methods_section29.json').write_text(json.dumps(d,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'docx':d['docx'],'english_word_count':d['english_word_count'],'paragraph_word_counts':d['paragraph_word_counts'],'prior_files':len(a['frozen_file_sha256'])},indent=2))

def verify():
    from pypdf import PdfReader
    d=json.loads((REPORT/'methods_section29.json').read_text())
    for p,h in d['source_audit']['frozen_file_sha256'].items():assert sha(ROOT/p)==h,p
    path=ROOT/d['docx'];assert sha(path)==d['docx_sha256'];texts=[p.text for p in Document(path).paragraphs]
    for p in PARAGRAPHS:
        assert p['english'] in texts and p['chinese'] in texts and texts.index(p['chinese'])==texts.index(p['english'])+1
        assert not re.search(r'[\u3400-\u9fff，。；：（）【】、]',p['english'])
    assert sum(wc(p['english']) for p in PARAGRAPHS)==d['english_word_count']
    ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
    with zipfile.ZipFile(path) as z:
        t=etree.fromstring(z.read('word/document.xml'));assert len(t.xpath('//m:oMath',namespaces=ns))==0
        assert not any(n.startswith('word/media/') for n in z.namelist())
    pages=len(PdfReader(OUT/'docx_render/UCSD_Manuscript_Methods_2_9_EN_ZH.pdf').pages);assert pages==3
    v={'success':True,'stage':31,'date':'2026-10-29','subsection':d['subsection'],'paragraph_ids':['M17','M18','M19'],
      'english_word_count':d['english_word_count'],'paragraph_word_counts':d['paragraph_word_counts'],'matched_chinese_paragraphs':3,
      'word_pages':pages,'visual_review_pages':[1,2,3],'native_editable_equations':0,'equation_numbers':[],'equation_images':0,
      'prior_files_unchanged':len(d['source_audit']['frozen_file_sha256']),'boundary_cases_checked':182,'repeated_manifest_rows_checked':3640,'folds_checked':105,'saved_successful_fold_model_records_checked':420,
      'repeated_auc_values_checked':80,'repeated_paired_differences_checked':80,'descriptive_summary_rows_checked':36,'single_scan_ledger_rows_checked':182,'single_scan_selected':135,
      'single_bootstrap_paired_values_checked':8000,'single_bootstrap_interval_endpoints_checked':40,'models_refitted':0,'new_hypothesis_tests':0,'new_manuscript_references':0,'next_section':d['next_section'],
      'sha256':{str(p.relative_to(ROOT)):sha(p) for p in OUT.iterdir() if p.is_file() and p.suffix in {'.docx','.txt'}}}
    (REPORT/'artifact_verification.json').write_text(json.dumps(v,ensure_ascii=False,indent=2)+'\n');print(json.dumps(v,indent=2))

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--verify',action='store_true');args=parser.parse_args()
    verify() if args.verify else build()
