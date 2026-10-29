"""First editorial revision of the assembled paper; existing analyses remain frozen."""
from pathlib import Path
from copy import deepcopy
import sys,json,re,hashlib,csv,zipfile,argparse
from lxml import etree
from docx import Document
from docx.oxml import OxmlElement
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
import assemble_ucsd_manuscript55 as base
OUT=ROOT/'outputs/ucsd-manuscript-revision1-20261030'
REPORT=ROOT/'reports/ucsd_manuscript_step56'
OLD=json.loads((ROOT/'reports/ucsd_manuscript_step55/assembly_receipt.json').read_text())
TEXT=[]
def put(pid,en,zh):TEXT.append({'id':pid,'english':en,'chinese':zh})

put('I1',
'''Magnetic resonance imaging (MRI) is central to follow-up after treatment for high-grade glioma. Contemporary glioblastoma management combines surgery, radiotherapy and systemic treatment, with subsequent imaging guiding assessment of disease status and further management (1). New or increasing enhancement can reflect either tumor activity or treatment effects, making their differentiation a recurring diagnostic problem. Response Assessment in Neuro-Oncology (RANO) 2.0 incorporates standardized response assessment and, in relevant early post-radiotherapy settings, confirmation of progression by repeat imaging or unequivocal histopathological evidence (2). Clinical history and the timing of imaging therefore remain important when assessing a suspicious examination. Quantitative MRI offers an additional way to summarize lesion appearance and examine its contribution to clinical assessment.''',
'''磁共振成像（MRI）是高级别胶质瘤治疗后随访的重要手段。当前胶质母细胞瘤治疗综合采用手术、放疗及系统性治疗，后续影像用于评估疾病状态并指导进一步管理（1）。治疗后新出现或增加的增强既可能反映肿瘤活动，也可能与治疗效应有关，二者的鉴别是反复遇到的诊断问题。神经肿瘤疗效评估（RANO）2.0采用标准化疗效评价，并在适用的放疗后早期场景中，通过重复影像或明确的组织病理证据确认进展（2）。因此，解读可疑检查时需要结合临床病史与影像检查时间。定量MRI可进一步概括病灶表现，并用于考察影像信息对临床评估的贡献。''')
put('I2',
'''Treatment-related abnormalities include pseudoprogression, radiation necrosis and nonspecific post-treatment changes (1,3). Pseudoprogression commonly presents as radiological worsening followed by stabilization or regression, whereas radiation necrosis is often considered a later form of treatment injury. RANO implementation guidance also distinguishes enhancing, nonenhancing and mixed presentations and specifies circumstances requiring confirmation scans (4). These diagnostic categories and their reference evidence shape the clinical question addressed by an imaging model. The public University of California San Diego Post-Treatment Glioblastoma (UCSD-PTGBM) collection contains residual or recurrent tumor and several categories of treatment-related change, supported by pathological, clinical and imaging assessments (5). It provides a resource for investigating their differentiation at the examination level.''',
'''治疗相关异常包括假性进展、放射性坏死及非特异性治疗后改变（1,3）。假性进展通常表现为影像恶化后趋于稳定或消退，放射性坏死则常被视为较晚期的治疗损伤。RANO实施指南还区分增强、非增强及混合表现，并规定需要确认扫描的具体情形（4）。这些诊断类别及其参考证据共同决定影像模型所回答的临床问题。公开的加州大学圣迭戈分校治疗后胶质母细胞瘤（UCSD-PTGBM）数据集包含残余或复发肿瘤及多类治疗相关改变，其诊断依据包括病理、临床与影像评估（5），为研究检查层面的鉴别提供了数据基础。''')
put('I3',
'''Quantitative image features describe both signal distribution and spatial organization. Intensity statistics summarize regional signal levels and dispersion, while texture measures characterize gray-level relationships between neighboring voxels. Radiomics studies have investigated these features for differentiating progression from treatment effects in glioma (3). The Image Biomarker Standardization Initiative (IBSI) provides common definitions for feature calculation (6). Preprocessing also influences the resulting measurements: Foltyn-Dumitru et al. showed that intensity normalization affected external generalization in a study of molecular glioma subtype classification (7). Together, this work supports examining a clearly specified imaging pipeline. A compact panel of fluid-attenuated inversion recovery (FLAIR) and postcontrast T1 features offers a practical representation for evaluating additional information beyond clinical variables.''',
'''定量影像特征能够描述信号分布及空间组织。强度统计概括区域信号水平和离散程度，纹理指标则刻画相邻体素之间的灰度关系。影像组学研究已探索利用这些特征鉴别胶质瘤进展与治疗效应（3）。影像生物标志物标准化倡议（IBSI）提供了共同的特征计算定义（6）。预处理也会影响相应测量：Foltyn-Dumitru等在胶质瘤分子亚型分类研究中发现，强度归一化影响模型的外部推广表现（7）。这些研究支持采用定义清晰的影像处理流程。由液体衰减反转恢复（FLAIR）与增强T1特征构成的精简面板，可用于评价影像在临床变量基础上提供的额外信息。''')
put('I4',
'''Clinical and imaging variables provide complementary descriptions of a post-treatment examination. Age, treatment history and the interval from an initial clinical event characterize its clinical context, while MRI features quantify regional appearance. Zhou et al. reported favorable results from a model combining radiomics with clinical, molecular and advanced MRI information (8). McKenney et al. reported improved classification accuracy when imaging features were combined with O6-methylguanine-DNA methyltransferase (MGMT) promoter methylation status in a histopathologically confirmed cohort (9). Leone et al., however, found no reliable improvement from adding diffusion, perfusion and radiomic features to their clinical reference model (10). The new-cohort evaluation by de Godoy et al. also illustrates the relevance of outcome grouping to model performance (11). This varied evidence makes a direct comparison between a fusion model and its clinical baseline particularly useful.''',
'''临床与影像变量从不同方面描述治疗后的检查。年龄、治疗史及初始临床事件至检查的间隔反映临床背景，MRI特征则量化区域影像表现。Zhou等报告，将影像组学与临床、分子及高级MRI信息结合的模型具有较好的表现（8）。McKenney等在经组织病理确认的队列中发现，将影像特征与O6-甲基鸟嘌呤-DNA甲基转移酶（MGMT）启动子甲基化状态结合能够提高分类准确率（9）。但Leone等未发现将弥散、灌注与影像组学特征加入临床参考模型可带来可靠改善（10）。de Godoy等的新队列评价也说明，结局分组与模型性能密切相关（11）。现有证据并不一致，因此，直接比较融合模型与其临床基线具有研究价值。''')
put('I5',
'''A reproducible comparison requires explicit feature definitions, consistent preprocessing and appropriate separation of training and validation data. CLEAR and PROBAST+AI emphasize the reporting and appraisal of these components (12,13), while prediction-model guidance distinguishes internal evaluation from evaluation in independent data (14). For patients with repeated examinations, grouping by patient preserves separation between training and validation folds. The importance of this separation is illustrated by experiments with duplicated scans in connectome-based prediction (15). Applying a common cohort, fold structure and evaluation framework to each candidate model makes the incremental comparison easier to interpret and reproduce. TRIPOD+AI provides a reporting framework for presenting the study design and performance estimates (16).''',
'''可复现的比较需要明确的特征定义、一致的预处理，以及合理的训练与验证数据划分。CLEAR与PROBAST+AI强调对这些环节进行报告和评价（12,13），预测模型指南则区分内部评价与独立数据评价（14）。对于有重复检查的患者，按患者分组能够保持训练折与验证折之间的隔离。基于脑连接组预测的重复扫描实验展示了这种隔离的重要性（15）。对各候选模型采用共同队列、分折结构和评价框架，有助于解释及复现其增量比较。TRIPOD+AI为呈现研究设计和性能估计提供了报告框架（16）。''')
put('I6',
'''We evaluated whether eight fixed MRI features improved discrimination between residual or recurrent tumor and treatment-related change beyond available clinical inputs in public UCSD-PTGBM data (5). Four models were compared: age and sex, expanded clinical inputs, MRI features, and clinical-MRI feature fusion. All models used the same imaging cohort and patient-grouped validation. The primary comparison was the paired difference in area under the receiver operating characteristic curve (AUC) between Fusion and the expanded clinical model. Calibration, probability error and classification performance complemented this comparison. Boundary erosion, repeated grouped partitions and one-examination-per-patient analyses were used to assess the consistency of the findings.''',
'''本研究利用公开UCSD-PTGBM数据（5），评价八项固定MRI特征能否在可用临床输入基础上，提高残余或复发肿瘤与治疗相关改变的鉴别能力。研究比较年龄与性别、扩展临床输入、MRI特征及临床–MRI特征融合四个模型，统一采用同一影像队列和患者分组验证。主要比较为Fusion与扩展临床模型之间的受试者工作特征曲线下面积（AUC）配对差值，同时评价校准、概率误差及分类表现。通过边界腐蚀、重复分组分折和每位患者保留一次检查的分析，进一步考察结果的一致性。''')

put('M1',
'''We conducted a retrospective secondary analysis of public, de-identified clinical records and MRI examinations to assess the incremental discrimination provided by MRI features. The outcome was residual or recurrent tumor versus treatment-related change at each examination. Clinical, MRI and fusion models were evaluated in a common cohort using patient-grouped cross-validation. The dataset authors reported institutional review board approval and a waiver of informed consent for the original collection (5). Reporting followed the organizational guidance of TRIPOD+AI and the Checklist for Artificial Intelligence in Medical Imaging (CLAIM) (16,17).''',
'''本研究对公开、去标识的临床记录和MRI检查进行回顾性二次分析，评价MRI特征提供的增量鉴别信息。结局为每次检查中的残余或复发肿瘤与治疗相关改变。临床、MRI及融合模型在共同队列中采用患者分组交叉验证。数据集作者报告，原始资料收集已获机构伦理审查批准并豁免知情同意（5）。论文组织参照TRIPOD+AI及医学影像人工智能报告清单CLAIM的指导（16,17）。''')
put('M2',
'''The UCSD-PTGBM dataset is distributed through The Cancer Imaging Archive under a Creative Commons Attribution 4.0 license (5). We used version 3 of its public clinical workbook and the accompanying list of examinations with treatment-related changes. FLAIR images, postcontrast T1-weighted images and standard BraTS masks were obtained as NIfTI files from the public data-nih/tcia release mirror on GitHub. Clinical and BraTS identifiers were linked to the corresponding image files. Download addresses, file sizes and SHA256 checksums were recorded for input verification and reproducibility.''',
'''UCSD-PTGBM数据集通过癌症影像档案库发布，采用知识共享署名4.0许可（5）。本研究使用公开的第3版临床工作簿及配套的治疗相关改变检查名单。FLAIR图像、增强T1加权图像及标准BraTS掩膜来自GitHub公开的data-nih/tcia发布镜像，文件格式为NIfTI。研究将临床与BraTS标识连接至对应影像文件，并记录下载地址、文件大小及SHA256校验值，用于输入核验和研究复现。''')
put('M3',
'''The source collection contained 243 examinations from 178 patients (5). Five examinations from two patients with discordant clinical and BraTS identifier mappings were excluded during linkage verification, leaving 238 examinations from 176 patients. Complete FLAIR, postcontrast T1-weighted and mask inputs were available for 182 examinations from 135 patients. This imaging cohort comprised 141 tumor and 41 treatment-related-change examinations and was used for all four primary models. Cohort assembly is shown in Figure 1; clinical characteristics according to imaging availability are summarized in Supplementary Table S1.''',
'''源数据集包含178名患者的243次检查（5）。连接核查时排除2名患者的5次临床与BraTS标识对应不一致的检查，保留176名患者的238次检查。其中，135名患者的182次检查具有完整的FLAIR、增强T1加权图像及掩膜输入。该影像队列包括141次肿瘤检查和41次治疗相关改变检查，四个主要模型均使用此队列。队列构建见Figure 1，按影像可获得性汇总的临床特征见Supplementary Table S1。''')
put('M4',
'''Diagnostic labels followed the published UCSD-PTGBM classification (5). The 51 examinations listed in the public treatment-related-change workbook formed the treatment-related-change class; the other 192 examinations formed the residual or recurrent tumor class in the full collection. Labels were retained after linkage verification and restriction to the imaging cohort. The source diagnoses incorporated pathological evidence and clinical-imaging assessment, with pathological confirmation reported for 45 of the 192 tumor examinations in the full collection. Treatment-related change included pseudoprogression, radiation necrosis and nonspecific post-treatment changes. The analysis used the binary diagnosis assigned to each examination.''',
'''诊断标签依据UCSD-PTGBM已发表的分类（5）。公开治疗相关改变工作簿所列的51次检查构成治疗相关改变类，完整数据集中其余192次检查构成残余或复发肿瘤类。连接核查及限定影像队列后沿用这些标签。源诊断结合病理证据与临床影像评估，完整数据集中192次肿瘤检查有45次获得病理确认。治疗相关改变包括假性进展、放射性坏死及非特异性治疗后改变。本分析采用每次检查所对应的二分类诊断。''')
put('M5',
'''Clinical predictors were age at imaging, sex at birth, the interval since the recorded initial clinical event, and other treatment before imaging, as provided in the public workbook (5). A missing-interval indicator was included, giving five clinical inputs for M1 and Fusion. M0 used age and sex. Predictor definitions and processing are summarized in Supplementary Table S2.''',
'''临床预测变量包括公开工作簿记录的检查时年龄、出生性别、初始临床事件至检查的间隔，以及扫描前其他治疗（5）。另纳入间隔缺失指示变量，使M1与Fusion共包含五项临床输入；M0采用年龄与性别。变量定义及处理方法见Supplementary Table S2。''')
put('M6',
'''The interval since the recorded initial clinical event was expressed in days and transformed as ln(1 + days). Missing intervals were imputed using the median transformed value in each training fold. Age and the transformed interval were then standardized using training-fold parameters, which were applied to the corresponding validation fold. The imaging cohort contained 36 examinations with a missing interval; age, sex and the other-treatment field were complete.''',
'''记录的初始临床事件至检查的间隔以天表示，并采用ln(1＋天数)转换。缺失间隔使用各训练折中转换后间隔的中位数填补，随后按训练折参数对年龄及转换后的间隔进行标准化，并将相同参数应用于对应验证折。影像队列中有36次检查的间隔缺失；年龄、性别及扫描前其他治疗字段均完整。''')
put('M7',
'''Each included examination comprised a FLAIR volume, a postcontrast T1-weighted volume and a standard BraTS mask, totaling 546 compressed NIfTI files. Input verification covered file integrity, finite voxel values, image dimensions and affine coordinates. Image and mask arrays shared a 256 × 256 × 256 voxel grid and were stored as unsigned 8-bit integers. Features were calculated on the supplied grids. The region of interest (ROI) was the union of positive labels in the supplied standard BraTS mask for each examination.''',
'''每次纳入检查均包括FLAIR图像、增强T1加权图像及标准BraTS掩膜，共546份压缩NIfTI文件。输入核验涵盖文件完整性、体素数值有限性、图像维度及仿射坐标。图像与掩膜采用一致的256×256×256体素网格，以无符号8位整数存储。特征在所提供的网格上计算，每次检查的感兴趣区域（ROI）由现有标准BraTS掩膜中的正标签区域合并得到。''')
put('M9',
'''FLAIR and postcontrast T1 images were normalized separately. For each image, the first and 99th percentiles were calculated from finite positive voxels using linear percentile interpolation. These image-specific reference values defined the affine scaling in Equation 1, where x is the supplied intensity and u is the scaled intensity. The transformation was applied to all ROI voxels, retaining the full range of scaled values.''',
'''FLAIR与增强T1图像分别进行归一化。对每幅图像，采用线性百分位插值，从有限正值体素计算第1与第99百分位数，并以其作为Equation 1仿射缩放的图像特异性参考值，其中x为原始输入强度，u为缩放后强度。该变换应用于全部ROI体素，保留缩放后的完整数值范围。''')
put('M10',
'''Four features were extracted from each sequence: mean, population standard deviation, Shannon entropy and three-dimensional gray-level co-occurrence matrix (GLCM) contrast. Mean and standard deviation were calculated from scaled ROI intensities using Equation 2, with N denoting the voxel count. Entropy and GLCM calculations used fixed-bin-number discretization based on IBSI definitions (6). For each sequence and examination, the minimum and maximum scaled ROI intensities, a and b, defined 32 bins according to Equation 3. Bin probabilities were the proportions of ROI voxels in each bin. Shannon entropy was calculated in bits using base-2 logarithms over occupied bins (Equation 4).''',
'''每个序列提取四项特征：均值、总体标准差、Shannon熵及三维灰度共生矩阵（GLCM）对比度。均值与标准差依据缩放后的ROI强度，按Equation 2计算，其中N为体素数。熵与GLCM计算采用基于IBSI定义的固定箱数离散化（6）。对每个序列和检查，以缩放后ROI强度的最小值a与最大值b确定32个箱，计算方法见Equation 3。各箱概率为落入该箱的ROI体素比例，采用以2为底的对数，对有体素分布的箱计算Shannon熵，单位为比特（Equation 4）。''')
put('M11',
'''GLCMs were constructed from neighboring voxel pairs within the ROI using 13 undirected offsets spanning the 26 immediate grid neighbors. The directional count matrix C recorded bin-index pairs i and j for each offset in D. Directional counts were merged and symmetrized to obtain S, then normalized to obtain P (Equation 5). Contrast was calculated from this pooled matrix using Equation 6 (6), producing one GLCM contrast value for each sequence.''',
'''GLCM基于ROI内的相邻体素对构建，采用覆盖网格上26个直接邻居的13个无向偏移。对D中的每个偏移，方向计数矩阵C记录箱索引i与j的组合。将各方向计数合并并对称化得到S，再归一化得到P（Equation 5）。按Equation 6从合并矩阵计算对比度（6），每个序列产生一个GLCM对比度值。''')
put('M12',
'''Four logistic regression models with an intercept were fitted. M0 used age and sex; M1 used the five clinical inputs; MRI used the eight imaging features; and Fusion concatenated clinical and imaging inputs into a 13-variable feature vector. The corresponding input counts were 2, 5, 8 and 13. Equation 7 gives the tumor probability for processed feature vector z, with m denoting the input count and the beta terms denoting the intercept and coefficients. All models used L2 regularization with C = 1, the lbfgs optimizer, a maximum of 2000 iterations and a convergence tolerance of 0.0001. The common regression framework allowed direct comparison of the four input sets.''',
'''拟合四个包含截距的Logistic回归模型。M0使用年龄与性别，M1使用五项临床输入，MRI使用八项影像特征，Fusion将临床与影像输入拼接为13变量特征向量，对应输入数分别为2、5、8和13。Equation 7给出处理后特征向量z的肿瘤概率，其中m为输入数，beta项为截距与系数。全部模型采用L2正则化，C＝1，优化器为lbfgs，最大迭代次数为2000，收敛容差为0.0001。共同的回归框架用于直接比较四组输入。''')
put('M13',
'''All models shared five patient-grouped folds in the 182-examination cohort. StratifiedGroupKFold was used with shuffling and random seed 20261029. All examinations from the same patient were assigned to one fold, with separation checked against both clinical and BraTS patient identifiers. Continuous predictors were standardized within each training fold; interval imputation followed Section 2.3. The fitted preprocessing and regression were applied to the held-out fold, giving one out-of-fold (OOF) tumor probability per examination and model. The four input panels and model settings were fixed throughout validation.''',
'''全部模型在182次检查队列中共用五个患者分组折，采用StratifiedGroupKFold，启用随机打乱并设随机种子为20261029。同一患者的全部检查分配至同一折，并通过临床及BraTS患者标识核查患者隔离。连续预测变量在各训练折内标准化，间隔填补按Section 2.3执行。将拟合的预处理与回归模型应用于留出折，使每次检查在每个模型中获得一个折外（OOF）肿瘤概率。四组输入及模型设置在验证过程中保持固定。''')
put('M14',
'''Performance was evaluated at the examination level using pooled OOF probabilities. The primary comparison was the Fusion-minus-M1 difference in ROC AUC on the same examinations (Equation 8). Additional measures included average precision (AP), Brier score, logarithmic loss and classification performance. Treatment-related-change AP used one minus the tumor probability as its score. Lower Brier scores and logarithmic losses indicated better probability accuracy. At a fixed threshold of 0.5, probabilities of at least 0.5 were classified as tumor. We calculated confusion counts, tumor sensitivity, treatment-change recognition, precision, balanced accuracy and overall accuracy.''',
'''基于汇总的OOF概率，在检查层面评价性能。主要比较为相同检查中Fusion减去M1的ROC AUC差值（Equation 8）。其他指标包括平均精确率（AP）、Brier评分、对数损失及分类表现。治疗相关改变AP以1减去肿瘤概率作为评分。Brier评分与对数损失越低，表示概率准确性越好。采用固定0.5阈值，将大于等于0.5的概率判为肿瘤，并计算混淆计数、肿瘤敏感度、治疗相关改变识别率、精确率、平衡准确率及总体准确率。''')
put('M15',
'''Uncertainty was estimated with 2000 patient-cluster bootstrap attempts using random seed 20261030. Each attempt sampled 135 patient groups with replacement and included all examinations belonging to each sampled group. The same sampled indices were applied to every model to obtain paired metric differences. Bootstrap calculations used the fixed OOF probabilities. The 2.5th and 97.5th percentiles defined the 95% intervals for AUC, treatment-related-change AP, Brier score, logarithmic loss and their differences. These intervals quantified patient-sampling uncertainty conditional on the saved predictions. A positive AUC or AP difference favored Fusion; a negative Brier-score or logarithmic-loss difference favored Fusion.''',
'''采用2000次患者聚类bootstrap估计不确定性，随机种子为20261030。每次有放回抽取135个患者组，并纳入每个抽中患者组的全部检查。各模型使用相同的抽样索引，以获得配对指标差值。bootstrap基于固定的OOF概率计算，以第2.5与第97.5百分位数构成AUC、治疗相关改变AP、Brier评分、对数损失及其差值的95%区间。这些区间量化条件于保存预测的患者抽样不确定性。AUC或AP差值为正表示有利于Fusion，Brier评分或对数损失差值为负表示有利于Fusion。''')
put('M16',
'''Calibration was assessed from pooled OOF tumor probabilities using five quantile groups, retaining the nonempty groups after duplicate boundaries were removed. Groups were summarized by sample size, mean predicted probability and observed tumor proportion. Calibration-in-the-large was estimated as an intercept with the predicted-log-odds coefficient fixed at one (Equation 9). A separate unpenalized logistic diagnostic regression jointly estimated a free intercept and calibration slope (Equation 10). Here, p denotes the OOF probability and q the fitted diagnostic probability. The logit calculation bounded p by machine epsilon and one minus machine epsilon. Calibration estimates were reported as descriptive diagnostics of the evaluated predictions.''',
'''采用汇总的OOF肿瘤概率评价校准，设置五个分位数组，去除重复边界后保留非空组。各组汇总样本数、平均预测概率及实际肿瘤比例。整体校准以预测对数优势系数固定为1的截距估计（Equation 9）；另采用无惩罚Logistic诊断回归联合估计自由截距及校准斜率（Equation 10）。其中p为OOF概率，q为拟合的诊断概率。计算logit时，将p限定在机器精度epsilon与1减epsilon之间。校准估计作为所评价预测的描述性诊断结果报告。''')
put('M17',
'''Boundary sensitivity was assessed by one iteration of ROI erosion using a three-dimensional six-neighbor structuring element on the supplied voxel grid. The same eight features were recalculated within the eroded region, retaining the original image-percentile references and full-ROI discretization boundaries. MRI and Fusion models were fitted with the original cases, patient groups, fold assignments and model settings. Paired patient-cluster bootstrap comparisons assessed eroded Fusion versus M1 and eroded versus original Fusion. The full ROI defined the primary analysis.''',
'''边界敏感性分析采用三维六邻域结构元素，在所提供的体素网格上对ROI执行一次腐蚀。在腐蚀后的区域内重新计算相同的八项特征，保留原始图像百分位参考及完整ROI确定的离散化边界。MRI与Fusion模型沿用原有检查、患者组、分折分配及模型设置拟合。通过配对患者聚类bootstrap比较腐蚀后Fusion与M1，以及腐蚀后与原始Fusion。主分析采用完整ROI。''')
put('M18',
'''Split sensitivity was evaluated using 20 patient-grouped five-fold partitions, with consecutive seeds 20261029–20261048. The first partition reproduced the primary analysis. Within each repetition, the four models shared the same cases and folds, and preprocessing and regression were fitted within training folds. Performance and Fusion-minus-M1 differences were summarized by the median, quartiles, observed range and direction of the differences. These summaries described variation across partitions of the same cohort. Model settings and the classification threshold were held constant.''',
'''分折敏感性分析采用20套患者分组五折划分，使用20261029至20261048的连续随机种子。第一套划分复现主分析。每次重复中，四个模型共用相同检查及分折，并在训练折内拟合预处理与回归。性能及Fusion减去M1的差值采用中位数、四分位数、实际范围及差值方向汇总，用于描述同一队列不同划分中的变化。模型设置与分类阈值保持固定。''')
put('M19',
'''The single-scan sensitivity analysis selected one available examination per patient using the smallest numeric session suffix in the imaging subset, with the complete examination identifier resolving ties. Patient-grouped five-fold partitions were rebuilt using seed 20261029, and the four models were fitted within the new training folds. AUC, treatment-related-change AP, Brier score, logarithmic loss and their Fusion-minus-M1 differences were assessed using 2000 patient-bootstrap attempts with seed 20261101 and fixed OOF predictions.''',
'''单扫描敏感性分析按影像子集中最小的数字会话后缀，为每位患者选择一次可用检查；出现并列时按完整检查标识确定。采用随机种子20261029重新构建患者分组五折划分，在新训练折内拟合四个模型。基于固定OOF预测，采用2000次患者bootstrap评价AUC、治疗相关改变AP、Brier评分、对数损失及其Fusion减去M1的差值，抽样随机种子为20261101。''')
put('M20',
'''Analyses used Python 3 with scikit-learn 1.8.0 (18), NumPy 2.3.5, pandas 2.2.3, SciPy 1.17.0 and Matplotlib 3.10.8. Jupyter notebooks documented feature extraction, model evaluation and sensitivity analyses. The public repository (https://github.com/BTomoto/rn-mri-clinical-fusion) contains analysis code, derived feature tables, OOF predictions, fold assignments and analysis locks. Source manifests record image acquisition URLs and checksums, and restoration scripts verify archived inputs. The sample comprised eligible linked examinations with complete imaging inputs.''',
'''分析使用Python 3，主要软件包包括scikit-learn 1.8.0（18）、NumPy 2.3.5、pandas 2.2.3、SciPy 1.17.0及Matplotlib 3.10.8。Jupyter notebook记录特征提取、模型评价及敏感性分析过程。公开仓库（https://github.com/BTomoto/rn-mri-clinical-fusion）保存分析代码、派生特征表、OOF预测、分折分配及分析锁定记录。源清单记录影像获取URL及校验值，恢复脚本用于核验归档输入。研究样本由完成连接且影像输入完整的符合条件检查构成。''')

put('R1',
'''The source cohort contained 243 examinations from 178 patients, including 192 tumor and 51 treatment-related-change examinations (Figure 1). Linkage verification excluded five examinations from two patients, leaving 238 examinations from 176 patients. Complete imaging inputs were available for 182 examinations from 135 patients, representing 76.5% of the retained clinical examinations. All 182 yielded the eight MRI features and entered the four-model comparison. The other 56 examinations came from 41 patients and comprised 46 tumor and 10 treatment-related-change examinations. The imaging cohort contained 141 tumor examinations (77.5%) and 41 treatment-related-change examinations (22.5%). At the patient level, 114 patients contributed tumor examinations and 29 contributed treatment-related-change examinations, including eight patients who contributed both. Table 1 summarizes the imaging cohort, and Supplementary Table S1 presents clinical characteristics by imaging coverage.''',
'''源队列包含178名患者的243次检查，其中192次为肿瘤、51次为治疗相关改变（Figure 1）。连接核查排除2名患者的5次检查后，保留176名患者的238次检查。135名患者的182次检查具有完整影像输入，占保留临床检查的76.5%；全部获得八项MRI特征并进入四模型比较。其余56次检查来自41名患者，包括46次肿瘤和10次治疗相关改变检查。影像队列中141次为肿瘤（77.5%），41次为治疗相关改变（22.5%）。按患者统计，114名患者贡献肿瘤检查，29名患者贡献治疗相关改变检查，其中8名患者同时贡献两类检查。影像队列特征见Table 1，按影像覆盖情况汇总的临床特征见Supplementary Table S1。''')
put('R2',
'''Median age in the imaging cohort was 56.5 years [49.0, 66.0], and 132 examinations were from male patients (72.5%; Table 1). Age, sex and other treatment before imaging were complete. Among 146 examinations with an observed initial-event interval, the median was 154.0 days [40.5, 359.5], with medians of 96.0 days for tumor and 346.0 days for treatment-related change. Intervals were missing for 36 examinations (19.8%), including 33 tumor and three treatment-related-change examinations. Other treatment before imaging was recorded for 25 examinations (13.7%). In the group without imaging coverage, median age was 53.5 years [45.0, 62.0], 60.7% of examinations were from male patients and 82.1% carried tumor labels (Supplementary Table S1). Brackets indicate the first and third quartiles; interval summaries use observed values before imputation.''',
'''影像队列年龄中位数为56.5岁［49.0，66.0］，132次检查来自男性患者（72.5%；Table 1）。年龄、性别及扫描前其他治疗记录完整。146次初始事件间隔已知的检查中，间隔中位数为154.0天［40.5，359.5］，肿瘤与治疗相关改变检查分别为96.0天和346.0天。36次检查的间隔缺失（19.8%），其中33次为肿瘤、3次为治疗相关改变。25次检查记录了扫描前其他治疗（13.7%）。影像未覆盖组年龄中位数为53.5岁［45.0，62.0］，60.7%的检查来自男性患者，82.1%为肿瘤标签（Supplementary Table S1）。方括号表示第一与第三四分位数，间隔汇总采用填补前的已知值。''')
put('R3',
'''All four models generated OOF probabilities for the 182 examinations using the shared patient-grouped five-fold partition. ROC AUCs were 0.526 for M0 (95% patient-bootstrap percentile interval, 0.404–0.644), 0.672 for M1 (0.553–0.778), 0.762 for MRI (0.683–0.838) and 0.782 for Fusion (0.698–0.859; Table 2 and Figure 2A). Treatment-related-change AP was 0.261, 0.354, 0.444 and 0.505, respectively, against a cohort prevalence of 0.225 (Figure 2B). Corresponding Brier scores were 0.177, 0.163, 0.151 and 0.144. Fusion had the highest AUC and treatment-related-change AP and the lowest Brier-score point estimate among the four models.''',
'''四个模型均在共同的患者分组五折划分中，为182次检查生成OOF概率。ROC AUC分别为M0的0.526（95%患者bootstrap百分位区间0.404—0.644）、M1的0.672（0.553—0.778）、MRI的0.762（0.683—0.838）及Fusion的0.782（0.698—0.859；Table 2和Figure 2A）。治疗相关改变AP依次为0.261、0.354、0.444和0.505，队列相应类别比例为0.225（Figure 2B）。Brier评分依次为0.177、0.163、0.151和0.144。四个模型中，Fusion的AUC与治疗相关改变AP最高，Brier评分点估计最低。''')
put('R4',
'''The primary Fusion-minus-M1 AUC difference was 0.110 (95% paired patient-bootstrap percentile interval, 0.021–0.214; Table 3). The treatment-related-change AP difference was 0.152 (0.026–0.268). Both intervals were above zero. Brier-score and logarithmic-loss differences favored Fusion in their point estimates: -0.019 (-0.040 to 0.002) and -0.037 (-0.107 to 0.039), respectively; their intervals included zero. The paired AUC comparison therefore supported improved discrimination from adding MRI features to the clinical inputs, whereas the probability-error differences were less conclusive.''',
'''主要比较中，Fusion减去M1的AUC差值为0.110（95%配对患者bootstrap百分位区间0.021—0.214；Table 3），治疗相关改变AP差值为0.152（0.026—0.268），两项区间均高于零。Brier评分与对数损失差值的点估计也有利于Fusion，分别为−0.019（−0.040至0.002）及−0.037（−0.107至0.039），但区间包含零。因此，配对AUC比较支持在临床输入基础上加入MRI特征可改善鉴别能力，概率误差差异的证据则相对不充分。''')
put('R5',
'''Fusion's descriptive calibration offset intercept was -0.121 and its calibration slope was 0.667 (Figure 2C; Supplementary Table S3). At the fixed 0.5 threshold, Fusion correctly classified 130/141 tumor examinations (92.2%) and 12/41 treatment-related-change examinations (29.3%). Eleven tumor examinations were classified as treatment-related change and 29 treatment-related-change examinations as tumor. Overall accuracy was 142/182 (78.0%), compared with 141/182 (77.5%) for the always-tumor reference, and balanced accuracy was 0.607. M1 correctly recognized two treatment-related-change examinations; MRI and Fusion each recognized 12. MRI correctly classified 128 tumor examinations, compared with 130 for Fusion.''',
'''Fusion的描述性校准offset截距为−0.121，校准斜率为0.667（Figure 2C；Supplementary Table S3）。在固定0.5阈值下，Fusion正确识别130/141次肿瘤检查（92.2%）及12/41次治疗相关改变检查（29.3%）；11次肿瘤检查被判为治疗相关改变，29次治疗相关改变检查被判为肿瘤。总体准确率为142/182（78.0%），全部判为肿瘤的参考规则为141/182（77.5%），平衡准确率为0.607。M1正确识别2次治疗相关改变检查，MRI与Fusion各识别12次。MRI正确分类128次肿瘤检查，Fusion为130次。''')
put('R6',
'''All 182 examinations yielded complete features after erosion. The eroded MRI model had an AUC of 0.765 versus 0.762 for the original MRI model (Supplementary Table S4; Figure S1). Eroded Fusion had an AUC of 0.788 (95% percentile interval, 0.706–0.867), compared with 0.782 for original Fusion. The paired eroded-minus-original Fusion AUC difference was 0.0062 (-0.0051 to 0.0183). Eroded Fusion had a Brier score of 0.141, treatment-related-change AP of 0.529 and treatment-change recognition of 14/41 at the fixed threshold. Corresponding original Fusion values were 0.144, 0.505 and 12/41. The small AUC difference indicated similar performance under this boundary perturbation.''',
'''腐蚀后182次检查均获得完整特征。腐蚀后MRI模型AUC为0.765，原始MRI模型为0.762（Supplementary Table S4；Figure S1）。腐蚀后Fusion的AUC为0.788（95%百分位区间0.706—0.867），原始Fusion为0.782；配对AUC差值为0.0062（−0.0051至0.0183）。腐蚀后Fusion的Brier评分为0.141，治疗相关改变AP为0.529，固定阈值下识别14/41次治疗相关改变检查；原始Fusion相应数值为0.144、0.505和12/41。AUC差异较小，显示该次边界扰动下模型表现相近。''')
put('R7',
'''All 20 patient-grouped five-fold repetitions were completed. Median AUCs were 0.556 for M0, 0.667 for M1, 0.775 for MRI and 0.784 for Fusion (Supplementary Table S5A; Figure 3A). Fusion's AUC quartiles were 0.775 and 0.795, with an observed range of 0.759–0.809. The median Fusion-minus-M1 AUC difference was 0.125, ranging from 0.103 to 0.150, and was positive in all 20 partitions (Supplementary Table S5B; Figure 3B). AP, Brier-score and logarithmic-loss differences also favored Fusion in all repetitions. These descriptive summaries showed a consistent direction of benefit across partitions of the same cohort. The first partition reproduced the primary Fusion AUC of 0.782.''',
'''20次患者分组五折重复均完成。AUC中位数分别为M0的0.556、M1的0.667、MRI的0.775及Fusion的0.784（Supplementary Table S5A；Figure 3A）。Fusion的AUC第一与第三四分位数为0.775和0.795，实际范围为0.759—0.809。Fusion减去M1的AUC差值中位数为0.125，范围为0.103—0.150，20套划分中均为正（Supplementary Table S5B；Figure 3B）。AP、Brier评分及对数损失差值也在全部重复中有利于Fusion。这些描述性结果显示，同一队列不同划分中的增益方向一致。第一套划分复现主分析Fusion AUC为0.782。''')
put('R8',
'''Selecting one available examination per patient retained 135 examinations, including 113 tumor and 22 treatment-related-change examinations. AUCs were 0.501 for M0, 0.609 for M1, 0.729 for MRI and 0.730 for Fusion; Fusion's 95% percentile interval was 0.606–0.833 (Supplementary Table S6; Figure S2A). The Fusion-minus-M1 AUC difference was 0.120 (-0.015 to 0.252; Supplementary Table S7). Its interval and those for AP, Brier-score and logarithmic-loss differences included zero. At the fixed threshold, Fusion correctly recognized 109/113 tumor examinations (96.5%) and 2/22 treatment-related-change examinations (9.1%; Figure S2B). Its descriptive calibration slope was 0.539. The AUC increment remained positive in its point estimate, with wider uncertainty in this selected cohort.''',
'''每位患者保留一次可用检查后，共纳入135次检查，包括113次肿瘤及22次治疗相关改变检查。AUC分别为M0的0.501、M1的0.609、MRI的0.729及Fusion的0.730，Fusion的95%百分位区间为0.606—0.833（Supplementary Table S6；Figure S2A）。Fusion减去M1的AUC差值为0.120（−0.015至0.252；Supplementary Table S7），该区间以及AP、Brier评分和对数损失差值区间均包含零。固定阈值下，Fusion正确识别109/113次肿瘤检查（96.5%）及2/22次治疗相关改变检查（9.1%；Figure S2B），描述性校准斜率为0.539。该选定队列中的AUC增量点估计仍为正，但不确定性更大。''')

put('D1',
'''Combining eight MRI features with five clinical inputs improved discrimination between residual or recurrent tumor and treatment-related change in this public cohort. Fusion achieved an OOF AUC of 0.782, compared with 0.672 for the expanded clinical model, with a paired increment of 0.110 (95% percentile interval, 0.021–0.214). Evaluating the models on identical examinations and patient-grouped folds provided a direct assessment of the information added by MRI. The positive increment across all 20 repeated partitions further supported the consistency of this finding (Supplementary Table S5B). Fusion also had the highest treatment-related-change AP and the lowest Brier-score point estimate among the four primary models. Together, these results support combining clinical context with quantitative image descriptors when assessing post-treatment glioma examinations.''',
'''在本公开队列中，将八项MRI特征与五项临床输入结合，提高了残余或复发肿瘤与治疗相关改变的鉴别能力。Fusion的OOF AUC为0.782，扩展临床模型为0.672，配对增量为0.110（95%百分位区间0.021—0.214）。各模型使用相同检查及患者分组分折，可直接评价MRI带来的额外信息。全部20套重复划分的增量均为正，进一步支持了这一发现的一致性（Supplementary Table S5B）。四个主要模型中，Fusion的治疗相关改变AP最高，Brier评分点估计最低。这些结果支持在评估胶质瘤治疗后检查时，将临床背景与定量影像描述结合。''')
put('D2',
'''Our findings add to evidence on integrating imaging and clinical information in post-treatment glioma assessment. Zhou et al. combined a radiomic score with relative cerebral blood volume, apparent diffusion coefficient, neutrophil-to-lymphocyte ratio, MGMT status and temozolomide treatment, reporting validation AUCs of 0.811 for the integrated model and 0.744 for the model excluding radiomics in a stratified internal split (8). McKenney et al. studied 74 patients undergoing repeat surgery and found higher classification accuracy when contrast-enhanced T1 and perfusion radiomics were combined with MGMT status (9). Conversely, Leone et al. reported AUCs of 0.69 for a clinical-MGMT reference model and 0.71 after adding median diffusion and perfusion measures, with no statistically significant improvement (10). De Godoy et al. reported an AUC of 0.757 in a new cohort using a diffusion-perfusion model, with mixed-tumor cases grouped differently between development and validation (11). Related work by Reecher et al. also used the UCSD-PTGBM collection (19). Across these studies, the incremental value of imaging depends on the clinical comparator, selected sequences and diagnostic categories. Our paired analysis extends this evidence using a compact, fixed set of FLAIR and postcontrast T1 features.''',
'''本研究补充了治疗后胶质瘤评估中融合影像与临床信息的证据。Zhou等将影像组学评分与相对脑血容量、表观扩散系数、中性粒细胞与淋巴细胞比值、MGMT状态及替莫唑胺治疗结合，在分层内部划分中，整合模型与去除影像组学的模型验证AUC分别为0.811和0.744（8）。McKenney等研究了74名再次手术患者，发现将增强T1及灌注影像组学与MGMT状态结合可提高分类准确率（9）。相比之下，Leone等的临床–MGMT参考模型AUC为0.69，加入弥散及灌注中位数后为0.71，改善未达到统计学显著性（10）。de Godoy等将弥散–灌注模型用于新队列，报告AUC为0.757，其开发与验证阶段对混合肿瘤病例采用了不同分组（11）。Reecher等的相关研究同样使用UCSD-PTGBM数据集（19）。这些研究表明，影像的增量价值与临床对照、所选序列及诊断类别有关。本研究采用精简且固定的FLAIR与增强T1特征，通过配对分析进一步扩展了相关证据。''')
put('D3',
'''The computational contribution is a reproducible evaluation of feature-level clinical-MRI fusion. The 13-input model combines an explicit clinical baseline with eight features whose calculations can be traced to intensity distributions and spatial gray-level relationships. A shared L2-regularized regression framework links the comparison directly to the information supplied by each input panel. Training-fold preprocessing, patient-grouped validation and retained OOF predictions make the relationship between inputs and evaluation outputs inspectable. The repository provides notebooks, derived feature tables, fold assignments and analysis locks, allowing model evaluation to be repeated from the saved features. This organization follows methodological priorities emphasized in METRICS, including documented image processing and model evaluation (20). It offers a practical computational workflow for studying MRI's additional contribution in public clinical-imaging data.''',
'''本研究在计算方法方面的贡献，是对临床–MRI特征级融合进行可复现的评价。13项输入模型将明确的临床基线与八项影像特征结合，其计算过程可追溯至强度分布及空间灰度关系。共同的L2正则化回归框架，使模型比较直接对应各输入面板提供的信息。训练折内预处理、患者分组验证及保存的OOF预测，使输入与评价输出之间的关系可以核查。仓库提供notebook、派生特征表、分折分配及分析锁定记录，可利用保存特征重复模型评价。这一组织方式与METRICS强调的影像处理及模型评价记录等方法学重点一致（20），为利用公开临床影像数据研究MRI的额外贡献提供了可操作的计算流程。''')
put('D4',
'''The selected features summarize complementary properties of regional MRI appearance. Mean intensity describes the average signal level after percentile scaling, standard deviation reflects dispersion, and Shannon entropy summarizes the distribution of discretized intensities. GLCM contrast incorporates spatial information by weighting squared gray-level differences between neighboring voxels (6). Together, these measurements capture regional signal and heterogeneity from FLAIR and postcontrast T1 images. Processing definitions are relevant to their interpretation: Foltyn-Dumitru et al. found that normalization influenced external generalization in molecular glioma classification (7), and Duenweg et al. showed that normalization affected radiomic comparability across acquisition settings (21). The present pipeline used image-specific percentile references and explicitly defined ROI discretization. This provides a consistent numerical basis for comparing clinical and imaging input sets.''',
'''所选特征概括了区域MRI表现的不同属性。均值反映百分位缩放后的平均信号水平，标准差反映离散程度，Shannon熵概括离散化强度的分布。GLCM对比度通过对相邻体素灰度差的平方加权，纳入空间信息（6）。这些测量共同描述FLAIR与增强T1图像的区域信号及异质性。处理定义与特征解释密切相关：Foltyn-Dumitru等发现，归一化影响胶质瘤分子分类的外部推广表现（7）；Duenweg等显示，归一化影响不同采集条件下影像组学特征的可比性（21）。本流程采用图像特异性百分位参考，并明确定义ROI离散化方法，为比较临床与影像输入提供一致的数值基础。''')
put('D5',
'''Boundary and split analyses examined two sources of variation in the fusion result. Eroding the supplied ROI produced a small Fusion AUC change of 0.0062, with a paired interval spanning zero, while retaining the original normalization and discretization references. This suggests that the observed performance was relatively stable under the tested boundary perturbation. Measurement repeatability is an important topic in glioma radiomics; Martin et al. investigated scan-rescan repeatability of diffusion and perfusion features (22). Our repeated grouped splits complemented the boundary analysis by assessing sensitivity to patient allocation across folds. Fusion-minus-M1 AUC differences were positive in all 20 partitions, with an observed range of 0.103–0.150. Maintaining patient separation throughout these comparisons also addressed the dependency concerns illustrated in scan-duplication experiments by Rosenblatt et al. (15). These findings support the internal consistency of the incremental discrimination.''',
'''边界及分折分析考察了融合结果的两类变化来源。保留原归一化和离散化参考，对现有ROI腐蚀后，Fusion的AUC变化较小，为0.0062，配对区间跨零，提示模型表现在所测试的边界扰动下相对稳定。测量可重复性是胶质瘤影像组学的重要问题，Martin等曾研究弥散与灌注特征的扫描—重扫描可重复性（22）。重复分组分折从患者在各折之间的分配角度补充了边界分析。20套划分中Fusion减去M1的AUC差值均为正，实际范围为0.103—0.150。各项比较始终保持患者隔离，也回应了Rosenblatt等扫描复制实验所展示的依赖性问题（15）。这些结果支持增量鉴别表现的内部一致性。''')
put('D6',
'''The examination-level and single-scan analyses offer complementary views of model performance. The primary analysis used all 182 eligible examinations, allowing repeated post-treatment observations from the same patient while keeping that patient within one validation fold. The single-scan analysis gave each of the 135 patients one examination and yielded a Fusion AUC of 0.730. Its paired AUC increment was 0.120, close in magnitude to the primary increment of 0.110, although the interval of -0.015 to 0.252 included zero. Selecting one examination also reduced treatment-related-change observations from 41 to 22 and altered the training sample and class distribution. The positive point estimates across both analyses favor an imaging contribution, with stronger precision in the primary examination-level analysis. Reporting both makes the influence of examination selection visible alongside the main comparison.''',
'''检查层面与单扫描分析从两个角度描述模型表现。主分析使用全部182次符合条件的检查，纳入同一患者治疗后的重复观察，同时将该患者保持在同一验证折内。单扫描分析为135名患者各保留一次检查，Fusion AUC为0.730，配对AUC增量为0.120，与主分析0.110的增量接近，但−0.015至0.252的区间包含零。选择一次检查也使治疗相关改变观察从41次减少至22次，并改变了训练样本及类别分布。两项分析的增量点估计均为正，有利于支持影像的贡献，而主分析的估计更精确。同时报告两项分析，使检查选择对结果的影响得到呈现。''')
put('D7',
'''Discrimination, calibration and classification describe different aspects of an imaging model, as emphasized in Metrics Reloaded (25). Fusion's AUC increment indicates improved ranking of tumor and treatment-related-change examinations relative to the clinical comparator. At the fixed 0.5 threshold, it correctly identified 130/141 tumor examinations and 12/41 treatment-related-change examinations, while its calibration slope was 0.667. The paired Brier-score and logarithmic-loss intervals included zero, so the clearest improvement concerned discrimination. This distinction helps identify the next stage of evaluation. A defined clinical use would require selecting a threshold according to the consequences of each classification and evaluating performance for both diagnostic categories. Independent assessment of calibration and decision thresholds would clarify how the additional MRI information could contribute to post-treatment assessment.''',
'''鉴别、校准及分类反映影像模型的不同方面，Metrics Reloaded也强调对这些维度分别评价（25）。Fusion的AUC增量表明，相较临床对照，其对肿瘤及治疗相关改变检查的排序能力有所提高。在固定0.5阈值下，模型正确识别130/141次肿瘤检查及12/41次治疗相关改变检查，校准斜率为0.667。配对Brier评分及对数损失区间包含零，因此改善证据主要集中于鉴别能力。这一区别为下一阶段评价提供了方向：针对明确的临床用途，结合各类分类结果的后果选择阈值，并评价两类诊断的表现。对校准及决策阈值开展独立评价，有助于明确额外MRI信息如何参与治疗后评估。''')
put('D8',
'''Two limitations merit consideration. The analysis used the available imaging subset of a single public collection and its existing diagnostic labels and annotations; the source diagnoses combined pathological and clinical-imaging evidence, and the information available during correction of the supplied masks was incompletely documented (5,17). This makes independently adjudicated diagnoses and documented annotation procedures priorities for replication. In addition, performance was assessed internally in 135 patients, including 41 treatment-related-change examinations in the primary analysis. The bootstrap intervals condition on the saved OOF predictions, and the single-scan increment remained imprecise. Evaluation in a larger independent cohort would help establish the generalizability and precision of the observed improvement.''',
'''本研究有两点局限需要考虑。研究采用单一公开数据集的可用影像子集及既有诊断标签和标注；源诊断结合病理与临床影像证据，现有掩膜修正过程中可获得的信息记录尚不完整（5,17）。后续复现宜采用独立裁定的诊断并完整记录标注流程。此外，性能评价基于135名患者的内部验证，主分析包含41次治疗相关改变检查；bootstrap区间以保存的OOF预测为条件，单扫描增量的估计仍不够精确。在更大的独立队列中评价，有助于明确所观察改善的可推广性及估计精度。''')
put('D9',
'''Further evaluation should retain the paired comparison between the clinical and fusion pipelines while broadening the patient population and acquisition settings. External-validation guidance recommends representative cohorts and sample-size planning for precise performance estimation (23,24). A multicenter study with documented acquisition metadata, consistent annotation procedures and adjudicated outcomes would provide a suitable setting. Treatment-related-change recognition and calibration should be evaluated alongside AUC, with thresholds linked to a defined clinical action. Additional sequences or molecular predictors could then be assessed for their incremental contribution using the same comparison framework. The present public-data workflow provides reusable feature definitions, code and validation outputs to support this progression.''',
'''后续评价可在保留临床与融合流程配对比较的基础上，拓展患者人群及影像采集条件。外部验证指南建议采用具有代表性的队列，并为精确估计性能规划样本量（23,24）。具备完整采集元数据、一致标注流程及经裁定结局的多中心研究，可为此提供合适场景。在评价AUC的同时，应关注治疗相关改变识别及校准，并使阈值对应明确的临床处置。随后可沿用同一比较框架，评价额外序列或分子预测变量的增量贡献。现有公开数据工作流程提供了可复用的特征定义、代码及验证输出，为后续研究奠定基础。''')
put('C1',
'''Clinical-MRI feature fusion improved discrimination between residual or recurrent tumor and treatment-related change in public UCSD-PTGBM data. In 182 examinations from 135 patients, Fusion achieved an OOF AUC of 0.782 and a paired increment of 0.110 over the expanded clinical model (95% percentile interval, 0.021–0.214). The increment was positive across 20 patient-grouped partitions. The reproducible workflow supports evaluating a compact MRI feature panel alongside clinical information, with independent cohort assessment as the next step toward clinical application.''',
'''在公开UCSD-PTGBM数据中，临床–MRI特征融合提高了残余或复发肿瘤与治疗相关改变的鉴别能力。在135名患者的182次检查中，Fusion的OOF AUC为0.782，较扩展临床模型的配对增量为0.110（95%百分位区间0.021—0.214），20套患者分组划分中增量均为正。可复现的工作流程支持将精简MRI特征面板与临床信息结合进行评价，后续可通过独立队列评估进一步推进临床应用研究。''')

ABSTRACT=deepcopy(OLD['abstract_paragraphs'])
ABSTRACT[-1]['english']='Clinical-MRI feature fusion improved discrimination over the clinical model in this public cohort. The reproducible approach merits evaluation in independent cohorts, with attention to treatment-change recognition and calibration.'
ABSTRACT[-1]['chinese']='在本公开队列中，临床–MRI特征融合提高了相较临床模型的鉴别能力。这一可复现方法值得在独立队列中进一步评价，并关注治疗相关改变识别及校准。'
DISCUSSION_HEADINGS=[
 '4.1 Principal findings','4.2 Comparison with related studies','4.3 Computational contribution',
 '4.4 Interpretation of MRI features','4.5 Boundary sensitivity',
 '4.6 Clinical and imaging information','4.7 Clinical implications',
 '4.8 Limitations','4.9 Further evaluation']
SUPP_KEYS=['TableS1','TableS2','TableS3','TableS4','TableS5A','TableS5B','TableS6','TableS7']

# Additional user instruction: remove the two standalone result subsections
# from the main text while retaining the analyses in supplementary material.
SUPPLEMENTARY_ANALYSES=[deepcopy(p)for p in TEXT if p['id']in ['R7','R8']]
SUPPLEMENTARY_ANALYSES[0]['english']=SUPPLEMENTARY_ANALYSES[0]['english'].replace('Figure 3A','Figure S3A').replace('Figure 3B','Figure S3B')
TEXT=[p for p in TEXT if p['id']not in ['R7','R8']]
def replace_text(pid,en,zh):
    p=next(p for p in TEXT if p['id']==pid);p.update(english=en,chinese=zh)
replace_text('D5',
'''The boundary analysis examined how a modest change to the supplied annotation region affected the extracted features and model performance. Fusion's AUC changed from 0.782 to 0.788 after erosion, with a paired difference of 0.0062 (-0.0051 to 0.0183). Retaining the original normalization and discretization references allowed the comparison to focus on the change in the sampled region. The resulting AUC estimates were similar, suggesting relative stability under this specific perturbation. Measurement reproducibility is also being investigated in glioma radiomics: Martin et al. evaluated scan-rescan repeatability of diffusion and perfusion features (22). The present boundary analysis provides complementary information about feature extraction within the supplied lesion region and supports further investigation of how annotation choices affect clinical-MRI fusion.''',
'''边界分析考察了对现有标注区域进行适度改变后，所提取特征及模型表现的变化。腐蚀后Fusion的AUC由0.782变为0.788，配对差值为0.0062（−0.0051至0.0183）。保留原归一化及离散化参考，使比较聚焦于取样区域的改变。两项AUC估计相近，提示模型在此次特定扰动下相对稳定。测量可重复性也是胶质瘤影像组学的研究重点，Martin等评价了弥散与灌注特征的扫描—重扫描可重复性（22）。本研究的边界分析补充了所提供病灶区域内特征提取的相关信息，并为进一步研究标注选择如何影响临床–MRI融合提供依据。''')
replace_text('D6',
'''The four-model comparison helps clarify how clinical and imaging inputs contribute to discrimination. Expanding the age-and-sex model increased the AUC point estimate from 0.526 to 0.672, while the MRI model achieved 0.762 and Fusion achieved 0.782. Age, other treatment before imaging and the initial-event interval describe the examination context, whereas intensity and texture features summarize the regional MRI appearance. Their combination places these two information sources in a common regression model. The primary paired comparison showed that MRI added discriminatory information to the expanded clinical baseline. This design provides a clear starting point for evaluating additional clinical variables or imaging sequences by comparing their contribution against the same reference pipeline.''',
'''四模型比较有助于理解临床与影像输入对鉴别能力的贡献。扩展年龄与性别模型后，AUC点估计由0.526提高至0.672；MRI模型为0.762，Fusion为0.782。年龄、扫描前其他治疗及初始事件间隔描述检查的临床背景，强度与纹理特征则概括区域MRI表现。融合模型将这两类信息纳入同一回归模型。主要配对比较显示，MRI在扩展临床基线之上提供了额外鉴别信息。这一设计也为后续纳入其他临床变量或影像序列提供了清晰起点，可依照相同参考流程比较新增信息的贡献。''')
replace_text('D8',
'''Two limitations merit consideration. The analysis used the available imaging subset of a single public collection and its existing diagnostic labels and annotations; source diagnoses combined pathological and clinical-imaging evidence, and the information available during correction of the supplied masks was incompletely documented (5,17). Independently adjudicated diagnoses and documented annotation procedures would strengthen replication. In addition, performance was evaluated internally in 135 patients, including 41 treatment-related-change examinations in the primary analysis. Bootstrap intervals condition on fixed OOF predictions, and the supplementary single-scan AUC-increment interval included zero (Supplementary Table S7). A larger independent cohort would help establish the precision and generalizability of the observed improvement.''',
'''本研究有两点局限需要考虑。研究采用单一公开数据集的可用影像子集及既有诊断标签和标注；源诊断结合病理与临床影像证据，现有掩膜修正过程中可获得的信息记录尚不完整（5,17）。采用独立裁定的诊断并完整记录标注流程，可加强后续复现。此外，性能评价基于135名患者的内部验证，主分析包含41次治疗相关改变检查；bootstrap区间以固定OOF预测为条件，补充材料中的单扫描AUC增量区间包含零（Supplementary Table S7）。更大的独立队列有助于明确所观察改善的估计精度及可推广性。''')
ABSTRACT[1]['english']='This secondary analysis of the public UCSD-PTGBM collection included 182 examinations from 135 patients. Four fixed L2-penalized logistic models used age and sex, five clinical inputs, eight MRI descriptors, or their fusion. MRI features summarized fluid-attenuated inversion recovery (FLAIR) and postcontrast T1 intensity and texture within supplied annotation regions. Five-fold patient-grouped validation generated out-of-fold predictions. The primary comparison was the Fusion-minus-expanded-clinical area under the receiver operating characteristic curve (AUC) difference. Intervals used 2000 patient-cluster bootstrap attempts conditional on fixed predictions. Sensitivity analyses examined annotation boundaries and examination selection.'
ABSTRACT[1]['chinese']='本项基于公开UCSD-PTGBM数据集的二次分析纳入135名患者的182次检查。四个固定的L2惩罚Logistic模型分别采用年龄与性别、五项临床输入、八项MRI描述量及其融合。MRI特征汇总所提供标注区域内液体衰减反转恢复（FLAIR）与增强T1的强度及纹理。五折患者分组验证生成折外预测，主要比较为Fusion减去扩展临床模型的受试者工作特征曲线下面积（AUC）差值。采用2000次患者聚类bootstrap构建条件于固定预测的区间，敏感性分析考察标注边界与检查选择。'
ABSTRACT[2]['english']='Fusion achieved an AUC of 0.782 versus 0.672 for the expanded clinical model and 0.762 for MRI alone. The paired AUC increment was 0.110 (95% percentile interval, 0.021 to 0.214). Treatment-related-change average precision was 0.505 for Fusion and 0.354 for the clinical model. At the fixed 0.5 threshold, Fusion recognized 130/141 tumor examinations and 12/41 treatment-related-change examinations. Its descriptive calibration slope was 0.667. Brier-score and logarithmic-loss difference intervals included zero.'
ABSTRACT[2]['chinese']='Fusion的AUC为0.782，扩展临床模型为0.672，仅MRI模型为0.762。配对AUC增量为0.110（95%百分位区间0.021至0.214）。Fusion与临床模型的治疗相关改变平均精确率分别为0.505和0.354。在固定0.5阈值下，Fusion识别出130/141次肿瘤检查及12/41次治疗相关改变检查，描述性校准斜率为0.667。Brier评分及对数损失差值的区间包含零。'

def update_objects():
    tables=deepcopy(OLD['tables']);figs=deepcopy(OLD['figures'])
    common='Intervals are the 2.5th and 97.5th percentiles of 2000 paired patient-cluster bootstrap attempts, conditional on fixed OOF predictions.'
    definitions='M0: age and sex; M1: five clinical inputs; MRI: eight MRI features; Fusion: clinical and MRI inputs.'
    ap='AP denotes average precision for treatment-related change, scored by one minus the tumor probability. Correct treatment-change counts use tumor probability below 0.5.'
    directions='Differences are Fusion minus M1. Positive AUC or AP and negative Brier-score or logarithmic-loss differences favor Fusion.'
    tables['Table1']['notes']=[tables['Table1']['notes'][0], 'Patient counts are deduplicated within each column; eight patients contributed examinations in both diagnostic categories. The initial event is recorded surgery, treatment or diagnosis. Interval summaries use observed values.']
    tables['TableS1']['notes']=[tables['TableS1']['notes'][0],'Coverage refers to availability in the accessed image mirror. Coverage groups contain 135 and 41 distinct patients, respectively. Interval summaries use observed values for recorded initial surgery, treatment or diagnosis.']
    tables['Table2']['notes']=[definitions+' All models used the same 182 examinations and patient-grouped folds.',common,ap+' Treatment-change prevalence was 41/182 (0.225). Lower Brier scores indicate smaller probability error. Values are rounded to three decimals.']
    tables['Table3']['notes']=[directions+' All contrasts used 182 examinations from 135 patients.',common,'ROC AUC is the primary incremental metric. Other metrics are auxiliary. Differences use unrounded metrics and are displayed to three decimals.']
    tables['TableS2']['notes']=[definitions+' Imputation and standardization parameters were estimated within training folds.','The initial event is recorded initial surgery, other initial treatment or diagnosis. Elapsed days were transformed as ln(1 + days).','MRI features use the positive-label union of the supplied standard BraTS mask. Texture distances are expressed in voxel-grid units.']
    tables['TableS3']['notes']=['The cohort contains 182 examinations from 135 patients. A tumor probability of at least 0.5 defines tumor. Balanced accuracy is the mean of tumor sensitivity and treatment-change recognition.','The calibration offset intercept fixes the predicted-log-odds coefficient at one. The calibration slope is jointly estimated with a free intercept in a separate unpenalized logistic diagnostic regression.','Calibration estimates describe the pooled OOF predictions. Metrics are rounded to three decimals; classification counts retain exact denominators.']
    tables['TableS4']['notes']=[tables['TableS4']['notes'][0],common+' The eroded-minus-original Fusion AUC difference was 0.0062 (-0.0051 to 0.0183).',ap+' The original ROI defines the primary analysis. Metrics use three decimals; the paired AUC difference uses four.']
    tables['TableS5A']['notes']=[tables['TableS5A']['notes'][0],'Q1 and Q3 are quartiles across the 20 partitions. Ranges are observed minima and maxima in this cohort. The primary Fusion AUC is 0.782 from partition 1.']
    tables['TableS5B']['notes']=['Fusion-minus-M1 differences were calculated within each repetition before summarizing their median and quartiles.',directions+' Q1 and Q3 are quartiles; ranges are observed minima and maxima.','Favorable-repetition counts describe the direction of paired estimates across the 20 partitions of the same cohort. Values are displayed to three decimals.']
    tables['TableS6']['notes']=['The selected cohort contains 135 examinations from 135 patients: 113 tumor and 22 treatment-related-change examinations. Each patient contributes the smallest available numeric session in the imaging subset.','Four models share the rebuilt patient-grouped five-fold partition. '+common,ap+' Values are rounded to three decimals.']
    tables['TableS7']['notes']=[directions+' The comparison uses the same 135 selected examinations and rebuilt folds.',common+' All four intervals include zero.','AP denotes average precision for treatment-related change. Differences use unrounded metrics and are displayed to four decimals.']
    figs['Figure 1']['caption']='Figure 1. Cohort assembly and model comparison. The imaging cohort contains 182 examinations from 135 patients, including 141 tumor and 41 treatment-related-change examinations. Clinical inputs and MRI features enter four fixed models evaluated using the same patient-grouped five-fold partition. M0: age and sex; M1: expanded clinical model; MRI: imaging features; Fusion: clinical and MRI inputs.'
    figs['Figure 2']['caption']='Figure 2. Patient-grouped internal validation in 182 examinations from 135 patients. A, ROC curves based on pooled OOF probabilities. B, Precision-recall curves for treatment-related change, with prevalence marked by the dashed horizontal line. C, Descriptive calibration by probability quantile groups, comparing mean predicted probability with observed tumor proportion. The diagonal lines indicate chance discrimination in A and agreement in C.'
    figs['Figure 3']['caption']='Figure 3. Performance across 20 patient-grouped five-fold partitions of the same 182 examinations from 135 patients. A, Pooled OOF AUC for each model within each partition. B, Paired Fusion-minus-M1 AUC differences, with zero marked by the dashed line. Partition 1 reproduces the primary analysis; the displayed variation describes split sensitivity within this cohort.'
    figs['Figure S1']['caption']='Figure S1. Boundary sensitivity in 182 examinations from 135 patients. Points show OOF AUCs for the original and eroded MRI and Fusion models. Error bars are 95% patient-cluster bootstrap percentile intervals conditional on fixed OOF probabilities. Erosion used one iteration of the six-neighbor operation on the supplied voxel grid. The paired Fusion difference is reported in Table S4.'
    figs['Figure S2']['caption']='Figure S2. Sensitivity analysis using one available examination per patient, comprising 113 tumor and 22 treatment-related-change examinations. A, OOF AUCs and patient-bootstrap 95% percentile intervals conditional on fixed predictions. B, Treatment-change recognition at the fixed tumor-probability threshold of 0.5, with exact correct/total counts. Selection used the smallest available numeric session, followed by rebuilt patient-grouped folds.'
    figs['Figure S3']=figs.pop('Figure 3');figs['Figure S3']['caption']=figs['Figure S3']['caption'].replace('Figure 3.','Figure S3.')
    return tables,figs

def revise_flow_figure(figs):
    # Redraw the cohort flow with aligned centers and orthogonal connectors.
    # Only presentation changes; all cohort counts come from frozen results.
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyBboxPatch,FancyArrowPatch
    folder=OUT/'figures';folder.mkdir(parents=True,exist_ok=True)
    plt.rcParams.update({'font.family':'Nimbus Roman','font.size':11,'pdf.fonttype':42,'svg.fonttype':'none'})
    fig,ax=plt.subplots(figsize=(8,8.8));fig.subplots_adjust(left=.025,right=.975,top=.98,bottom=.02)
    ax.set(xlim=(0,1),ylim=(0,1));ax.axis('off')
    nodes=[];segments=[]
    def box(key,cx,cy,w,h,label,fill='#F3F5F7',fontsize=10.5):
        left,bottom=cx-w/2,cy-h/2
        ax.add_patch(FancyBboxPatch((left,bottom),w,h,boxstyle='round,pad=0.005,rounding_size=0.008',facecolor=fill,edgecolor='#607583',linewidth=.9))
        text=ax.text(cx,cy,label,ha='center',va='center',fontsize=fontsize,linespacing=1.25)
        nodes.append({'key':key,'center':[cx,cy],'width':w,'height':h,'label':label,'_text':text})
    def line(points,arrow=False):
        for a,b in zip(points[:-2],points[1:-1]):ax.plot([a[0],b[0]],[a[1],b[1]],color='#607583',lw=.9)
        a,b=points[-2:]
        if arrow:ax.add_patch(FancyArrowPatch(a,b,arrowstyle='-|>',mutation_scale=11,color='#607583',linewidth=.9,shrinkA=0,shrinkB=2.5))
        else:ax.plot([a[0],b[0]],[a[1],b[1]],color='#607583',lw=.9)
        segments.append({'points':points,'arrow':arrow})
    box('source',.60,.92,.52,.08,'Public UCSD-PTGBM collection\n243 examinations / 178 patients')
    box('clinical',.60,.74,.52,.10,'Linked clinical cohort\n238 examinations / 176 patients')
    box('linkage_exclusions',.16,.74,.26,.10,'Excluded at linkage\n5 examinations / 2 patients',fontsize=9.5)
    line([(.60,.88),(.60,.79)],True)
    line([(.60,.835),(.16,.835),(.16,.79)],True)
    box('imaging',.60,.56,.52,.13,'Complete imaging cohort\n182 examinations / 135 patients\n141 tumor / 41 treatment-related change')
    box('other_clinical',.16,.56,.26,.13,'Clinical records retained\n56 examinations / 41 patients\nCoverage summary (Table S1)',fontsize=9.3)
    line([(.60,.69),(.60,.625)],True)
    line([(.60,.66),(.16,.66),(.16,.625)],True)
    box('inputs',.60,.40,.52,.10,'Clinical and MRI inputs\n5 clinical inputs + 8 MRI features')
    line([(.60,.495),(.60,.45)],True)
    centers=[.14,.38,.62,.86]
    labels=['M0\nAge and sex\n2 inputs','M1\nExpanded clinical\n5 inputs','MRI\nImage features\n8 inputs','Fusion\nClinical + MRI\n13 inputs']
    line([(.60,.35),(.60,.32)]);line([(.14,.32),(.86,.32)])
    for x,label in zip(centers,labels):
        box(label.split('\n')[0],x,.235,.20,.10,label,fill='#ECF0F3',fontsize=10)
        line([(x,.32),(x,.285)],True);line([(x,.185),(x,.15)])
    line([(.14,.15),(.86,.15)])
    box('validation',.50,.075,.90,.10,'Patient-grouped five-fold validation\nPaired comparison of Fusion versus the expanded clinical model')
    line([(.50,.15),(.50,.125)],True)
    fig.canvas.draw();renderer=fig.canvas.get_renderer()
    for node in nodes:
        bb=node['_text'].get_window_extent(renderer).transformed(ax.transData.inverted())
        cx,cy=node['center'];w,h=node['width'],node['height']
        assert bb.x0>=cx-w/2 and bb.x1<=cx+w/2 and bb.y0>=cy-h/2 and bb.y1<=cy+h/2,node['key']
        node.pop('_text');node['text_bbox']=[bb.x0,bb.y0,bb.x1,bb.y1]
    for ext in ['png','pdf','svg']:
        target=folder/f'Figure1_Cohort_Flow.{ext}';fig.savefig(target,format=ext,dpi=600)
        if ext=='svg':target.write_text('\n'.join(line.rstrip()for line in target.read_text().splitlines())+'\n')
    plt.close(fig)
    (REPORT/'figure1_layout.json').write_text(json.dumps({'nodes':nodes,'connectors':segments,'all_text_inside_nodes':True,'orthogonal_connectors':True,'cohort_counts_unchanged':True},indent=2)+'\n')
    figs['Figure 1']['path']=str((folder/'Figure1_Cohort_Flow.png').relative_to(ROOT))
    figs['Figure 1']['sha256']=base.sha(folder/'Figure1_Cohort_Flow.png')

def citation_numbers(s,chinese=False):
    pat=r'（([0-9]+(?:\s*[,，]\s*[0-9]+)*)）' if chinese else r'\(([0-9]+(?:\s*,\s*[0-9]+)*)\)'
    return [int(n.strip()) for group in re.findall(pat,s) for n in re.split('[,，]',group)]

def renumber():
    numbering={}
    for p in TEXT:
        en=citation_numbers(p['english']);zh=citation_numbers(p['chinese'],True)
        assert en==zh,(p['id'],en,zh)
        for n in en:
            if n not in numbering:numbering[n]=len(numbering)+1
    assert set(numbering)==set(range(1,26)),numbering
    paragraphs=deepcopy(TEXT)
    for p in paragraphs:
        for language,pat in [('english',r'\(([0-9]+(?:\s*,\s*[0-9]+)*)\)'),('chinese',r'（([0-9]+(?:\s*[,，]\s*[0-9]+)*)）')]:
            left,right=('(',')') if language=='english' else ('（','）')
            p[language]=re.sub(pat,lambda m:left+','.join(str(numbering[int(x.strip())]) for x in re.split('[,，]',m.group(1)))+right,p[language])
        p['citation_ids']=[OLD['references'][n-1]['id'] for n in citation_numbers(next(x['english']for x in TEXT if x['id']==p['id']))]
    refs=[dict(OLD['references'][n-1],number=new) for n,new in numbering.items()]
    return paragraphs,refs,numbering

def sections():
    _,source,_,_,_=base.load_sources()
    result=[]
    for d in source:
        if d['stage'] in [39,40]:continue
        if d['stage'] in range(41,53):continue
        ids=[p['id']for p in d['paragraphs']if p['id']!='M8']
        # Put the nine discussion sections immediately before the conclusion.
        if d['stage']==53:
            for n,title in enumerate(DISCUSSION_HEADINGS,1):result.append({'heading':title,'ids':[f'D{n}'],'major':'4 Discussion' if n==1 else None})
        result.append({'heading':d['subsection'],'ids':ids,'major':{23:'2 Materials and methods',33:'3 Results'}.get(d['stage'])})
    return result

def build():
    OUT.mkdir(parents=True,exist_ok=True);REPORT.mkdir(parents=True,exist_ok=True)
    paragraphs,refs,numbering=renumber();lookup={p['id']:p for p in paragraphs};assert len(lookup)==41
    tables,figs=update_objects();revise_flow_figure(figs)
    _,source,_,_,_=base.load_sources();eq,_,_=base.objects(source)
    # All equation objects retain their original XML and numbering.
    base.ABSTRACT=ABSTRACT
    outputs=[]
    abstract=base.document('Abstract');abstract.add_paragraph(OLD['title'],style='Title');base.add_abstract(abstract,True)
    path=OUT/'Abstract_Revision1_EN_ZH.docx';abstract.save(path);outputs.append(path)
    (OUT/'Abstract_Revision1_EN.txt').write_text('\n\n'.join(p['label']+': '+p['english']for p in ABSTRACT)+'\n\nKeywords: '+'; '.join(OLD['keywords'])+'\n')
    for bilingual,name in [(True,'Manuscript_Revision1_EN_ZH.docx'),(False,'Manuscript_Revision1_EN.docx')]:
        doc=base.document(OLD['title']);doc.add_paragraph(OLD['title'],style='Title');base.add_abstract(doc,bilingual)
        for section in sections():
            if section['major']:doc.add_heading(section['major'],1)
            level=1 if section['heading'].startswith(('1 Introduction','5 Conclusion'))else 2
            head=doc.add_heading(section['heading'],level)
            if section['heading'].startswith('1 Introduction'):head.paragraph_format.page_break_before=True
            for pid in section['ids']:
                p=lookup[pid];base.paragraph(doc,p['english'])
                if bilingual:base.paragraph(doc,p['chinese'],True)
                for form in eq.get(pid,[]):doc.element.body.insert(len(doc.element.body)-1,deepcopy(form))
        h=doc.add_heading('References',1);h.paragraph_format.page_break_before=True
        for r in refs:
            p=base.paragraph(doc,str(r['number'])+'. '+r['citation']);p.paragraph_format.keep_together=True
        for key in ['Table1','Table2','Table3']:base.add_table(doc,tables[key])
        for key in ['Figure 1','Figure 2']:base.add_figure(doc,key,figs[key],key!='Figure 1')
        base.lines(doc);endmatter=False
        for p in doc.paragraphs:
            if p.text==tables['Table1']['title']:endmatter=True
            if endmatter:p._p.get_or_add_pPr().append(OxmlElement('w:suppressLineNumbers'))
        path=OUT/name;doc.save(path);outputs.append(path)
    supp=base.document('Supplementary material');supp.add_paragraph('Supplementary material',style='Title')
    supp.add_heading('Additional sensitivity analyses',1)
    for item in SUPPLEMENTARY_ANALYSES:
        supp.add_heading('Repeated patient-grouped partitions' if item['id']=='R7' else 'Single-scan analysis',2)
        base.paragraph(supp,item['english'])
    for i,key in enumerate(SUPP_KEYS):
        base.add_table(supp,tables[key],key=='TableS2' or len(tables[key]['headers'])>5)
    for key in ['Figure S1','Figure S2','Figure S3']:base.add_figure(supp,key,figs[key],True)
    path=OUT/'Supplementary_Revision1_EN.docx';supp.save(path);outputs.append(path)
    previous={str(p.relative_to(ROOT)):base.sha(p) for p in [ROOT/'reports/ucsd_manuscript_step55/assembly_receipt.json',ROOT/'scripts/assemble_ucsd_manuscript55.py',*[ROOT/p for p in OLD['outputs']]]}
    previous.update(OLD['source_files_sha256'])
    previous.update({f['path']:f['sha256']for f in OLD['figures'].values()})
    previous.update({str(p.relative_to(ROOT)):base.sha(p)for p in ROOT.glob('data/processed/ucsd*.csv')})
    for rel,h in previous.items():assert base.sha(ROOT/rel)==h,rel
    counts={key:sum(base.wc(p['english'])for p in paragraphs if p['id'].startswith(key))for key in ['I','M','R','D','C']}
    receipt={'stage':56,'date':'2026-10-30','revision':'First editorial revision','title':OLD['title'],'main_paragraphs':paragraphs,'supplementary_analyses':SUPPLEMENTARY_ANALYSES,'sections':sections(),'abstract_paragraphs':ABSTRACT,'section_words':counts,'main_english_words':sum(counts.values()),'abstract_english_words':sum(base.wc(p['english'])for p in ABSTRACT),'previous_main_english_words':8200,'discussion_parts':9,'limitations_paragraphs':1,'limitations_items':2,'references':refs,'citation_mapping_previous_to_revision':numbering,'tables':tables,'figures':figs,'source_files_sha256':previous,'outputs':{str(p.relative_to(ROOT)):base.sha(p)for p in outputs},'model_refits':0,'new_predictions':0,'new_bootstrap_draws':0,'numeric_table_cells_changed':0,'figure1_redrawn_with_orthogonal_alignment':True,'Figure3_moved_to_supplement_as_FigureS3':True,'post_conclusion_declarations_deferred':True}
    (REPORT/'revision_receipt.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n')
    with (REPORT/'paragraph_comparison.csv').open('w') as f:
        w=csv.writer(f,lineterminator='\n');w.writerow(['Section','Previous words','Revised words'])
        for key,old in [('I',1054),('M',3131),('R',1446),('D',2424),('C',145)]:w.writerow([key,old,counts[key]])
    print(json.dumps({'outputs':list(receipt['outputs']),'main_words':sum(counts.values()),'section_words':counts,'abstract_words':receipt['abstract_english_words']},ensure_ascii=False))

def verify():
    from pypdf import PdfReader
    d=json.loads((REPORT/'revision_receipt.json').read_text())
    for rel,h in d['source_files_sha256'].items():assert base.sha(ROOT/rel)==h,rel
    assert len(d['main_paragraphs'])==41 and len([s for s in d['sections']if s['heading'].startswith('4.')])==9
    assert len([p for p in d['main_paragraphs']if p['id']=='D8'])==1
    # Source analyses are rechecked read-only; the saved result values remain authoritative.
    base.abstract_data_audit()
    metrics=json.loads((ROOT/'reports/ucsd_fusion/metrics.json').read_text())
    for value in [metrics['metrics'][k]['roc_auc']for k in ['Fusion','M1','MRI']]+[metrics['metrics'][k]['ap_treatment_change']for k in ['Fusion','M1']]+[metrics['calibration']['Fusion']['calibration_slope']]:
        assert f'{value:.3f}'in ABSTRACT[2]['english']
    _,source,_,_,_=base.load_sources();eq,_,_=base.objects(source)
    math=[etree.tostring(o,method='c14n',exclusive=True)for v in eq.values()for obj in v for o in obj.xpath('.//m:oMath')]
    docs={}
    for rel,h in d['outputs'].items():
        p=ROOT/rel;assert base.sha(p)==h;doc=Document(p);texts=[x.text for x in doc.paragraphs]
        ismain=p.name.startswith('Manuscript');isabstract=p.name.startswith('Abstract');isenglish=p.name.endswith('_EN.docx')
        if ismain or isabstract:
            for item in ABSTRACT:
                assert item['label']+': '+item['english']in texts
                if not isenglish:assert texts.index(item['chinese'])==texts.index(item['label']+': '+item['english'])+1
        if ismain:
            for item in d['main_paragraphs']:
                assert texts.count(item['english'])==1,item['id']
                if not isenglish:assert texts.index(item['chinese'])==texts.index(item['english'])+1
            assert [x.text for x in doc.paragraphs if x.style.name=='Heading 2' and x.text.startswith('4.')]==DISCUSSION_HEADINGS
            assert len(doc.tables)==3
            for r in d['references']:assert str(r['number'])+'. '+r['citation']in texts
            actual=[etree.tostring(o,method='c14n',exclusive=True)for o in doc.element.xpath('//m:oMath')]
            assert actual==math
            assert not any(t.startswith(('3.7 ','3.8 '))for t in texts)
            keys=['Table1','Table2','Table3'];figkeys=['Figure 1','Figure 2']
        elif isabstract:keys=[];figkeys=[]
        else:
            keys=SUPP_KEYS;figkeys=['Figure S1','Figure S2','Figure S3'];assert len(doc.tables)==8
            for item in SUPPLEMENTARY_ANALYSES:assert texts.count(item['english'])==1
        for table,key in zip(doc.tables,keys):assert [[c.text for c in row.cells]for row in table.rows]==[OLD['tables'][key]['headers'],*OLD['tables'][key]['rows']],key
        with zipfile.ZipFile(p)as z:
            images=[hashlib.sha256(z.read(n)).hexdigest()for n in z.namelist()if n.startswith('word/media/')]
            assert sorted(images)==sorted(d['figures'][k]['sha256']for k in figkeys)
        if isenglish:
            alltext='\n'.join(texts+[cell.text for t in doc.tables for row in t.rows for cell in row.cells])
            assert not re.search(r'[\u3400-\u9fff]',alltext)
        pdf=OUT/(p.stem+'_render')/(p.stem+'.pdf');pages=len(PdfReader(pdf).pages)
        docs[p.name]={'pages':pages,'editable_tables':len(keys),'native_equation_groups':len(doc.element.xpath('//m:oMath')),'figures':len(figkeys)}
    order=[]
    for p in d['main_paragraphs']:
        for n in citation_numbers(p['english']):
            if n not in order:order.append(n)
    assert order==list(range(1,26))
    assert {r['id']for r in d['references']}=={r['id']for r in OLD['references']}
    for key in OLD['tables']:assert d['tables'][key]['rows']==OLD['tables'][key]['rows']
    assert 'M8'not in [p['id']for p in d['main_paragraphs']]
    for bad in ['processing failure','not interpreted as','did not establish','did not constitute','not used to select','not interpreted as causal','without outcome-based']:
        assert not any(bad in p['english'] for p in d['main_paragraphs']),bad
    v={'success':True,'stage':56,'main_english_words':d['main_english_words'],'previous_main_english_words':8200,'abstract_english_words':d['abstract_english_words'],'discussion_parts':9,'limitations_paragraphs':1,'limitations_items':2,'numeric_table_cells_unchanged':True,'native_formulas_preserved':10,'references_all_retained':25,'citations_first_appearance_order':True,'prior_files_unchanged':len(d['source_files_sha256']),'source_data_and_OOF_results_unchanged':True,'main_result_subsections_removed':['3.7','3.8'],'sensitivity_details_retained_in_supplement':True,'Figure3_moved_as_FigureS3':True,'Figure1_alignment_verified':True,'statistical_uncertainty_and_classification_results_retained':True,'models_refitted':0,'new_bootstrap_draws':0,'documents':docs,'sha256':{str(p.relative_to(ROOT)):base.sha(p)for p in [*[p for p in OUT.iterdir()if p.is_file()],*sorted((OUT/'figures').glob('*'))]if p.is_file()},'declarations_still_deferred':True}
    (REPORT/'artifact_verification.json').write_text(json.dumps(v,ensure_ascii=False,indent=2)+'\n');print(json.dumps(v,ensure_ascii=False,indent=2))

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--verify',action='store_true');args=parser.parse_args();verify()if args.verify else build()
