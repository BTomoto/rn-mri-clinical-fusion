"""Draft M20 from saved software, notebook and source-recovery records."""
from pathlib import Path
import argparse,hashlib,json,re,sys,zipfile
import numpy as np
import pandas as pd
from lxml import etree
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'outputs/ucsd-manuscript-step32-20261029'
REPORT=ROOT/'reports/ucsd_manuscript_step32'
PARAGRAPHS=[
 {'id':'M20','citation_ids':['S40'],
  'english': 'Analyses were implemented in Python 3 using scikit-learn 1.8.0 (5), NumPy 2.3.5, pandas 2.2.3, SciPy 1.17.0 and Matplotlib 3.10.8. Jupyter notebooks for feature extraction, model evaluation and sensitivity analyses were executed through an in-process ipykernel 7.4.0 kernel, with code hashes and execution receipts retained. The public project repository (https://github.com/BTomoto/rn-mri-clinical-fusion) contains code, derived data, OOF probabilities, fold assignments and analysis locks linking inputs, protocols and code by SHA256 checksums. Model analyses can be rerun from the saved feature tables without downloading images again. Re-extracting features requires the image files: 18 inputs are included as repository samples, whereas 528 additional inputs were archived separately. Source manifests provide acquisition URLs and checksums, and an archive restoration script verifies the archive and its extracted members. The saved restoration receipt records successful verification of all 528 additional files; access to the stored archive requires its holder’s account, with public acquisition URLs providing an alternative route. Sample size was determined by eligible linked examinations with available inputs, without a prospective sample-size calculation.',
  'chinese': '分析采用Python 3，使用scikit-learn 1.8.0（5）、NumPy 2.3.5、pandas 2.2.3、SciPy 1.17.0和Matplotlib 3.10.8。特征提取、模型评价和敏感性分析的Jupyter Notebook通过进程内ipykernel 7.4.0内核执行，保留代码哈希及执行回执。公开项目仓库（https://github.com/BTomoto/rn-mri-clinical-fusion）包含代码、派生数据、OOF概率、分折划分，以及通过SHA256校验值关联输入、方案与代码的分析锁。模型分析可直接使用已保存的特征表重新运行，无需再次下载图像。重新提取特征则需要影像文件：其中18份输入作为仓库样本保存，另外528份输入单独归档。源清单提供获取URL与校验值，归档恢复脚本会核验压缩包及解出的成员。既有恢复回执记录528份追加文件均核验成功；获取所存归档需要持有者的账户权限，也可通过公开获取URL重新取得。样本规模由具备可用输入且符合条件的关联检查决定，未进行前瞻性样本量计算。'}
]

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def wc(s):return len(re.findall(r"[A-Za-z0-9]+(?:['/-][A-Za-z0-9]+)*",s))
def close(a,b):assert np.isclose(a,b,rtol=1e-12,atol=1e-12),(a,b)

def audit():
    old=json.loads((ROOT/'reports/ucsd_manuscript_step31/methods_section29.json').read_text());frozen=old['source_audit']['frozen_file_sha256'].copy()
    frozen.update(json.loads((ROOT/'reports/ucsd_manuscript_step31/artifact_verification.json').read_text())['sha256'])
    for p,h in frozen.items():assert sha(ROOT/p)==h,p
    extra=['requirements.txt','requirements-imaging.txt','pyproject.toml','scripts/run_notebooks_kernel.py','src/rn_public/notebook_kernel.py',
      'scripts/restore_ucsd_batch_sources.py','scripts/run_ucsd_batch_features.py','scripts/run_ucsd_fusion.py','scripts/run_ucsd_stability.py',
      'reports/ucsd_batch_features/source_archive_receipt.json','reports/ucsd_batch_features/source_archive_access.json','reports/ucsd_batch_features/source_restore_verification.json']
    receipts=[];total_cells=0
    for number,name in [(17,'17_ucsd_batch_image_features'),(18,'18_ucsd_clinical_mri_fusion'),(19,'19_ucsd_model_stability')]:
        receipt_path=f'reports/kernel_execution_{number}.json';path=f'notebooks/{name}.ipynb'
        extra.extend([receipt_path,path]);r=json.loads((ROOT/receipt_path).read_text());n=json.loads((ROOT/path).read_text())
        assert r['success'] and r['mode']=='in_process_ipykernel' and r['browser_frontend_tested'] is False and r['separate_kernel_server_tested'] is False
        assert r['versions']['ipykernel']=='7.4.0' and len(r['notebooks'])==1
        code=[c for c in n['cells'] if c['cell_type']=='code']
        text='\n'.join(''.join(c['source']) if isinstance(c['source'],list) else c['source'] for c in code)
        recorded=r['notebooks'][0];assert recorded['path']==path and recorded['code_sha256']==hashlib.sha256(text.encode()).hexdigest()
        assert len(code)==recorded['code_cells']==5 and recorded['errors']==0
        assert not any(o['output_type']=='error' for c in code for o in c.get('outputs',[]))
        assert n['metadata']['execution_success'] and n['metadata']['kernelspec']['language']=='python'
        receipts.append({'path':path,'code_sha256':recorded['code_sha256'],'code_cells':len(code),'recorded_errors':0});total_cells+=len(code)
    versions=json.loads((ROOT/'reports/ucsd_fusion/metrics.json').read_text())['versions']
    required={'scikit-learn':'1.8.0','numpy':'2.3.5','pandas':'2.2.3','scipy':'1.17.0','matplotlib':'3.10.8'}
    assert versions==required
    requirements=(ROOT/'requirements.txt').read_text().splitlines()
    for package,version in required.items():assert f'{package}=={version}' in requirements
    assert 'ipykernel==7.4.0' in requirements
    archive=json.loads((ROOT/'reports/ucsd_batch_features/source_archive_receipt.json').read_text());access=json.loads((ROOT/'reports/ucsd_batch_features/source_archive_access.json').read_text())
    restored=json.loads((ROOT/'reports/ucsd_batch_features/source_restore_verification.json').read_text())
    manifest=pd.read_csv(ROOT/'reports/ucsd_batch_features/source_manifest.csv')
    assert len(manifest)==546 and len(manifest[manifest.storage.eq('existing_git_sample')])==18 and len(manifest[manifest.storage.eq('batch_package')])==528
    assert archive['source_manifest_sha256']==sha(ROOT/'reports/ucsd_batch_features/source_manifest.csv')
    assert archive['additional_source_files']==restored['verified']==restored['restored']==528 and restored['reused']==0
    assert archive['archive_sha256']==access['archive_sha256'] and archive['archive_bytes']==access['archive_bytes']
    assert archive['batch_source_archive_in_Git'] is False and access['save_status']=='succeeded'
    assert manifest.url.str.startswith('https://').all() and manifest.sha256.str.fullmatch('[0-9a-f]{64}').all()
    # The 18 repository image samples are small enough to verify directly; the separate ZIP is not restored again.
    for row in manifest[manifest.storage.eq('existing_git_sample')].itertuples():
        p=ROOT/'data/samples/ucsd_mirror'/row.mirror_filename;assert p.stat().st_size==row.bytes and sha(p)==row.sha256
    lock_checks=[]
    for folder in ['ucsd_fusion','ucsd_stability']:
        lock=json.loads((ROOT/f'reports/{folder}/analysis_lock.json').read_text())
        for p,h in lock['input_sha256'].items():assert sha(ROOT/p)==h,p
        assert sha(ROOT/f'reports/{folder}_protocol.md')==lock['protocol_sha256']
        assert sha(ROOT/f'src/rn_public/{folder}.py')==lock['analysis_code_sha256']
        lock_checks.append(folder)
    ref=next(r for r in json.loads((ROOT/'reports/ucsd_manuscript_step23/reference_library.json').read_text())['references'] if r['id']=='S40')
    assert ref['publication_year']==2011 and ref['pages_or_article']=='2825-2830' and ref['read_level']=='Publisher metadata and abstract reviewed'
    frozen.update({p:sha(ROOT/p) for p in extra})
    return {'frozen_file_sha256':frozen,'saved_model_package_versions':versions,'pinned_package_entries_checked':6,
      'python_version_claim':'Python 3 only; historical Python minor/patch version not recorded in the reviewed receipts',
      'notebook_execution_receipts_checked':receipts,'notebook_code_cells_checked':total_cells,'notebook_errors':0,
      'notebook_browser_frontend_claimed_tested':False,'analysis_locks_checked':lock_checks,
      'image_source_manifest_rows_checked':546,'repository_image_sample_files_rehashed':18,'additional_archive_files_verified_in_saved_restoration_receipt':528,
      'archive_restored_this_stage':False,'public_source_urls_retried_this_stage':False,
      'archive_access_limitation':'Stored archive is holder-account restricted; public source URLs are listed, with no guarantee of future availability',
      'models_refitted':0,'features_reextracted':0,'reference_record':ref,
      'verification_scope':'Inspect historical software and execution receipts, verify notebook code hashes and analysis locks, rehash repository samples and check saved archive-recovery evidence; no model or imaging rerun'}

def build():
    OUT.mkdir(parents=True,exist_ok=True);REPORT.mkdir(parents=True,exist_ok=True);a=audit()
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(.79)
    for name in ['Normal','Title','Heading 1']:
        st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        rf=st.element.get_or_add_rPr().rFonts;rf.set(qn('w:eastAsia'),'Noto Serif CJK SC')
        for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:rf.attrib.pop(qn('w:'+key),None)
    normal=doc.styles['Normal'];normal.font.size=Pt(12);normal.paragraph_format.line_spacing=1.;normal.paragraph_format.space_after=Pt(9)
    doc.styles['Title'].font.size=Pt(16);doc.styles['Title'].paragraph_format.space_after=Pt(12)
    doc.styles['Heading 1'].font.size=Pt(13);doc.styles['Heading 1'].paragraph_format.space_before=Pt(8);doc.styles['Heading 1'].paragraph_format.space_after=Pt(9)
    doc.core_properties.title='Materials and methods 2.10 Reproducibility and software';doc.core_properties.author='Tianyu Feng'
    doc.add_paragraph('Materials and methods',style='Title');doc.add_paragraph('2.10 Reproducibility and software',style='Heading 1')
    for item in PARAGRAPHS:
        en=doc.add_paragraph(item['english']);en.paragraph_format.page_break_before=False
        en.paragraph_format.keep_together=True;en.paragraph_format.keep_with_next=True
        zh=doc.add_paragraph(item['chinese']);zh.paragraph_format.keep_together=True
        for r in zh.runs:r.font.size=Pt(11)
    doc.add_paragraph('References',style='Heading 1')
    ref=doc.add_paragraph('5. '+a['reference_record']['citation']);ref.paragraph_format.keep_together=True
    for e in [*doc.styles.element.xpath('.//w:pBdr'),*doc.element.xpath('.//w:pBdr')]:e.getparent().remove(e)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    path=OUT/'UCSD_Manuscript_Methods_2_10_EN_ZH.docx';doc.save(path)
    (OUT/'UCSD_Manuscript_Methods_2_10_EN.txt').write_text('\n\n'.join(p['english'] for p in PARAGRAPHS)+'\n')
    d={'stage':32,'date':'2026-10-29','subsection':'2.10 Reproducibility and software','v2_paragraph_ids':['M20'],'paragraphs':PARAGRAPHS,
      'english_word_count':sum(wc(p['english']) for p in PARAGRAPHS),'paragraph_word_counts':{p['id']:wc(p['english']) for p in PARAGRAPHS},
      'word_count_scope':'English narrative; excludes equation, headings and Chinese translation','v2_approximate_target_words':150,
      'equations':{'count':0,'numbers':[],'cumulative_previous_numbered_groups':10},
      'citation_mapping':{'5':'S40'},'citation_numbering':'Provisional first-appearance number 5; regenerate globally when Introduction is assembled','new_manuscript_references':1,'cumulative_unique_cited_sources':5,'source_audit':a,'models_refitted':0,'new_hypothesis_tests':0,
      'next_section':'Results 3.1 Cohort flow (R1)','docx':str(path.relative_to(ROOT)),'docx_sha256':sha(path)}
    (REPORT/'methods_section210.json').write_text(json.dumps(d,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'docx':d['docx'],'english_word_count':d['english_word_count'],'paragraph_word_counts':d['paragraph_word_counts'],'prior_files':len(a['frozen_file_sha256'])},indent=2))

def verify():
    from pypdf import PdfReader
    d=json.loads((REPORT/'methods_section210.json').read_text())
    for p,h in d['source_audit']['frozen_file_sha256'].items():assert sha(ROOT/p)==h,p
    path=ROOT/d['docx'];assert sha(path)==d['docx_sha256'];texts=[p.text for p in Document(path).paragraphs]
    for p in PARAGRAPHS:
        assert p['english'] in texts and p['chinese'] in texts and texts.index(p['chinese'])==texts.index(p['english'])+1
        assert not re.search(r'[\u3400-\u9fff，。；：（）【】、]',p['english'])
    assert sum(wc(p['english']) for p in PARAGRAPHS)==d['english_word_count']
    assert any(t.startswith('5. Pedregosa F') and ':2825-2830.' in t for t in texts)
    ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
    with zipfile.ZipFile(path) as z:
        t=etree.fromstring(z.read('word/document.xml'));assert len(t.xpath('//m:oMath',namespaces=ns))==0
        assert not any(n.startswith('word/media/') for n in z.namelist())
    pages=len(PdfReader(OUT/'docx_render/UCSD_Manuscript_Methods_2_10_EN_ZH.pdf').pages);assert pages==1
    methods=[json.loads(p.read_text()) for p in ROOT.glob('reports/ucsd_manuscript_step*/methods_section*.json')]
    ids=[p['id'] for section in methods for p in section['paragraphs']]
    assert len(methods)==10 and len(ids)==20 and set(ids)=={f'M{i}' for i in range(1,21)}
    assert sum(section['english_word_count'] for section in methods)==3131
    assert sum(section.get('equations',{}).get('count',0) for section in methods)==10
    v={'success':True,'stage':32,'date':'2026-10-29','subsection':d['subsection'],'paragraph_ids':['M20'],
      'english_word_count':d['english_word_count'],'paragraph_word_counts':d['paragraph_word_counts'],'matched_chinese_paragraphs':1,
      'word_pages':pages,'visual_review_pages':[1],'native_editable_equations':0,'equation_numbers':[],'equation_images':0,
      'prior_files_unchanged':len(d['source_audit']['frozen_file_sha256']),'pinned_package_entries_checked':6,'notebook_execution_receipts_checked':3,'notebook_code_cells_checked':15,'notebook_errors':0,
      'analysis_locks_checked':2,'image_source_manifest_rows_checked':546,'repository_image_sample_files_rehashed':18,'saved_archive_restore_verified_files':528,'models_refitted':0,'new_hypothesis_tests':0,'new_manuscript_references':1,'current_section_citations':1,'cumulative_unique_cited_sources':5,'citation_mapping':{'5':'S40'},'methods_v2_paragraphs_complete':True,'completed_methods_subsections':10,'completed_methods_english_word_count':3131,'completed_methods_chinese_paragraphs':20,'completed_methods_equation_groups':10,'next_section':d['next_section'],
      'sha256':{str(p.relative_to(ROOT)):sha(p) for p in OUT.iterdir() if p.is_file() and p.suffix in {'.docx','.txt'}}}
    (REPORT/'artifact_verification.json').write_text(json.dumps(v,ensure_ascii=False,indent=2)+'\n');print(json.dumps(v,indent=2))

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--verify',action='store_true');args=parser.parse_args()
    verify() if args.verify else build()
