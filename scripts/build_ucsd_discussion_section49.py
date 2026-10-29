"""Draft D9 from frozen image QC and normalization amendment records."""
from pathlib import Path
import argparse, hashlib, json, re, zipfile
import pandas as pd
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'outputs/ucsd-manuscript-step49-20261030'
REPORT = ROOT/'reports/ucsd_manuscript_step49'
EVIDENCE = REPORT/'literature_evidence.json'
MAPPING = {'12': 'S39'}
PARAGRAPHS = [{'id': 'D9', 'citation_ids': ['S39'], 'english': 'Image processing limits the portability and biological interpretation of the observed signal. The supplied structural volumes were stored as unsigned 8-bit integers, so subsequent scaling could not recover distinctions absent from those derived values. Unknown spatial units also prevented physical-volume interpretation and defining texture neighborhoods in verified millimeters. Matching image and mask grids supported computation at corresponding voxel indices, but did not independently establish anatomical registration. Our positive-label union pooled the supplied annotation compartments, potentially combining different tissue and treatment effects within a single summary. No compartment-specific diagnostic claim follows from its predictive performance. We replaced percentile clipping with unclipped affine scaling after technical diagnostics revealed additional tail collapse. This documented engineering revision avoided further clipping, but was neither preregistered nor independently shown to be optimal. Evidence that normalization choices influence MRI feature comparability across acquisition settings reinforces the need for evaluation under other processing conditions (12). It does not validate our transform, and ROI-range fixed-bin-number texture features need not change under every positive affine transformation. Finally, the eight descriptors used only FLAIR and postcontrast T1 images. Quantitative ADC and perfusion information were not evaluated. The findings therefore concern this particular derived-image representation; confirmation with preserved acquisition metadata, verified spatial scales and standardized processing remains necessary.', 'chinese': '影像处理限制了所观察信号的可迁移性及生物学解释。所用结构影像以无符号8位整数存储，后续缩放无法恢复这些派生数值中已经不存在的差别。空间单位未知，也使我们无法按物理体积解释结果，或以经核实的毫米距离定义纹理邻域。图像与掩膜网格一致，支持在对应体素索引处计算，但并未独立证实解剖配准准确。我们将全部正标签取并集，合并了所提供的标注区域，可能在同一汇总中混合不同组织及治疗效应；其预测表现不能支持区域成分特异性的诊断结论。在技术诊断发现额外的尾部灰度合并后，我们将百分位裁切改为不裁切的仿射缩放。这项有记录的工程修订避免了进一步裁切，但并非预先注册，也未被独立证明为最优方案。有研究发现，强度归一化的选择会影响不同采集条件下MRI特征的可比性，因此有必要在其他处理条件下进行评价（12）。这一证据并未验证本研究的变换，而且按ROI范围进行固定箱数离散化的纹理特征，不一定随每一种正仿射变换而变化。最后，八项描述量仅使用FLAIR及增强T1图像，未评价定量ADC或灌注信息。因此，当前发现对应这一特定的派生影像表达；仍需在保留采集元数据、核实空间尺度并采用标准化处理的条件下进一步确认。'}]

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def wc(text):
    return len(re.findall(r"[A-Za-z0-9]+(?:['/-][A-Za-z0-9]+)*", text))

def audit():
    previous=json.loads((ROOT/'reports/ucsd_manuscript_step48/discussion_section48.json').read_text())
    frozen=previous['source_audit']['frozen_file_sha256'].copy()
    frozen.update(json.loads((ROOT/'reports/ucsd_manuscript_step48/artifact_verification.json').read_text())['sha256'])
    for rel in ['reports/ucsd_roi_pilot/normalization_choice.json','reports/ucsd_roi_pilot_intensity_amendment.md','reports/ucsd_manuscript_step27/methods_section25.json','reports/ucsd_manuscript_step45/literature_evidence.json']:
        if rel not in frozen: frozen[rel]=sha(ROOT/rel)
    frozen[str(EVIDENCE.relative_to(ROOT))]=sha(EVIDENCE)
    for p,h in frozen.items(): assert sha(ROOT/p)==h,p
    evidence=json.loads(EVIDENCE.read_text()); record=evidence['studies'][0]
    refs={r['id']:r for r in json.loads((ROOT/'reports/ucsd_manuscript_step23/reference_library.json').read_text())['references']}
    ref=refs['S39'];assert ref['doi']==record['doi'] and record['provisional_citation_number']==12
    selected=[{'number':12,'reference_id':'S39','citation':ref['citation'],'doi':ref['doi'],'read_level_this_stage':record['read_level_this_stage']}]
    qc=pd.read_csv(ROOT/'reports/ucsd_batch_features/actual_file_QC.csv')
    assert len(qc)==546 and qc.ID.nunique()==182 and not qc.duplicated(['ID','sequence']).any()
    assert set(qc.sequence)=={'FLAIR','T1post','BraTS_tumor_seg'} and qc.groupby('ID').size().eq(3).all()
    assert qc.dtype.eq('uint8').all() and qc.spatial_units.eq('unknown').all() and qc['shape'].eq('[256, 256, 256]').all()
    assert qc.form_code.eq(1).all() and qc.slope.eq(1).all() and qc.intercept.eq(0).all()
    assert not qc.physical_scale_confirmed.any()
    for c in ['decoded','gzip_crc_verified','digest_matches','relative_grid_match']: assert qc[c].all()
    choice=json.loads((ROOT/'reports/ucsd_roi_pilot/normalization_choice.json').read_text())
    assert choice['recommended_next_stage_method']=='per_image_p01_p99_affine_unclipped_then_ROI_FBN32'
    assert choice['added_after_observing_clipping'] and not choice['uses_clinical_class'] and not choice['model_performance_used']
    assert choice['original_clipping_outputs_preserved'] and not choice['independent_optimality_validated']
    assert not choice['quantitative_ADC_allowed'] and not choice['physical_volume_allowed']
    diag=pd.read_csv(ROOT/'reports/ucsd_batch_features/normalization_and_texture_diagnostics.csv')
    assert len(diag)==364 and diag.ID.nunique()==182 and set(diag.sequence)=={'FLAIR','T1post'}
    assert not diag.duplicated(['ID','sequence']).any() and diag.p99.gt(diag.p01).all()
    assert diag.normalized_maximum.gt(1).any() and diag.erosion_error.isna().all()
    dictionary=pd.read_csv(ROOT/'reports/ucsd_batch_features/feature_dictionary.csv')
    expected={f'{seq}_{suffix}' for seq in ['FLAIR','T1post'] for suffix in ['mean','std','entropy','glcm_contrast']}
    assert len(dictionary)==8 and set(dictionary.feature)==expected and dictionary.allowed_model_predictor.all()
    assert dictionary.ROI.eq('ordinary_BraTS_positive_label_union').all() and not dictionary.physical_distance_assumed.any()
    assert dictionary.normalization.eq('per_image_nonzero_finite_p1_p99_affine_no_clipping').all()
    prior_ids=set()
    for pattern in ['methods_section*.json','results_section*.json','discussion_section*.json']:
        for p in ROOT.glob('reports/ucsd_manuscript_step*/'+pattern):
            if int(p.parent.name.rsplit('step',1)[1])<49: prior_ids.update(r for item in json.loads(p.read_text())['paragraphs'] for r in item['citation_ids'])
    assert len(prior_ids)==16 and set(MAPPING.values()).issubset(prior_ids)
    return {'frozen_file_sha256':frozen,'reference_records':selected,'source_evidence_record':str(EVIDENCE.relative_to(ROOT)),
            'new_full_texts_checked_for_current_claims':[],'previously_reviewed_full_text_revisited':['S39'],
            'saved_file_QC_rows_checked':546,'structural_image_rows_checked':364,'saved_intensity_diagnostic_rows_checked':364,
            'eight_feature_definitions_checked':True,'uint8_and_unknown_spatial_units_confirmed':True,
            'relative_grid_checks_not_anatomical_validation':True,'normalization_revision_after_diagnostics_disclosed':True,
            'model_performance_used_for_normalization_selection':False,'normalization_preregistered_or_independently_optimal_claimed':False,
            'physical_volume_or_quantitative_ADC_included':False,'all_features_claimed_sensitive_to_every_affine_transform':False,
            'previous_unique_cited_sources':sorted(prior_ids),'models_refitted':0,'features_reextracted':0,'new_hypothesis_tests':0,
            'verification_scope':'Read saved file QC, normalization choice, image diagnostics and feature dictionary; no source-image reopening, feature extraction, fitting or new sensitivity experiment'}

def build():
    a = audit(); OUT.mkdir(parents=True, exist_ok=True); REPORT.mkdir(parents=True, exist_ok=True)
    doc = Document(); sec = doc.sections[0]; sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(.79)
    for name in ['Normal', 'Title', 'Heading 1']:
        st = doc.styles[name]; st.font.name = 'Times New Roman'; st.font.color.rgb = RGBColor(0, 0, 0)
        rf = st.element.get_or_add_rPr().rFonts; rf.set(qn('w:eastAsia'), 'Noto Serif CJK SC')
        for key in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']: rf.attrib.pop(qn('w:'+key), None)
    normal = doc.styles['Normal']; normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.; normal.paragraph_format.space_after = Pt(9)
    doc.styles['Title'].font.size = Pt(16); doc.styles['Title'].paragraph_format.space_after = Pt(12)
    heading = doc.styles['Heading 1']; heading.font.size = Pt(13)
    heading.paragraph_format.space_before = Pt(8); heading.paragraph_format.space_after = Pt(9)
    doc.core_properties.title = 'Discussion 4.9 Image-processing limitations'; doc.core_properties.author = 'Tianyu Feng'
    doc.add_paragraph('Discussion', style='Title'); doc.add_paragraph('4.9 Image-processing limitations', style='Heading 1')
    for item in PARAGRAPHS:
        en = doc.add_paragraph(item['english']); en.paragraph_format.keep_together = True; en.paragraph_format.keep_with_next = True
        zh = doc.add_paragraph(item['chinese']); zh.paragraph_format.keep_together = True
        for run in zh.runs: run.font.size = Pt(11)
    title = doc.add_paragraph('References', style='Heading 1'); title.paragraph_format.page_break_before = True
    for ref in a['reference_records']:
        p = doc.add_paragraph(str(ref['number'])+'. '+ref['citation']); p.paragraph_format.keep_together = True
    for e in [*doc.styles.element.xpath('.//w:pBdr'), *doc.element.xpath('.//w:pBdr')]: e.getparent().remove(e)
    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    path = OUT/'UCSD_Manuscript_Discussion_4_9_EN_ZH.docx'; doc.save(path)
    (OUT/'UCSD_Manuscript_Discussion_4_9_EN.txt').write_text('\n\n'.join(p['english'] for p in PARAGRAPHS)+'\n')
    d = {'stage': 49, 'date': '2026-10-30', 'subsection': '4.9 Image-processing limitations', 'v2_paragraph_ids': ['D9'],
         'paragraphs': PARAGRAPHS, 'english_word_count': sum(wc(p['english']) for p in PARAGRAPHS),
         'paragraph_word_counts': {p['id']: wc(p['english']) for p in PARAGRAPHS},
         'word_count_scope': 'English narrative including in-text citation numbers; excludes headings, reference list and Chinese translation',
         'v2_approximate_target_words': 200, 'equations': {'count': 0, 'numbers': [], 'cumulative_previous_numbered_groups': 10},
         'citation_mapping': MAPPING, 'citation_numbering': 'Existing number 12 preserved; renumber globally when Introduction is assembled',
         'new_manuscript_references': 0, 'cumulative_unique_cited_sources': 16,
         'source_audit': a, 'models_refitted': 0, 'new_hypothesis_tests': 0,
         'next_section': 'Discussion 4.10 Sampling and model uncertainty (D10)',
         'docx': str(path.relative_to(ROOT)), 'docx_sha256': sha(path)}
    (REPORT/'discussion_section49.json').write_text(json.dumps(d, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps({'docx': d['docx'], 'english_word_count': d['english_word_count'], 'frozen_files': len(a['frozen_file_sha256'])}))

def verify():
    from pypdf import PdfReader
    d = json.loads((REPORT/'discussion_section49.json').read_text()); a = audit(); assert a == d['source_audit']
    path = ROOT/d['docx']; assert sha(path) == d['docx_sha256']; doc = Document(path); texts = [p.text for p in doc.paragraphs]
    for p in PARAGRAPHS:
        assert p['english'] in texts and p['chinese'] in texts and texts.index(p['chinese']) == texts.index(p['english'])+1
        assert not re.search(r'[\u3400-\u9fff，。；：（）【】、]', p['english'])
    assert sum(wc(p['english']) for p in PARAGRAPHS) == d['english_word_count']
    assert re.findall(r'\((\d+)\)', PARAGRAPHS[0]['english']) == ['12']
    for ref in a['reference_records']: assert str(ref['number'])+'. '+ref['citation'] in texts
    assert not doc.tables and len(doc.sections) == 1
    with zipfile.ZipFile(path) as z:
        xml = etree.fromstring(z.read('word/document.xml'))
        assert not xml.xpath('//m:oMath', namespaces={'m': 'http://schemas.openxmlformats.org/officeDocument/2006/math'})
        assert not [n for n in z.namelist() if n.startswith('word/media/')]
    pdf = PdfReader(OUT/'docx_render/UCSD_Manuscript_Discussion_4_9_EN_ZH.pdf'); pages = len(pdf.pages); assert pages == 2
    assert '4.9 Image' in pdf.pages[0].extract_text() and 'References' in pdf.pages[1].extract_text()
    v = {'success': True, 'stage': 49, 'date': '2026-10-30', 'subsection': d['subsection'], 'paragraph_ids': ['D9'],
         'english_word_count': d['english_word_count'], 'matched_chinese_paragraphs': 1, 'word_pages': pages,
         'visual_review_pages': [1, 2], 'native_editable_equations': 0, 'equation_images': 0,
         'frozen_files_unchanged': len(a['frozen_file_sha256']), 'prior_files_unchanged': len(a['frozen_file_sha256'])-1,
         'new_full_texts_checked_for_current_claims': [], 'saved_file_QC_rows_checked': 546,
         'structural_image_rows_checked': 364, 'saved_intensity_diagnostic_rows_checked': 364,
         'eight_feature_definitions_checked': True, 'uint8_and_unknown_spatial_units_confirmed': True,
         'relative_grid_checks_not_anatomical_validation': True, 'normalization_revision_after_diagnostics_disclosed': True,
         'model_performance_used_for_normalization_selection': False,
         'normalization_preregistered_or_independently_optimal_claimed': False,
         'physical_volume_or_quantitative_ADC_included': False, 'features_reextracted': 0,
         'all_features_claimed_sensitive_to_every_affine_transform': False,
         'new_manuscript_references': 0, 'provisional_citation_mapping': MAPPING, 'cumulative_unique_cited_sources': 16,
         'models_refitted': 0, 'new_hypothesis_tests': 0, 'next_section': d['next_section'],
         'sha256': {str(p.relative_to(ROOT)): sha(p) for p in OUT.iterdir() if p.is_file() and p.suffix in {'.docx', '.txt'}}}
    (REPORT/'artifact_verification.json').write_text(json.dumps(v, ensure_ascii=False, indent=2)+'\n'); print(json.dumps(v, indent=2))

if __name__ == '__main__':
    parser = argparse.ArgumentParser(); parser.add_argument('--verify', action='store_true'); args = parser.parse_args()
    verify() if args.verify else build()
