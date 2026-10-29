"""Write M16 with read-only checks of saved calibration diagnostics and bins."""
from pathlib import Path
import argparse,hashlib,json,re,sys,zipfile
import numpy as np
import pandas as pd
from lxml import etree
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from build_ucsd_methods_section25 import el,mr,sub,bracket,function

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'outputs/ucsd-manuscript-step30-20261029'
REPORT=ROOT/'reports/ucsd_manuscript_step30'
PARAGRAPHS=[
 {'id':'M16','citation_ids':[],
  'english': 'Calibration was assessed descriptively using the pooled OOF tumor probabilities. For each model, we requested five probability quantile groups, dropping duplicate boundaries and retaining the actual nonempty groups. Each group was summarized by its examination count, mean predicted probability, observed tumor proportion and tumor count. Calibration-in-the-large was estimated as an intercept with the coefficient of the predicted log odds fixed at one (Equation 9). A separate unpenalized logistic diagnostic regression estimated a free intercept and calibration slope jointly (Equation 10). Here, p denotes the saved OOF probability and q the fitted diagnostic probability. For logit calculations only, p was clipped to machine epsilon and one minus machine epsilon. The offset intercept was distinguished from the joint intercept accompanying the slope. Failed or separated fits were retained as missing estimates with their recorded errors. These diagnostics used the same OOF predictions being evaluated and did not constitute independent validation. We did not use them to recalibrate predictions, alter the fixed classification threshold or replace any evaluated OOF probabilities.',
  'chinese': '使用汇总的OOF肿瘤概率进行描述性校准评估。对每个模型，设定五个概率分位组，删除重复边界并保留实际非空组。每组汇总检查数、平均预测概率、实际肿瘤比例和肿瘤数。总体校准偏差通过固定预测对数优势系数为1的截距估计（Equation 9）。另以未惩罚的Logistic诊断回归联合估计自由截距与校准斜率（Equation 10）。其中，p表示已保存的OOF概率，q表示诊断拟合概率。仅在计算logit时，将p限制在机器精度epsilon与1减epsilon之间。固定offset的截距与伴随斜率联合估计的截距分别记录。拟合失败或发生分离时，估计值保留为缺失，并记录错误原因。这些诊断使用正在评价的同一批OOF预测，不构成独立验证。我们未据此重新校准预测、改变固定分类阈值或替换任何被评价的OOF概率。'}
]

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def wc(s):return len(re.findall(r"[A-Za-z0-9]+(?:['/-][A-Za-z0-9]+)*",s))
def close(a,b):assert np.isclose(a,b,rtol=1e-12,atol=1e-12),(a,b)

def audit():
    old=json.loads((ROOT/'reports/ucsd_manuscript_step29/methods_section27.json').read_text());frozen=old['source_audit']['frozen_file_sha256'].copy()
    frozen.update(json.loads((ROOT/'reports/ucsd_manuscript_step29/artifact_verification.json').read_text())['sha256'])
    extra=['reports/ucsd_fusion/calibration_bins.csv','outputs/ucsd-paper-materials-20261029/tables/TableS3.csv']
    for p,h in frozen.items():assert sha(ROOT/p)==h,p
    frozen.update({p:sha(ROOT/p) for p in extra})
    from scipy.special import expit,logit
    pred=pd.read_csv(ROOT/'data/processed/ucsd_fusion_oof.csv');y=pred.tumor_label.to_numpy()
    metrics=json.loads((ROOT/'reports/ucsd_fusion/metrics.json').read_text())
    bins=pd.read_csv(ROOT/'reports/ucsd_fusion/calibration_bins.csv');assert len(pred)==182 and len(bins)==30
    checks={};rows_checked=0
    for name,r in metrics['calibration'].items():
        assert r['error'] is None
        p=pred[name+'_oof_probability'].to_numpy();z=logit(np.clip(p,np.finfo(float).eps,1-np.finfo(float).eps))
        offset_res=float(abs(np.sum(expit(z+r['calibration_intercept_offset'])-y)))
        x=np.column_stack([np.ones(len(y)),z]);fit=expit(x@np.array([r['calibration_joint_intercept'],r['calibration_slope']]))
        score=float(np.max(np.abs(x.T@(fit-y))));mineig=float(np.linalg.eigvalsh(x.T@((fit*(1-fit))[:,None]*x)).min())
        assert offset_res<1e-8 and score<1e-5 and mineig>=1e-8
        frame=pd.DataFrame(dict(y=y,p=p,bin=pd.qcut(pd.Series(p),5,duplicates='drop')))
        saved=bins.loc[bins.model==name].reset_index(drop=True);n=0
        for k,(_,g) in enumerate(frame.groupby('bin',observed=True),1):
            row=saved.iloc[k-1];assert row['bin']==k and row.n==len(g) and row.tumor_count==int(g.y.sum())
            close(row.mean_predicted,float(g.p.mean()));close(row.observed_tumor,float(g.y.mean()));n+=1;rows_checked+=1
        assert n==len(saved) and saved.n.sum()==182 and saved.tumor_count.sum()==141
        checks[name]={'offset_score_absolute_residual':offset_res,'joint_score_maximum_absolute_residual':score,'joint_hessian_minimum_eigenvalue':mineig,'nonempty_quantile_groups':n,'diagnostic_fit_error':r['error']}
    table=pd.read_csv(ROOT/'outputs/ucsd-paper-materials-20261029/tables/TableS3.csv');assert len(table)==4
    for (_,r),name in zip(table.iterrows(),['M0','M1','MRI','Fusion']):
        close(r.Calibration_offset,metrics['calibration'][name]['calibration_intercept_offset'])
        close(r.Calibration_slope,metrics['calibration'][name]['calibration_slope'])
    return {'frozen_file_sha256':frozen,'calibration_variants_checked':6,'saved_diagnostic_score_and_curvature_checks':checks,
      'calibration_bin_rows_checked':rows_checked,'bin_summary_numeric_fields_checked':rows_checked*5,'table_s3_calibration_values_checked':8,
      'calibration_diagnostics_refitted':0,'prediction_models_refitted':0,'probabilities_replaced':0,
      'verification_scope':'Evaluate score residuals and curvature at saved estimates, reproduce quantile-bin summaries, and check existing Table S3; no new fitting or recalibration'}

def equations():
    def logit(symbol):return function(mr('logit',True),bracket(sub(mr(symbol),'i')))
    def named_a(name):return el('sSub',el('e',mr('a')),el('sub',mr(name,True)))
    return {9:[logit('q'),mr(' = '),named_a('offset'),mr(' + '),logit('p')],
            10:[logit('q'),mr(' = '),named_a('joint'),mr(' + b '),logit('p')]}

def build():
    OUT.mkdir(parents=True,exist_ok=True);REPORT.mkdir(parents=True,exist_ok=True);a=audit()
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(.79)
    for name in ['Normal','Title','Heading 1']:
        st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        rf=st.element.get_or_add_rPr().rFonts;rf.set(qn('w:eastAsia'),'Noto Serif CJK SC')
        for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:rf.attrib.pop(qn('w:'+key),None)
    normal=doc.styles['Normal'];normal.font.size=Pt(12);normal.paragraph_format.line_spacing=1.;normal.paragraph_format.space_after=Pt(9)
    doc.styles['Title'].font.size=Pt(16);doc.styles['Title'].paragraph_format.space_after=Pt(12)
    doc.styles['Heading 1'].font.size=Pt(13);doc.styles['Heading 1'].paragraph_format.space_before=Pt(8);doc.styles['Heading 1'].paragraph_format.space_after=Pt(9)
    doc.core_properties.title='Materials and methods 2.8 Calibration assessment';doc.core_properties.author='Tianyu Feng'
    doc.add_paragraph('Materials and methods',style='Title');doc.add_paragraph('2.8 Calibration assessment',style='Heading 1')
    for item in PARAGRAPHS:
        en=doc.add_paragraph(item['english']);en.paragraph_format.page_break_before=False
        en.paragraph_format.keep_together=True;en.paragraph_format.keep_with_next=True
        zh=doc.add_paragraph(item['chinese']);zh.paragraph_format.keep_together=True
        for r in zh.runs:r.font.size=Pt(11)
        for number,nodes in equations().items():
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_together=True
            p.paragraph_format.space_before=Pt(6);p._p.append(el('oMath',nodes));p.add_run(f'    ({number})')
    for e in [*doc.styles.element.xpath('.//w:pBdr'),*doc.element.xpath('.//w:pBdr')]:e.getparent().remove(e)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    path=OUT/'UCSD_Manuscript_Methods_2_8_EN_ZH.docx';doc.save(path)
    (OUT/'UCSD_Manuscript_Methods_2_8_EN.txt').write_text('\n\n'.join(p['english'] for p in PARAGRAPHS)+'\n')
    d={'stage':30,'date':'2026-10-29','subsection':'2.8 Calibration assessment','v2_paragraph_ids':['M16'],'paragraphs':PARAGRAPHS,
      'english_word_count':sum(wc(p['english']) for p in PARAGRAPHS),'paragraph_word_counts':{p['id']:wc(p['english']) for p in PARAGRAPHS},
      'word_count_scope':'English narrative; excludes equation, headings and Chinese translation','v2_approximate_target_words':150,
      'equations':{'count':2,'format':'native editable OMML','numbers':[9,10],'meaning':'Offset calibration intercept at fixed slope one, and separate joint intercept and slope diagnostic model'},
      'citation_mapping':{},'new_manuscript_references':0,'cumulative_unique_cited_sources':4,'source_audit':a,'models_refitted':0,'new_hypothesis_tests':0,
      'next_section':'Sensitivity analyses (M17-M19)','docx':str(path.relative_to(ROOT)),'docx_sha256':sha(path)}
    (REPORT/'methods_section28.json').write_text(json.dumps(d,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'docx':d['docx'],'english_word_count':d['english_word_count'],'paragraph_word_counts':d['paragraph_word_counts'],'prior_files':len(a['frozen_file_sha256'])},indent=2))

def verify():
    from pypdf import PdfReader
    d=json.loads((REPORT/'methods_section28.json').read_text())
    for p,h in d['source_audit']['frozen_file_sha256'].items():assert sha(ROOT/p)==h,p
    path=ROOT/d['docx'];assert sha(path)==d['docx_sha256'];texts=[p.text for p in Document(path).paragraphs]
    for p in PARAGRAPHS:
        assert p['english'] in texts and p['chinese'] in texts and texts.index(p['chinese'])==texts.index(p['english'])+1
        assert not re.search(r'[\u3400-\u9fff，。；：（）【】、]',p['english'])
    assert sum(wc(p['english']) for p in PARAGRAPHS)==d['english_word_count'] and all(f'    ({i})' in texts for i in [9,10])
    ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
    with zipfile.ZipFile(path) as z:
        t=etree.fromstring(z.read('word/document.xml'));assert len(t.xpath('//m:oMath',namespaces=ns))==2
        assert len(t.xpath('//m:sSub',namespaces=ns))==6 and not any(n.startswith('word/media/') for n in z.namelist())
    pages=len(PdfReader(OUT/'docx_render/UCSD_Manuscript_Methods_2_8_EN_ZH.pdf').pages);assert pages==1
    v={'success':True,'stage':30,'date':'2026-10-29','subsection':d['subsection'],'paragraph_ids':['M16'],
      'english_word_count':d['english_word_count'],'paragraph_word_counts':d['paragraph_word_counts'],'matched_chinese_paragraphs':1,
      'word_pages':pages,'visual_review_pages':[1],'native_editable_equations':2,'equation_numbers':[9,10],'equation_images':0,
      'prior_files_unchanged':len(d['source_audit']['frozen_file_sha256']),'calibration_variants_checked':6,'calibration_bin_rows_checked':30,'bin_summary_numeric_fields_checked':150,'table_s3_calibration_values_checked':8,
      'calibration_diagnostics_refitted':0,'models_refitted':0,'new_hypothesis_tests':0,'new_manuscript_references':0,'next_section':d['next_section'],
      'sha256':{str(p.relative_to(ROOT)):sha(p) for p in OUT.iterdir() if p.is_file() and p.suffix in {'.docx','.txt'}}}
    (REPORT/'artifact_verification.json').write_text(json.dumps(v,ensure_ascii=False,indent=2)+'\n');print(json.dumps(v,indent=2))

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--verify',action='store_true');args=parser.parse_args()
    verify() if args.verify else build()
