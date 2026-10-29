"""Draft R8 from saved single-scan selection, OOF metrics and paired intervals."""
from pathlib import Path
import argparse,hashlib,json,re,sys,zipfile
import numpy as np
import pandas as pd
from lxml import etree
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION,WD_ORIENT

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'outputs/ucsd-manuscript-step40-20261029'
REPORT=ROOT/'reports/ucsd_manuscript_step40'
PARAGRAPHS=[
 {'id':'R8','citation_ids':[],
  'english': 'Selecting the smallest available numeric session per patient retained 135 examinations from 135 patients, including 113 tumor and 22 treatment-related-change examinations. This selection did not establish the true first clinical visit. OOF AUCs were 0.501 for M0, 0.609 for M1, 0.729 for MRI and 0.730 for Fusion (95% patient-bootstrap percentile interval for Fusion, 0.606 to 0.833; Supplementary Table S6 and Figure S2A). The Fusion-minus-M1 AUC difference was 0.120, with a paired interval of -0.015 to 0.252 that included zero (Supplementary Table S7). The auxiliary AP, Brier-score and logarithmic-loss difference intervals also included zero. At the fixed threshold, Fusion recognized 109/113 tumor examinations (96.5%) and 2/22 treatment-related-change examinations (9.1%; Figure S2B). Its descriptive calibration slope was 0.539, without subsequent recalibration. The selected analysis differed from the primary 182-examination analysis in examination composition, class distribution and training data. Consequently, the difference between their AUC estimates could not be attributed solely to removing repeated scans. Incremental discrimination remained uncertain in the single-scan analysis, which was retained as a sensitivity analysis without replacing the primary comparison.',
  'chinese': '按每位患者最小可用数字会话编号选择后，保留135名患者的135次检查，包括113次肿瘤检查和22次治疗相关改变检查。该选择不能确定真实首次临床就诊。OOF AUC分别为M0的0.501、M1的0.609、MRI的0.729和Fusion的0.730；Fusion的95%患者bootstrap百分位区间为0.606至0.833（Supplementary Table S6和Figure S2A）。Fusion减去M1的AUC差值为0.120，配对区间为−0.015至0.252，包含零（Supplementary Table S7）。辅助AP、Brier评分与对数损失差值区间也均包含零。在固定阈值下，Fusion正确识别109/113次肿瘤检查（96.5%）和2/22次治疗相关改变检查（9.1%；Figure S2B）。其描述性校准斜率为0.539，未据此重新校准。所选分析与182次检查的主分析在检查构成、类别分布和训练数据方面不同。因此，两者AUC估计的差异不能仅归因于去除重复扫描。单扫描分析中的增量鉴别仍存在不确定性，该分析保留为敏感性分析，不替代主要比较。'}
]
MODELS=['M0','M1','MRI','Fusion']
METRICS=['roc_auc','ap_treatment_change','brier','log_loss']
METRIC_LABELS=['ROC AUC','Treatment-related-change AP','Brier score','Logarithmic loss']
FIGURE=ROOT/'outputs/ucsd-paper-materials-20261029/figures/FigureS2_Single_Scan.png'
FIGURE_CAPTION='Figure S2. Sensitivity analysis using one available examination per patient. The selected cohort contains 135 examinations, including 113 tumor and 22 treatment-related-change examinations. A, Pooled OOF AUCs with patient-bootstrap 95% percentile intervals conditional on fixed predictions without refitting. B, Treatment-change recognition at the fixed tumor-probability threshold of 0.5, with exact correct/total counts. Selection used the smallest numeric session available in the imaging subset and does not establish the true first clinical visit. The analysis used rebuilt grouped folds and does not replace the primary 182-examination analysis.'
TABLE_SPECS=[
 {'key':'TableS6','title':'Table S6. Performance in the one-examination-per-patient analysis',
  'headers':['Model','ROC AUC (95% interval)','Brier score','AP for treatment change','Treatment change correct, n/N'],
  'notes':['The selected cohort contains 135 examinations from 135 patients, including 113 tumor and 22 treatment-related-change examinations. Each patient contributes the smallest available numeric session in the imaging subset; this is not necessarily the first clinical examination.',
  'Four models share the rebuilt patient-grouped five-fold partition. Intervals use 2000 patient-bootstrap attempts with fixed OOF probabilities and no model refitting. They do not capture full retraining or partition variability.',
  'AP denotes average precision for treatment-related change, scored by one minus the tumor probability. Correct treatment-change counts use tumor probability below 0.5. Values are displayed to three decimals. The cohort has different examination composition, class distribution and training data from the primary analysis.']},
 {'key':'TableS7','title':'Table S7. Paired Fusion-minus-M1 differences in the single-scan analysis',
  'headers':['Metric','Fusion minus M1','95% paired interval'],
  'notes':['Differences use the same 135 selected examinations and rebuilt grouped folds. Positive AUC or treatment-change AP differences favor Fusion; negative Brier-score or logarithmic-loss differences favor Fusion.',
  'Intervals are the 2.5th and 97.5th percentiles of 2000 paired patient-bootstrap attempts with fixed OOF predictions and no refitting. All four intervals include zero. The result does not replace the primary comparison.',
  'AP denotes average precision for treatment-related change. Differences are calculated from unrounded metrics and displayed to four decimals to preserve the small positive Brier difference. This is a sensitivity comparison within the selected cohort, not a paired contrast between the 135- and 182-examination analyses.']}
]

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def wc(s):return len(re.findall(r"[A-Za-z0-9]+(?:['/-][A-Za-z0-9]+)*",s))
def close(a,b):assert np.isclose(a,b,rtol=1e-12,atol=1e-12),(a,b)

def audit():
    old=json.loads((ROOT/'reports/ucsd_manuscript_step39/results_section37.json').read_text());frozen=old['source_audit']['frozen_file_sha256'].copy()
    frozen.update(json.loads((ROOT/'reports/ucsd_manuscript_step39/artifact_verification.json').read_text())['sha256'])
    for p,h in frozen.items():assert sha(ROOT/p)==h,p
    extra=['outputs/ucsd-paper-materials-20261029/tables/TableS6.csv','outputs/ucsd-paper-materials-20261029/tables/TableS7.csv',
      str(FIGURE.relative_to(ROOT)),str(FIGURE.with_suffix('.svg').relative_to(ROOT))]
    frozen.update({p:sha(ROOT/p) for p in extra})
    directory=ROOT/'reports/ucsd_stability';saved=json.loads((directory/'metrics.json').read_text())['single_scan'];lock=json.loads((directory/'analysis_lock.json').read_text())
    full=pd.read_csv(ROOT/'data/processed/ucsd_clinical_image_features.csv');ledger=pd.read_csv(directory/'single_scan_selection.csv')
    identifiers=full[['ID','patient_id']].copy();identifiers['session']=identifiers.ID.str.rsplit('_',n=1).str[-1].astype(int)
    chosen=set(identifiers.sort_values(['patient_id','session','ID']).drop_duplicates('patient_id').ID)
    pred=pd.read_csv(ROOT/'data/processed/ucsd_stability_single_scan_oof.csv');manifest=pd.read_csv(directory/'single_scan_fold_manifest.csv')
    assert len(ledger)==182 and set(ledger.loc[ledger.selected,'ID'])==chosen==set(pred.ID) and len(pred)==135
    assert pred.patient_id.nunique()==pred.brats_patient_id.nunique()==pred.validation_group.nunique()==135
    assert set(pred.patient_id)==set(full.patient_id)
    pd.testing.assert_frame_equal(pred[manifest.columns],manifest)
    y=pred.tumor_label.to_numpy();assert y.sum()==113 and (y==0).sum()==22 and saved['complete']
    assert saved['scans']==saved['patients']==135 and saved['tumor_scans']==113 and saved['treatment_change_scans']==22
    assert lock['single_scan_seed']==20261029 and lock['single_scan_bootstrap_seed']==20261101
    foldaudit=json.loads((directory/'single_scan_fold_audit.json').read_text());assert foldaudit['complete'] and not foldaudit['errors'] and len(foldaudit['folds'])==5
    seen=[]
    for f in foldaudit['folds']:
        tr=set(f['train_IDs']);te=set(f['test_IDs']);assert not tr&te and tr|te==chosen and te==set(pred.loc[pred.validation_fold.eq(f['fold']),'ID'])
        train=pred[pred.ID.isin(tr)];test=pred[pred.ID.isin(te)]
        for k in ['patient_id','brats_patient_id','validation_group']:assert not set(train[k])&set(test[k])
        assert train.tumor_label.nunique()==test.tumor_label.nunique()==2 and f['valid'] and not f['errors']
        assert set(f['models'])==set(MODELS) and all(r['success'] for r in f['models'].values());seen.extend(te)
    assert len(seen)==135 and set(seen)==chosen
    boot=pd.read_csv(directory/'single_scan_patient_bootstrap.csv');assert boot.shape==(2000,21) and saved['bootstrap_valid']==2000 and saved['bootstrap_skipped']==0
    for key in boot.columns[1:]:
        close(float(boot[key].quantile(.025)),saved['bootstrap_95ci'][key]['lower']);close(float(boot[key].quantile(.975)),saved['bootstrap_95ci'][key]['upper'])
    def manual_values(p):
        diff=p[y==1][:,None]-p[y==0][None,:];auc=float(np.mean((diff>0)+.5*(diff==0)))
        labels=1-y;score=1-p;order=np.argsort(-score,kind='stable');labels=labels[order];score=score[order]
        ends=np.flatnonzero(np.r_[score[:-1]!=score[1:],True]);tp_sum=np.cumsum(labels)[ends]
        ap=float(np.sum(np.diff(np.r_[0,tp_sum])/labels.sum()*tp_sum/(ends+1)))
        positive=p>=.5;tp=int((positive&(y==1)).sum());tn=int((~positive&(y==0)).sum());fp=int((positive&(y==0)).sum());fn=int((~positive&(y==1)).sum())
        return dict(roc_auc=auc,ap_treatment_change=ap,brier=float(np.mean((p-y)**2)),log_loss=float(-np.mean(y*np.log(p)+(1-y)*np.log1p(-p))),
          tp=tp,tn=tn,fp=fp,fn=fn,sensitivity_tumor=tp/113,specificity_treatment_change=tn/22,balanced_accuracy=.5*(tp/113+tn/22),accuracy=(tp+tn)/135)
    original_a=pd.read_csv(ROOT/'outputs/ucsd-paper-materials-20261029/tables/TableS6.csv');original_b=pd.read_csv(ROOT/'outputs/ucsd-paper-materials-20261029/tables/TableS7.csv')
    assert original_a.shape==(4,8) and original_b.shape==(4,4)
    records_a=[];rows_a=[];records_b=[];rows_b=[];values={}
    for i,name in enumerate(MODELS):
        p=pred[name+'_oof_probability'].to_numpy();assert np.isfinite(p).all() and np.all((p>0)&(p<1));v=manual_values(p);values[name]=v
        for key,value in v.items():close(value,saved['metrics'][name][key])
        lo=saved['bootstrap_95ci'][name+'__roc_auc']['lower'];hi=saved['bootstrap_95ci'][name+'__roc_auc']['upper']
        for key,value in [('AUC',v['roc_auc']),('CI_lower',lo),('CI_upper',hi),('Brier',v['brier']),('AP_TC',v['ap_treatment_change']),('TC_correct',v['tn']),('TC_total',22)]:close(value,original_a.iloc[i][key])
        records_a.append({'Model':name,'ROC_AUC':float(original_a.iloc[i].AUC),'AUC_interval_lower':float(original_a.iloc[i].CI_lower),'AUC_interval_upper':float(original_a.iloc[i].CI_upper),
          'Brier_score':float(original_a.iloc[i].Brier),'AP_treatment_change':float(original_a.iloc[i].AP_TC),'Treatment_change_correct':int(original_a.iloc[i].TC_correct),'Treatment_change_total':22})
        rows_a.append([name,f"{v['roc_auc']:.3f} [{lo:.3f}, {hi:.3f}]",f"{v['brier']:.3f}",f"{v['ap_treatment_change']:.3f}",f"{v['tn']}/22"])
    for i,metric in enumerate(METRICS):
        difference=values['Fusion'][metric]-values['M1'][metric];key='Fusion_minus_M1__'+metric
        np.testing.assert_allclose(boot[key],boot['Fusion__'+metric]-boot['M1__'+metric],rtol=1e-12,atol=1e-12)
        lo=saved['bootstrap_95ci'][key]['lower'];hi=saved['bootstrap_95ci'][key]['upper'];assert lo<0<hi
        close(difference,saved['paired_differences'][metric])
        for col,value in [('差值',difference),('CI_lower',lo),('CI_upper',hi)]:close(value,original_b.iloc[i][col])
        records_b.append({'Metric':METRIC_LABELS[i],'Difference_Fusion_minus_M1':float(original_b.iloc[i]['差值']),'Interval_lower':float(original_b.iloc[i].CI_lower),'Interval_upper':float(original_b.iloc[i].CI_upper)})
        rows_b.append([METRIC_LABELS[i],f'{difference:+.4f}',f'[{lo:.4f}, {hi:.4f}]'])
    from scipy.special import expit,logit
    c=saved['metrics']['Fusion'];assert c['error'] is None
    z=logit(np.clip(pred.Fusion_oof_probability,np.finfo(float).eps,1-np.finfo(float).eps));x=np.column_stack([np.ones(len(y)),z])
    fit=expit(x@np.array([c['calibration_joint_intercept'],c['calibration_slope']]))
    off_score=float(abs(np.sum(expit(z+c['calibration_intercept_offset'])-y)));joint_score=float(np.max(np.abs(x.T@(fit-y))))
    mineig=float(np.linalg.eigvalsh(x.T@((fit*(1-fit))[:,None]*x)).min());assert off_score<1e-8 and joint_score<1e-5 and mineig>=1e-8
    assert f"{c['calibration_slope']:.3f}"=='0.539' and values['Fusion']['tp']==109 and values['Fusion']['tn']==2
    assert f"{100*values['Fusion']['sensitivity_tumor']:.1f}"=='96.5' and f"{100*values['Fusion']['specificity_treatment_change']:.1f}"=='9.1'
    receipt=json.loads((ROOT/'reports/ucsd_paper_materials/generation_receipt.json').read_text());assert receipt['output_sha256'][str(FIGURE.relative_to(ROOT))]==sha(FIGURE)
    svg=etree.parse(str(FIGURE.with_suffix('.svg')));texts=svg.xpath('//s:text/text()',namespaces={'s':'http://www.w3.org/2000/svg'})
    assert all(name in texts for name in MODELS) and texts.count('0/22')==texts.count('2/22')==2
    specs=[dict(**TABLE_SPECS[0],rows=rows_a),dict(**TABLE_SPECS[1],rows=rows_b)]
    for spec in specs:assert not re.search(r'[\u3400-\u9fff]',' '.join(spec['headers']+sum(spec['rows'],[])+spec['notes']))
    assert not re.search(r'[\u3400-\u9fff]',' '.join(texts)+FIGURE_CAPTION)
    return {'frozen_file_sha256':frozen,'translated_tables':specs,'full_precision_table_records':{'TableS6':records_a,'TableS7':records_b},
      'selection_ledger_rows_checked':182,'selected_examinations_checked':135,'selected_patients':135,'tumor_examinations':113,'treatment_change_examinations':22,
      'folds_checked':5,'saved_successful_fold_model_records_checked':20,'fixed_oof_metric_and_count_values_reconstructed':48,
      'table_s6_numeric_values_checked':28,'table_s7_numeric_values_checked':12,'bootstrap_paired_difference_values_checked':8000,'bootstrap_interval_endpoints_checked':40,
      'fusion_calibration_diagnostic_checked':{'offset_score_absolute_residual':off_score,'joint_score_maximum_absolute_residual':joint_score,'joint_hessian_minimum_eigenvalue':mineig},
      'calibration_diagnostics_refitted':0,'figure_unchanged_from_saved_materials':True,'numeric_values_changed':0,'models_refitted':0,'new_hypothesis_tests':0,
      'verification_scope':'Reconstruct identifier-only selection and fixed-OOF metrics, inspect saved fold records, verify all existing bootstrap endpoints and paired differences, and retain the unchanged figure without model fits or new resampling'}

def add_table(doc,spec):
    title=doc.add_paragraph(spec['title'],style='Heading 1');title.paragraph_format.page_break_before=True
    title.paragraph_format.keep_with_next=True
    table=doc.add_table(rows=1,cols=len(spec['headers']));table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    widths=[.64,2.12,.86,1.44,1.86] if spec['key']=='TableS6' else [2.60,1.50,2.82]
    for col,width in zip(table.columns,widths):col.width=Inches(width)
    for cell,width in zip(table.rows[0].cells,widths):cell.width=Inches(width)
    for cell,label in zip(table.rows[0].cells,spec['headers']):cell.text=label
    for values in spec['rows']:
        for cell,value,width in zip(table.add_row().cells,values,widths):cell.text=value;cell.width=Inches(width)
    props=table._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side)
        for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:e.set(qn('w:'+k),v)
        borders.append(e)
    props.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for side in ['top','left','bottom','right']:
        e=OxmlElement('w:'+side);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
    props.append(margins)
    for ri,row in enumerate(table.rows):
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if ri==0:trpr.append(OxmlElement('w:tblHeader'))
        for ci,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7EBEF');cell._tc.get_or_add_tcPr().append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:run.font.size=Pt(10.5);run.font.bold=ri==0
    for note in spec['notes']:
        p=doc.add_paragraph(note);p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(4)
        p.paragraph_format.keep_together=True
        for r in p.runs:r.font.size=Pt(10)

def build():
    OUT.mkdir(parents=True,exist_ok=True);REPORT.mkdir(parents=True,exist_ok=True);a=audit()
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(.79)
    for name in ['Normal','Title','Heading 1']:
        st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        rf=st.element.get_or_add_rPr().rFonts;rf.set(qn('w:eastAsia'),'Noto Serif CJK SC')
        for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:rf.attrib.pop(qn('w:'+key),None)
    normal=doc.styles['Normal'];normal.font.size=Pt(12);normal.paragraph_format.line_spacing=1.;normal.paragraph_format.space_after=Pt(9)
    doc.styles['Title'].font.size=Pt(16);doc.styles['Title'].paragraph_format.space_after=Pt(12)
    doc.styles['Heading 1'].font.size=Pt(13);doc.styles['Heading 1'].paragraph_format.space_before=Pt(8);doc.styles['Heading 1'].paragraph_format.space_after=Pt(9)
    doc.core_properties.title='Results 3.8 Single-scan sensitivity';doc.core_properties.author='Tianyu Feng'
    doc.add_paragraph('Results',style='Title');doc.add_paragraph('3.8 Single-scan sensitivity',style='Heading 1')
    for item in PARAGRAPHS:
        en=doc.add_paragraph(item['english']);en.paragraph_format.page_break_before=False
        en.paragraph_format.keep_together=True;en.paragraph_format.keep_with_next=True
        zh=doc.add_paragraph(item['chinese']);zh.paragraph_format.keep_together=True
        for r in zh.runs:r.font.size=Pt(11)
    for spec in a['translated_tables']:
        add_table(doc,spec)
        pd.DataFrame(a['full_precision_table_records'][spec['key']]).to_csv(OUT/(spec['key']+'_EN.csv'),index=False)
    landscape=doc.add_section(WD_SECTION.NEW_PAGE);landscape.orientation=WD_ORIENT.LANDSCAPE
    landscape.page_width=Inches(11);landscape.page_height=Inches(8.5)
    landscape.top_margin=landscape.bottom_margin=landscape.left_margin=landscape.right_margin=Inches(.79)
    doc.add_paragraph('Figure S2',style='Heading 1');doc.add_picture(str(FIGURE),width=Inches(9.3))
    p=doc.add_paragraph(FIGURE_CAPTION);p.paragraph_format.keep_together=True
    for r in p.runs:r.font.size=Pt(10)
    for e in [*doc.styles.element.xpath('.//w:pBdr'),*doc.element.xpath('.//w:pBdr')]:e.getparent().remove(e)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    path=OUT/'UCSD_Manuscript_Results_3_8_EN_ZH.docx';doc.save(path)
    (OUT/'UCSD_Manuscript_Results_3_8_EN.txt').write_text('\n\n'.join(p['english'] for p in PARAGRAPHS)+'\n')
    d={'stage':40,'date':'2026-10-29','subsection':'3.8 Single-scan sensitivity','v2_paragraph_ids':['R8'],'paragraphs':PARAGRAPHS,
      'english_word_count':sum(wc(p['english']) for p in PARAGRAPHS),'paragraph_word_counts':{p['id']:wc(p['english']) for p in PARAGRAPHS},
      'word_count_scope':'English narrative only; excludes headings, tables/notes, figure/caption and Chinese translation','v2_approximate_target_words':200,
      'equations':{'count':0,'numbers':[],'cumulative_previous_numbered_groups':10},
      'citation_mapping':{},'new_manuscript_references':0,'cumulative_unique_cited_sources':5,'editable_english_tables':['Table S6','Table S7'],'table_numeric_values_changed':0,'included_figure':str(FIGURE.relative_to(ROOT)),'figure_sha256':sha(FIGURE),'figure_caption':FIGURE_CAPTION,'figure_reused_without_regeneration':True,'source_audit':a,'models_refitted':0,'new_hypothesis_tests':0,
      'next_section':'Discussion 4.1 Principal findings (D1)','docx':str(path.relative_to(ROOT)),'docx_sha256':sha(path)}
    (REPORT/'results_section38.json').write_text(json.dumps(d,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'docx':d['docx'],'english_word_count':d['english_word_count'],'paragraph_word_counts':d['paragraph_word_counts'],'prior_files':len(a['frozen_file_sha256'])},indent=2))

def verify():
    from pypdf import PdfReader
    d=json.loads((REPORT/'results_section38.json').read_text())
    for p,h in d['source_audit']['frozen_file_sha256'].items():assert sha(ROOT/p)==h,p
    path=ROOT/d['docx'];assert sha(path)==d['docx_sha256'];texts=[p.text for p in Document(path).paragraphs]
    for p in PARAGRAPHS:
        assert p['english'] in texts and p['chinese'] in texts and texts.index(p['chinese'])==texts.index(p['english'])+1
        assert not re.search(r'[\u3400-\u9fff，。；：（）【】、]',p['english'])
    assert sum(wc(p['english']) for p in PARAGRAPHS)==d['english_word_count']
    doc=Document(path);assert len(doc.tables)==2
    for table,spec in zip(doc.tables,d['source_audit']['translated_tables']):
        assert [[c.text for c in row.cells] for row in table.rows]==[spec['headers'],*spec['rows']]
        assert all(not re.search(r'[\u3400-\u9fff]',c.text) for row in table.rows for c in row.cells)
        for note in spec['notes']:assert note in texts
        full=pd.read_csv(OUT/(spec['key']+'_EN.csv'));assert len(full)==4
        for (_,row),r in zip(full.iterrows(),d['source_audit']['full_precision_table_records'][spec['key']]):
            for key,value in r.items():
                if isinstance(value,str):assert row[key]==value
                else:close(row[key],value)
    assert len(doc.sections)==2 and doc.sections[1].orientation==WD_ORIENT.LANDSCAPE and FIGURE_CAPTION in texts
    ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
    with zipfile.ZipFile(path) as z:
        t=etree.fromstring(z.read('word/document.xml'));assert len(t.xpath('//m:oMath',namespaces=ns))==0
        media=[n for n in z.namelist() if n.startswith('word/media/')];assert len(media)==1
        assert hashlib.sha256(z.read(media[0])).hexdigest()==sha(FIGURE)
    pages=len(PdfReader(OUT/'docx_render/UCSD_Manuscript_Results_3_8_EN_ZH.pdf').pages);assert pages==4
    results=[json.loads(p.read_text()) for p in ROOT.glob('reports/ucsd_manuscript_step*/results_section*.json')]
    ids=[p['id'] for section in results for p in section['paragraphs']]
    assert len(results)==8 and len(ids)==8 and set(ids)=={f'R{i}' for i in range(1,9)}
    result_words=sum(section['english_word_count'] for section in results)
    v={'success':True,'stage':40,'date':'2026-10-29','subsection':d['subsection'],'paragraph_ids':['R8'],
      'english_word_count':d['english_word_count'],'paragraph_word_counts':d['paragraph_word_counts'],'matched_chinese_paragraphs':1,
      'word_pages':pages,'visual_review_pages':[1,2,3,4],'native_editable_equations':0,'equation_numbers':[],'equation_images':0,
      'prior_files_unchanged':len(d['source_audit']['frozen_file_sha256']),'selection_ledger_rows_checked':182,'selected_examinations_checked':135,'selected_patients':135,'tumor_examinations':113,'treatment_change_examinations':22,
      'folds_checked':5,'saved_successful_fold_model_records_checked':20,'fixed_oof_metric_and_count_values_reconstructed':48,'table_s6_numeric_values_checked':28,'table_s7_numeric_values_checked':12,
      'bootstrap_paired_difference_values_checked':8000,'bootstrap_interval_endpoints_checked':40,'calibration_diagnostics_refitted':0,'editable_english_tables':2,'table_numeric_values_changed':0,'paired_intervals_include_zero_disclosed':True,'unchanged_figure_embedded':True,'figure_page_landscape':True,
      'models_refitted':0,'new_hypothesis_tests':0,'new_manuscript_references':0,'cumulative_unique_cited_sources':5,'results_v2_paragraphs_complete':True,'completed_results_subsections':8,'completed_results_english_word_count':result_words,'next_section':d['next_section'],
      'sha256':{str(p.relative_to(ROOT)):sha(p) for p in OUT.iterdir() if p.is_file() and p.suffix in {'.docx','.txt','.csv'}}}
    (REPORT/'artifact_verification.json').write_text(json.dumps(v,ensure_ascii=False,indent=2)+'\n');print(json.dumps(v,indent=2))

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--verify',action='store_true');args=parser.parse_args()
    verify() if args.verify else build()
