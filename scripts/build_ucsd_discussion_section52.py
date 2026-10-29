"""Draft D12 as a prospective synthesis of the completed Discussion limitations."""
from pathlib import Path
import argparse, hashlib, json, re, zipfile
import pandas as pd
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'outputs/ucsd-manuscript-step52-20261030'
REPORT = ROOT/'reports/ucsd_manuscript_step52'
EVIDENCE = REPORT/'literature_evidence.json'
MAPPING = {'15': 'S25', '16': 'S26'}
PARAGRAPHS = [{'id': 'D12', 'citation_ids': ['S25', 'S26'], 'english': 'Further work should prioritize independently adjudicated examination diagnoses and documented annotation procedures, including the information available to annotators. Acquisition metadata and spatial scales should be verified, with preprocessing specified before evaluation. The locked clinical and fusion pipelines should be evaluated in an independent cohort representing the intended population, with sample size planned for precise performance estimation (15,16). Assessment should give particular attention to treatment-related-change recognition and calibration, alongside discrimination and probability error. Any clinical-utility study would additionally require a defined action and justified thresholds. These studies would provide evidence on whether the observed imaging increment persists beyond the present dataset.', 'chinese': '后续工作应优先取得经过独立裁定的检查诊断，并记录标注流程，包括标注者能够获得的信息。应核实采集元数据及空间尺度，并在评价前确定预处理流程。应在代表目标人群的独立队列中评价已锁定的临床与融合流程，并以精确估计性能为目标规划样本量（15,16）。除鉴别能力及概率误差外，评价尤其应关注治疗相关改变的识别与校准。任何临床效用研究还需要明确相应处置，并为阈值提供依据。这些研究将为所观察到的影像增量能否延续至当前数据集之外提供证据。'}]

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def wc(text):
    return len(re.findall(r"[A-Za-z0-9]+(?:['/-][A-Za-z0-9]+)*", text))

def audit():
    previous=json.loads((ROOT/'reports/ucsd_manuscript_step51/discussion_section51.json').read_text())
    frozen=previous['source_audit']['frozen_file_sha256'].copy()
    frozen.update(json.loads((ROOT/'reports/ucsd_manuscript_step51/artifact_verification.json').read_text())['sha256'])
    discussion=[]
    for stage in range(41,52):
        rel=f'reports/ucsd_manuscript_step{stage}/discussion_section{stage}.json'
        if rel not in frozen:frozen[rel]=sha(ROOT/rel)
        d=json.loads((ROOT/rel).read_text());discussion.extend(d['paragraphs'])
    frozen[str(EVIDENCE.relative_to(ROOT))]=sha(EVIDENCE)
    for p,h in frozen.items():assert sha(ROOT/p)==h,p
    assert [p['id'] for p in discussion]==[f'D{i}' for i in range(1,12)]
    assert sum(wc(p['english']) for p in discussion)==2323
    evidence=json.loads(EVIDENCE.read_text());records={s['reference_id']:s for s in evidence['studies']}
    refs={s['id']:s for s in json.loads((ROOT/'reports/ucsd_manuscript_step23/reference_library.json').read_text())['references']}
    selected=[]
    for number,rid in MAPPING.items():
        ref=refs[rid];rec=records[rid];assert ref['doi']==rec['doi'] and int(number)==rec['provisional_citation_number']
        selected.append({'number':int(number),'reference_id':rid,'citation':ref['citation'],'doi':ref['doi'],'read_level_this_stage':rec['read_level_this_stage']})
    refs_to_limits={
        'independent_diagnoses_and_annotation_records':'reports/ucsd_manuscript_step48/discussion_section48.json',
        'metadata_scales_and_processing':'reports/ucsd_manuscript_step49/discussion_section49.json',
        'independent_evaluation_and_precision_planning':'reports/ucsd_manuscript_step50/discussion_section50.json',
        'minority_class_calibration_and_clinical_utility':'reports/ucsd_manuscript_step51/discussion_section51.json'}
    assert all(rel in frozen for rel in refs_to_limits.values())
    m=json.loads((ROOT/'reports/ucsd_fusion/metrics.json').read_text())
    assert not m['independent_external_validation'] and not m['source_annotation_blinding_confirmed'] and not m['physical_scale_confirmed']
    lock=json.loads((ROOT/'reports/ucsd_fusion/analysis_lock.json').read_text())
    assert lock['model_features']==m['models'] and len(lock['model_features']['Fusion'])==13
    prior_ids=set()
    for pattern in ['methods_section*.json','results_section*.json','discussion_section*.json']:
        for p in ROOT.glob('reports/ucsd_manuscript_step*/'+pattern):
            if int(p.parent.name.rsplit('step',1)[1])<52:prior_ids.update(r for item in json.loads(p.read_text())['paragraphs'] for r in item['citation_ids'])
    assert len(prior_ids)==17 and set(MAPPING.values()).issubset(prior_ids)
    return {'frozen_file_sha256':frozen,'reference_records':selected,'source_evidence_record':str(EVIDENCE.relative_to(ROOT)),
            'new_full_texts_checked_for_current_claims':[],'previous_full_texts_revisited':['S25','S26'],
            'prior_discussion_paragraphs_checked':11,'prior_discussion_english_words':2323,'recommendations_to_documented_limitations':refs_to_limits,
            'locked_model_inputs_unchanged':True,'future_recommendations_not_completed_work':True,
            'new_external_cohort_obtained':False,'independent_diagnoses_or_annotation_blinding_newly_verified':False,
            'physical_scale_newly_verified':False,'new_sample_size_target_calculated':False,'clinical_net_benefit_calculated':False,
            'independent_probability_recalibration_performed':False,'new_threshold_recommended':False,
            'previous_unique_cited_sources':sorted(prior_ids),'models_refitted':0,'new_hypothesis_tests':0,
            'verification_scope':'Check all eleven prior Discussion paragraphs, map recommendations to their frozen limitation records, preserve locked model inputs and use previously reviewed primary guidance; no new validation, fitting or sample-size calculation'}

def build():
    a = audit(); OUT.mkdir(parents=True, exist_ok=True); REPORT.mkdir(parents=True, exist_ok=True)
    doc = Document(); sec = doc.sections[0]; sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(.79)
    for name in ['Normal', 'Title', 'Heading 1']:
        st = doc.styles[name]; st.font.name = 'Times New Roman'; st.font.color.rgb = RGBColor(0, 0, 0)
        rf = st.element.get_or_add_rPr().rFonts; rf.set(qn('w:eastAsia'), 'Noto Serif CJK SC')
        for key in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']: rf.attrib.pop(qn('w:'+key), None)
    normal = doc.styles['Normal']; normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.; normal.paragraph_format.space_after = Pt(9)
    doc.styles['Title'].font.size = Pt(16); doc.styles['Title'].paragraph_format.space_after = Pt(12)
    heading = doc.styles['Heading 1']; heading.font.size = Pt(13)
    heading.paragraph_format.space_before = Pt(8); heading.paragraph_format.space_after = Pt(9)
    doc.core_properties.title = 'Discussion 4.12 Priorities for further evaluation'; doc.core_properties.author = 'Tianyu Feng'
    doc.add_paragraph('Discussion', style='Title'); doc.add_paragraph('4.12 Priorities for further evaluation', style='Heading 1')
    for item in PARAGRAPHS:
        en = doc.add_paragraph(item['english']); en.paragraph_format.keep_together = True; en.paragraph_format.keep_with_next = True
        zh = doc.add_paragraph(item['chinese']); zh.paragraph_format.keep_together = True
        for run in zh.runs: run.font.size = Pt(11)
    if a['reference_records']:
        title = doc.add_paragraph('References', style='Heading 1'); title.paragraph_format.page_break_before = False
    for ref in a['reference_records']:
        p = doc.add_paragraph(str(ref['number'])+'. '+ref['citation']); p.paragraph_format.keep_together = True
    for e in [*doc.styles.element.xpath('.//w:pBdr'), *doc.element.xpath('.//w:pBdr')]: e.getparent().remove(e)
    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    path = OUT/'UCSD_Manuscript_Discussion_4_12_EN_ZH.docx'; doc.save(path)
    (OUT/'UCSD_Manuscript_Discussion_4_12_EN.txt').write_text('\n\n'.join(p['english'] for p in PARAGRAPHS)+'\n')
    d = {'stage': 52, 'date': '2026-10-30', 'subsection': '4.12 Priorities for further evaluation', 'v2_paragraph_ids': ['D12'],
         'paragraphs': PARAGRAPHS, 'english_word_count': sum(wc(p['english']) for p in PARAGRAPHS),
         'paragraph_word_counts': {p['id']: wc(p['english']) for p in PARAGRAPHS},
         'word_count_scope': 'English narrative including in-text citation numbers; excludes headings, reference list and Chinese translation',
         'v2_approximate_target_words': 100, 'equations': {'count': 0, 'numbers': [], 'cumulative_previous_numbered_groups': 10},
         'citation_mapping': MAPPING, 'citation_numbering': 'Existing numbers 15 and 16 preserved; renumber globally when Introduction is assembled',
         'new_manuscript_references': 0, 'cumulative_unique_cited_sources': 17,
         'source_audit': a, 'models_refitted': 0, 'new_hypothesis_tests': 0,
         'next_section': 'Conclusion (C1)',
         'docx': str(path.relative_to(ROOT)), 'docx_sha256': sha(path)}
    (REPORT/'discussion_section52.json').write_text(json.dumps(d, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps({'docx': d['docx'], 'english_word_count': d['english_word_count'], 'frozen_files': len(a['frozen_file_sha256'])}))

def verify():
    from pypdf import PdfReader
    d = json.loads((REPORT/'discussion_section52.json').read_text()); a = audit(); assert a == d['source_audit']
    path = ROOT/d['docx']; assert sha(path) == d['docx_sha256']; doc = Document(path); texts = [p.text for p in doc.paragraphs]
    for p in PARAGRAPHS:
        assert p['english'] in texts and p['chinese'] in texts and texts.index(p['chinese']) == texts.index(p['english'])+1
        assert not re.search(r'[\u3400-\u9fff，。；：（）【】、]', p['english'])
    assert sum(wc(p['english']) for p in PARAGRAPHS) == d['english_word_count']
    assert re.findall(r'\((\d+),(\d+)\)', PARAGRAPHS[0]['english']) == [('15','16')]
    for ref in a['reference_records']: assert str(ref['number'])+'. '+ref['citation'] in texts
    assert not doc.tables and len(doc.sections) == 1
    with zipfile.ZipFile(path) as z:
        xml = etree.fromstring(z.read('word/document.xml'))
        assert not xml.xpath('//m:oMath', namespaces={'m': 'http://schemas.openxmlformats.org/officeDocument/2006/math'})
        assert not [n for n in z.namelist() if n.startswith('word/media/')]
    pdf = PdfReader(OUT/'docx_render/UCSD_Manuscript_Discussion_4_12_EN_ZH.pdf'); pages = len(pdf.pages); assert pages == 1
    assert '4.12 Priorities' in pdf.pages[0].extract_text() and 'References' in pdf.pages[0].extract_text()
    v = {'success': True, 'stage': 52, 'date': '2026-10-30', 'subsection': d['subsection'], 'paragraph_ids': ['D12'],
         'english_word_count': d['english_word_count'], 'matched_chinese_paragraphs': 1, 'word_pages': pages,
         'visual_review_pages': [1], 'native_editable_equations': 0, 'equation_images': 0,
         'frozen_files_unchanged': len(a['frozen_file_sha256']), 'prior_files_unchanged': len(a['frozen_file_sha256'])-1,
         'new_full_texts_checked_for_current_claims': [], 'prior_discussion_paragraphs_checked': 11,
         'recommendations_mapped_to_frozen_limitations': True, 'locked_model_inputs_unchanged': True,
         'future_recommendations_not_completed_work': True, 'discussion_v2_paragraphs_complete': True,
         'discussion_english_word_count': 2323+d['english_word_count'],
         'new_external_cohort_obtained': False, 'independent_diagnoses_or_annotation_blinding_newly_verified': False,
         'physical_scale_newly_verified': False, 'new_sample_size_target_calculated': False,
         'clinical_net_benefit_calculated': False, 'independent_probability_recalibration_performed': False,
         'new_threshold_recommended': False,
         'new_manuscript_references': 0, 'provisional_citation_mapping': MAPPING, 'cumulative_unique_cited_sources': 17,
         'models_refitted': 0, 'new_hypothesis_tests': 0, 'next_section': d['next_section'],
         'sha256': {str(p.relative_to(ROOT)): sha(p) for p in OUT.iterdir() if p.is_file() and p.suffix in {'.docx', '.txt'}}}
    (REPORT/'artifact_verification.json').write_text(json.dumps(v, ensure_ascii=False, indent=2)+'\n'); print(json.dumps(v, indent=2))

if __name__ == '__main__':
    parser = argparse.ArgumentParser(); parser.add_argument('--verify', action='store_true'); args = parser.parse_args()
    verify() if args.verify else build()
