"""Write M12-M13 and verify fixed models and OOF probabilities without fitting."""
from pathlib import Path
from collections import Counter
import argparse,hashlib,json,re,sys,zipfile
import numpy as np
import pandas as pd
from lxml import etree
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from build_ucsd_methods_section25 import el,mr,frac,sub,bracket,summation,function

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'outputs/ucsd-manuscript-step28-20261029'
REPORT=ROOT/'reports/ucsd_manuscript_step28'
PARAGRAPHS=[
 {'id':'M12','citation_ids':[],
  'english': 'We fitted four binary logistic regression models with an intercept. M0 included age and sex; M1 included the five clinical inputs defined above and served as the principal clinical comparator. The MRI model used the eight fixed imaging features, and Fusion concatenated the five clinical inputs and eight MRI features into one 13-variable predictor vector. This was feature-level fusion within a single model. Equation 7 gives the tumor probability for the processed feature vector z, where m denotes the input count and the beta terms denote the intercept and coefficients. The four models contained 2, 5, 8 and 13 nonintercept terms, respectively. All used L2 regularization with inverse regularization strength C set to 1, the lbfgs optimizer, a maximum of 2000 iterations and convergence tolerance 0.0001. No class weighting was applied. A common regression framework provided a consistent comparison of the fixed input sets. Coefficients were used for prediction and were not interpreted as causal effects.',
  'chinese': '我们拟合了四个含截距的二分类Logistic回归模型。M0纳入年龄与性别；M1纳入前述五个临床输入，作为主要临床对照。MRI模型使用八项固定影像特征；Fusion将五个临床输入与八项MRI特征拼接为一个含13项变量的预测向量。这属于单一模型内的特征级融合。Equation 7给出处理后特征向量z对应的肿瘤概率，其中m表示输入数量，beta项表示截距与系数。四个模型分别含2、5、8和13个非截距项。全部采用L2正则化，逆正则化强度C设为1，优化器为lbfgs，最大迭代次数2000，收敛容差0.0001。不使用类别加权。共同的回归框架用于一致地比较固定输入集合。系数用于预测，不解释为因果效应。'},
 {'id':'M13','citation_ids':[],
  'english': 'All four models were evaluated in the same 182-examination cohort using the same five patient-grouped folds. We used StratifiedGroupKFold with shuffling enabled and random seed 20261029, grouping examinations by the audited patient-group identifier. All examinations from a given patient were kept in one fold, and separation was checked using both the source clinical and BraTS patient identifiers. The case-to-fold manifest was fixed before model fitting and was not reselected after viewing performance. Age, the transformed interval and MRI features were standardized within training folds wherever included; interval imputation followed Section 2.3. Binary inputs were not scaled. Fitted preprocessing and regression parameters were applied to the held-out fold without using its outcomes to train or select the model. Each examination received one out-of-fold (OOF) tumor probability per model. There was no hyperparameter search, feature selection, dimensionality reduction, oversampling or interaction-term expansion. Validation was internal to the public collection.',
  'chinese': '四个模型均在同一包含182次检查的队列中，使用相同的五个患者分组折进行评价。采用StratifiedGroupKFold，启用随机打乱，随机种子为20261029，按已核查的患者分组标识划分检查。同一患者的全部检查保持在同一折，并分别使用临床源患者标识和BraTS患者标识检查训练与验证的隔离。病例与折的对应清单在模型拟合前固定，未根据观察到的性能重新选择。年龄、变换后的间隔和MRI特征在各自被纳入时，均仅在训练折内进行标准化；间隔填补按照Section 2.3执行。二分类输入不缩放。已拟合的预处理和回归参数用于留出的验证折，不使用该折结局训练或选择模型。每次检查在每个模型中获得一个折外（OOF）肿瘤概率。未进行超参数搜索、特征筛选、降维、过采样或交互项扩展。验证范围限于该公开数据集内部。'}
]

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def wc(s):return len(re.findall(r"[A-Za-z0-9]+(?:['/-][A-Za-z0-9]+)*",s))
def audit():
    old=json.loads((ROOT/'reports/ucsd_manuscript_step27/methods_section25.json').read_text())
    frozen=old['source_audit']['frozen_file_sha256'].copy()
    frozen.update(json.loads((ROOT/'reports/ucsd_manuscript_step27/artifact_verification.json').read_text())['sha256'])
    for p,h in frozen.items():assert sha(ROOT/p)==h,p
    lock=json.loads((ROOT/'reports/ucsd_fusion/analysis_lock.json').read_text())
    for p,h in lock['input_sha256'].items():assert sha(ROOT/p)==h,p
    assert sha(ROOT/'src/rn_public/ucsd_fusion.py')==lock['analysis_code_sha256']
    assert sha(ROOT/'reports/ucsd_fusion_protocol.md')==lock['protocol_sha256']
    manifest_path=ROOT/'reports/ucsd_fusion/cohort_and_fold_manifest.csv'
    assert sha(manifest_path)==lock['fold_manifest_sha256']
    sys.path.insert(0,str(ROOT/'src'))
    from rn_public.ucsd_fusion import FEATURES,make_model,SEED
    from sklearn.model_selection import StratifiedGroupKFold
    assert SEED==lock['seed']==20261029 and lock['n_splits']==5
    sizes={k:len(v) for k,v in FEATURES.items()};assert sizes=={'M0':2,'M1':5,'MRI':8,'Fusion':13}
    assert FEATURES==lock['model_features'] and FEATURES['Fusion']==FEATURES['M1']+FEATURES['MRI']
    factory={}
    for name in FEATURES:
        model=make_model(name);p=model.named_steps['logistic'].get_params()
        expected={'C':1,'l1_ratio':0,'solver':'lbfgs','max_iter':2000,'tol':.0001,'class_weight':None,'random_state':SEED,'fit_intercept':True}
        assert all(p[k]==v for k,v in expected.items())
        assert not hasattr(model.named_steps['logistic'],'coef_')
        factory[name]=expected
    data=pd.read_csv(ROOT/'data/processed/ucsd_clinical_image_features.csv').sort_values('ID').reset_index(drop=True)
    manifest=pd.read_csv(manifest_path);oof=pd.read_csv(ROOT/'data/processed/ucsd_fusion_oof.csv').set_index('ID')
    assert len(data)==len(manifest)==len(oof)==182 and data.patient_id.nunique()==135
    assert set(data.ID)==set(manifest.ID)==set(oof.index)
    folds=json.loads((ROOT/'reports/ucsd_fusion/fold_audit.json').read_text());assert len(folds)==5
    cv=StratifiedGroupKFold(n_splits=5,shuffle=True,random_state=SEED);derived=np.zeros(len(data),dtype=int)
    for fold,(_,test) in enumerate(cv.split(data,data.tumor_label,data.validation_group),1):derived[test]=fold
    expected_manifest=data[['ID','patient_id','brats_patient_id','validation_group','tumor_label']].copy();expected_manifest['validation_fold']=derived
    pd.testing.assert_frame_equal(expected_manifest,manifest)
    records=[];covered=Counter();maximum_error=0.;coefficient_terms=0
    indexed=data.set_index('ID')
    for f in folds:
        train=indexed.loc[f['train_IDs']];test=indexed.loc[f['test_IDs']]
        assert set(train.index).isdisjoint(test.index) and set(train.index)|set(test.index)==set(indexed.index)
        assert train.tumor_label.nunique()==test.tumor_label.nunique()==2
        assert all(int(oof.loc[i,'validation_fold'])==f['fold'] for i in test.index)
        for namespace in ['patient_id','brats_patient_id','validation_group']:
            assert set(train[namespace]).isdisjoint(test[namespace]) and f[namespace+'_overlap']==0
        covered.update(test.index)
        for name in FEATURES:
            p=f['models'][name]['parameters'];assert p['classes']==[0,1]
            assert len(p['transformed_names'])==len(p['coefficients'])==sizes[name]
            assert [a.split('__',1)[1] for a in p['transformed_names']]==FEATURES[name]
            assert p['iterations']<2000
            normalized={}
            for block,v in p['continuous'].items():
                raw_train=train[v['columns']].to_numpy(dtype=float);raw_test=test[v['columns']].to_numpy(dtype=float)
                if block=='interval':
                    median=float(np.nanmedian(raw_train));assert np.isclose(median,v['impute_median'],rtol=1e-12,atol=1e-12)
                    raw_train=np.where(np.isnan(raw_train),median,raw_train);raw_test=np.where(np.isnan(raw_test),median,raw_test)
                np.testing.assert_allclose(raw_train.mean(axis=0),v['mean'],rtol=1e-12,atol=1e-12)
                np.testing.assert_allclose(raw_train.std(axis=0,ddof=0),v['scale'],rtol=1e-12,atol=1e-12)
                values=(raw_test-np.asarray(v['mean']))/np.asarray(v['scale'])
                for col,j in zip(v['columns'],range(values.shape[1])):normalized[block+'__'+col]=values[:,j]
            for col in p['binary']:normalized['binary__'+col]=test[col].to_numpy(dtype=float)
            z=np.column_stack([normalized[col] for col in p['transformed_names']])
            linear=p['intercept']+z@np.asarray(p['coefficients']);probability=1/(1+np.exp(-linear))
            saved=oof.loc[test.index,name+'_oof_probability'].to_numpy(dtype=float)
            error=float(np.max(np.abs(probability-saved)));assert error<1e-12
            maximum_error=max(maximum_error,error);coefficient_terms+=len(p['coefficients'])
            records.append({'fold':f['fold'],'model':name,'input_count':sizes[name],'test_examinations':len(test),'iterations_recorded':p['iterations'],'maximum_probability_error':error})
    assert set(covered)==set(data.ID) and set(covered.values())=={1}
    assert len(records)==20 and coefficient_terms==140
    metrics=json.loads((ROOT/'reports/ucsd_fusion/metrics.json').read_text())
    assert metrics['same_cases_and_folds'] and not metrics['hyperparameter_search'] and not metrics['feature_selection']
    assert metrics['patient_overlap']==metrics['brats_patient_overlap']==0 and not metrics['full_fit_training_metrics_reported']
    extra=['reports/ucsd_fusion/analysis_lock.json','reports/ucsd_fusion_protocol.md','scripts/build_ucsd_methods_section25.py']
    frozen.update({p:sha(ROOT/p) for p in extra})
    return {'frozen_file_sha256':frozen,'model_sizes':sizes,'fixed_factory_parameters':factory,'saved_fold_models_checked':20,'saved_coefficient_terms_checked':140,
      'saved_OOF_probabilities_reconstructed':728,'maximum_absolute_probability_error':maximum_error,'fold_assignments_reproduced':182,
      'source_patients':135,'all_examinations_validated_exactly_once':True,'all_three_patient_namespaces_disjoint':True,
      'training_preprocessing_parameters_match':True,'reconstruction_records':records,'models_fitted_this_stage':0,'new_hypothesis_tests':0,
      'documentation':['https://scikit-learn.org/1.8/modules/generated/sklearn.linear_model.LogisticRegression.html','https://scikit-learn.org/1.8/modules/generated/sklearn.model_selection.StratifiedGroupKFold.html']}

def equation():
    linear=[sub(mr('β'),'0'),mr(' + '),summation('j=1','m',[sub(mr('β'),'j'),sub(mr('z'),'j')])]
    denominator=[mr('1 + '),function(mr('exp',True),bracket([mr('−'),bracket(linear)]))]
    return [mr('p'),bracket(mr('z')),mr(' = '),frac(mr('1'),denominator)]

def build():
    OUT.mkdir(parents=True,exist_ok=True);REPORT.mkdir(parents=True,exist_ok=True);a=audit()
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(.79)
    for name in ['Normal','Title','Heading 1']:
        st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        rf=st.element.get_or_add_rPr().rFonts;rf.set(qn('w:eastAsia'),'Noto Serif CJK SC')
        for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:rf.attrib.pop(qn('w:'+key),None)
    normal=doc.styles['Normal'];normal.font.size=Pt(12);normal.paragraph_format.line_spacing=1.;normal.paragraph_format.space_after=Pt(9)
    doc.styles['Title'].font.size=Pt(16);doc.styles['Title'].paragraph_format.space_after=Pt(12)
    doc.styles['Heading 1'].font.size=Pt(13);doc.styles['Heading 1'].paragraph_format.space_before=Pt(8);doc.styles['Heading 1'].paragraph_format.space_after=Pt(9)
    doc.core_properties.title='Materials and methods 2.6 Model specification and grouped training';doc.core_properties.author='Tianyu Feng'
    doc.add_paragraph('Materials and methods',style='Title');doc.add_paragraph('2.6 Model specification and grouped training',style='Heading 1')
    for item in PARAGRAPHS:
        en=doc.add_paragraph(item['english']);en.paragraph_format.page_break_before=item['id']=='M13'
        en.paragraph_format.keep_together=True;en.paragraph_format.keep_with_next=True
        zh=doc.add_paragraph(item['chinese']);zh.paragraph_format.keep_together=True
        for r in zh.runs:r.font.size=Pt(11)
        if item['id']=='M12':
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_together=True
            p.paragraph_format.space_before=Pt(6);p._p.append(el('oMath',equation()));p.add_run('    (7)')
    for e in [*doc.styles.element.xpath('.//w:pBdr'),*doc.element.xpath('.//w:pBdr')]:e.getparent().remove(e)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    path=OUT/'UCSD_Manuscript_Methods_2_6_EN_ZH.docx';doc.save(path)
    (OUT/'UCSD_Manuscript_Methods_2_6_EN.txt').write_text('\n\n'.join(p['english'] for p in PARAGRAPHS)+'\n')
    d={'stage':28,'date':'2026-10-29','subsection':'2.6 Model specification and grouped training','v2_paragraph_ids':['M12','M13'],'paragraphs':PARAGRAPHS,
      'english_word_count':sum(wc(p['english']) for p in PARAGRAPHS),'paragraph_word_counts':{p['id']:wc(p['english']) for p in PARAGRAPHS},
      'word_count_scope':'English narrative; excludes equation, headings and Chinese translation','v2_approximate_target_words':300,
      'equations':{'count':1,'format':'native editable OMML','numbers':[7],'meaning':'Tumor probability from a processed feature vector; intercept and additive linear terms'},
      'citation_mapping':{},'new_manuscript_references':0,'cumulative_unique_cited_sources':4,
      'source_audit':a,'models_refitted':0,'new_hypothesis_tests':0,'next_section':'Performance metrics and conditional uncertainty (M14-M15)',
      'docx':str(path.relative_to(ROOT)),'docx_sha256':sha(path)}
    (REPORT/'methods_section26.json').write_text(json.dumps(d,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'docx':d['docx'],'english_word_count':d['english_word_count'],'paragraph_word_counts':d['paragraph_word_counts'],'maximum_probability_error':a['maximum_absolute_probability_error'],'prior_files':len(a['frozen_file_sha256'])},indent=2))

def verify():
    from pypdf import PdfReader
    d=json.loads((REPORT/'methods_section26.json').read_text())
    for p,h in d['source_audit']['frozen_file_sha256'].items():assert sha(ROOT/p)==h,p
    path=ROOT/d['docx'];assert sha(path)==d['docx_sha256'];texts=[p.text for p in Document(path).paragraphs]
    for p in PARAGRAPHS:
        assert p['english'] in texts and p['chinese'] in texts and texts.index(p['chinese'])==texts.index(p['english'])+1
        assert not re.search(r'[\u3400-\u9fff，。；：（）【】、]',p['english'])
    assert sum(wc(p['english']) for p in PARAGRAPHS)==d['english_word_count'] and '    (7)' in texts
    ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
    with zipfile.ZipFile(path) as z:
        t=etree.fromstring(z.read('word/document.xml'));assert len(t.xpath('//m:oMath',namespaces=ns))==1
        assert len(t.xpath('//m:f',namespaces=ns))==len(t.xpath('//m:nary',namespaces=ns))==1
        assert not any(n.startswith('word/media/') for n in z.namelist())
    pages=len(PdfReader(OUT/'docx_render/UCSD_Manuscript_Methods_2_6_EN_ZH.pdf').pages);assert pages==2
    v={'success':True,'stage':28,'date':'2026-10-29','subsection':d['subsection'],'paragraph_ids':['M12','M13'],
      'english_word_count':d['english_word_count'],'paragraph_word_counts':d['paragraph_word_counts'],'matched_chinese_paragraphs':2,
      'word_pages':pages,'visual_review_pages':[1,2],'native_editable_equations':1,'equation_numbers':[7],'equation_images':0,
      'prior_files_unchanged':len(d['source_audit']['frozen_file_sha256']),'model_sizes':d['source_audit']['model_sizes'],
      'fold_assignments_checked':182,'saved_fold_models_checked':20,'OOF_probabilities_reconstructed':728,
      'maximum_absolute_probability_error':d['source_audit']['maximum_absolute_probability_error'],'all_patient_namespaces_disjoint':True,
      'models_refitted':0,'new_hypothesis_tests':0,'new_manuscript_references':0,'next_section':d['next_section'],
      'sha256':{str(p.relative_to(ROOT)):sha(p) for p in OUT.iterdir() if p.is_file() and p.suffix in {'.docx','.txt'}}}
    (REPORT/'artifact_verification.json').write_text(json.dumps(v,ensure_ascii=False,indent=2)+'\n');print(json.dumps(v,indent=2))

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--verify',action='store_true');args=parser.parse_args()
    verify() if args.verify else build()
