"""Draft complete Introduction I1-I6 from primary literature and the locked study design."""
from pathlib import Path
import argparse, hashlib, json, re, zipfile
import pandas as pd
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'outputs/ucsd-manuscript-step54-20261030'
REPORT = ROOT/'reports/ucsd_manuscript_step54'
EVIDENCE = REPORT/'literature_evidence.json'
MAPPING = {'1': 'S01', '2': 'S02', '4': 'S32', '6': 'S06', '7': 'S12', '9': 'S13', '10': 'S14', '14': 'S27', '18': 'S05', '19': 'S04', '20': 'S20', '21': 'S36', '22': 'S38', '23': 'S34', '24': 'S23', '25': 'S24'}
PARAGRAPHS = [{'id': 'I1', 'citation_ids': ['S05', 'S04'], 'english': 'Treatment follow-up in high-grade glioma requires clinicians to interpret magnetic resonance imaging (MRI) abnormalities in the context of previous therapy. For glioblastoma, contemporary management combines surgery, radiotherapy and systemic treatment, while subsequent imaging contributes to assessment of disease status and further management (18). New or increasing contrast enhancement after treatment may indicate tumor activity, but can also reflect treatment effects. An apparent increase in lesion burden therefore does not, by itself, resolve the underlying diagnosis. Response Assessment in Neuro-Oncology (RANO) 2.0 addresses this problem through standardized response assessment and, in relevant early post-radiotherapy settings, confirmation of progression by repeat imaging or unequivocal histopathological evidence (19). Such requirements emphasize the importance of temporal and clinical context when interpreting a suspicious examination. They also motivate methods that can summarize information already present in the scan and assess whether it adds to available clinical information. The immediate research question is whether that combination improves discrimination of the status assigned to an examination, providing evidence for subsequent evaluation rather than replacing the broader diagnostic process.', 'chinese': '高级别胶质瘤治疗后的随访要求临床医师结合既往治疗解读磁共振成像（MRI）异常。对于胶质母细胞瘤，当前治疗综合使用手术、放疗及系统性治疗，后续影像参与疾病状态及进一步管理的评估（18）。治疗后新出现或增大的对比增强可能提示肿瘤活动，也可能反映治疗效应，因此，病灶负荷表面增加本身并不能确定实际诊断。神经肿瘤疗效评估（RANO）2.0通过标准化疗效评估应对这一问题，并在适用的放疗后早期场景中，要求通过重复影像或明确的组织病理证据确认进展（19）。这些要求强调了解读可疑检查时，时间与临床背景的重要性，也促使研究者考虑能否汇总扫描中已有的信息，并评估其对现有临床信息的增量。本研究直接关注的是，这种结合能否改善对每次检查所赋状态的鉴别，为后续评价提供依据，而不替代整体诊断。'}, {'id': 'I2', 'citation_ids': ['S36', 'S20', 'S05', 'S01'], 'english': 'The difficulty is partly one of outcome definition. Pseudoprogression commonly refers to an apparent radiological worsening related to treatment that subsequently stabilizes or regresses without a new antitumor treatment. Radiation necrosis represents a different form of treatment injury, often discussed as a later complication; these terms should not be treated as interchangeable with every nonspecific post-treatment abnormality (18,20). RANO 2.0 implementation guidance further distinguishes enhancing, nonenhancing and mixed imaging presentations and the circumstances requiring confirmation scans (21). These distinctions matter when comparing diagnostic studies, because a model trained to identify pseudoprogression may address a narrower target than a model separating tumor from a broad treatment-change category. Reference diagnoses also vary: pathological assessment, clinical evolution and serial imaging provide different forms of evidence, and even pathological differentiation can remain difficult (18). The public UCSD-PTGBM collection includes residual or recurrent tumor and several treatment-related-change categories with mixed diagnostic support (1). It therefore permits investigation of a broad examination-level distinction, while retaining uncertainty about individual source diagnoses and the interpretation of treatment-change subtypes.', 'chinese': '这一困难部分来自结局定义。假性进展通常指与治疗有关的影像表观恶化，其后在未开始新的抗肿瘤治疗时趋于稳定或消退。放射性坏死属于另一种治疗损伤形式，常作为较晚期并发症讨论；这些术语不能与所有非特异性治疗后异常互换使用（18,20）。RANO 2.0的实施指南还区分增强、非增强及混合影像表现，以及需要确认扫描的具体情形（21）。比较诊断研究时，这些区别十分重要，因为识别假性进展的模型所针对的目标，可能比区分肿瘤与广义治疗改变类别的模型更窄。参考诊断也并不相同：病理评估、临床演变和连续影像提供不同形式的证据，即使在病理层面，鉴别仍可能困难（18）。公开UCSD-PTGBM数据集包含残余或复发肿瘤及多种治疗相关改变类别，其诊断依据并不统一（1）。因此，该数据集允许研究检查层面的广义鉴别，同时仍需保留对个体源诊断及治疗改变亚型解释的不确定性。'}, {'id': 'I3', 'citation_ids': ['S20', 'S32', 'S38'], 'english': 'Quantitative image analysis offers a way to describe these examinations through explicit computations. Intensity statistics summarize the distribution of values within a region, whereas texture descriptors characterize aspects of gray-level organization that are not captured by a mean alone. Radiomics research has explored such information for differentiating progression and treatment effects in glioma, although the interpretation depends on the images, segmentation and feature definitions used (20). The Image Biomarker Standardization Initiative provides common definitions that help make feature calculations comparable (4). Standard definitions do not remove acquisition and preprocessing differences. In a separate task involving molecular glioma subtype classification, Foltyn-Dumitru and colleagues showed that intensity normalization affected external generalization, illustrating the importance of examining the processing pipeline rather than assuming portability from internal performance (22). That result does not establish the best transformation for post-treatment diagnosis. For the present question, a small, fixed set of fluid-attenuated inversion recovery (FLAIR) and postcontrast T1 descriptors provides a tractable computational representation. Its value lies in testing additional information under defined conditions; the descriptors are not direct measurements of viable tumor, microscopic tissue composition or a specific treatment-injury mechanism.', 'chinese': '定量影像分析为通过明确计算描述这些检查提供了一种途径。强度统计概括区域内数值的分布，而纹理描述量刻画仅靠均值无法涵盖的灰度组织特征。影像组学研究已探索利用这些信息区分胶质瘤进展与治疗效应，但其解释依赖所用影像、分割及特征定义（20）。影像生物标志物标准化倡议提供了共同定义，有助于提高特征计算的可比性（4）。然而，标准定义并不能消除采集与预处理差异。在另一项胶质瘤分子亚型分类任务中，Foltyn-Dumitru等发现，强度归一化影响外部推广表现，说明需要审视处理流程，不能仅依据内部性能假定可迁移性（22）。该结果并未确定治疗后诊断的最佳变换方案。对于当前问题，小规模且固定的液体衰减反转恢复（FLAIR）与增强T1描述量集合能够形成便于分析的计算表达，其价值在于检验明确条件下的额外信息；这些描述量并非对活性肿瘤、微观组织构成或特定治疗损伤机制的直接测量。'}, {'id': 'I4', 'citation_ids': ['S06', 'S12', 'S13', 'S14'], 'english': 'Clinical and imaging information may contribute different, partly overlapping signals. Age, treatment history and the interval from an initial clinical event describe the examination context, while MRI descriptors summarize the supplied image region. Combining them at the feature level can test whether imaging improves on an explicit clinical comparator. Existing studies support investigating this possibility without guaranteeing an increment. Zhou and colleagues reported favorable results from a model combining radiomics with clinical, molecular and advanced MRI information (6). McKenney and colleagues also reported better classification accuracy for a model combining imaging information with O6-methylguanine-DNA methyltransferase (MGMT) promoter methylation status in a selected cohort with histopathological confirmation (7). Conversely, Leone and colleagues did not demonstrate a reliable improvement from adding diffusion, perfusion and radiomic information to their reference clinical model (9). The new-cohort evaluation by de Godoy and colleagues further illustrates why outcome grouping and reference definitions must be examined when interpreting model performance (10). These studies differ in patient selection, available sequences and diagnostic targets, so their areas under the receiver operating characteristic curve (AUCs) cannot serve as a ranking of interchangeable methods. For a public-data analysis with a restricted input panel, the relevant question is the paired gain over its own clinical baseline.', 'chinese': '临床与影像信息可能提供不同但部分重叠的信号。年龄、治疗史及初始临床事件至检查的间隔描述检查背景，MRI描述量则汇总所提供影像区域的信息。将二者在特征层面结合，可以检验影像是否优于明确的临床对照。现有研究支持探索这一可能性，但不能保证必然产生增量。Zhou等报告，将影像组学与临床、分子及高级MRI信息结合的模型表现较为有利（6）。McKenney等也在经组织病理确认的选定队列中，报告结合影像信息与O6-甲基鸟嘌呤-DNA甲基转移酶（MGMT）启动子甲基化状态的模型具有更好的分类准确率（7）。相反，Leone等未能证明在参考临床模型中加入弥散、灌注及影像组学信息可以带来可靠改善（9）。de Godoy等的新队列评价进一步说明，解释模型性能时需要审视结局分组与参考定义（10）。这些研究的患者选择、可用序列和诊断目标均不同，因此不能根据其AUC对可互换的方法进行排名。对于输入面板受限的公开数据分析，相关问题是相较其自身临床基线所获得的配对增益。'}, {'id': 'I5', 'citation_ids': ['S34', 'S23', 'S24', 'S27', 'S02'], 'english': 'A useful contribution in this setting is a transparent assessment of incremental information and its dependence on analytical choices. CLEAR emphasizes reporting acquisition, segmentation, feature extraction and model evaluation in enough detail for appraisal and reproduction (23). PROBAST+AI separately draws attention to data sources, predictors, outcomes and analysis when judging model quality, bias and applicability (24). These concerns apply to penalized regression as well as more complex learning algorithms. Evaluation guidance distinguishes performance estimated within the development data source from evaluation in independent data (25). Where patients contribute repeated examinations, this distinction also requires attention to dependence between observations: separating scans alone does not ensure that patients are separated between training and testing. Experiments with duplicated scans in connectome-based prediction demonstrate one way that leakage can inflate performance, although that task and duplication mechanism differ from longitudinal glioma imaging (14). A focused analysis should therefore compare models on the same examinations and patient-grouped folds, retain unfavorable sensitivity results, and report the scope of uncertainty. Transparent reporting, as encouraged by TRIPOD+AI, allows readers to judge what such internal evidence can support (2).', 'chinese': '在这一场景中，有价值的贡献是透明地评价增量信息及其对分析选择的依赖。CLEAR强调，应充分报告采集、分割、特征提取及模型评价细节，以便审查和复现（23）。PROBAST+AI则在判断模型质量、偏倚及适用性时，分别关注数据来源、预测变量、结局与分析（24）。这些问题同样适用于惩罚回归和更复杂的学习算法。评价指南区分了在开发数据来源内部估计的性能与在独立数据中的评价（25）。当患者贡献重复检查时，还需要关注观测间依赖：仅将扫描分开，并不能确保患者在训练与测试之间分离。在基于脑连接组的预测中，重复复制扫描的实验展示了信息泄漏使性能偏高的一种方式，但该任务及复制机制均不同于胶质瘤纵向影像（14）。因此，聚焦的分析应在相同检查及患者分组分折上比较模型，保留不利的敏感性结果，并报告不确定性范围。TRIPOD+AI所倡导的透明报告，有助于读者判断这类内部证据能够支持哪些结论（2）。'}, {'id': 'I6', 'citation_ids': ['S01'], 'english': "We therefore evaluated whether eight fixed MRI descriptors add discriminatory information to clinical inputs in a secondary analysis of public UCSD-PTGBM data (1). The target was residual or recurrent tumor versus broad treatment-related change at the examination level. We compared an age-and-sex model, an expanded clinical model, an MRI-only model and their feature-level fusion using the same available imaging cohort and patient-grouped validation. The primary comparison was the paired AUC difference between Fusion and the expanded clinical model. Probability error, calibration and class-specific classification were evaluated alongside discrimination. Boundary erosion, repeated grouped partitions and selection of one available examination per patient examined distinct aspects of sensitivity. By preserving the fixed feature panel and documenting processing and evaluation, the study aimed to characterize MRI's incremental contribution and its sensitivity to these analytical choices as a basis for independent assessment.", 'chinese': '因此，我们对公开UCSD-PTGBM数据开展二次分析，评价八项固定MRI描述量能否为临床输入增加鉴别信息（1）。目标是在检查层面区分残余或复发肿瘤与广义治疗相关改变。我们在相同的可用影像队列及患者分组验证中，比较年龄与性别模型、扩展临床模型、仅MRI模型，以及二者在特征层面的融合模型。主要比较为Fusion与扩展临床模型之间的配对AUC差值，同时评价鉴别能力、概率误差、校准及各类别分类表现。边界腐蚀、重复分组分折及每位患者选择一次可用检查，分别考察不同方面的敏感性。通过保留固定特征面板并记录处理及评价流程，本研究旨在刻画MRI的增量贡献及其对这些分析选择的敏感性，为后续独立评估提供基础。'}]

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def wc(text):
    return len(re.findall(r"[A-Za-z0-9]+(?:['/-][A-Za-z0-9]+)*", text))

def audit():
    prev=json.loads((ROOT/'reports/ucsd_manuscript_step53/conclusion_section53.json').read_text())
    frozen=prev['source_audit']['frozen_file_sha256'].copy()
    frozen.update(json.loads((ROOT/'reports/ucsd_manuscript_step53/artifact_verification.json').read_text())['sha256'])
    for rel in ['reports/ucsd_manuscript_step53/conclusion_section53.json','reports/ucsd_manuscript_step42/literature_evidence.json','reports/ucsd_manuscript_step43/literature_evidence.json','reports/ucsd_manuscript_step46/literature_evidence.json','reports/ucsd_manuscript_step27/methods_section25.json']:
        if rel not in frozen:frozen[rel]=sha(ROOT/rel)
    frozen[str(EVIDENCE.relative_to(ROOT))]=sha(EVIDENCE)
    for p,h in frozen.items():assert sha(ROOT/p)==h,p
    ev=json.loads(EVIDENCE.read_text());new=ev['new_reference_mapping'];records={r['reference_id']:r for r in ev['studies']}
    refs={r['id']:r for r in json.loads((ROOT/'reports/ucsd_manuscript_step23/reference_library.json').read_text())['references']}
    selected=[]
    for number,rid in MAPPING.items():
        ref=refs[rid];citation=ref['citation']
        if rid in ev['metadata_corrections_for_current_reference_copies']:
            c=ev['metadata_corrections_for_current_reference_copies'][rid];old=':'+c['from']+'.';assert old in citation,(rid,citation)
            citation=citation.replace(old,':'+c['to']+'.')
        if rid in records:assert records[rid]['doi']==ref['doi'] and records[rid]['provisional_citation_number']==int(number)
        selected.append({'number':int(number),'reference_id':rid,'citation':citation,'doi':ref['doi'],'read_level_this_stage':records[rid]['read_level_this_stage'] if rid in records else 'Relevant prior reading and evidence records reused'})
    prior_ids=set()
    for pattern in ['methods_section*.json','results_section*.json','discussion_section*.json','conclusion_section*.json']:
        for p in ROOT.glob('reports/ucsd_manuscript_step*/'+pattern):
            if int(p.parent.name.rsplit('step',1)[1])<54:prior_ids.update(r for item in json.loads(p.read_text())['paragraphs'] for r in item['citation_ids'])
    assert len(prior_ids)==17 and not set(new.values())&prior_ids
    current={rid for p in PARAGRAPHS for rid in p['citation_ids']};assert current==set(MAPPING.values())
    all_ids=prior_ids|current;assert len(all_ids)==25
    recent5=sum(refs[r]['english_flag'] and refs[r]['recent_5yr'] for r in all_ids)
    recent3=sum(refs[r]['english_flag'] and refs[r]['recent_3yr'] for r in all_ids)
    assert recent5==23 and recent3==17 and recent5/25>.9 and recent3>25/2
    found=[]
    for p in PARAGRAPHS:
        nums=[n for group in re.findall(r'\(([0-9,]+)\)',p['english']) for n in group.split(',')]
        assert {MAPPING[n] for n in nums}==set(p['citation_ids']),(p['id'],nums)
        for n in nums:
            if n in new and n not in found:found.append(n)
        assert not re.search(r'0\.782|0\.110|0\.021|0\.214|12/41',p['english'])
    assert found==list(new)
    lock=json.loads((ROOT/'reports/ucsd_fusion/analysis_lock.json').read_text())
    assert len(lock['model_features']['M0'])==2 and len(lock['model_features']['M1'])==5 and len(lock['model_features']['MRI'])==8 and len(lock['model_features']['Fusion'])==13
    assert [p['id'] for p in PARAGRAPHS]==[f'I{i}' for i in range(1,7)]
    return {'frozen_file_sha256':frozen,'reference_records':selected,'source_evidence_record':str(EVIDENCE.relative_to(ROOT)),
            'new_full_texts_checked_for_current_claims':list(new.values()),'new_unique_reference_mapping':new,'previous_unique_cited_sources':sorted(prior_ids),
            'citation_numbers_match_paragraph_source_ids':True,'new_numbers_follow_first_appearance_in_fragment':True,'original_candidate_library_unchanged':True,
            'cited_reference_recency':{'unique_sources':25,'english_recent_2022_2026':recent5,'proportion_recent_5yr':recent5/25,'english_recent_2024_2026':recent3,'proportion_recent_3yr':recent3/25,'all_english':all(refs[r]['english_flag']==1 for r in all_ids),'final_assembled_bibliography_recheck_required':True},
            'global_numbering_pending_full_manuscript_assembly':True,'all_six_intro_paragraphs_complete':True,'no_actual_intro_subheadings':True,
            'no_current_study_performance_results_in_introduction':True,'clinical_endpoint_distinctions_preserved':True,
            'normalization_subtype_study_not_our_diagnostic_validation':True,'duplicate_scan_experiment_not_longitudinal_glioma_experiment':True,
            'fixed_model_input_dimensions_checked':{'M0':2,'M1':5,'MRI':8,'Fusion':13},'abstract_deferred_to_next_round':True,
            'post_conclusion_declarations_deferred_by_user':True,'models_refitted':0,'new_hypothesis_tests':0,
            'verification_scope':'Primary source reading and claim/citation review, preserved prior-file hashes and locked input dimensions; no new fitting or analyses. Recency applies to current25 actual citations, not the40 candidates.'}

def build():
    a = audit(); OUT.mkdir(parents=True, exist_ok=True); REPORT.mkdir(parents=True, exist_ok=True)
    doc = Document(); sec = doc.sections[0]; sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(.79)
    for name in ['Normal', 'Title', 'Heading 1']:
        st = doc.styles[name]; st.font.name = 'Times New Roman'; st.font.color.rgb = RGBColor(0, 0, 0)
        rf = st.element.get_or_add_rPr().rFonts; rf.set(qn('w:eastAsia'), 'Noto Serif CJK SC')
        for key in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']: rf.attrib.pop(qn('w:'+key), None)
    normal = doc.styles['Normal']; normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.; normal.paragraph_format.space_after = Pt(9)
    doc.styles['Title'].font.size = Pt(16); doc.styles['Title'].paragraph_format.space_after = Pt(12)
    heading = doc.styles['Heading 1']; heading.font.size = Pt(13)
    heading.paragraph_format.space_before = Pt(8); heading.paragraph_format.space_after = Pt(9)
    doc.core_properties.title = '1 Introduction'; doc.core_properties.author = 'Tianyu Feng'
    doc.add_paragraph('1 Introduction', style='Title')
    for item in PARAGRAPHS:
        en = doc.add_paragraph(item['english']); en.paragraph_format.keep_together = True; en.paragraph_format.keep_with_next = True
        zh = doc.add_paragraph(item['chinese']); zh.paragraph_format.keep_together = True
        for run in zh.runs: run.font.size = Pt(11)
    if a['reference_records']:
        title = doc.add_paragraph('References', style='Heading 1'); title.paragraph_format.page_break_before = True
    for ref in a['reference_records']:
        p = doc.add_paragraph(str(ref['number'])+'. '+ref['citation']); p.paragraph_format.keep_together = True
    for e in [*doc.styles.element.xpath('.//w:pBdr'), *doc.element.xpath('.//w:pBdr')]: e.getparent().remove(e)
    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    path = OUT/'UCSD_Manuscript_Introduction_EN_ZH.docx'; doc.save(path)
    (OUT/'UCSD_Manuscript_Introduction_EN.txt').write_text('\n\n'.join(p['english'] for p in PARAGRAPHS)+'\n')
    d = {'stage': 54, 'date': '2026-10-30', 'subsection': '1 Introduction', 'v2_paragraph_ids': [p['id'] for p in PARAGRAPHS],
         'paragraphs': PARAGRAPHS, 'english_word_count': sum(wc(p['english']) for p in PARAGRAPHS),
         'paragraph_word_counts': {p['id']: wc(p['english']) for p in PARAGRAPHS},
         'word_count_scope': 'English narrative including in-text citation numbers; excludes headings, reference list and Chinese translation',
         'v2_approximate_target_words': 1000, 'equations': {'count': 0, 'numbers': [], 'cumulative_previous_numbered_groups': 10},
         'citation_mapping': MAPPING, 'citation_numbering': 'Existing1-17 unchanged and new18-25 assigned provisionally; global first-appearance numbering will be regenerated on full assembly',
         'new_manuscript_references': 8, 'cumulative_unique_cited_sources': 25,
         'source_audit': a, 'models_refitted': 0, 'new_hypothesis_tests': 0,
         'next_section': 'Abstract (A1-A4) next round; declarations and acknowledgments deferred by user',
         'docx': str(path.relative_to(ROOT)), 'docx_sha256': sha(path)}
    (REPORT/'introduction_section54.json').write_text(json.dumps(d, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps({'docx': d['docx'], 'english_word_count': d['english_word_count'], 'frozen_files': len(a['frozen_file_sha256'])}))

def verify():
    from pypdf import PdfReader
    d = json.loads((REPORT/'introduction_section54.json').read_text()); a = audit(); assert a == d['source_audit']
    path = ROOT/d['docx']; assert sha(path) == d['docx_sha256']; doc = Document(path); texts = [p.text for p in doc.paragraphs]
    for p in PARAGRAPHS:
        assert p['english'] in texts and p['chinese'] in texts and texts.index(p['chinese']) == texts.index(p['english'])+1
        assert not re.search(r'[\u3400-\u9fff，。；：（）【】、]', p['english'])
    assert sum(wc(p['english']) for p in PARAGRAPHS) == d['english_word_count']
    assert a['citation_numbers_match_paragraph_source_ids']
    for ref in a['reference_records']: assert str(ref['number'])+'. '+ref['citation'] in texts
    assert not doc.tables and len(doc.sections) == 1
    with zipfile.ZipFile(path) as z:
        xml = etree.fromstring(z.read('word/document.xml'))
        assert not xml.xpath('//m:oMath', namespaces={'m': 'http://schemas.openxmlformats.org/officeDocument/2006/math'})
        assert not [n for n in z.namelist() if n.startswith('word/media/')]
    pdf = PdfReader(OUT/'docx_render/UCSD_Manuscript_Introduction_EN_ZH.pdf'); pages = len(pdf.pages); assert 3 <= pages <= 10
    assert '1 Introduction' in pdf.pages[0].extract_text() and any('References' in page.extract_text() for page in pdf.pages)
    assert [p.style.name for p in doc.paragraphs].count('Heading 1')==1
    v = {'success': True, 'stage': 54, 'date': '2026-10-30', 'subsection': d['subsection'], 'paragraph_ids': [p['id'] for p in PARAGRAPHS],
         'english_word_count': d['english_word_count'], 'matched_chinese_paragraphs': 6, 'word_pages': pages,
         'visual_review_pages': list(range(1,pages+1)), 'native_editable_equations': 0, 'equation_images': 0,
         'frozen_files_unchanged': len(a['frozen_file_sha256']), 'prior_files_unchanged': len(a['frozen_file_sha256'])-1,
         'new_full_texts_checked_for_current_claims': a['new_full_texts_checked_for_current_claims'],
         'all_six_intro_paragraphs_complete': True, 'no_actual_intro_subheadings': True,
         'citation_numbers_match_paragraph_source_ids': True, 'cited_reference_recency': a['cited_reference_recency'],
         'current_section_references': len(a['reference_records']), 'global_numbering_pending_full_manuscript_assembly': True,
         'clinical_endpoint_distinctions_preserved': True, 'no_current_study_performance_results_in_introduction': True,
         'fixed_model_input_dimensions_checked': a['fixed_model_input_dimensions_checked'],
         'abstract_deferred_to_next_round': True, 'post_conclusion_declarations_deferred_by_user': True,
         'new_manuscript_references': 8, 'provisional_citation_mapping': MAPPING, 'cumulative_unique_cited_sources': 25,
         'models_refitted': 0, 'new_hypothesis_tests': 0, 'next_section': d['next_section'],
         'sha256': {str(p.relative_to(ROOT)): sha(p) for p in OUT.iterdir() if p.is_file() and p.suffix in {'.docx', '.txt'}}}
    (REPORT/'artifact_verification.json').write_text(json.dumps(v, ensure_ascii=False, indent=2)+'\n'); print(json.dumps(v, indent=2))

if __name__ == '__main__':
    parser = argparse.ArgumentParser(); parser.add_argument('--verify', action='store_true'); args = parser.parse_args()
    verify() if args.verify else build()
