"""Draft D4 and verify saved computational records without fitting models."""
from pathlib import Path
import argparse, hashlib, json, re, zipfile
import numpy as np
import pandas as pd
from scipy.special import expit
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'outputs/ucsd-manuscript-step44-20261030'
REPORT = ROOT/'reports/ucsd_manuscript_step44'
EVIDENCE = REPORT/'literature_evidence.json'
MAPPING = {'11': 'S33'}
PARAGRAPHS = [{
    'id': 'D4', 'citation_ids': ['S33'],
    'english': 'The computational contribution of this study is an auditable evaluation of low-dimensional feature-level fusion using public clinical and MRI data. Eight explicitly defined descriptors from FLAIR and postcontrast T1 images were concatenated with five clinical inputs and fitted with a fixed L2-regularized logistic model. This design makes the imaging contribution traceable to specified intensity and texture calculations and anchors the comparison to the same clinical input set. Evaluating all models on matched examinations with shared patient-grouped folds reduces the influence of case composition and partition differences on the incremental comparison. Training-fold preprocessing and retained OOF probabilities also make the relationship between inputs, fitted models and evaluation outputs inspectable. The public repository provides notebooks, derived feature tables, fold assignments and analysis locks, allowing model evaluation to be rerun from the saved features. Re-extracting features additionally requires the image files and their documented acquisition or restoration routes. This emphasis on processing details, comparator models, partitioning and code availability is consistent with methodological priorities described in METRICS (11). These contributions position the work as an application and evaluation study in medical image computing. They facilitate scrutiny of the reported increment, while the reference-standard and external-validation limitations still constrain its interpretation.',
    'chinese': '本研究在计算方法方面的贡献，是利用公开临床与MRI数据，对低维特征级融合开展可核查的评价。将来自FLAIR和增强T1图像的8项明确定义的描述特征与5项临床输入拼接，采用固定的L2正则化Logistic模型拟合。这一设计使影像信息的贡献能够追溯至具体的强度和纹理计算，并使比较基于相同的临床输入集合。所有模型在匹配检查和共同的患者分组折上评价，减少了病例构成及分折差异对增量比较的影响。训练折内预处理和保留的OOF概率，也使输入、拟合模型与评价输出之间的关系可以检查。公开仓库提供Notebook、派生特征表、分折分配及分析锁，可利用保存的特征重新运行模型评价。重新提取特征则还需要影像文件，并使用已记录的获取或恢复途径。这种对处理细节、比较模型、数据划分和代码可用性的重视，与METRICS所描述的方法学关注点一致（11）。这些贡献将本研究定位为医学图像计算中的应用与评价研究。它们有助于审视所报告的增量，但参考标准和缺乏外部验证等局限仍限制了对该增量的解释。'
}]

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def wc(text):
    return len(re.findall(r"[A-Za-z0-9]+(?:['/-][A-Za-z0-9]+)*", text))

def audit():
    previous = json.loads((ROOT/'reports/ucsd_manuscript_step43/discussion_section43.json').read_text())
    frozen = previous['source_audit']['frozen_file_sha256'].copy()
    frozen.update(json.loads((ROOT/'reports/ucsd_manuscript_step43/artifact_verification.json').read_text())['sha256'])
    frozen[str(EVIDENCE.relative_to(ROOT))] = sha(EVIDENCE)
    for p, h in frozen.items(): assert sha(ROOT/p) == h, p
    evidence = json.loads(EVIDENCE.read_text()); record = evidence['studies'][0]
    refs = {r['id']: r for r in json.loads((ROOT/'reports/ucsd_manuscript_step23/reference_library.json').read_text())['references']}
    ref = refs['S33']; assert record['reference_id'] == 'S33' and record['doi'] == ref['doi']
    assert record['provisional_citation_number'] == 11 and ref['english_flag'] == ref['recent_5yr'] == ref['recent_3yr'] == 1
    assert ref['eligibility_year'] == record['facts']['publication_year'] == 2024
    assert [r['number'] for r in record['facts']['relevant_items']] == [11, 13, 18, 24, 25, 29]
    main = json.loads((ROOT/'reports/ucsd_fusion/metrics.json').read_text())
    assert len(main['models']['MRI']) == 8 and len(main['models']['M1']) == 5
    assert main['models']['Fusion'] == main['models']['M1']+main['models']['MRI']
    assert main['same_cases_and_folds'] and main['C'] == 1 and not main['hyperparameter_search'] and not main['feature_selection']
    code = (ROOT/'src/rn_public/ucsd_fusion.py').read_text(); assert 'LogisticRegression(C=1,l1_ratio=0' in code
    data = pd.read_csv(ROOT/'data/processed/ucsd_clinical_image_features.csv').set_index('ID')
    oof = pd.read_csv(ROOT/'data/processed/ucsd_fusion_oof.csv').set_index('ID')
    assert len(data) == len(oof) == 182 and set(data.index) == set(oof.index)
    folds = json.loads((ROOT/'reports/ucsd_fusion/fold_audit.json').read_text()); assert len(folds) == 5
    checked = 0; maximum_error = 0.; seen = []; preprocessing_values = 0
    for fold in folds:
        train = data.loc[fold['train_IDs']]; test = data.loc[fold['test_IDs']]; seen.extend(test.index)
        assert set(train.index).isdisjoint(test.index) and set(train.index)|set(test.index) == set(data.index)
        for k in ['patient_id', 'brats_patient_id', 'validation_group']:
            assert set(train[k]).isdisjoint(test[k])
        assert oof.loc[test.index, 'validation_fold'].eq(fold['fold']).all()
        for name in ['M1', 'Fusion']:
            params = fold['models'][name]['parameters']; transformed = {}
            for prefix, part in params['continuous'].items():
                cols = part['columns']; tr = train[cols].to_numpy(float); te = test[cols].to_numpy(float)
                if 'impute_median' in part:
                    median = float(np.nanmedian(tr)); np.testing.assert_allclose(median, part['impute_median'], atol=1e-12, rtol=1e-12)
                    tr = np.where(np.isnan(tr), median, tr); te = np.where(np.isnan(te), median, te); preprocessing_values += 1
                np.testing.assert_allclose(tr.mean(axis=0), part['mean'], atol=1e-12, rtol=1e-12)
                scale = np.std(tr, axis=0); scale = np.where(scale == 0, 1, scale)
                np.testing.assert_allclose(scale, part['scale'], atol=1e-12, rtol=1e-12); preprocessing_values += 2*len(cols)
                te = (te-np.asarray(part['mean']))/np.asarray(part['scale'])
                for i, col in enumerate(cols): transformed[prefix+'__'+col] = te[:, i]
            for col in params['binary']: transformed['binary__'+col] = test[col].to_numpy(float)
            x = np.column_stack([transformed[k] for k in params['transformed_names']])
            assert x.shape[1] == len(main['models'][name])
            p = expit(x@np.asarray(params['coefficients'])+params['intercept'])
            saved = oof.loc[test.index, name+'_oof_probability'].to_numpy()
            error = float(np.max(np.abs(p-saved))); assert error < 1e-10
            maximum_error = max(maximum_error, error); checked += len(test)
    assert len(seen) == len(set(seen)) == 182 and checked == 364 and preprocessing_values == 130
    locks = []
    for name in ['ucsd_fusion', 'ucsd_stability']:
        lock = json.loads((ROOT/f'reports/{name}/analysis_lock.json').read_text())
        for p, h in lock['input_sha256'].items(): assert sha(ROOT/p) == h, p
        assert sha(ROOT/f'src/rn_public/{name}.py') == lock['analysis_code_sha256']
        assert sha(ROOT/f'reports/{name}_protocol.md') == lock['protocol_sha256']; locks.append(name)
    notebooks = []
    for number, name in [(17, '17_ucsd_batch_image_features'), (18, '18_ucsd_clinical_mri_fusion'), (19, '19_ucsd_model_stability')]:
        path = f'notebooks/{name}.ipynb'; n = json.loads((ROOT/path).read_text())
        r = json.loads((ROOT/f'reports/kernel_execution_{number}.json').read_text())
        cells = [c for c in n['cells'] if c['cell_type'] == 'code']
        source = '\n'.join(''.join(c['source']) if isinstance(c['source'], list) else c['source'] for c in cells)
        assert r['success'] and r['mode'] == 'in_process_ipykernel' and r['notebooks'][0]['errors'] == 0
        assert hashlib.sha256(source.encode()).hexdigest() == r['notebooks'][0]['code_sha256']; notebooks.append(path)
    availability = json.loads((ROOT/'reports/ucsd_manuscript_step32/methods_section210.json').read_text())['source_audit']
    assert availability['additional_archive_files_verified_in_saved_restoration_receipt'] == 528
    assert availability['repository_image_sample_files_rehashed'] == 18 and not availability['archive_restored_this_stage']
    prior = []
    for pattern in ['methods_section*.json', 'results_section*.json', 'discussion_section*.json']:
        for p in ROOT.glob('reports/ucsd_manuscript_step*/'+pattern):
            if int(p.parent.name.rsplit('step', 1)[1]) < 44: prior.append(json.loads(p.read_text()))
    prior_ids = {rid for s in prior for p in s['paragraphs'] for rid in p['citation_ids']}
    assert len(prior_ids) == 10 and 'S33' not in prior_ids
    return {'frozen_file_sha256': frozen, 'reference_records': [{'number': 11, 'reference_id': 'S33', 'citation': ref['citation'],
            'doi': ref['doi'], 'eligibility_year': 2024, 'read_level_this_stage': record['read_level_this_stage']}],
            'source_evidence_record': str(EVIDENCE.relative_to(ROOT)), 'full_texts_checked_for_current_claims': ['S33'],
            'METRICS_topics_checked': [11, 13, 18, 24, 25, 29], 'METRICS_score_calculated': False,
            'model_input_counts_checked': {'MRI': 8, 'M1': 5, 'Fusion': 13}, 'folds_checked': 5,
            'saved_fold_model_records_checked': 10, 'training_preprocessing_values_checked': 130,
            'fixed_oof_probabilities_reconstructed': checked, 'maximum_absolute_probability_difference': maximum_error,
            'analysis_locks_checked': locks, 'historical_notebook_execution_receipts_checked': notebooks,
            'notebooks_executed_this_stage': 0, 'images_restored_this_stage': 0,
            'all_raw_images_claimed_public_in_project_repository': False, 'new_algorithm_claimed': False,
            'independent_reproduction_performed': False, 'previous_unique_cited_sources': sorted(prior_ids),
            'models_refitted': 0, 'new_hypothesis_tests': 0,
            'verification_scope': 'Reconstruct saved M1/Fusion OOF probabilities and training preprocessing summaries, check existing locks and kernel receipts, and verify manually reviewed METRICS topics; no fitting or resampling'}

def build():
    a = audit(); OUT.mkdir(parents=True, exist_ok=True); REPORT.mkdir(parents=True, exist_ok=True)
    doc = Document(); sec = doc.sections[0]; sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(.79)
    for name in ['Normal', 'Title', 'Heading 1']:
        st = doc.styles[name]; st.font.name = 'Times New Roman'; st.font.color.rgb = RGBColor(0, 0, 0)
        rf = st.element.get_or_add_rPr().rFonts; rf.set(qn('w:eastAsia'), 'Noto Serif CJK SC')
        for key in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']: rf.attrib.pop(qn('w:'+key), None)
    normal = doc.styles['Normal']; normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.; normal.paragraph_format.space_after = Pt(9)
    doc.styles['Title'].font.size = Pt(16); doc.styles['Title'].paragraph_format.space_after = Pt(12)
    heading = doc.styles['Heading 1']; heading.font.size = Pt(13)
    heading.paragraph_format.space_before = Pt(8); heading.paragraph_format.space_after = Pt(9)
    doc.core_properties.title = 'Discussion 4.4 Computational contribution'; doc.core_properties.author = 'Tianyu Feng'
    doc.add_paragraph('Discussion', style='Title'); doc.add_paragraph('4.4 Computational contribution', style='Heading 1')
    for item in PARAGRAPHS:
        en = doc.add_paragraph(item['english']); en.paragraph_format.keep_together = True; en.paragraph_format.keep_with_next = True
        zh = doc.add_paragraph(item['chinese']); zh.paragraph_format.keep_together = True
        for run in zh.runs: run.font.size = Pt(11)
    doc.add_paragraph('References', style='Heading 1')
    for ref in a['reference_records']:
        p = doc.add_paragraph(str(ref['number'])+'. '+ref['citation']); p.paragraph_format.keep_together = True
    for e in [*doc.styles.element.xpath('.//w:pBdr'), *doc.element.xpath('.//w:pBdr')]: e.getparent().remove(e)
    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    path = OUT/'UCSD_Manuscript_Discussion_4_4_EN_ZH.docx'; doc.save(path)
    (OUT/'UCSD_Manuscript_Discussion_4_4_EN.txt').write_text('\n\n'.join(p['english'] for p in PARAGRAPHS)+'\n')
    d = {'stage': 44, 'date': '2026-10-30', 'subsection': '4.4 Computational contribution', 'v2_paragraph_ids': ['D4'],
         'paragraphs': PARAGRAPHS, 'english_word_count': sum(wc(p['english']) for p in PARAGRAPHS),
         'paragraph_word_counts': {p['id']: wc(p['english']) for p in PARAGRAPHS},
         'word_count_scope': 'English narrative including in-text citation number; excludes headings, reference list and Chinese translation',
         'v2_approximate_target_words': 200, 'equations': {'count': 0, 'numbers': [], 'cumulative_previous_numbered_groups': 10},
         'citation_mapping': MAPPING, 'citation_numbering': 'Provisional first-appearance number 11; renumber globally when Introduction is assembled',
         'new_manuscript_references': 1, 'cumulative_unique_cited_sources': 11,
         'source_audit': a, 'models_refitted': 0, 'new_hypothesis_tests': 0,
         'next_section': 'Discussion 4.5 Interpretation of image features (D5)',
         'docx': str(path.relative_to(ROOT)), 'docx_sha256': sha(path)}
    (REPORT/'discussion_section44.json').write_text(json.dumps(d, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps({'docx': d['docx'], 'english_word_count': d['english_word_count'], 'frozen_files': len(a['frozen_file_sha256'])}))

def verify():
    from pypdf import PdfReader
    d = json.loads((REPORT/'discussion_section44.json').read_text()); a = audit(); assert a == d['source_audit']
    path = ROOT/d['docx']; assert sha(path) == d['docx_sha256']; doc = Document(path); texts = [p.text for p in doc.paragraphs]
    for p in PARAGRAPHS:
        assert p['english'] in texts and p['chinese'] in texts and texts.index(p['chinese']) == texts.index(p['english'])+1
        assert not re.search(r'[\u3400-\u9fff，。；：（）【】、]', p['english'])
    assert sum(wc(p['english']) for p in PARAGRAPHS) == d['english_word_count']
    assert re.findall(r'\((\d+)\)', PARAGRAPHS[0]['english']) == list(MAPPING)
    for ref in a['reference_records']: assert str(ref['number'])+'. '+ref['citation'] in texts
    assert not doc.tables and len(doc.sections) == 1
    with zipfile.ZipFile(path) as z:
        xml = etree.fromstring(z.read('word/document.xml'))
        assert not xml.xpath('//m:oMath', namespaces={'m': 'http://schemas.openxmlformats.org/officeDocument/2006/math'})
        assert not [n for n in z.namelist() if n.startswith('word/media/')]
    pdf = PdfReader(OUT/'docx_render/UCSD_Manuscript_Discussion_4_4_EN_ZH.pdf'); pages = len(pdf.pages); assert pages == 1
    v = {'success': True, 'stage': 44, 'date': '2026-10-30', 'subsection': d['subsection'], 'paragraph_ids': ['D4'],
         'english_word_count': d['english_word_count'], 'matched_chinese_paragraphs': 1, 'word_pages': pages,
         'visual_review_pages': [1], 'native_editable_equations': 0, 'equation_images': 0,
         'frozen_files_unchanged': len(a['frozen_file_sha256']), 'prior_files_unchanged': len(a['frozen_file_sha256'])-1,
         'full_texts_checked_for_current_claims': ['S33'], 'new_manuscript_references': 1,
         'provisional_citation_mapping': MAPPING, 'cumulative_unique_cited_sources': 11,
         'model_input_counts_checked': 3, 'folds_checked': 5, 'saved_fold_model_records_checked': 10,
         'training_preprocessing_values_checked': 130, 'fixed_oof_probabilities_reconstructed': 364,
         'maximum_absolute_probability_difference': a['maximum_absolute_probability_difference'],
         'analysis_locks_checked': 2, 'historical_notebook_execution_receipts_checked': 3,
         'METRICS_topics_checked': 6, 'METRICS_score_calculated': False, 'models_refitted': 0, 'new_hypothesis_tests': 0,
         'next_section': d['next_section'],
         'sha256': {str(p.relative_to(ROOT)): sha(p) for p in OUT.iterdir() if p.is_file() and p.suffix in {'.docx', '.txt'}}}
    (REPORT/'artifact_verification.json').write_text(json.dumps(v, ensure_ascii=False, indent=2)+'\n'); print(json.dumps(v, indent=2))

if __name__ == '__main__':
    parser = argparse.ArgumentParser(); parser.add_argument('--verify', action='store_true'); args = parser.parse_args()
    verify() if args.verify else build()
