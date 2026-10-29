"""Draft D3 from checked primary studies and frozen own-study results."""
from pathlib import Path
import argparse, hashlib, json, re, zipfile
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'outputs/ucsd-manuscript-step43-20261030'
REPORT = ROOT/'reports/ucsd_manuscript_step43'
EVIDENCE = REPORT/'literature_evidence.json'
MAPPING = {'9': 'S13', '10': 'S14'}
PARAGRAPHS = [{
    'id': 'D3', 'citation_ids': ['S13', 'S14'],
    'english': 'Published evidence does not show a uniform benefit from adding imaging information. Leone et al. evaluated 105 patients with IDH-wildtype glioblastoma, including 20 with pseudoprogression, using five repetitions of five-fold cross-validation (9). A clinical-MGMT reference model achieved an AUC of 0.69; adding median diffusion and perfusion measures produced 0.71 without a statistically significant improvement. Models incorporating radiomic features yielded AUCs of 0.52 and 0.54. Their baseline included MGMT status and extent of resection, which were absent from our M1, and outcomes largely relied on subsequent imaging and clinical evolution. Consequently, the differing results do not isolate an effect of model complexity or establish why fusion did or did not help. A separate study by de Godoy et al. applied a previously developed diffusion-perfusion model to 56 new patients and reported an AUC of 0.757 (10). At their 0.5 probability threshold, it correctly identified 18/19 pseudoprogression cases and 21/37 cases in the progression-plus-mixed-tumor group. However, mixed tumors were grouped with pseudoprogression during development and with progression during validation, changing the classification task. This was evaluation in a new cohort, but it did not test our features or model. These studies show why discrimination, class-specific recognition and endpoint alignment require separate scrutiny. For our analysis, the uncertain single-scan increment and limited treatment-change recognition remain material findings, despite the favorable primary AUC difference. Further evaluation should preserve the intended reference categories and assess a locked model in an independent cohort.',
    'chinese': '已发表证据并未显示加入影像信息能够一致带来获益。Leone等采用5次重复的5折交叉验证，评估了105名IDH野生型胶质母细胞瘤患者，其中20名为假性进展（9）。临床-MGMT参考模型的AUC为0.69；加入扩散和灌注指标的中位数后为0.71，但改善未达到统计学显著性。纳入影像组学特征的模型AUC为0.52和0.54。其基线模型包含我们的M1未纳入的MGMT状态和切除范围，结局判断主要依赖后续影像及临床演变。因此，这些不同结果不能单独识别模型复杂度的影响，也不能确定融合产生或未产生获益的原因。另一项由de Godoy等开展的研究将既有扩散-灌注模型应用于56名新患者，报告AUC为0.757（10）。在其0.5概率阈值下，模型正确识别18/19例假性进展，以及进展及混合肿瘤组中的21/37例。然而，混合肿瘤在开发阶段归入假性进展组，在验证阶段归入进展组，因而分类任务发生了变化。这是在新队列中的评价，但并未检验我们的特征或模型。这些研究说明，鉴别能力、各类别的正确识别及终点一致性需要分别审视。对于本研究，尽管主要AUC差值有利，单扫描增量的不确定性和治疗相关改变识别不足仍是重要发现。后续评价应保持预期参考类别的一致，并在独立队列中评估锁定的模型。'
}]

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def wc(text):
    return len(re.findall(r"[A-Za-z0-9]+(?:['/-][A-Za-z0-9]+)*", text))

def audit():
    previous = json.loads((ROOT/'reports/ucsd_manuscript_step42/discussion_section42.json').read_text())
    frozen = previous['source_audit']['frozen_file_sha256'].copy()
    frozen.update(json.loads((ROOT/'reports/ucsd_manuscript_step42/artifact_verification.json').read_text())['sha256'])
    frozen[str(EVIDENCE.relative_to(ROOT))] = sha(EVIDENCE)
    for p, h in frozen.items(): assert sha(ROOT/p) == h, p
    evidence = json.loads(EVIDENCE.read_text()); studies = {r['reference_id']: r for r in evidence['studies']}
    assert set(studies) == set(MAPPING.values()) and evidence['cumulative_unique_cited_sources'] == 10
    refs = {r['id']: r for r in json.loads((ROOT/'reports/ucsd_manuscript_step23/reference_library.json').read_text())['references']}
    selected = []
    for number, rid in MAPPING.items():
        ref = refs[rid]; record = studies[rid]
        assert ref['doi'] == record['doi'] and record['provisional_citation_number'] == int(number)
        assert ref['english_flag'] == ref['recent_5yr'] == 1 and ref['eligibility_year'] == 2023
        selected.append({'number': int(number), 'reference_id': rid, 'citation': ref['citation'], 'doi': ref['doi'],
                         'eligibility_year': ref['eligibility_year'], 'read_level_this_stage': record['read_level_this_stage']})
    l = studies['S13']['facts']; g = studies['S14']['facts']
    assert (l['patients_included'], l['pseudoprogression_cases'], l['validation_repeats'], l['validation_folds']) == (105, 20, 5, 5)
    assert [l[k] for k in ['clinical_MGMT_auc', 'clinical_MGMT_plus_median_ADC_nrCBV_auc',
            'clinical_MGMT_maps_advanced_radiomics_auc', 'clinical_MGMT_maps_all_radiomics_auc']] == [.69, .71, .52, .54]
    assert l['reported_AUC_comparison_p_value'] == .416
    assert (g['new_cohort_patients'], g['new_cohort_auc_figure7'], g['fixed_probability_threshold']) == (56, .757, .5)
    assert [g[k] for k in ['pseudoprogression_correct', 'pseudoprogression_cases',
            'progression_plus_mixed_tumor_correct', 'progression_plus_mixed_tumor_cases']] == [18, 19, 21, 37]
    assert g['total_correct'] == 39 and g['development_grouping'] != g['validation_grouping']
    main = json.loads((ROOT/'reports/ucsd_fusion/metrics.json').read_text())
    single = json.loads((ROOT/'reports/ucsd_stability/metrics.json').read_text())['single_scan']
    ci = single['bootstrap_95ci']['Fusion_minus_M1__roc_auc']; assert ci['lower'] < 0 < ci['upper']
    pair = next(r for r in main['paired_comparisons'] if r['comparison'] == 'Fusion_minus_M1' and r['metric'] == 'roc_auc')
    assert pair['lower'] > 0 and pair['difference'] > 0
    assert main['metrics']['Fusion']['tn'] == 12 and single['metrics']['Fusion']['tn'] == 2
    assert not any('mgmt' in p.lower() or 'resection' in p.lower() for p in main['models']['M1'])
    assert not main['independent_external_validation']
    prior = []
    for pattern in ['methods_section*.json', 'results_section*.json', 'discussion_section*.json']:
        for p in ROOT.glob('reports/ucsd_manuscript_step*/'+pattern):
            if int(p.parent.name.rsplit('step', 1)[1]) < 43: prior.append(json.loads(p.read_text()))
    prior_ids = {rid for s in prior for p in s['paragraphs'] for rid in p['citation_ids']}
    assert len(prior_ids) == 8 and not prior_ids.intersection(MAPPING.values())
    return {'frozen_file_sha256': frozen, 'reference_records': selected,
            'source_evidence_record': str(EVIDENCE.relative_to(ROOT)), 'full_texts_checked_for_current_claims': ['S13', 'S14'],
            'published_numeric_values_checked': 13, 'published_resampling_design_counts_checked': 2,
            'new_cohort_AUC_not_misrepresented_as_classification_accuracy': True,
            'mixed_tumor_grouping_change_disclosed': True, 'new_cohort_study_is_validation_of_our_model': False,
            'causal_explanation_of_cross_study_differences_established': False, 'cross_paper_AUC_ranking_performed': False,
            'own_study_positive_primary_and_uncertain_single_scan_intervals_verified': True,
            'previous_unique_cited_sources': sorted(prior_ids), 'models_refitted': 0, 'new_hypothesis_tests': 0,
            'verification_scope': 'Validate manually reviewed primary-source claims and frozen own-study limits; no article reanalysis, feature extraction or model fitting'}

def build():
    a = audit(); OUT.mkdir(parents=True, exist_ok=True); REPORT.mkdir(parents=True, exist_ok=True)
    doc = Document(); sec = doc.sections[0]; sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(.79)
    for name in ['Normal', 'Title', 'Heading 1']:
        st = doc.styles[name]; st.font.name = 'Times New Roman'; st.font.color.rgb = RGBColor(0, 0, 0)
        rf = st.element.get_or_add_rPr().rFonts; rf.set(qn('w:eastAsia'), 'Noto Serif CJK SC')
        for key in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']: rf.attrib.pop(qn('w:'+key), None)
    normal = doc.styles['Normal']; normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.; normal.paragraph_format.space_after = Pt(9)
    doc.styles['Title'].font.size = Pt(16); doc.styles['Title'].paragraph_format.space_after = Pt(12)
    heading = doc.styles['Heading 1']; heading.font.size = Pt(13)
    heading.paragraph_format.space_before = Pt(8); heading.paragraph_format.space_after = Pt(9)
    doc.core_properties.title = 'Discussion 4.3 Conflicting evidence'; doc.core_properties.author = 'Tianyu Feng'
    doc.add_paragraph('Discussion', style='Title'); doc.add_paragraph('4.3 Conflicting evidence', style='Heading 1')
    for item in PARAGRAPHS:
        en = doc.add_paragraph(item['english']); en.paragraph_format.keep_together = True; en.paragraph_format.keep_with_next = True
        zh = doc.add_paragraph(item['chinese']); zh.paragraph_format.keep_together = True
        for run in zh.runs: run.font.size = Pt(11)
    title = doc.add_paragraph('References', style='Heading 1'); title.paragraph_format.page_break_before = True
    for ref in a['reference_records']:
        p = doc.add_paragraph(str(ref['number'])+'. '+ref['citation']); p.paragraph_format.keep_together = True
    for e in [*doc.styles.element.xpath('.//w:pBdr'), *doc.element.xpath('.//w:pBdr')]: e.getparent().remove(e)
    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    path = OUT/'UCSD_Manuscript_Discussion_4_3_EN_ZH.docx'; doc.save(path)
    (OUT/'UCSD_Manuscript_Discussion_4_3_EN.txt').write_text('\n\n'.join(p['english'] for p in PARAGRAPHS)+'\n')
    d = {'stage': 43, 'date': '2026-10-30', 'subsection': '4.3 Conflicting evidence', 'v2_paragraph_ids': ['D3'],
         'paragraphs': PARAGRAPHS, 'english_word_count': sum(wc(p['english']) for p in PARAGRAPHS),
         'paragraph_word_counts': {p['id']: wc(p['english']) for p in PARAGRAPHS},
         'word_count_scope': 'English narrative including in-text citation numbers; excludes headings, reference list and Chinese translation',
         'v2_approximate_target_words': 250, 'equations': {'count': 0, 'numbers': [], 'cumulative_previous_numbered_groups': 10},
         'citation_mapping': MAPPING, 'citation_numbering': 'Provisional first-appearance numbers 9-10; renumber globally when Introduction is assembled',
         'new_manuscript_references': 2, 'cumulative_unique_cited_sources': 10,
         'source_audit': a, 'models_refitted': 0, 'new_hypothesis_tests': 0,
         'next_section': 'Discussion 4.4 Computational contribution (D4)',
         'docx': str(path.relative_to(ROOT)), 'docx_sha256': sha(path)}
    (REPORT/'discussion_section43.json').write_text(json.dumps(d, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps({'docx': d['docx'], 'english_word_count': d['english_word_count'], 'frozen_files': len(a['frozen_file_sha256'])}))

def verify():
    from pypdf import PdfReader
    d = json.loads((REPORT/'discussion_section43.json').read_text()); a = audit(); assert a == d['source_audit']
    path = ROOT/d['docx']; assert sha(path) == d['docx_sha256']
    doc = Document(path); texts = [p.text for p in doc.paragraphs]
    for p in PARAGRAPHS:
        assert p['english'] in texts and p['chinese'] in texts and texts.index(p['chinese']) == texts.index(p['english'])+1
        assert not re.search(r'[\u3400-\u9fff，。；：（）【】、]', p['english'])
    assert sum(wc(p['english']) for p in PARAGRAPHS) == d['english_word_count']
    assert re.findall(r'\((\d+)\)', PARAGRAPHS[0]['english']) == list(MAPPING)
    for ref in a['reference_records']: assert str(ref['number'])+'. '+ref['citation'] in texts
    assert len(doc.tables) == 0 and len(doc.sections) == 1
    with zipfile.ZipFile(path) as z:
        xml = etree.fromstring(z.read('word/document.xml'))
        assert not xml.xpath('//m:oMath', namespaces={'m': 'http://schemas.openxmlformats.org/officeDocument/2006/math'})
        assert not [n for n in z.namelist() if n.startswith('word/media/')]
    pdf = PdfReader(OUT/'docx_render/UCSD_Manuscript_Discussion_4_3_EN_ZH.pdf'); pages = len(pdf.pages); assert pages == 2
    assert '4.3 Conflicting' in pdf.pages[0].extract_text() and 'References' in pdf.pages[1].extract_text()
    v = {'success': True, 'stage': 43, 'date': '2026-10-30', 'subsection': d['subsection'], 'paragraph_ids': ['D3'],
         'english_word_count': d['english_word_count'], 'matched_chinese_paragraphs': 1, 'word_pages': pages,
         'visual_review_pages': [1, 2], 'native_editable_equations': 0, 'equation_images': 0,
         'frozen_files_unchanged': len(a['frozen_file_sha256']), 'prior_files_unchanged': len(a['frozen_file_sha256'])-1,
         'full_texts_checked_for_current_claims': ['S13', 'S14'], 'published_numeric_values_checked': 13,
         'published_resampling_design_counts_checked': 2, 'new_manuscript_references': 2,
         'provisional_citation_mapping': MAPPING, 'cumulative_unique_cited_sources': 10,
         'mixed_tumor_grouping_change_disclosed': True, 'new_cohort_AUC_not_misrepresented_as_classification_accuracy': True,
         'cross_paper_AUC_ranking_performed': False, 'models_refitted': 0, 'new_hypothesis_tests': 0,
         'next_section': d['next_section'],
         'sha256': {str(p.relative_to(ROOT)): sha(p) for p in OUT.iterdir() if p.is_file() and p.suffix in {'.docx', '.txt'}}}
    (REPORT/'artifact_verification.json').write_text(json.dumps(v, ensure_ascii=False, indent=2)+'\n'); print(json.dumps(v, indent=2))

if __name__ == '__main__':
    parser = argparse.ArgumentParser(); parser.add_argument('--verify', action='store_true'); args = parser.parse_args()
    verify() if args.verify else build()
