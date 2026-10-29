"""Check literature counts, outline coverage, document, and frozen result anchors."""
import hashlib
import json
import re
import subprocess
from pathlib import Path
from docx import Document
import fitz

ROOT=Path(__file__).resolve().parents[1]
REPORT=ROOT/'reports/ucsd_writing_preparation'


def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    plan=json.loads((REPORT/'writing_outline.json').read_text())
    counts=json.loads((REPORT/'literature_structure_counts.json').read_text())
    receipt=json.loads((REPORT/'generation_receipt.json').read_text())
    path=ROOT/receipt['deliverable'];doc=Document(path)
    text='\n'.join(p.text for p in doc.paragraphs)
    assert sha(path)==receipt['sha256']
    expected={'abstract':4,'introduction':4,'methods':12,'results':6,'discussion':8,'conclusion':1}
    ids=[]
    for section,n in expected.items():
        rows=plan['paragraphs'][section];assert len(rows)==n
        for r in rows:
            ident=r[0];ids.append(ident)
            assert len([p for p in doc.paragraphs if re.match(r'^'+ident+r'\s',p.text)])==1,ident
    assert len(ids)==len(set(ids))==35
    assert sum(r[1] for k,v in plan['paragraphs'].items() if k!='abstract' for r in v)==4180
    assert sum(r[1] for r in plan['paragraphs']['abstract'])==250
    assert len(plan['references'])==8 and len(doc.tables)==3
    for r in plan['references']:assert r['doi'] in text
    countrows=doc.tables[0].rows
    for i,p in enumerate(counts['papers'],1):
        c=p['section_words'];assert sum(c.values())==p['main_prose_words']
        assert p['abstract_words']+p['main_prose_words']==p['abstract_plus_main_words']
        assert [c.text for c in countrows[i].cells]==list(map(str,[p['key'],p['abstract_words'],c['Introduction'],c['Methods'],c['Results'],c['Discussion'],c['Conclusion'] or 'in D',p['main_prose_words'],p['abstract_plus_main_words']]))
    # Reconcile quoted planning anchors to already independently verified Stage18-20 results.
    package=json.loads((ROOT/'reports/ucsd_paper_materials/table_package.json').read_text())
    specs={t['key']:t for t in package['tables']}
    main_metrics=json.loads((ROOT/'reports/ucsd_fusion/metrics.json').read_text())['metrics']
    for model in ['M1','MRI','Fusion']:assert f"{main_metrics[model]['roc_auc']:.3f}" in text
    single=json.loads((ROOT/'reports/ucsd_stability/metrics.json').read_text())['single_scan']['metrics']
    assert f"{single['Fusion']['roc_auc']:.3f}" in text
    for key in ['Table3','TableS7']:
        r=specs[key]['rows'][0]
        for x in r[1:]:assert f'{x:.3f}' in text
    prior=json.loads((ROOT/'reports/ucsd_paper_materials/generation_receipt.json').read_text())
    for rel,h in prior['input_sha256'].items():assert sha(ROOT/rel)==h,rel
    for rel,h in prior['output_sha256'].items():assert sha(ROOT/rel)==h,rel
    assert '12/41' in text and '2/22' in text and '36条间隔缺失' in text
    assert not doc.styles.element.xpath('.//w:pBdr') and not doc.element.xpath('.//w:pBdr')
    folder=path.parent/'docx_render_verified';pdf=fitz.open(folder/(path.stem+'.pdf'))
    assert len(pdf)==16
    for i,p in enumerate(pdf,1):
        assert len(p.get_text())>350
        assert (folder/f'page-{i}.png').exists()
    changed=subprocess.check_output(['git','diff','ba3ef15895d1d400d32f4482aa5e37284510a003','--name-only'],cwd=ROOT,text=True).splitlines()
    allowed={'README.md','.gitignore','reports/project_status.json'}
    for name in changed:
        existed=subprocess.run(['git','cat-file','-e','ba3ef15895d1d400d32f4482aa5e37284510a003:'+name],cwd=ROOT,capture_output=True).returncode==0
        assert not existed or name in allowed,name
    result={'success':True,'word_pages':16,'visual_review_pages':16,'main_paragraphs':31,'abstract_sections':4,
        'literature_word_count_rows_checked':4,'document_tables_checked':3,'references_checked':8,
        'main_word_target':4180,'abstract_word_target':250,'prior_input_and_output_hashes_preserved':True,
        'model_refits':0,'new_hypothesis_tests':0,'manuscript_body_drafted':False,'awaiting_user_template':True,
        'docx_sha256':sha(path),'scope':'literature learning and writing outline only'}
    (REPORT/'artifact_verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(result,ensure_ascii=False,indent=2))


if __name__=='__main__':main()
