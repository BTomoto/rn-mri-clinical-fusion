"""Draft R2 and editable English clinical tables from unchanged numeric records."""
from pathlib import Path
import argparse,hashlib,json,re,sys,zipfile
import numpy as np
import pandas as pd
from lxml import etree
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'outputs/ucsd-manuscript-step34-20261029'
REPORT=ROOT/'reports/ucsd_manuscript_step34'
PARAGRAPHS=[
 {'id':'R2','citation_ids':[],
  'english': 'At the examination level, median age in the imaging cohort was 56.5 years [49.0, 66.0], and male patients contributed 132 examinations (72.5%; Table 1). Age, sex and other treatment before imaging were complete. Among 146 examinations with a known interval since the initial event, the median was 154.0 days [40.5, 359.5]. The corresponding medians were 96.0 days for tumor examinations and 346.0 days for treatment-related-change examinations. Intervals were missing for 36 examinations (19.8%), comprising 33 tumor and three treatment-related-change examinations. Other treatment before imaging was recorded for 25 examinations (13.7%). In the group without imaging coverage, median age was 53.5 years [45.0, 62.0], 60.7% of examinations were from male patients and 82.1% carried tumor labels, compared with 77.5% in the imaging cohort (Supplementary Table S1). These summaries used examination records, with interval distributions restricted to observed values before imputation. Bracketed values indicate the first and third quartiles. No between-group significance tests were performed.',
  'chinese': '按检查记录统计，影像队列年龄中位数为56.5岁［49.0，66.0］，男性患者贡献132次检查（72.5%；Table 1）。年龄、性别和扫描前其他治疗信息均完整。在距初始事件间隔已知的146次检查中，间隔中位数为154.0天［40.5，359.5］；肿瘤检查和治疗相关改变检查的相应中位数分别为96.0天和346.0天。36次检查的间隔缺失（19.8%），其中33次为肿瘤检查、3次为治疗相关改变检查。25次检查记录了扫描前其他治疗（13.7%）。在没有影像覆盖的组中，年龄中位数为53.5岁［45.0，62.0］，男性患者贡献的检查占60.7%，肿瘤标签占82.1%，影像队列中的肿瘤标签占比为77.5%（Supplementary Table S1）。上述汇总以检查记录为单位，间隔分布仅使用填补前的已知值。方括号内为第一和第三四分位数。未进行组间显著性检验。'}
]
ROW_LABELS=['Examinations, n','Unique source patients, n','Age, years','Male sex at birth',
  'Tumor-labeled examinations','Treatment-related-change examinations','Days since initial event (known only)',
  'Known interval','Missing interval','Other treatment before imaging']
COMMON_NOTE='Values are counts or n/N (%), except continuous variables, which are median [Q1, Q3]. Percentages use the examination count in each column. Q1 and Q3 denote the first and third quartiles.'
TABLES=[
 {'key':'Table1','title':'Table 1. Clinical characteristics of the imaging cohort',
  'headers':['Characteristic','All imaging examinations','Tumor examinations','Treatment-related change'],
  'notes':[COMMON_NOTE,'Source-patient counts are deduplicated within each column; eight patients contributed examinations to both diagnostic categories. The initial event denotes the recorded initial surgery, treatment or diagnosis, not a uniform radiotherapy date. Interval summaries use known values without imputation.']},
 {'key':'TableS1','title':'Table S1. Clinical characteristics by imaging coverage',
  'headers':['Characteristic','Retained clinical cohort','Imaging coverage','No imaging coverage'],
  'notes':[COMMON_NOTE,'Coverage refers to the accessed image mirror. The 135 patients with coverage and 41 without coverage do not overlap in this cohort. The initial event denotes recorded surgery, treatment or diagnosis; interval summaries use known values without imputation. No between-group significance tests were performed.']}
]

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def wc(s):return len(re.findall(r"[A-Za-z0-9]+(?:['/-][A-Za-z0-9]+)*",s))
def close(a,b):assert np.isclose(a,b,rtol=1e-12,atol=1e-12),(a,b)

def audit():
    old=json.loads((ROOT/'reports/ucsd_manuscript_step33/results_section31.json').read_text());frozen=old['source_audit']['frozen_file_sha256'].copy()
    frozen.update(json.loads((ROOT/'reports/ucsd_manuscript_step33/artifact_verification.json').read_text())['sha256'])
    for p,h in frozen.items():assert sha(ROOT/p)==h,p
    paths=['data/processed/ucsd_clinical_revalidated_cohort.csv','data/processed/ucsd_clinical_image_features.csv',
      'outputs/ucsd-paper-materials-20261029/tables/Table1.csv','outputs/ucsd-paper-materials-20261029/tables/TableS1.csv',
      'outputs/ucsd-paper-materials-20261029/tables/baseline_numeric_details.csv']
    frozen.update({p:sha(ROOT/p) for p in paths})
    clinical=pd.read_csv(ROOT/paths[0]);image=pd.read_csv(ROOT/paths[1]);uncovered=clinical[~clinical.ID.isin(image.ID)]
    groups={'Table1':[image,image[image.tumor_label.eq(1)],image[image.tumor_label.eq(0)]],
            'TableS1':[clinical,image,uncovered]}
    numeric=pd.read_csv(ROOT/paths[4]);assert len(numeric)==60
    labels=['scans','patients','age_years','sex_male','tumor_label','treatment_change','days_since_initial_event','interval_observed','initial_event_missing','other_treatment_prior']
    payload=[];cells_checked=0;numeric_rows=0
    for spec in TABLES:
        original=pd.read_csv(ROOT/f"outputs/ucsd-paper-materials-20261029/tables/{spec['key']}.csv",dtype=str)
        assert original.shape==(10,4);rows=[]
        for ri,var in enumerate(labels):
            out=[ROW_LABELS[ri]]
            for ci,frame in enumerate(groups[spec['key']],1):
                n=len(frame);raw=numeric[(numeric.table_id==spec['key'])&(numeric.cohort==original.columns[ci])&(numeric.variable==var)]
                assert len(raw)==1;raw=raw.iloc[0];assert raw.scan_denominator==n and raw.patient_count==frame.patient_id.nunique()
                if var in ['scans','patients']:
                    value=n if var=='scans' else int(frame.patient_id.nunique());formatted=str(value);close(raw['count'],value)
                elif var in ['age_years','days_since_initial_event']:
                    values=np.sort(frame[var].dropna().to_numpy());q1,median,q3=np.quantile(values,[.25,.5,.75],method='linear')
                    assert raw.observed_n==len(values)
                    for key,v in [('median',median),('q25',q1),('q75',q3)]:close(raw[key],v)
                    formatted=f'{median:.1f} [{q1:.1f}, {q3:.1f}]'
                else:
                    values=(1-frame.tumor_label) if var=='treatment_change' else frame.initial_event_missing.eq(0) if var=='interval_observed' else frame[var]
                    count=int(values.sum());percent=100*count/n;close(raw['count'],count);close(raw.percent,percent);assert raw.observed_n==n
                    formatted=f'{count}/{n} ({percent:.1f}%)'
                assert formatted==original.iloc[ri,ci],(spec['key'],var,ci,formatted,original.iloc[ri,ci])
                out.append(formatted);cells_checked+=1;numeric_rows+=1
            rows.append(out)
        payload.append(dict(**spec,rows=rows))
    assert cells_checked==numeric_rows==60
    assert image[['age_years','sex_male','other_treatment_prior']].notna().all().all()
    assert image.initial_event_missing.sum()==36 and image.days_since_initial_event.notna().sum()==146
    assert image.loc[image.tumor_label.eq(1),'initial_event_missing'].sum()==33 and image.loc[image.tumor_label.eq(0),'initial_event_missing'].sum()==3
    missing=image[image.initial_event_missing.eq(1)].initial_event_offset_raw.value_counts().to_dict();assert missing=={'Unknown':35,'0':1}
    tumor=set(image.loc[image.tumor_label.eq(1),'patient_id']);tc=set(image.loc[image.tumor_label.eq(0),'patient_id']);assert len(tumor&tc)==8
    assert not set(image.patient_id)&set(uncovered.patient_id)
    for table in payload:
        alltext=' '.join(table['headers']+sum(table['rows'],[])+table['notes']+[table['title']]);assert not re.search(r'[\u3400-\u9fff]',alltext)
    return {'frozen_file_sha256':frozen,'original_table_numeric_cells_checked':cells_checked,'saved_numeric_detail_rows_checked':numeric_rows,
      'translated_tables':payload,'numeric_values_changed':0,'interval_missing_original_tokens':missing,
      'interval_missing_counts':{'all':36,'tumor':33,'treatment_change':3},'interval_known_count':146,
      'age_sex_other_treatment_complete':True,'diagnostic_patient_overlap':8,'coverage_patient_overlap':0,
      'models_refitted':0,'features_reextracted':0,'new_hypothesis_tests':0,
      'verification_scope':'Independently reconstruct the 60 original formatted table cells from fixed examination rows and compare unrounded statistics with saved numeric details; translate labels and notes without changing values'}

def add_table(doc,spec):
    title=doc.add_paragraph(spec['title'],style='Heading 1');title.paragraph_format.page_break_before=True
    title.paragraph_format.keep_with_next=True
    table=doc.add_table(rows=1,cols=4);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    widths=[2.62,1.43,1.43,1.44]
    for col,width in zip(table.columns,widths):col.width=Inches(width)
    for cell,width in zip(table.rows[0].cells,widths):cell.width=Inches(width)
    for cell,label in zip(table.rows[0].cells,spec['headers']):cell.text=label
    for values in spec['rows']:
        for cell,value,width in zip(table.add_row().cells,values,widths):cell.text=value;cell.width=Inches(width)
    props=table._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side)
        for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:e.set(qn('w:'+k),v)
        borders.append(e)
    props.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for side in ['top','left','bottom','right']:
        e=OxmlElement('w:'+side);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
    props.append(margins)
    for ri,row in enumerate(table.rows):
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if ri==0:trpr.append(OxmlElement('w:tblHeader'))
        for ci,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7EBEF');cell._tc.get_or_add_tcPr().append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:run.font.size=Pt(10.5);run.font.bold=ri==0
    for note in spec['notes']:
        p=doc.add_paragraph(note);p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(4)
        p.paragraph_format.keep_together=True
        for r in p.runs:r.font.size=Pt(10)

def build():
    OUT.mkdir(parents=True,exist_ok=True);REPORT.mkdir(parents=True,exist_ok=True);a=audit()
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(.79)
    for name in ['Normal','Title','Heading 1']:
        st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        rf=st.element.get_or_add_rPr().rFonts;rf.set(qn('w:eastAsia'),'Noto Serif CJK SC')
        for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:rf.attrib.pop(qn('w:'+key),None)
    normal=doc.styles['Normal'];normal.font.size=Pt(12);normal.paragraph_format.line_spacing=1.;normal.paragraph_format.space_after=Pt(9)
    doc.styles['Title'].font.size=Pt(16);doc.styles['Title'].paragraph_format.space_after=Pt(12)
    doc.styles['Heading 1'].font.size=Pt(13);doc.styles['Heading 1'].paragraph_format.space_before=Pt(8);doc.styles['Heading 1'].paragraph_format.space_after=Pt(9)
    doc.core_properties.title='Results 3.2 Clinical characteristics';doc.core_properties.author='Tianyu Feng'
    doc.add_paragraph('Results',style='Title');doc.add_paragraph('3.2 Clinical characteristics',style='Heading 1')
    for item in PARAGRAPHS:
        en=doc.add_paragraph(item['english']);en.paragraph_format.page_break_before=False
        en.paragraph_format.keep_together=True;en.paragraph_format.keep_with_next=True
        zh=doc.add_paragraph(item['chinese']);zh.paragraph_format.keep_together=True
        for r in zh.runs:r.font.size=Pt(11)
    for spec in a['translated_tables']:
        add_table(doc,spec)
        pd.DataFrame(spec['rows'],columns=spec['headers']).to_csv(OUT/(spec['key']+'_EN.csv'),index=False)
    for e in [*doc.styles.element.xpath('.//w:pBdr'),*doc.element.xpath('.//w:pBdr')]:e.getparent().remove(e)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    path=OUT/'UCSD_Manuscript_Results_3_2_EN_ZH.docx';doc.save(path)
    (OUT/'UCSD_Manuscript_Results_3_2_EN.txt').write_text('\n\n'.join(p['english'] for p in PARAGRAPHS)+'\n')
    d={'stage':34,'date':'2026-10-29','subsection':'3.2 Clinical characteristics','v2_paragraph_ids':['R2'],'paragraphs':PARAGRAPHS,
      'english_word_count':sum(wc(p['english']) for p in PARAGRAPHS),'paragraph_word_counts':{p['id']:wc(p['english']) for p in PARAGRAPHS},
      'word_count_scope':'English narrative only; excludes headings, tables, table titles/notes and Chinese translation','v2_approximate_target_words':150,
      'equations':{'count':0,'numbers':[],'cumulative_previous_numbered_groups':10},
      'citation_mapping':{},'new_manuscript_references':0,'cumulative_unique_cited_sources':5,'editable_english_tables':['Table 1','Table S1'],'table_numeric_values_changed':0,'source_audit':a,'models_refitted':0,'new_hypothesis_tests':0,
      'next_section':'Results 3.3 Four-model discrimination (R3)','docx':str(path.relative_to(ROOT)),'docx_sha256':sha(path)}
    (REPORT/'results_section32.json').write_text(json.dumps(d,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'docx':d['docx'],'english_word_count':d['english_word_count'],'paragraph_word_counts':d['paragraph_word_counts'],'prior_files':len(a['frozen_file_sha256'])},indent=2))

def verify():
    from pypdf import PdfReader
    d=json.loads((REPORT/'results_section32.json').read_text())
    for p,h in d['source_audit']['frozen_file_sha256'].items():assert sha(ROOT/p)==h,p
    path=ROOT/d['docx'];assert sha(path)==d['docx_sha256'];texts=[p.text for p in Document(path).paragraphs]
    for p in PARAGRAPHS:
        assert p['english'] in texts and p['chinese'] in texts and texts.index(p['chinese'])==texts.index(p['english'])+1
        assert not re.search(r'[\u3400-\u9fff，。；：（）【】、]',p['english'])
    assert sum(wc(p['english']) for p in PARAGRAPHS)==d['english_word_count']
    doc=Document(path);assert len(doc.tables)==2
    cells_checked=0
    for table,spec in zip(doc.tables,d['source_audit']['translated_tables']):
        expected=[spec['headers'],*spec['rows']]
        assert [[c.text for c in row.cells] for row in table.rows]==expected
        assert all(not re.search(r'[\u3400-\u9fff]',c.text) for row in table.rows for c in row.cells)
        for note in spec['notes']:assert note in texts
        exported=pd.read_csv(OUT/(spec['key']+'_EN.csv'),dtype=str);assert exported.columns.tolist()==spec['headers'] and exported.values.tolist()==spec['rows']
        cells_checked+=len(spec['rows'])*3
    assert cells_checked==60
    ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
    with zipfile.ZipFile(path) as z:
        t=etree.fromstring(z.read('word/document.xml'));assert len(t.xpath('//m:oMath',namespaces=ns))==0
        assert not any(n.startswith('word/media/') for n in z.namelist())
    pages=len(PdfReader(OUT/'docx_render/UCSD_Manuscript_Results_3_2_EN_ZH.pdf').pages);assert pages==3
    v={'success':True,'stage':34,'date':'2026-10-29','subsection':d['subsection'],'paragraph_ids':['R2'],
      'english_word_count':d['english_word_count'],'paragraph_word_counts':d['paragraph_word_counts'],'matched_chinese_paragraphs':1,
      'word_pages':pages,'visual_review_pages':[1,2,3],'native_editable_equations':0,'equation_numbers':[],'equation_images':0,
      'prior_files_unchanged':len(d['source_audit']['frozen_file_sha256']),'original_table_numeric_cells_checked':60,'saved_numeric_detail_rows_checked':60,'editable_english_tables':2,'table_numeric_values_changed':0,
      'models_refitted':0,'new_hypothesis_tests':0,'new_manuscript_references':0,'cumulative_unique_cited_sources':5,'next_section':d['next_section'],
      'sha256':{str(p.relative_to(ROOT)):sha(p) for p in OUT.iterdir() if p.is_file() and p.suffix in {'.docx','.txt','.csv'}}}
    (REPORT/'artifact_verification.json').write_text(json.dumps(v,ensure_ascii=False,indent=2)+'\n');print(json.dumps(v,indent=2))

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--verify',action='store_true');args=parser.parse_args()
    verify() if args.verify else build()
