"""Reconcile Stage20 baseline statistics and publication artifacts without fitting."""
from pathlib import Path
import base64,hashlib,importlib.util,json,math,subprocess
import nbformat,numpy as np,pandas as pd
from PIL import Image
from openpyxl import load_workbook
from docx import Document
ROOT=Path(__file__).resolve().parents[1]


def digest(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def main():
 out=ROOT/'reports/ucsd_paper_materials';folder=ROOT/'outputs/ucsd-paper-materials-20261029'
 package=json.loads((out/'table_package.json').read_text());receipt=json.loads((out/'generation_receipt.json').read_text())
 assert len(package['tables'])==11 and len(package['figures'])==5 and receipt['model_refits']==0 and receipt['new_hypothesis_tests']==0
 assert not package['manuscript_body_drafted'] and receipt['awaiting_user_writing_instructions']
 for p,h in receipt['input_sha256'].items():assert digest(ROOT/p)==h
 for p,h in receipt['output_sha256'].items():assert digest(ROOT/p)==h
 assert digest(ROOT/'scripts/make_ucsd_paper_materials.py')==receipt['analysis_code_sha256']
 c=pd.read_csv(ROOT/'data/processed/ucsd_clinical_revalidated_cohort.csv');d=pd.read_csv(ROOT/'data/processed/ucsd_clinical_image_features.csv')
 groups={'Table1':{'全部影像扫描':d,'肿瘤扫描':d[d.tumor_label.eq(1)],'治疗改变扫描':d[d.tumor_label.eq(0)]},
         'TableS1':{'临床全集':c,'影像覆盖':d,'影像未覆盖':c[~c.ID.isin(d.ID)]}}
 baseline=pd.read_csv(folder/'tables/baseline_numeric_details.csv')
 for r in baseline.itertuples():
  frame=groups[r.table_id][r.cohort];assert len(frame)==r.scan_denominator and frame.patient_id.nunique()==r.patient_count
  if r.statistic=='n':assert r.count==len(frame)
  elif r.statistic=='patients':assert r.count==frame.patient_id.nunique()
  elif r.statistic=='continuous':
   a=frame[r.variable].dropna().to_numpy();q=np.quantile(a,[.25,.5,.75]);assert len(a)==r.observed_n
   np.testing.assert_allclose(q,[r.q25,r.median,r.q75],atol=1e-12)
  else:
   a=1-frame.tumor_label if r.variable=='treatment_change' else frame.initial_event_missing.eq(0) if r.variable=='interval_observed' else frame[r.variable]
   assert r.count==int(a.sum());assert math.isclose(r.percent,100*a.sum()/len(frame),abs_tol=1e-12)
 specs={t['key']:t for t in package['tables']};main=json.loads((ROOT/'reports/ucsd_fusion/metrics.json').read_text());stable=json.loads((ROOT/'reports/ucsd_stability/metrics.json').read_text())
 for tkey,names,source,ci,total in [
  ('Table2',['M0','M1','MRI','Fusion'],main['metrics'],main['patient_cluster_bootstrap_95ci'],41),
  ('TableS4',['MRI','Fusion','MRI_eroded','Fusion_eroded'],main['metrics'],main['patient_cluster_bootstrap_95ci'],41),
  ('TableS6',['M0','M1','MRI','Fusion'],stable['single_scan']['metrics'],stable['single_scan']['bootstrap_95ci'],22)]:
  for row,name in zip(specs[tkey]['rows'],names):
   expected=[source[name]['roc_auc'],ci[name+'__roc_auc']['lower'],ci[name+'__roc_auc']['upper'],source[name]['brier'],source[name]['ap_treatment_change'],source[name]['tn'],total]
   np.testing.assert_allclose(row[1:],expected,atol=1e-12)
 for row,key in zip(specs['TableS7']['rows'],['roc_auc','ap_treatment_change','brier','log_loss']):
  ci=stable['single_scan']['bootstrap_95ci']['Fusion_minus_M1__'+key]
  np.testing.assert_allclose(row[1:],[stable['single_scan']['paired_differences'][key],ci['lower'],ci['upper']],atol=1e-12)
 datasets={t['key']:t for t in package['tables']};datasets['BaselineNumbers']=package['baseline_numeric']
 wb=load_workbook(folder/'UCSD_Paper_Numeric_Tables.xlsx',read_only=True,data_only=True);checked=0
 for name,data in datasets.items():
  actual=list(wb[name].iter_rows(min_row=1,max_row=len(data['rows'])+1,max_col=len(data['headers']),values_only=True));assert list(actual[0])==data['headers']
  for row,want in zip(actual[1:],data['rows']):
   for a,b in zip(row,want):
    if b is None or b=='':assert a is None or a==''
    elif isinstance(b,(int,float)):assert isinstance(a,(int,float)) and math.isclose(a,b,abs_tol=1e-12,rel_tol=1e-12),(name,a,b)
    else:assert a==b,(name,a,b)
   checked+=1
 assert math.isclose(wb['Index']['C19'].value,main['metrics']['Fusion']['roc_auc'],abs_tol=1e-12)
 assert wb['Index']['C21'].value==2 and wb['Index']['C22'].value==182
 assert not [(sh.title,cell.coordinate) for sh in wb for row in sh for cell in row if cell.data_type=='e']
 spec=importlib.util.spec_from_file_location('word_tables',ROOT/'scripts/build_ucsd_paper_tables.py');module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
 doc=Document(folder/'UCSD_Paper_Tables.docx');assert len(doc.tables)==11
 ordered=['Table1','Table2','Table3','TableS1','TableS2','TableS3','TableS4','TableS5A','TableS5B','TableS6','TableS7']
 for table,key in zip(doc.tables,ordered):
  headers,rows,width=module.shown_table(specs[key]);assert [[c.text for c in row.cells] for row in table.rows]==[headers,*[[str(x) for x in r] for r in rows]]
 assert not doc.styles.element.xpath('.//w:pBdr') and not doc.element.xpath('.//w:pBdr')
 text='\n'.join(p.text for p in doc.paragraphs);assert not any(title in text for title in ['摘要','引言','材料与方法','讨论'])
 pages=list((folder/'docx_render').glob('page-*.png'));assert len(pages)==7
 for name in package['figures']:
  for suffix in ['png','pdf','svg']:assert (folder/'figures'/f'{name}.{suffix}').stat().st_size>1000
  with Image.open(folder/'figures'/f'{name}.png') as img:assert min(img.info['dpi'])>599 and img.width>3000 and img.height>2000
 nb=nbformat.read(ROOT/'notebooks/20_ucsd_paper_tables_and_figures.ipynb',as_version=4);code=[c for c in nb.cells if c.cell_type=='code'];kr=json.loads((ROOT/'reports/kernel_execution_20.json').read_text())
 assert kr['success'] and len(code)==5 and [c.execution_count for c in code]==[1,2,3,4,5]
 assert hashlib.sha256('\n'.join(c.source for c in code).encode()).hexdigest()==kr['notebooks'][0]['code_sha256']
 assert not [o for c in code for o in c.outputs if o.output_type=='error']
 images=[base64.b64decode(o.data['image/png']) for c in code for o in c.outputs if o.output_type in ['display_data','execute_result'] and 'image/png' in o.data];assert len(images)==2
 for b,name in zip(images,['Figure1_Cohort_Flow','Figure2_Main_Comparison']):assert b==(folder/'figures'/f'{name}.png').read_bytes()
 changed=subprocess.check_output(['git','diff','28c758b0226c8cdb52772387a822f70727459dfe','--name-only'],cwd=ROOT,text=True).splitlines()
 for p in changed:
  existed=subprocess.run(['git','cat-file','-e','28c758b0226c8cdb52772387a822f70727459dfe:'+p],cwd=ROOT,capture_output=True).returncode==0
  assert not existed or p in ['README.md','.gitignore','reports/project_status.json'],p
 result=dict(success=True,tables_checked=11,baseline_numeric_rows_independently_checked=len(baseline),workbook_data_rows_checked=checked,
  word_pages_rendered=7,word_tables_editable=11,figure_sets_checked=5,figure_png_dpi=600,vector_formats=['PDF','SVG'],
  notebook_code_cells=5,notebook_errors=0,notebook_embedded_figures_byte_identical=2,prior_stage_artifacts_preserved=True,
  model_refits=0,new_hypothesis_tests=0,manuscript_body_drafted=False,awaiting_user_writing_instructions=True)
 (out/'artifact_verification.json').write_text(json.dumps(result,indent=2)+'\n');print(json.dumps(result,indent=2))


if __name__=='__main__':main()
