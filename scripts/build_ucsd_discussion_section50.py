"""Draft D10 from frozen cohort, model and uncertainty records."""
from pathlib import Path
import argparse, hashlib, json, re, zipfile
import pandas as pd
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'outputs/ucsd-manuscript-step50-20261030'
REPORT = ROOT/'reports/ucsd_manuscript_step50'
EVIDENCE = REPORT/'literature_evidence.json'
MAPPING = {}
PARAGRAPHS = [{'id': 'D10', 'citation_ids': [], 'english': 'Availability-based sampling restricts the population represented by this analysis. MRI inputs were available for 182 of 238 eligible clinical examinations, leaving 56 examinations outside the paired model comparison. The included 135 patients and 41 treatment-related-change examinations provide limited coverage of the diagnostic spectrum; performance may differ in unavailable cases or other clinical settings. Repeated examinations increased observations without creating independent patients. Because fitting and pooled performance gave each examination equal weight, patients with more examinations contributed more than patients with a single scan, even though grouping prevented patient overlap between folds. Fusion also estimated 13 input coefficients from a modest cohort. Fixed regularization and the absence of tuning reduced analytical flexibility but did not establish that estimation was sufficiently stable. The initial-event interval was missing for 36 examinations and referred to heterogeneous starting events. Training-fold median imputation and a missingness indicator could not recover that history or establish an ignorable missingness mechanism. Finally, the patient-cluster bootstrap resampled fixed out-of-fold predictions without refitting. Its intervals therefore describe patient-sampling uncertainty conditional on those predictions, excluding variation from repeating model development. Repeated grouped splits assessed another aspect of instability within the same data; neither procedure substitutes for independent evaluation.', 'chinese': '基于可获得性的抽样限制了本分析所代表的人群。238次符合条件的临床检查中，182次具有可用MRI输入，其余56次未进入配对模型比较。纳入的135名患者及41次治疗相关改变检查，对诊断谱的覆盖有限；未能纳入的病例或其他临床场景中的表现可能不同。重复检查增加了观测次数，但没有增加独立患者数。模型拟合及汇总性能对每次检查赋予相同权重，因此，检查次数更多的患者比仅有一次扫描的患者贡献更大，尽管分组避免了患者在不同折之间重叠。Fusion还在规模有限的队列中估计了13项输入的系数。固定正则化且不进行调参减少了分析灵活性，但不能证明估计已足够稳定。36次检查的初始事件间隔缺失，且该间隔对应不同类型的起始事件。训练折内中位数填补及缺失指示变量无法恢复这些既往信息，也不能确认缺失机制可被忽略。最后，患者聚类bootstrap对固定的折外预测进行重采样，没有重新拟合。因此，其区间描述的是条件于这些预测的患者抽样不确定性，不包含重复模型开发产生的变异。重复分组分折在同一数据内评估了不稳定性的另一个方面；两种程序均不能替代独立评价。'}]

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def wc(text):
    return len(re.findall(r"[A-Za-z0-9]+(?:['/-][A-Za-z0-9]+)*", text))

def audit():
    previous=json.loads((ROOT/'reports/ucsd_manuscript_step49/discussion_section49.json').read_text())
    frozen=previous['source_audit']['frozen_file_sha256'].copy()
    frozen.update(json.loads((ROOT/'reports/ucsd_manuscript_step49/artifact_verification.json').read_text())['sha256'])
    for rel in ['reports/ucsd_manuscript_step25/methods_section23.json','reports/ucsd_manuscript_step29/methods_section27.json','reports/ucsd_fusion/paired_patient_bootstrap.csv']:
        if rel not in frozen: frozen[rel]=sha(ROOT/rel)
    frozen[str(EVIDENCE.relative_to(ROOT))]=sha(EVIDENCE)
    for p,h in frozen.items(): assert sha(ROOT/p)==h,p
    clinical=pd.read_csv(ROOT/'data/processed/ucsd_clinical_revalidated_cohort.csv')
    image=pd.read_csv(ROOT/'data/processed/ucsd_clinical_image_features.csv')
    unavailable=clinical.loc[~clinical.ID.isin(image.ID)]
    assert clinical.ID.is_unique and image.ID.is_unique and set(image.ID).issubset(clinical.ID)
    assert len(clinical)==238 and len(image)==182 and len(unavailable)==56
    assert image.patient_id.nunique()==135 and image.tumor_label.eq(0).sum()==41
    assert unavailable.patient_id.nunique()==41 and not set(unavailable.patient_id)&set(image.patient_id)
    contribution={str(int(k)):int(v) for k,v in image.groupby('patient_id').size().value_counts().sort_index().items()}
    assert contribution=={'1':99,'2':27,'3':7,'4':2}
    assert int(image.initial_event_missing.sum())==36 and int(image.log_initial_event_days.isna().sum())==36
    lock=json.loads((ROOT/'reports/ucsd_fusion/analysis_lock.json').read_text())
    metrics=json.loads((ROOT/'reports/ucsd_fusion/metrics.json').read_text())
    assert len(lock['model_features']['Fusion'])==13 and lock['model_features']==metrics['models']
    assert metrics['C']==1 and metrics['class_weight'] is None and not metrics['hyperparameter_search'] and not metrics['feature_selection']
    assert metrics['interval_missing_scans']==36 and not metrics['bootstrap_models_refitted']
    assert metrics['bootstrap_attempts']==metrics['bootstrap_valid']==2000 and metrics['bootstrap_one_class_skipped']==0
    draws=pd.read_csv(ROOT/'reports/ucsd_fusion/paired_patient_bootstrap.csv')
    assert len(draws)==2000 and draws.draw.tolist()==list(range(1,2001))
    assert draws.resampled_scan_count.between(135,540).all() and draws.resampled_scan_count.nunique()>1
    folds=json.loads((ROOT/'reports/ucsd_fusion/fold_audit.json').read_text());assert len(folds)==5
    for f in folds:
        for key in ['patient_id_overlap','brats_patient_id_overlap','validation_group_overlap']:assert f[key]==0
        assert len(f['models']['Fusion']['parameters']['coefficients'])==13
        assert len(f['models']['Fusion']['parameters']['transformed_names'])==13
    stability=json.loads((ROOT/'reports/ucsd_stability/metrics.json').read_text())
    assert not stability['external_validation']
    prior_ids=set()
    for pattern in ['methods_section*.json','results_section*.json','discussion_section*.json']:
        for p in ROOT.glob('reports/ucsd_manuscript_step*/'+pattern):
            if int(p.parent.name.rsplit('step',1)[1])<50:prior_ids.update(r for item in json.loads(p.read_text())['paragraphs'] for r in item['citation_ids'])
    assert len(prior_ids)==16
    return {'frozen_file_sha256':frozen,'reference_records':[],'source_evidence_record':str(EVIDENCE.relative_to(ROOT)),
            'new_full_texts_checked_for_current_claims':[],'clinical_examinations':238,'imaged_examinations':182,'unavailable_examinations':56,
            'imaged_patients':135,'treatment_change_examinations':41,'patient_examination_frequency':contribution,
            'interval_missing_examinations':36,'fusion_input_coefficients_per_saved_fold':13,'saved_folds_checked':5,
            'examination_weighting_verified_by_source_code_review':True,'fixed_regularization_without_tuning_checked':True,
            'saved_bootstrap_draws_checked':2000,'bootstrap_models_refitted':False,
            'bootstrap_excludes_full_development_variation_disclosed':True,'missingness_mechanism_claimed_established':False,
            'sample_size_adequacy_claimed_established':False,'retrospective_power_calculated':False,'external_validation_claimed':False,
            'previous_unique_cited_sources':sorted(prior_ids),'models_refitted':0,'new_hypothesis_tests':0,'new_bootstrap_draws':0,
            'verification_scope':'Read-only cohort accounting, patient contributions, missingness count, saved coefficient dimensions, bootstrap records and source code; no fitting, resampling, selection-bias test or new sample-size calculation'}

def build():
    a = audit(); OUT.mkdir(parents=True, exist_ok=True); REPORT.mkdir(parents=True, exist_ok=True)
    doc = Document(); sec = doc.sections[0]; sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(.79)
    for name in ['Normal', 'Title', 'Heading 1']:
        st = doc.styles[name]; st.font.name = 'Times New Roman'; st.font.color.rgb = RGBColor(0, 0, 0)
        rf = st.element.get_or_add_rPr().rFonts; rf.set(qn('w:eastAsia'), 'Noto Serif CJK SC')
        for key in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']: rf.attrib.pop(qn('w:'+key), None)
    normal = doc.styles['Normal']; normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.; normal.paragraph_format.space_after = Pt(9)
    doc.styles['Title'].font.size = Pt(16); doc.styles['Title'].paragraph_format.space_after = Pt(12)
    heading = doc.styles['Heading 1']; heading.font.size = Pt(13)
    heading.paragraph_format.space_before = Pt(8); heading.paragraph_format.space_after = Pt(9)
    doc.core_properties.title = 'Discussion 4.10 Sampling and model uncertainty'; doc.core_properties.author = 'Tianyu Feng'
    doc.add_paragraph('Discussion', style='Title'); doc.add_paragraph('4.10 Sampling and model uncertainty', style='Heading 1')
    for item in PARAGRAPHS:
        en = doc.add_paragraph(item['english']); en.paragraph_format.keep_together = True; en.paragraph_format.keep_with_next = True
        zh = doc.add_paragraph(item['chinese']); zh.paragraph_format.keep_together = True
        for run in zh.runs: run.font.size = Pt(11)
    if a['reference_records']:
        title = doc.add_paragraph('References', style='Heading 1'); title.paragraph_format.page_break_before = True
    for ref in a['reference_records']:
        p = doc.add_paragraph(str(ref['number'])+'. '+ref['citation']); p.paragraph_format.keep_together = True
    for e in [*doc.styles.element.xpath('.//w:pBdr'), *doc.element.xpath('.//w:pBdr')]: e.getparent().remove(e)
    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    path = OUT/'UCSD_Manuscript_Discussion_4_10_EN_ZH.docx'; doc.save(path)
    (OUT/'UCSD_Manuscript_Discussion_4_10_EN.txt').write_text('\n\n'.join(p['english'] for p in PARAGRAPHS)+'\n')
    d = {'stage': 50, 'date': '2026-10-30', 'subsection': '4.10 Sampling and model uncertainty', 'v2_paragraph_ids': ['D10'],
         'paragraphs': PARAGRAPHS, 'english_word_count': sum(wc(p['english']) for p in PARAGRAPHS),
         'paragraph_word_counts': {p['id']: wc(p['english']) for p in PARAGRAPHS},
         'word_count_scope': 'English narrative including in-text citation numbers; excludes headings, reference list and Chinese translation',
         'v2_approximate_target_words': 200, 'equations': {'count': 0, 'numbers': [], 'cumulative_previous_numbered_groups': 10},
         'citation_mapping': MAPPING, 'citation_numbering': 'No citations added; existing manuscript reference numbering unchanged',
         'new_manuscript_references': 0, 'cumulative_unique_cited_sources': 16,
         'source_audit': a, 'models_refitted': 0, 'new_hypothesis_tests': 0,
         'next_section': 'Discussion 4.11 Clinical interpretation (D11)',
         'docx': str(path.relative_to(ROOT)), 'docx_sha256': sha(path)}
    (REPORT/'discussion_section50.json').write_text(json.dumps(d, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps({'docx': d['docx'], 'english_word_count': d['english_word_count'], 'frozen_files': len(a['frozen_file_sha256'])}))

def verify():
    from pypdf import PdfReader
    d = json.loads((REPORT/'discussion_section50.json').read_text()); a = audit(); assert a == d['source_audit']
    path = ROOT/d['docx']; assert sha(path) == d['docx_sha256']; doc = Document(path); texts = [p.text for p in doc.paragraphs]
    for p in PARAGRAPHS:
        assert p['english'] in texts and p['chinese'] in texts and texts.index(p['chinese']) == texts.index(p['english'])+1
        assert not re.search(r'[\u3400-\u9fff，。；：（）【】、]', p['english'])
    assert sum(wc(p['english']) for p in PARAGRAPHS) == d['english_word_count']
    assert re.findall(r'\((\d+)\)', PARAGRAPHS[0]['english']) == []
    for ref in a['reference_records']: assert str(ref['number'])+'. '+ref['citation'] in texts
    assert not doc.tables and len(doc.sections) == 1
    with zipfile.ZipFile(path) as z:
        xml = etree.fromstring(z.read('word/document.xml'))
        assert not xml.xpath('//m:oMath', namespaces={'m': 'http://schemas.openxmlformats.org/officeDocument/2006/math'})
        assert not [n for n in z.namelist() if n.startswith('word/media/')]
    pdf = PdfReader(OUT/'docx_render/UCSD_Manuscript_Discussion_4_10_EN_ZH.pdf'); pages = len(pdf.pages); assert pages == 1
    assert '4.10 Sampling' in pdf.pages[0].extract_text() and 'References' not in texts
    v = {'success': True, 'stage': 50, 'date': '2026-10-30', 'subsection': d['subsection'], 'paragraph_ids': ['D10'],
         'english_word_count': d['english_word_count'], 'matched_chinese_paragraphs': 1, 'word_pages': pages,
         'visual_review_pages': [1], 'native_editable_equations': 0, 'equation_images': 0,
         'frozen_files_unchanged': len(a['frozen_file_sha256']), 'prior_files_unchanged': len(a['frozen_file_sha256'])-1,
         'new_full_texts_checked_for_current_claims': [], 'clinical_examinations': 238, 'imaged_examinations': 182,
         'unavailable_examinations': 56, 'imaged_patients': 135, 'treatment_change_examinations': 41,
         'patient_examination_frequency': a['patient_examination_frequency'], 'interval_missing_examinations': 36,
         'fusion_input_coefficients_per_saved_fold': 13, 'saved_folds_checked': 5,
         'examination_weighting_verified_by_source_code_review': True, 'saved_bootstrap_draws_checked': 2000,
         'bootstrap_models_refitted': False, 'bootstrap_excludes_full_development_variation_disclosed': True,
         'missingness_mechanism_claimed_established': False, 'sample_size_adequacy_claimed_established': False,
         'retrospective_power_calculated': False, 'external_validation_claimed': False, 'new_bootstrap_draws': 0,
         'new_manuscript_references': 0, 'provisional_citation_mapping': MAPPING, 'cumulative_unique_cited_sources': 16,
         'models_refitted': 0, 'new_hypothesis_tests': 0, 'next_section': d['next_section'],
         'sha256': {str(p.relative_to(ROOT)): sha(p) for p in OUT.iterdir() if p.is_file() and p.suffix in {'.docx', '.txt'}}}
    (REPORT/'artifact_verification.json').write_text(json.dumps(v, ensure_ascii=False, indent=2)+'\n'); print(json.dumps(v, indent=2))

if __name__ == '__main__':
    parser = argparse.ArgumentParser(); parser.add_argument('--verify', action='store_true'); args = parser.parse_args()
    verify() if args.verify else build()
