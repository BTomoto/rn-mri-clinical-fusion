"""Draft V2 M4 from the frozen reference labels without refitting any model."""
from pathlib import Path
from collections import Counter
import argparse,csv,hashlib,json,re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'outputs/ucsd-manuscript-step24-20261029'
REPORT=ROOT/'reports/ucsd_manuscript_step24'
ENGLISH=(
 'The examination-level outcome was coded as 1 for residual or recurrent tumor and 0 for treatment-related change. '
 'We matched the 51 entries in the public negative-case workbook to the clinical records by examination identifier and assigned them to class 0. '
 'Class 1 comprised the remaining 192 examinations, using the published partition of the full 243-examination collection (1). '
 'These class assignments were retained after identifier quarantine and restriction to examinations with available MRI inputs. '
 'The source reference standard combined pathological evidence with clinical and imaging assessment; the dataset authors reported pathological confirmation for 45 of the 192 tumor examinations in the full collection (1). '
 'Treatment-related change encompassed pseudoprogression, radiation necrosis and nonspecific post-treatment changes. '
 'Subtype counts in the frozen workbook differed from the published summary, although the overall binary counts agreed. We therefore retained the original subtype entries for audit and modeled the broad binary outcome. '
 'We carried forward the source classifications without independent diagnostic adjudication. '
 'This outcome represented the status assigned to the examination and did not define subsequent recurrence or a future radiation-necrosis event.'
)
CHINESE=(
 '检查层面的结局编码为1表示残余或复发肿瘤，编码为0表示治疗相关改变。'
 '我们依据检查标识，将公开negative-case工作簿中的51条记录与临床记录匹配，并将其赋为0类。'
 '依据已发表的完整243次检查的分组，其余192次检查归为1类（1）。'
 '在隔离标识冲突记录、并限定为具有可用MRI输入的检查后，仍沿用这些类别赋值。'
 '源数据的参考标准结合了病理证据、临床及影像评估；数据集作者报告，完整数据集中192次肿瘤检查有45次获得病理证实（1）。'
 '治疗相关改变涵盖假性进展、放射性坏死和非特异性治疗后改变。'
 '冻结工作簿的亚型数量与已发表汇总不同，但二分类总数一致。因此，我们保留原始亚型条目供核查，并针对广义二分类结局建模。'
 '本研究沿用源数据分类，未开展独立诊断裁定。'
 '该结局表示每次检查被赋予的状态，不用于界定后续复发或未来放射性坏死事件。'
)

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def wc(s):return len(re.findall(r"[A-Za-z0-9]+(?:['/-][A-Za-z0-9]+)*",s))

def source_audit():
    receipt=json.loads((ROOT/'reports/ucsd_paper_materials/generation_receipt.json').read_text())
    frozen={**receipt['input_sha256'],**receipt['output_sha256']}
    previous=json.loads((ROOT/'reports/ucsd_manuscript_step23/artifact_verification.json').read_text())
    frozen.update(previous['sha256'])
    for p in ['reports/ucsd_writing_requirements_v2/requirements_and_outline.json','reports/ucsd_manuscript_step23/reference_library.json','reports/ucsd_manuscript_step23/methods_section21.json']:
        frozen[p]=sha(ROOT/p)
    assert all(sha(ROOT/p)==h for p,h in frozen.items())
    source=json.loads((ROOT/'data/metadata/public_label_audit_20261029/downloaded_sources.json').read_text())
    for s in source:assert sha(ROOT/s['local_path'])==s['sha256']
    wb=load_workbook(ROOT/'data/raw/reference/ucsd_ptgbm_negative_cases.xlsx',read_only=True,data_only=True)
    rows=list(wb['all_negative_cases_TCIA'].iter_rows(values_only=True));wb.close()
    negatives=[dict(zip(rows[0],r)) for r in rows[1:]]
    negative_ids={r['ID'] for r in negatives};assert len(negative_ids)==len(negatives)==51
    subtypes=dict(Counter(r['Dx'] for r in negatives))
    assert subtypes=={'NS':29,'PsP (clinically confirmed)':9,'Radiation necrosis':7,'PsP (bx proven)':6}
    counts={}
    for key,file in [('source','ucsd_clinical_baseline_cohort.csv'),('clinical','ucsd_clinical_revalidated_cohort.csv'),('imaging','ucsd_clinical_image_features.csv')]:
        with (ROOT/'data/processed'/file).open() as f:records=list(csv.DictReader(f))
        for r in records:assert int(r['tumor_label'])==int(r['ID'] not in negative_ids)
        counts[key]={'examinations':len(records),'patients':len({r['patient_id'] for r in records}),'tumor':sum(int(r['tumor_label']) for r in records),'treatment_change':sum(r['ID'] in negative_ids for r in records)}
    assert counts=={'source':{'examinations':243,'patients':178,'tumor':192,'treatment_change':51},'clinical':{'examinations':238,'patients':176,'tumor':187,'treatment_change':51},'imaging':{'examinations':182,'patients':135,'tumor':141,'treatment_change':41}}
    return {'frozen_file_sha256':frozen,'raw_source_files':source,'subtype_counts_in_frozen_workbook':subtypes,
      'subtype_counts_in_published_full_collection':{'NS':28,'PsP (clinically confirmed)':9,'Radiation necrosis':10,'PsP (bx proven)':4},
      'label_counts':counts,'all_labels_reconstructed_and_matched':True,'label_rows_checked':sum(c['examinations'] for c in counts.values()),
      'pathology_confirmed_tumor_examinations_in_published_full_collection':45,
      'pathology_confirmed_count_in_182_imaging_subset':'not inferred from aggregate source data'}

def build():
    OUT.mkdir(parents=True,exist_ok=True);REPORT.mkdir(parents=True,exist_ok=True)
    audit=source_audit()
    refs=json.loads((ROOT/'reports/ucsd_manuscript_step23/reference_library.json').read_text())['references']
    reference=next(r for r in refs if r['id']=='S01')
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(.79)
    for name in ['Normal','Title','Heading 1']:
        st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        rf=st.element.get_or_add_rPr().rFonts;rf.set(qn('w:eastAsia'),'Noto Serif CJK SC')
        for a in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:rf.attrib.pop(qn('w:'+a),None)
    normal=doc.styles['Normal'];normal.font.size=Pt(12);normal.paragraph_format.line_spacing=1.;normal.paragraph_format.space_after=Pt(10)
    doc.styles['Title'].font.size=Pt(16);doc.styles['Title'].paragraph_format.space_after=Pt(12)
    doc.styles['Heading 1'].font.size=Pt(13);doc.styles['Heading 1'].paragraph_format.space_before=Pt(8);doc.styles['Heading 1'].paragraph_format.space_after=Pt(9)
    doc.core_properties.title='Materials and methods 2.2 Reference standard';doc.core_properties.author='Tianyu Feng'
    doc.add_paragraph('Materials and methods',style='Title')
    doc.add_paragraph('2.2 Reference standard',style='Heading 1')
    p=doc.add_paragraph(ENGLISH);p.paragraph_format.keep_together=True
    p=doc.add_paragraph(CHINESE);p.paragraph_format.keep_together=True
    for r in p.runs:r.font.size=Pt(11)
    doc.add_paragraph('References',style='Heading 1')
    p=doc.add_paragraph('1. '+reference['citation']);p.paragraph_format.keep_together=True
    for r in p.runs:r.font.size=Pt(11)
    for e in [*doc.styles.element.xpath('.//w:pBdr'),*doc.element.xpath('.//w:pBdr')]:e.getparent().remove(e)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    p=OUT/'UCSD_Manuscript_Methods_2_2_EN_ZH.docx';doc.save(p)
    (OUT/'UCSD_Manuscript_Methods_2_2_EN.txt').write_text(ENGLISH+'\n')
    data={'stage':24,'date':'2026-10-29','subsection':'2.2 Reference standard','v2_paragraph_ids':['M4'],
      'paragraphs':[{'id':'M4','english':ENGLISH,'chinese':CHINESE,'citation_ids':['S01']}],
      'english_word_count':wc(ENGLISH),'v2_target_words':150,'word_target_is_approximate':True,
      'citation_mapping':{'1':'S01'},'citation_numbering':'S01 retains number 1 from section 2.1; regenerate globally when Introduction is assembled',
      'models_refitted':0,'new_hypothesis_tests':0,'next_section':'Clinical predictor definitions (M5-M6)',
      'source_audit':audit,'claim_sources':[
        {'claim':'Binary outcome construction and preservation through the existing cohorts','sources':['src/rn_public/ucsd_clinical_baseline.py','reports/ucsd_clinical_baseline_protocol.md','data/raw/reference/ucsd_ptgbm_negative_cases.xlsx'],'independent_readjudication':False},
        {'claim':'Published full collection partition, pathological confirmation and broad treatment-change categories','citation_id':'S01','url':'https://pmc.ncbi.nlm.nih.gov/articles/PMC12886825/','section':'Methods / Data collection','checked_date':'2026-10-29'},
        {'claim':'Subtype discrepancy and retention of broad grouping','sources':['data/processed/ucsd_ptgbm_subtype_count_comparison.csv','reports/public_label_audit_step10_20261029.md']}],
      'docx':str(p.relative_to(ROOT)),'docx_sha256':sha(p)}
    (REPORT/'methods_section22.json').write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'docx':data['docx'],'english_word_count':data['english_word_count'],'labels_checked':audit['label_rows_checked'],'prior_files_checked':len(audit['frozen_file_sha256'])},indent=2))

def verify():
    from pypdf import PdfReader
    d=json.loads((REPORT/'methods_section22.json').read_text())
    for p,h in d['source_audit']['frozen_file_sha256'].items():assert sha(ROOT/p)==h,p
    document=ROOT/d['docx'];assert sha(document)==d['docx_sha256']
    paragraphs=[p.text for p in Document(document).paragraphs]
    assert paragraphs[2:4]==[ENGLISH,CHINESE]
    assert paragraphs[4]=='References' and paragraphs[5].startswith('1. Gagnon L')
    assert not re.search(r'[\u3400-\u9fff，。；：（）【】、]',ENGLISH)
    assert wc(ENGLISH)==d['english_word_count']
    assert set(re.findall(r'\((\d+)\)',ENGLISH))=={'1'}
    pages=len(PdfReader(OUT/'docx_render/UCSD_Manuscript_Methods_2_2_EN_ZH.pdf').pages);assert pages==1
    p={'success':True,'stage':24,'date':'2026-10-29','subsection':d['subsection'],'paragraph_ids':['M4'],
       'english_word_count':d['english_word_count'],'matched_chinese_paragraphs':1,'word_pages':pages,'current_section_citations':1,
       'existing_citation_number_1_preserved':True,'english_cjk_check_passed':True,'visual_review_pages':[1],
       'prior_files_unchanged':len(d['source_audit']['frozen_file_sha256']),
       'label_records_checked':d['source_audit']['label_rows_checked'],'all_labels_match_source_mapping':True,
       'subtype_discrepancy_disclosed':True,'pathology_count_not_extrapolated_to_imaging_subset':True,
       'models_refitted':0,'new_hypothesis_tests':0,'next_section':d['next_section'],
       'sha256':{str(f.relative_to(ROOT)):sha(f) for f in OUT.iterdir() if f.is_file() and f.suffix in {'.txt','.docx'}}}
    (REPORT/'artifact_verification.json').write_text(json.dumps(p,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(p,ensure_ascii=False,indent=2))

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--verify',action='store_true');args=parser.parse_args()
    verify() if args.verify else build()
