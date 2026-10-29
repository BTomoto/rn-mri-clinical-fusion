"""Write R1 and attach the unchanged cohort figure using read-only cohort checks."""
from pathlib import Path
import argparse,hashlib,json,re,sys,zipfile
import numpy as np
import pandas as pd
from lxml import etree
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'outputs/ucsd-manuscript-step33-20261029'
REPORT=ROOT/'reports/ucsd_manuscript_step33'
PARAGRAPHS=[
 {'id':'R1','citation_ids':[],
  'english': 'The source cohort contained 243 examinations from 178 patients, including 192 examinations labeled as tumor and 51 labeled as treatment-related change (Figure 1). Identifier quarantine removed five examinations from two patients, leaving 238 examinations from 176 patients in the clinical cohort. The retained cohort contained 187 tumor and 51 treatment-related-change examinations. Complete FLAIR, postcontrast T1-weighted and annotation-mask inputs were available for 182 examinations from 135 patients, representing 76.5% of the retained clinical examinations. All 182 yielded the eight MRI features and entered the common comparison of M0, M1, MRI and Fusion. The remaining 56 examinations from 41 patients were not covered by the accessed mirror and comprised 46 tumor and 10 treatment-related-change examinations. These records remained in the coverage ledger and did not enter the imaging comparison. The imaging cohort included 141 tumor examinations (77.5%) and 41 treatment-related-change examinations (22.5%). At the patient level, 114 patients contributed tumor examinations and 29 contributed treatment-related-change examinations; eight patients contributed examinations in both categories. These category-specific patient counts therefore overlapped. Table 1 summarizes the imaging cohort, and Supplementary Table S1 describes the retained clinical cohort according to imaging coverage. The same 182 examinations were retained for the primary comparison and the erosion sensitivity analysis.',
  'chinese': '源队列包含178名患者的243次检查，其中192次标记为肿瘤，51次标记为治疗相关改变（Figure 1）。标识隔离排除了2名患者的5次检查，临床队列保留176名患者的238次检查，包括187次肿瘤检查和51次治疗相关改变检查。其中，135名患者的182次检查具有完整的FLAIR、增强T1加权图像及标注掩膜输入，占保留临床检查的76.5%。全部182次检查均获得八项MRI特征，并进入M0、M1、MRI与Fusion的共同比较。其余41名患者的56次检查未被所访问的镜像覆盖，包括46次肿瘤检查和10次治疗相关改变检查。这些记录保留在覆盖清单中，未进入影像比较。影像队列包含141次肿瘤检查（77.5%）和41次治疗相关改变检查（22.5%）。按患者统计，114名患者贡献过肿瘤检查，29名贡献过治疗相关改变检查，其中8名患者在两类中均有检查。因此，两类患者人数存在重叠。Table 1汇总影像队列特征，Supplementary Table S1按影像覆盖情况描述保留的临床队列。主比较与腐蚀敏感性分析均保留相同的182次检查。'}
]
FIGURE=ROOT/'outputs/ucsd-paper-materials-20261029/figures/Figure1_Cohort_Flow.png'
FIGURE_CAPTION='Figure 1. Cohort assembly and analysis populations. Counts refer to examinations and source patients. After identifier quarantine, the imaging cohort included 182 examinations from 135 patients. The one-scan-per-patient sensitivity analysis selected the smallest available session number within this imaging cohort. It does not necessarily represent the first examination in the clinical course.'

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def wc(s):return len(re.findall(r"[A-Za-z0-9]+(?:['/-][A-Za-z0-9]+)*",s))
def close(a,b):assert np.isclose(a,b,rtol=1e-12,atol=1e-12),(a,b)

def audit():
    old=json.loads((ROOT/'reports/ucsd_manuscript_step32/methods_section210.json').read_text());frozen=old['source_audit']['frozen_file_sha256'].copy()
    frozen.update(json.loads((ROOT/'reports/ucsd_manuscript_step32/artifact_verification.json').read_text())['sha256'])
    for p,h in frozen.items():assert sha(ROOT/p)==h,p
    files=['data/processed/ucsd_clinical_baseline_cohort.csv','data/processed/ucsd_clinical_revalidated_cohort.csv',
      'data/processed/ucsd_clinical_image_features.csv','reports/ucsd_clinical_revalidation/excluded_identity_conflicts.csv',
      'reports/ucsd_batch_features/all_clinical_case_status.csv','outputs/ucsd-paper-materials-20261029/tables/Table1.csv',
      'outputs/ucsd-paper-materials-20261029/tables/TableS1.csv',str(FIGURE.relative_to(ROOT))]
    frozen.update({p:sha(ROOT/p) for p in files})
    source=pd.read_csv(ROOT/files[0]);clinical=pd.read_csv(ROOT/files[1]);image=pd.read_csv(ROOT/files[2]);excluded=pd.read_csv(ROOT/files[3]);status=pd.read_csv(ROOT/files[4])
    def counts(d):
        assert not d.ID.duplicated().any()
        return {'examinations':len(d),'patients':int(d.patient_id.nunique()),'tumor_examinations':int(d.tumor_label.eq(1).sum()),'treatment_change_examinations':int(d.tumor_label.eq(0).sum())}
    expected={'source':{'examinations':243,'patients':178,'tumor_examinations':192,'treatment_change_examinations':51},
      'quarantined':{'examinations':5,'patients':2,'tumor_examinations':5,'treatment_change_examinations':0},
      'clinical':{'examinations':238,'patients':176,'tumor_examinations':187,'treatment_change_examinations':51},
      'imaging':{'examinations':182,'patients':135,'tumor_examinations':141,'treatment_change_examinations':41},
      'uncovered':{'examinations':56,'patients':41,'tumor_examinations':46,'treatment_change_examinations':10}}
    uncovered=clinical[~clinical.ID.isin(image.ID)];frames=dict(source=source,quarantined=excluded,clinical=clinical,imaging=image,uncovered=uncovered)
    actual={k:counts(v) for k,v in frames.items()};assert actual==expected
    assert set(source.ID)==set(clinical.ID)|set(excluded.ID) and not set(clinical.ID)&set(excluded.ID)
    assert not set(clinical.patient_id)&set(excluded.patient_id)
    assert set(clinical.ID)==set(image.ID)|set(uncovered.ID) and not set(image.ID)&set(uncovered.ID)
    assert not set(image.patient_id)&set(uncovered.patient_id)
    for d in [clinical,image,excluded]:
        pd.testing.assert_series_equal(d.set_index('ID').tumor_label.sort_index(),source.set_index('ID').loc[d.ID,'tumor_label'].sort_index())
    assert set(status.ID)==set(clinical.ID) and set(status.loc[status.status.eq('complete'),'ID'])==set(image.ID)
    assert set(status.loc[status.status.eq('not_listed_in_current_mirror'),'ID'])==set(uncovered.ID)
    assert set(status.status)=={'complete','not_listed_in_current_mirror'}
    features=['FLAIR_'+s for s in ['mean','std','entropy','glcm_contrast']]+['T1post_'+s for s in ['mean','std','entropy','glcm_contrast']]
    assert np.isfinite(image[features].to_numpy()).all()
    manifest=pd.read_csv(ROOT/'reports/ucsd_batch_features/source_manifest.csv');assert len(manifest)==546
    assert set(manifest.ID)==set(image.ID) and manifest.groupby('ID').size().eq(3).all()
    assert manifest.groupby('ID').sequence.nunique().eq(3).all()
    for file in ['data/processed/ucsd_image_features_eroded.csv','data/processed/ucsd_fusion_oof.csv']:
        d=pd.read_csv(ROOT/file);assert len(d)==182 and set(d.ID)==set(image.ID)
    pred=pd.read_csv(ROOT/'data/processed/ucsd_fusion_oof.csv')
    for name in ['M0','M1','MRI','Fusion','MRI_eroded','Fusion_eroded']:
        assert np.isfinite(pred[name+'_oof_probability']).all()
    tumor=set(image.loc[image.tumor_label.eq(1),'patient_id']);tc=set(image.loc[image.tumor_label.eq(0),'patient_id'])
    assert len(tumor)==114 and len(tc)==29 and len(tumor&tc)==8 and len(tumor|tc)==135
    pct={'clinical_examination_imaging_coverage':100*len(image)/len(clinical),'imaging_tumor':100*image.tumor_label.eq(1).mean(),'imaging_treatment_change':100*image.tumor_label.eq(0).mean()}
    assert [f'{p:.1f}' for p in pct.values()]==['76.5','77.5','22.5']
    table=pd.read_csv(ROOT/files[5],dtype=str).set_index('指标');s1=pd.read_csv(ROOT/files[6],dtype=str).set_index('指标')
    for col,n,patient in [('全部影像扫描',182,135),('肿瘤扫描',141,114),('治疗改变扫描',41,29)]:
        assert table.loc['扫描记录数',col]==str(n) and table.loc['源患者数',col]==str(patient)
    for col,n,patient in [('临床全集',238,176),('影像覆盖',182,135),('影像未覆盖',56,41)]:
        assert s1.loc['扫描记录数',col]==str(n) and s1.loc['源患者数',col]==str(patient)
    prior=json.loads((ROOT/'reports/ucsd_manuscript_step23/methods_section21.json').read_text());assert sha(FIGURE)==prior['figure_sha256']
    return {'frozen_file_sha256':frozen,'cohort_counts_checked':actual,'cohort_count_fields_checked':20,
      'source_to_clinical_partition_checked':True,'clinical_to_imaging_partition_checked':True,
      'label_rows_compared_to_source':len(clinical)+len(image)+len(excluded),'clinical_status_rows_checked':238,
      'imaging_file_sets_checked':182,'finite_mri_feature_values_checked':1456,'main_and_eroded_case_sets_identical':True,
      'class_specific_patient_counts':{'tumor':114,'treatment_change':29,'both_labels_on_different_examinations':8,'union':135},
      'percentages_unrounded':pct,'table_count_cells_checked':12,'figure_identical_to_confirmed_methods_copy':True,
      'models_refitted':0,'features_reextracted':0,'new_hypothesis_tests':0,
      'verification_scope':'Count and partition existing cohort tables, verify label preservation, saved processing status and complete MRI/OOF case sets; inspect unchanged Figure 1; no new model analysis'}

def build():
    OUT.mkdir(parents=True,exist_ok=True);REPORT.mkdir(parents=True,exist_ok=True);a=audit()
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(.79)
    for name in ['Normal','Title','Heading 1']:
        st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        rf=st.element.get_or_add_rPr().rFonts;rf.set(qn('w:eastAsia'),'Noto Serif CJK SC')
        for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:rf.attrib.pop(qn('w:'+key),None)
    normal=doc.styles['Normal'];normal.font.size=Pt(12);normal.paragraph_format.line_spacing=1.;normal.paragraph_format.space_after=Pt(9)
    doc.styles['Title'].font.size=Pt(16);doc.styles['Title'].paragraph_format.space_after=Pt(12)
    doc.styles['Heading 1'].font.size=Pt(13);doc.styles['Heading 1'].paragraph_format.space_before=Pt(8);doc.styles['Heading 1'].paragraph_format.space_after=Pt(9)
    doc.core_properties.title='Results 3.1 Cohort flow';doc.core_properties.author='Tianyu Feng'
    doc.add_paragraph('Results',style='Title');doc.add_paragraph('3.1 Cohort flow',style='Heading 1')
    for item in PARAGRAPHS:
        en=doc.add_paragraph(item['english']);en.paragraph_format.page_break_before=False
        en.paragraph_format.keep_together=True;en.paragraph_format.keep_with_next=True
        zh=doc.add_paragraph(item['chinese']);zh.paragraph_format.keep_together=True
        for r in zh.runs:r.font.size=Pt(11)
    p=doc.add_paragraph('Figure 1',style='Heading 1');p.paragraph_format.page_break_before=True
    doc.add_picture(str(FIGURE),width=Inches(6.7))
    p=doc.add_paragraph(FIGURE_CAPTION);p.paragraph_format.keep_together=True
    for r in p.runs:r.font.size=Pt(10)
    for e in [*doc.styles.element.xpath('.//w:pBdr'),*doc.element.xpath('.//w:pBdr')]:e.getparent().remove(e)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    path=OUT/'UCSD_Manuscript_Results_3_1_EN_ZH.docx';doc.save(path)
    (OUT/'UCSD_Manuscript_Results_3_1_EN.txt').write_text('\n\n'.join(p['english'] for p in PARAGRAPHS)+'\n')
    d={'stage':33,'date':'2026-10-29','subsection':'3.1 Cohort flow','v2_paragraph_ids':['R1'],'paragraphs':PARAGRAPHS,
      'english_word_count':sum(wc(p['english']) for p in PARAGRAPHS),'paragraph_word_counts':{p['id']:wc(p['english']) for p in PARAGRAPHS},
      'word_count_scope':'English narrative only; excludes headings, figure text/caption and Chinese translation','v2_approximate_target_words':200,
      'equations':{'count':0,'numbers':[],'cumulative_previous_numbered_groups':10},
      'citation_mapping':{},'new_manuscript_references':0,'cumulative_unique_cited_sources':5,'included_figure':str(FIGURE.relative_to(ROOT)),'figure_sha256':sha(FIGURE),'figure_caption':FIGURE_CAPTION,'figure_reused_without_regeneration':True,'source_audit':a,'models_refitted':0,'new_hypothesis_tests':0,
      'next_section':'Results 3.2 Clinical characteristics (R2)','docx':str(path.relative_to(ROOT)),'docx_sha256':sha(path)}
    (REPORT/'results_section31.json').write_text(json.dumps(d,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'docx':d['docx'],'english_word_count':d['english_word_count'],'paragraph_word_counts':d['paragraph_word_counts'],'prior_files':len(a['frozen_file_sha256'])},indent=2))

def verify():
    from pypdf import PdfReader
    d=json.loads((REPORT/'results_section31.json').read_text())
    for p,h in d['source_audit']['frozen_file_sha256'].items():assert sha(ROOT/p)==h,p
    path=ROOT/d['docx'];assert sha(path)==d['docx_sha256'];texts=[p.text for p in Document(path).paragraphs]
    for p in PARAGRAPHS:
        assert p['english'] in texts and p['chinese'] in texts and texts.index(p['chinese'])==texts.index(p['english'])+1
        assert not re.search(r'[\u3400-\u9fff，。；：（）【】、]',p['english'])
    assert sum(wc(p['english']) for p in PARAGRAPHS)==d['english_word_count']
    assert FIGURE_CAPTION in texts and not re.search(r'[\u3400-\u9fff]',FIGURE_CAPTION)
    ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
    with zipfile.ZipFile(path) as z:
        t=etree.fromstring(z.read('word/document.xml'));assert len(t.xpath('//m:oMath',namespaces=ns))==0
        media=[n for n in z.namelist() if n.startswith('word/media/')];assert len(media)==1
        assert hashlib.sha256(z.read(media[0])).hexdigest()==sha(FIGURE)
    pages=len(PdfReader(OUT/'docx_render/UCSD_Manuscript_Results_3_1_EN_ZH.pdf').pages);assert pages==2
    v={'success':True,'stage':33,'date':'2026-10-29','subsection':d['subsection'],'paragraph_ids':['R1'],
      'english_word_count':d['english_word_count'],'paragraph_word_counts':d['paragraph_word_counts'],'matched_chinese_paragraphs':1,
      'word_pages':pages,'visual_review_pages':[1,2],'native_editable_equations':0,'equation_numbers':[],'equation_images':0,
      'prior_files_unchanged':len(d['source_audit']['frozen_file_sha256']),'cohort_count_fields_checked':20,'label_rows_compared_to_source':425,'clinical_status_rows_checked':238,'imaging_file_sets_checked':182,
      'finite_mri_feature_values_checked':1456,'table_count_cells_checked':12,'overlapping_diagnostic_patient_groups_disclosed':True,
      'unchanged_figure_embedded':True,'models_refitted':0,'new_hypothesis_tests':0,'new_manuscript_references':0,'cumulative_unique_cited_sources':5,'next_section':d['next_section'],
      'sha256':{str(p.relative_to(ROOT)):sha(p) for p in OUT.iterdir() if p.is_file() and p.suffix in {'.docx','.txt'}}}
    (REPORT/'artifact_verification.json').write_text(json.dumps(v,ensure_ascii=False,indent=2)+'\n');print(json.dumps(v,indent=2))

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--verify',action='store_true');args=parser.parse_args()
    verify() if args.verify else build()
