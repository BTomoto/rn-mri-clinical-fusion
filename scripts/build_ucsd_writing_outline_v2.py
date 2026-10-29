"""Stage22: map collaborator requirements to an 8,000-word manuscript outline.

Only the outline is written; no manuscript narrative or new model is produced.
"""
from pathlib import Path
import hashlib,json,math
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'outputs/ucsd-writing-outline-v2-20261029'
REPORT=ROOT/'reports/ucsd_writing_requirements_v2'
INPUT=ROOT.parent/'upload'

# Each row: stable V2 paragraph ID, planned English words, topic, task, evidence.
I=[
 ('I1',160,'Clinical setting','介绍胶质瘤治疗后复查场景与当前状态鉴别问题，迅速聚焦残余或复发肿瘤和治疗相关改变。','近期临床与影像研究'),
 ('I2',170,'Diagnostic uncertainty','解释两类MRI表现重叠、混合病理和临床判断困难，区分治疗相关改变、假性进展与放射性坏死。','近期终点及影像评估文献'),
 ('I3',180,'Quantitative imaging','介绍常规FLAIR、增强T1的强度和纹理计算为何可能补充视觉信息；明确本研究仅8项固定特征。','近期常规MRI及影像计算研究'),
 ('I4',180,'Clinical and imaging fusion','解释临床背景与MRI描述的互补性，综述融合获益与未获益证据，不能假定增加特征必然提高性能。','正向与无获益的融合研究'),
 ('I5',180,'Gap in reproducible evaluation','将缺口限定为公开派生影像中低维融合的增量与稳定性；说明同病例对照及患者隔离的重要性。','近期验证方法与复现研究'),
 ('I6',130,'Study objective','提出扫描级鉴别目标、四模型和Fusion-M1主要比较；简述边界、重复分组及单扫描复核。','第18至19步方案'),
]
M=[
 ('M1',150,'Study design','界定公开去标识资料的回顾性二次计算研究，说明分析单位和内部验证目的，不写成新招募研究。','第18步方案与源数据论文'),
 ('M2',150,'Public sources','说明UCSD-PTGBM临床表、公开镜像、版本及来源分工；区分源采集流程与本研究实际取得的派生文件。','第14至17步来源清单'),
 ('M3',150,'Cohort linkage','写明243扫描/178患者、ID冲突隔离后238/176、影像队列182/135及未覆盖56扫描；区分患者与扫描。','Figure 1和Table S1'),
 ('M4',150,'Reference labels','定义1为肿瘤、0为治疗改变，说明公开阴性表与冻结总体分组补集，以及混合临床/病理参考标准。','标签审计与数据描述'),
 ('M5',150,'Clinical inputs','列明年龄、性别、初始事件对数间隔、扫描前其他治疗和间隔缺失标记；解释各输入在分析中的角色。','Table S2与变量字典'),
 ('M6',150,'Time and missingness','解释初始事件可能为手术/治疗/诊断；负偏移映射为正天数，其他异常按冻结规则缺失，36条间隔折内填补。','临床重验证及第18步代码'),
 ('M7',150,'MRI data quality checks','说明FLAIR、增强T1和现成标注的546文件检查、uint8派生强度、256立方数组与空间单位unknown。','第15至17步文件与数组QC'),
 ('M8',150,'Annotation region','ROI为普通BraTS全部正标签并集，保留各连通分量；说明现成标注来源及盲法未知，非本研究新训练分割。','既有ROI协议'),
 ('M9',150,'Normalization','介绍有限正值支持区域的p1/p99线性变换且不截断，ROI内0保留；交代不截断为观察后的工程修订。','第16步修订及第17步实现'),
 ('M10',150,'Feature definitions','定义两序列各自均值、标准差与32级直方图熵，说明ROI范围离散化及最大值分箱规则。','8项固定特征字典'),
 ('M11',150,'Three-dimensional texture','解释13个无向单原生体素偏移、对称计数汇总再归一化的GLCM对比度；公式用可编辑数学对象。','原GLCM代码与补充方法'),
 ('M12',150,'Fixed models and fusion','列出M0、M1、MRI、Fusion的2/5/8/13输入与特征级拼接，给出L2 Logistic核心参数及模型选择理由。','第18步固定模型工厂'),
 ('M13',150,'Training and partitions','交代同病例同五折、固定种子、源/BraTS患者隔离，训练折学习填补和标准化；没有调参、筛选或类别加权。','折清单与fold_audit'),
 ('M14',150,'Discrimination and errors','定义主比较AUC差及两类AP、Brier、log loss、固定0.5阈值混淆矩阵；治疗改变PR使用1-p。','第18步指标定义'),
 ('M15',150,'Conditional uncertainty','说明患者聚类配对bootstrap 2000次及百分位区间；固定OOF概率不重拟合，不能覆盖全部训练和分折变异。','bootstrap回执与抽样指标'),
 ('M16',150,'Calibration assessment','说明校准分组、斜率及offset截距的诊断定义；未据此重校准或更改阈值。','原校准协议和Table S3'),
 ('M17',150,'Boundary sensitivity','一次六邻域单原生体素腐蚀，保留主归一化和分箱边界，在同病例同折复核MRI及Fusion。','Table S4和Figure S1'),
 ('M18',150,'Repeated grouped validation','20套事先固定种子的患者分组五折，描述中位数、四分位和范围；不选最优重复或构建额外集成。','Tables S5A及S5B和Figure 3'),
 ('M19',150,'One scan per patient','按当前可用影像内最小数字会话选135扫描，独立重建验证与bootstrap；不是整个病程的真实首次就诊。','选择清单及Tables S6和S7'),
 ('M20',150,'Reproducibility','说明Python、Notebook、软件版本、逐例OOF、源文件哈希和归档恢复路径；交代样本量由可用数据决定。','项目仓库及复现材料'),
]
R=[
 ('R1',200,'Cohort flow','先写源队列、隔离与覆盖流程，分别报告扫描数、患者数和182扫描中141肿瘤/41治疗改变。','Figure 1和Table 1'),
 ('R2',150,'Clinical characteristics','报告扫描口径年龄等分布、36条时间缺失和影像覆盖情况；患者可在两标签间重叠，不作未经计算的组间检验。','Table 1和Table S1'),
 ('R3',200,'Four-model discrimination','按M0、M1、MRI、Fusion顺序呈现固定主五折；M1 0.672、MRI 0.762、Fusion 0.782，附准确区间。','Table 2和Figure 2'),
 ('R4',200,'Paired incremental comparison','Fusion-M1 AUC差0.110，区间0.021至0.214；Brier差-0.019且区间跨0，治疗改变AP按原结果报告。','Table 3'),
 ('R5',150,'Calibration and errors','写主Fusion斜率0.667，肿瘤正确130/141、治疗改变12/41；整体准确率78.0%与全判肿瘤77.5%参考并列。','Figure 2和Table S3'),
 ('R6',150,'Boundary sensitivity','报告腐蚀Fusion AUC 0.788及原/腐蚀差值区间跨0，不以较高版本替换主分析。','Table S4和Figure S1'),
 ('R7',150,'Split stability','20套完整，Fusion AUC中位数0.784、范围0.759至0.809；Fusion-M1差全部为正，范围不是置信区间。','Figure 3和Tables S5A及S5B'),
 ('R8',200,'Single-scan sensitivity','135扫描/135患者，113肿瘤/22治疗改变；Fusion AUC 0.730，增量区间-0.015至0.252跨0，只识别2/22治疗改变。','Tables S6和S7及Figure S2'),
]
D=[
 ('D1',200,'Principal findings','开头回答增量问题，同时交代主分析有利信号、分折方向一致与单扫描增量不确定。','对应R3至R8'),
 ('D2',250,'Comparison with fusion studies','比较近期融合研究的信息类型、样本、模型和内部验证设计；解释与本研究观察的联系，不按跨论文AUC排名。','近期融合原始研究'),
 ('D3',250,'Conflicting evidence','讨论未获益、较低性能及新队列验证研究，重点比较终点、参考标准、成像和抽样差异。','负向及独立验证研究'),
 ('D4',200,'Computational contribution','阐明固定MRI特征、特征级融合、同条件模型比较及公开复现的贡献，准确定位为应用验证。','本研究流程与近期复现文献'),
 ('D5',200,'Interpretation of image features','解释强度与空间灰度变化所描述的信息，区分统计关联与病理机制；不把系数写成独立危险因素。','特征定义和有依据的近期解释文献'),
 ('D6',200,'Boundary and split robustness','解释腐蚀和重复分折分别回答的问题，说明稳定性有利但不能验证源标签或外部推广。','M17至M18及对应结果'),
 ('D7',200,'Meaning of the single-scan result','讨论更少治疗改变样本下区间宽及识别不足，承认人群、类别组成和训练集不同，不能直接归因于重复扫描。','M19和R8'),
 ('D8',250,'Reference-standard limitations','优先讨论标签补集、混合参考、治疗改变异质性和现成标注盲法未确认；区分可能泄漏与已证实泄漏。','标签审计与源数据论文'),
 ('D9',200,'Image-processing limitations','讨论uint8、unknown空间单位、全正标签ROI及工程修订的适用范围；没有形状、定量ADC或独立重画实验。','原影像与特征QC'),
 ('D10',200,'Sampling and model uncertainty','说明覆盖选择、135患者、少数类有限、重复扫描及扫描权重、13输入和缺失时间变量，区间为条件bootstrap区间。','队列审计与近期验证文献'),
 ('D11',200,'Clinical interpretation','将AUC排序能力、校准和阈值识别分开，解释不能据内部AUC证明临床获益；没有外部验证或决策曲线结论。','R4至R5及近期临床转化研究'),
 ('D12',100,'Next validation priorities','按独立参考与标注、标准化影像、外部队列和少数类/校准评估的优先顺序提出后续方向。','从本研究局限推导'),
]
C=[('C1',150,'Conclusion','简洁回到研究目标，保留主内部验证增量信号与单扫描不确定性，并限定适用范围及独立验证需求。','对应I6和D1')]
A=[
 ('A1',35,'Background','临床鉴别问题及MRI对临床信息的增量目标。','不引用文献、图或表'),
 ('A2',85,'Methods','182扫描/135患者、固定特征及Logistic，患者分组验证和主要比较；简述敏感性。','与正式方法一致'),
 ('A3',100,'Results','主要AUC、增量及区间，加单扫描不确定与少数类识别不足。','不改用最佳重复结果'),
 ('A4',30,'Conclusions','内部增量信号及独立验证需求。','与结论一致'),
]

PDF_RULES=[
 ('P1','PDF p1, item 1','使用通用模板并保持标题逻辑','使用Title和分级Heading样式；引言无小标题，方法/结果分组标题对应流程。'),
 ('P2','PDF p1, item 2','正文清楚引用每幅图和每张表','建立首次出现顺序清单，主文图表与补充图表分别编号并逐一对应。'),
 ('P3','PDF p1, item 3','英文稿中不出现中文字符','纯英文投稿版扫描汉字及中文标点；中文仅在对照审阅稿的译文区域保留。'),
 ('P4','PDF p2, item 4','英文标点及空格规范','常规逗号、句号等紧跟前词，后接一个空格；小数、缩写、DOI及数学运算按各自规则。'),
 ('P5','PDF p2, item 5','行内与独立公式规范','Word公式编辑器的可编辑对象；上下标、希腊符号、变量和公式编号统一，不能用普通文字拼装。'),
 ('P6','PDF p2-p3, item 6','翻译后核对语法','先按英文论文逻辑写作，再逐段忠实中译，核对术语、数值、否定和不确定性。'),
 ('P7','PDF p3, item 7','图片清晰可见','沿用真实数据图的600 dpi及矢量格式，并检查最终尺寸下文字可读；不以提高dpi掩盖模糊。'),
 ('P8','PDF p3, item 8','参考文献格式和顺序一致','按医学模板先采用Vancouver数字引文与出现顺序；目标期刊确定后统一具体标点。'),
]
TEMPLATE_RULES=[
 ('T1','Title and author block','题名、作者、单位、通讯信息按实际资料填写；示例作者和邮箱全部清理。'),
 ('T2','Abstract and keywords','摘要不引用文献；初稿规划250词、4个逻辑单元，最终是否显示标签以目标期刊为准；关键词5至8个。'),
 ('T3','Headings and style','实际投稿保留研究章节，不复制Article types等模板说明章节；英文正文Times New Roman 12 pt、单倍行距，设置页码和行号。'),
 ('T4','Figures and tables','主图单独供稿，图注及可编辑表置稿件末尾，均在正文按首次出现顺序引用；补充材料独立。'),
 ('T5','Reference system','本研究按医学领域采用Vancouver序号，先使用模板医学圆括号形式；若目标刊另有要求则整篇一致调整。'),
 ('T6','Declarations','纳入数据可用性、利益冲突、作者贡献、资助、致谢及适用伦理说明；根据真实信息填写，不能直接套用无冲突声明。'),
 ('T7','Template cleanup','删除指导语、占位文本和旧公式示例；修复示例中可见的域代码/编号残留，不将其复制进成稿。'),
]


def p(doc,text,bold=False):
    para=doc.add_paragraph()
    if bold and '：' in text:
        a,b=text.split('：',1);para.add_run(a+'：').bold=True;para.add_run(b)
    else:para.add_run(text)
    return para


def h(doc,text,level=1):doc.add_paragraph(text,style=f'Heading {level}')
def page(doc,title):
    para=doc.add_paragraph(title,style='Heading 1')
    para.paragraph_format.page_break_before=True


def table(doc,headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    for j,x in enumerate(headers):t.rows[0].cells[j].text=str(x)
    for row in rows:
        cells=t.add_row().cells
        for j,x in enumerate(row):cells[j].text=str(x)
    for i,row in enumerate(t.rows):
        tr=row._tr.get_or_add_trPr();tr.append(OxmlElement('w:cantSplit'))
        if not i:tr.append(OxmlElement('w:tblHeader'))
        for j,c in enumerate(row.cells):
            c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');mar.append(e)
            pr.append(mar);borders=OxmlElement('w:tcBorders')
            for side in ['top','bottom','left','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            pr.append(borders)
            if not i:
                e=OxmlElement('w:shd');e.set(qn('w:fill'),'EDF0F2');pr.append(e)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=1.05
                para.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
                for r in para.runs:r.font.size=Pt(10);r.bold=i==0


def items(doc,rows,compact=False):
    for ident,words,title,task,evidence in rows:
        h(doc,f'{ident}  {title}  {words} words',2)
        if compact:
            p(doc,task+' 依据：'+evidence)
        else:
            p(doc,task)
            para=p(doc,'依据：'+evidence)
            for r in para.runs:r.font.size=Pt(9.5)


def build():
    OUT.mkdir(parents=True,exist_ok=True);REPORT.mkdir(parents=True,exist_ok=True)
    groups={'abstract':A,'introduction':I,'methods':M,'results':R,'discussion':D,'conclusion':C}
    targets={k:sum(x[1] for x in v) for k,v in groups.items()}
    assert targets=={'abstract':250,'introduction':1000,'methods':3000,'results':1400,'discussion':2450,'conclusion':150},targets
    inputs=[]
    for filename in ['编辑部稿件要求【第一版】.pdf','SCI论文模板（Word）.docx']:
        b=(INPUT/filename).read_bytes();inputs.append({'filename':filename,'bytes':len(b),'sha256':hashlib.sha256(b).hexdigest()})
    plan={'stage':22,'version':'V2','date':'2026-10-29','source_documents':inputs,
          'count_scope':'main English narrative from Introduction through Conclusion, including in-text citations; excludes abstract, headings, captions, table contents, reference list, declarations, Chinese translation and supplement; assumption based on supplied template link, not author-confirmed custom counting rule',
          'word_target':8000,'soft_upper_bound':10000,'section_targets':targets,
          'paragraphs':groups,'pdf_requirements':PDF_RULES,'template_requirements':TEMPLATE_RULES,
          'bibliography':{'minimum_total':25,'planned_total':40,'recent_5yr_english_min':37,'recent_3yr_english_min':25,
             'calendar_5yr':[2022,2026],'calendar_3yr':[2024,2026],'older_max_at_planned_total':3,
             'language_not_author_nationality':True,'ratio_denominator':'all unique references','completed':False,
             'next_task':'expand and verify reference evidence table before drafting'},
          'manuscript_body_drafted':False,'models_refitted':0,'old_outline_replaced':False,
          'official_guideline_url':'https://www.frontiersin.org/guidelines/author-guidelines',
          'future_documents':['English paragraphs with matching Chinese translation for review','English-only submission manuscript'],
          'next_stage':'reference evidence table and first authorized writing section'}
    (REPORT/'requirements_and_outline.json').write_text(json.dumps(plan,ensure_ascii=False,indent=2)+'\n')
    doc=Document();s=doc.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11)
    s.top_margin=Inches(.65);s.bottom_margin=Inches(.65);s.left_margin=Inches(.8);s.right_margin=Inches(.8)
    for name in ['Normal','Title','Heading 1','Heading 2']:
        st=doc.styles[name];st.font.name='Nimbus Roman';st.font.color.rgb=RGBColor(0,0,0)
        st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC')
    st=doc.styles['Normal'];st.font.size=Pt(11);st.paragraph_format.line_spacing=1.12;st.paragraph_format.space_after=Pt(5)
    doc.styles['Title'].font.size=Pt(20);doc.styles['Title'].paragraph_format.space_after=Pt(10)
    doc.styles['Heading 1'].font.size=Pt(14);doc.styles['Heading 1'].paragraph_format.space_after=Pt(8)
    st=doc.styles['Heading 2'];st.font.size=Pt(11.5);st.paragraph_format.space_before=Pt(7);st.paragraph_format.space_after=Pt(3)
    doc.core_properties.title='UCSD study writing requirements and revised outline V2';doc.core_properties.author='Tianyu Feng'
    doc.add_paragraph('写作要求理解与修订大纲 V2',style='Title')
    p(doc,'胶质瘤治疗后鉴别研究  2026年10月29日')
    p(doc,'本版依据合作者提供的3页编辑部要求、Word通用模板及补充的篇幅/文献要求，调整第21步写作大纲。研究目标、队列、输入、模型与主结果沿用已完成分析。')
    h(doc,'关键执行规则')
    for t in [
        '篇幅：英文正文以约8,000词为目标，尽量不超过10,000词。方法与讨论扩充已有流程、证据比较和局限，不以重复背景增加篇幅。',
        '计数口径：先按模板链接的正文口径规划，摘要、图表注释、表内容、参考文献、声明、补充材料与中文译文另计。最终同时给出正文及全稿英文词数，便于按合作者口径调整。',
        '文献：至少25篇；先按40篇规划，其中至少37篇为近5年英文文献、至少25篇为近3年英文文献，分别为92.5%及62.5%。这是待完成的检索目标。',
        '时间窗口：按2026年写作年暂取2022至2026年和2024至2026年；以核实的正式发表日期记录，存在在线与卷期日期时均保留。投稿年份变化后重新核对。',
        '语言：对照审阅稿每个英文段落后附对应中文译文；投稿稿仅英文。图表及其题目与注释仅英文，英文区域不混入汉字或中文标点。',
        '模板：识别为Frontiers通用格式模板，先用于版式与栏目约束。没有据此选定具体期刊；正式投稿前再核对目标期刊和文章类型要求。']:
        p(doc,t,True)
    p(doc,'目前没有影响大纲调整的阻塞问题。8,000词的计数范围为可调整假定；如果合作者指包括参考文献在内的全部英文内容，需要相应缩减正文预算。')

    page(doc,'1  编辑部八项要求的落实')
    for ident,source,rule,action in PDF_RULES:
        h(doc,f'{ident}  {rule}',2);p(doc,action);para=p(doc,'来源：'+source)
        for r in para.runs:r.font.size=Pt(9.5)

    page(doc,'2  Word模板的适用规则')
    for ident,source,action in TEMPLATE_RULES:
        h(doc,f'{ident}  {source}',2);p(doc,action)
    p(doc,'模板的图注和表格置后要求适用于投稿稿。对照审阅时可提供定位提示，但提交版会按模板组织独立图件、文末图注/表和单独补充材料。模板中的软件、动物命名等示例不自动成为本研究必需栏目。')

    page(doc,'3  篇幅与参考文献配置')
    table(doc,['Section','Target words','Paragraphs'],[[k.title(),targets[k],len(groups[k])] for k in groups]+[['Main text',8000,47]],[3.25,1.5,1.2])
    p(doc,'摘要250词是写作规划，不是已经选定期刊的上限。各段词数可在节内调配；若准确、充分表达后正文略少于8,000词，优先保留信息质量，不为凑字数增加新结果或未经执行的方法。')
    table(doc,['Reference category','Target count','Share of all'],[['All unique references',40,'100%'],['English, 2022-2026',37,'at least 92.5%'],['English, 2024-2026',25,'at least 62.5%'],['Older than 2022',3,'at most 7.5%']],[3.25,1.5,1.2])
    p(doc,'近3年文献包含在近5年文献中，不能把37与25相加。比例分母为完整去重文献表；“外文”按文献发表语言判断，不按作者国籍判断。经典原始方法或较早研究最多占少量额度，不能为满足年份而改写原始文献日期。')
    p(doc,'第21步的8项来源用于学习写作，并不满足正式论文的数量和年份结构。本轮确定配额，下一阶段补充检索并核验作者、年份、题名、DOI及其支持的具体陈述，达标后才进入正文引用。')
    p(doc,'重点扩展证据：近期治疗后终点/评估标准、MRI计算与临床融合、阴性或外部验证、影像特征稳健性、患者分组与模型评价、报告规范及数据来源。不能用大量低相关文献凑够数量。')

    page(doc,'4  摘要与引言逐段安排')
    p(doc,'以下ID仅用于V2写作管理，不是实际投稿稿的小标题。引言最终为连续段落，无小标题；正文的分节标题在后页统一安排。')
    h(doc,'Abstract  250 words',2)
    for ident,words,title,task,evidence in A:p(doc,f'{ident} {title} {words}词：{task}')
    h(doc,'Introduction  1000 words')
    items(doc,I,compact=True)

    page(doc,'5  方法段落一 研究资料与临床输入')
    items(doc,M[:7])
    page(doc,'6  方法段落二 影像计算与模型训练')
    items(doc,M[7:14])
    page(doc,'7  方法段落三 不确定性与稳健性')
    items(doc,M[14:])
    p(doc,'方法正式小节安排：2.1 Study design and data sources（M1-M3）；2.2 Reference standard（M4）；2.3 Clinical inputs（M5-M6）；2.4 MRI processing and features（M7-M11）；2.5 Model development and validation（M12-M13）；2.6 Performance and uncertainty（M14-M16）；2.7 Sensitivity analyses（M17-M19）；2.8 Reproducibility（M20）。')

    page(doc,'8  结果段落一 队列与主要比较')
    items(doc,R[:4])
    page(doc,'9  结果段落二 错误与敏感性')
    items(doc,R[4:])
    p(doc,'结果按队列、四模型主比较、概率及类别错误、敏感性组织。逐段报告实际数字与区间，不把表格逐格重述成长段，不在结果中加入未经验证的生物机制。')

    page(doc,'10  讨论段落一 证据与计算意义')
    items(doc,D[:6])
    page(doc,'11  讨论段落二 局限与结论')
    items(doc,D[6:])
    items(doc,C)

    page(doc,'12  图表与后续写作顺序')
    table(doc,['Item','Main purpose','V2 paragraphs'],[
        ['Table 1 / Figure 1','Cohort and baseline','M3, R1-R2'],['Table 2 / Figure 2','OOF model performance','M14-M16, R3-R5'],
        ['Table 3','Fusion minus M1 differences','R4'],['Figure 3 / Tables S5A-S5B','Repeated grouped validation','M18, R7'],
        ['Table S1','Imaging coverage','M3, R2'],['Table S2','Input definitions','M5-M12'],['Table S3','Calibration and class errors','R5'],
        ['Table S4 / Figure S1','Boundary sensitivity','M17, R6'],['Tables S6-S7 / Figure S2','One scan per patient','M19, R8']],[1.75,3.2,1.3])
    h(doc,'下一阶段')
    p(doc,'先补齐可核验的参考文献证据表并检查数量、年份、语言与主题覆盖，再从方法开始按本大纲逐段写作，英文在前、对应中译在后；每完成一个授权阶段总结问题与下一步。后续整合引言、讨论、结论与摘要，统一引文、图表顺序及声明。')
    p(doc,'正文与附加材料将按Title、Author information、Abstract、Keywords、Introduction、Materials and methods、Results、Discussion、Conclusion、Declarations、References、Figure legends和Tables组织。补充材料单独，正式参考文献顺序由正文首次引用确定。')
    h(doc,'最终核对项目')
    p(doc,'分别检查英文正文/全稿词数、文献总量与两个年份窗口比例、数字引文顺序、图表首次出现、英文区中文字符、公式可编辑性、语法和双语一致性。英文稿与对照稿共享同一英文原文及相同数据结果。')
    p(doc,'资料依据为合作者两份文件、2026年10月29日补充要求及模板链接的Frontiers作者指南。两份输入保留原样，文件哈希与逐条规则另存于项目核对记录。官方指南：https://www.frontiersin.org/guidelines/author-guidelines。')
    p(doc,'本轮完成要求理解和V2大纲，尚未起草论文正文，也未将40篇文献目标标记为已完成。作者、单位、排序、贡献、资助及适用声明将在正式组稿时按实际资料核对。')

    for e in [*doc.styles.element.xpath('.//w:pBdr'),*doc.element.xpath('.//w:pBdr')]:e.getparent().remove(e)
    footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    path=OUT/'UCSD_Writing_Requirements_and_Outline_V2.docx';doc.save(path)
    receipt={'stage':22,'version':'V2','deliverable':str(path.relative_to(ROOT)),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),
        'source_documents':inputs,'main_word_target':8000,'soft_upper_bound':10000,'main_paragraphs':47,'abstract_sections':4,
        'pdf_rules':8,'template_rules':7,'reference_target':40,'reference_5yr_min':37,'reference_3yr_min':25,
        'bibliography_completed':False,'manuscript_body_drafted':False,'models_refitted':0}
    (REPORT/'generation_receipt.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(receipt,ensure_ascii=False,indent=2))


if __name__=='__main__':build()
