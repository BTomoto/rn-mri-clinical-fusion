"""Audit public labels only; no cohort switch, image processing or model fitting."""
from pathlib import Path
import argparse
import collections
import gzip
import hashlib
import io
import json
import re
import urllib.request
import xml.etree.ElementTree as ET
import zipfile
import pandas as pd
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
META=ROOT/'data/metadata/public_label_audit_20261029'
CACHE=ROOT/'data/raw/reference/_label_audit_cache'
SOURCES=[
 dict(filename='ucsd_ptgbm_clinical_v3.xlsx',url='https://www.cancerimagingarchive.net/wp-content/uploads/UCSD_PTGBM-clinical-information_v3_2026-12-Mar.xlsx',sha256='bbf40bbdd289266924869052806de4d7e8fc9ee73286cdc0cc066672ef5bb51d'),
 dict(filename='ucsd_ptgbm_negative_cases.xlsx',url='https://www.cancerimagingarchive.net/wp-content/uploads/details_of_negative_cases_TCIA.xlsx',sha256='6d99db8ac1feac569e50c18ed214a1a761858fb8bfa614a29ec5f036946b96e4')]
SUPPLEMENT_URL='https://www.ebi.ac.uk/europepmc/webservices/rest/PMC13490108/supplementaryFiles?includeInlineImage=false'
SUPPLEMENT_SHA='b69dc975548e3bf52c886aea0b9a9a57b78404720808c98571d2572ac7e94ed4'


def read_verified(path,url,sha,refresh=False):
    if refresh or not path.exists():
        b=urllib.request.urlopen(urllib.request.Request(url,headers={'User-Agent':'RNPublicDataAudit/1.0'}),timeout=60).read()
        if hashlib.sha256(b).hexdigest()!=sha:raise ValueError('Public source changed; preserve pinned version and review: '+url)
        path.parent.mkdir(exist_ok=True,parents=True);path.write_bytes(b)
    b=path.read_bytes()
    if hashlib.sha256(b).hexdigest()!=sha:raise ValueError('Local source SHA mismatch: '+str(path))
    return b


def run(refresh=False):
    META.mkdir(exist_ok=True,parents=True);CACHE.mkdir(exist_ok=True,parents=True)
    receipts=[]
    for source in SOURCES:
        path=ROOT/'data/raw/reference'/source['filename'];b=read_verified(path,source['url'],source['sha256'],refresh)
        receipts.append(dict(source,local_path=str(path.relative_to(ROOT)),size_bytes=len(b),checked_date='2026-10-29',
                             role='feasibility_only_not_merged_into_OpenBTAI',license='CC BY 4.0 per TCIA collection'))
    (META/'downloaded_sources.json').write_text(json.dumps(receipts,indent=2))
    bundle=read_verified(CACHE/'rnd3_supplements.zip',SUPPLEMENT_URL,SUPPLEMENT_SHA,refresh)
    with zipfile.ZipFile(io.BytesIO(bundle)) as z:doc=z.read('Data_Sheet_1.DOCX')
    with zipfile.ZipFile(io.BytesIO(doc)) as z:
        root=ET.fromstring(z.read('word/document.xml'));ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        paragraphs=[''.join(t.text or '' for t in p.findall('.//w:t',ns)) for p in root.findall('.//w:p',ns)]
        text='\n'.join(paragraphs)
        media=[n for n in z.namelist() if n.startswith('word/media/')]
        supplement=dict(url=SUPPLEMENT_URL,bundle_sha256=SUPPLEMENT_SHA,docx_sha256=hashlib.sha256(doc).hexdigest(),
            docx_bytes=len(doc),table_count=len(root.findall('.//w:tbl',ns)),embedded_image_count=len(media),
            text_molab_id_pattern_hits=sorted(set(re.findall(r'\b(?:100\d{2}|200\d{2}|40\d{3})\b',text))),
            manually_reviewed_table_roles=['aggregate_cohort_characteristics','model_architecture','hyperparameters'],
            manually_reviewed_figure_roles=['training_curves','classifier_ROC_curves'],
            embedded_image_sha256={n:hashlib.sha256(z.read(n)).hexdigest() for n in media},
            lesion_level_joinable_label_table_observed=False,previous_access_gap_closed=True,
            interpretation='No usable per-lesion registry in this reviewed supplement; not a universal absence claim')
    (META/'rnd3_supplement_review.json').write_text(json.dumps(supplement,indent=2))
    w=load_workbook(ROOT/'data/raw/reference/ucsd_ptgbm_clinical_v3.xlsx',data_only=True,read_only=True)
    rows=list(w['Cinical info'].iter_rows(values_only=True));w.close()
    clinical=pd.DataFrame([r[:36] for r in rows[1:]],columns=[str(x).strip() for x in rows[0][:36]])
    clinical.insert(0,'source_row_clinical',range(2,len(clinical)+2))
    w=load_workbook(ROOT/'data/raw/reference/ucsd_ptgbm_negative_cases.xlsx',data_only=True,read_only=True)
    rows=list(w.active.iter_rows(values_only=True));w.close()
    negative=pd.DataFrame(rows[1:],columns=rows[0]);negative['source_row_negative_list']=range(2,len(negative)+2)
    if clinical.ID.duplicated().any() or negative.ID.duplicated().any():raise ValueError('Duplicate public timepoint identifiers')
    if not clinical.ID.str.fullmatch(r'UCSD-PTGBM-\d{4}_\d{2}').all():raise ValueError('Unexpected public ID syntax')
    if not negative.ID.isin(clinical.ID).all():raise ValueError('Negative-list ID not in current clinical version')
    allowed={'NS','PsP (clinically confirmed)','PsP (bx proven)','Radiation necrosis'}
    if not set(negative.Dx).issubset(allowed):raise ValueError('New subtype requires source review')
    if len(clinical)!=243 or len(negative)!=51:raise ValueError('Published partition no longer matches source totals')
    table=clinical.merge(negative.rename(columns={'Dx':'negative_list_dx'}),on='ID',how='left',validate='one_to_one')
    table['patient_id']=table.ID.str.rsplit('_',n=1).str[0]
    table['broad_tumor_label_candidate']=table.negative_list_dx.isna().astype(int)
    table['broad_label_source']=table.negative_list_dx.map(lambda x:'explicit_negative_case_list' if pd.notna(x) else 'derived_from_published_243_equals_192_plus_51_partition')
    table['rn_status_for_audit']=table.negative_list_dx.map(lambda x:'explicit_RN_record' if x=='Radiation necrosis' else 'not_adjudicated_as_RN_negative')
    table['analysis_question_approved']=0;table['use_for_model']=0
    cols=['ID','patient_id','source_row_clinical','source_row_negative_list','negative_list_dx','broad_tumor_label_candidate',
          'broad_label_source','rn_status_for_audit',"Patient's Age",'Sex at birth','Primary Diagnosis','analysis_question_approved','use_for_model']
    table[cols].to_csv(ROOT/'data/processed/ucsd_ptgbm_label_feasibility.csv',index=False)
    raw_counts=negative.Dx.value_counts().to_dict()
    published_counts={'NS':28,'PsP (clinically confirmed)':9,'PsP (bx proven)':4,'Radiation necrosis':10}
    comparison=pd.DataFrame([dict(subtype=k,published_description_count=published_counts[k],downloaded_file_count=raw_counts.get(k,0),difference=raw_counts.get(k,0)-published_counts[k]) for k in published_counts])
    comparison.to_csv(ROOT/'data/processed/ucsd_ptgbm_subtype_count_comparison.csv',index=False)
    rn=table[table.negative_list_dx=='Radiation necrosis'];change=table[table.broad_tumor_label_candidate==0];tumor=table[table.broad_tumor_label_candidate==1]
    summary=dict(step='public_label_feasibility_only',date='2026-10-29',rnd3_supplement_access_gap_closed=True,
        rnd3_joinable_label_table_found=False,original_OpenBTAI_model_ready=False,
        ucsd_timepoints=len(table),ucsd_patients=int(table.patient_id.nunique()),
        explicit_rn_timepoints=len(rn),explicit_rn_patients=int(rn.patient_id.nunique()),
        downloaded_subtype_counts=raw_counts,published_subtype_counts=published_counts,subtype_count_discrepancy=True,
        all_negative_list_ids_match_current_clinical=True,tumor_candidate_timepoints=len(tumor),tumor_candidate_patients=int(tumor.patient_id.nunique()),
        treatment_change_timepoints=len(change),treatment_change_patients=int(change.patient_id.nunique()),
        patients_with_both_timepoint_classes=int(table.groupby('patient_id').broad_tumor_label_candidate.nunique().gt(1).sum()),
        broad_label_limit='Tumor labels are a documented-partition derivation, not individual pathology adjudication; confirm protocol before use',
        rn_only_model_recommended=False,broad_diagnostic_question_is_a_candidate=True,
        scope_change_approved=False,models_fitted=False,mri_downloaded_for_new_cohort=False,
        next_step_requires_user_approval=True,
        proposed_next_step='If approved, change to high-grade-glioma post-treatment tumor-versus-treatment-change diagnosis; verify patient-grouped eligibility and run a prespecified simple clinical logistic baseline, stopping if label checks fail',
        new_question_is_not_brain_metastasis_future_RN_risk=True)
    (ROOT/'reports/public_label_feasibility_20261029.json').write_text(json.dumps(summary,indent=2,allow_nan=False))
    print(json.dumps(summary,indent=2),flush=True);return summary


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--refresh',action='store_true')
    run(parser.parse_args().refresh)
