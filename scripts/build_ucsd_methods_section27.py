"""Write M14-M15 with a read-only audit of saved metrics and bootstrap records."""
from pathlib import Path
import argparse,hashlib,json,re,sys,zipfile
import numpy as np
import pandas as pd
from lxml import etree
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from build_ucsd_methods_section25 import el,mr

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'outputs/ucsd-manuscript-step29-20261029'
REPORT=ROOT/'reports/ucsd_manuscript_step29'
PARAGRAPHS=[
 {'id':'M14','citation_ids':[],
  'english': 'Performance was calculated at the examination level by pooling OOF probabilities from all five folds. The primary incremental comparison was the Fusion-minus-M1 difference in area under the receiver operating characteristic curve (ROC AUC), evaluated on the same examinations (Equation 8). We also reported ROC AUC for each model, average precision (AP) for each diagnostic class, Brier score and logarithmic loss. Treatment-related-change AP used treatment-related change as the positive class and one minus the tumor probability as its score. AP was computed from precision weighted by increments in recall. Brier score and logarithmic loss summarized probability error, with lower values indicating better performance. At the fixed threshold of 0.5, probabilities at least 0.5 were classified as tumor. Confusion counts, tumor sensitivity, recognition of treatment-related change, precision for each class, balanced accuracy and overall accuracy were calculated without threshold optimization.',
  'chinese': '将五折的OOF概率汇总，在检查层面计算性能。主要增量比较为同一批检查中Fusion减去M1的受试者工作特征曲线下面积（ROC AUC）差值（Equation 8）。同时报告每个模型的ROC AUC、两类诊断各自的平均精确率（AP）、Brier评分及对数损失。治疗相关改变的AP以治疗相关改变为正类，以1减去肿瘤概率作为评分。AP由按召回率增量加权的精确率计算。Brier评分与对数损失概括概率误差，数值越低表示表现越好。在固定的0.5阈值下，大于或等于0.5的概率归为肿瘤。计算混淆计数、肿瘤敏感度、治疗相关改变识别率、各类精确率、平衡准确率及总体准确率，不进行阈值优化。'},
 {'id':'M15','citation_ids':[],
  'english': 'Uncertainty was summarized using 2000 patient-cluster bootstrap attempts with random seed 20261030. Each attempt sampled 135 patient groups with replacement and retained all available examinations for each selected group, including repeated copies when a group was sampled more than once. All models used identical resampled examination indices, enabling paired metric differences. The saved OOF probabilities remained fixed; prediction models and their preprocessing were not refitted within resampling. Attempts containing only one diagnostic class were ineligible and were not replaced by additional draws. For ROC AUC, treatment-related-change AP, Brier score and logarithmic loss, and their paired differences, 95% percentile intervals were defined by the 2.5th and 97.5th percentiles of eligible bootstrap values. These intervals represented patient-sampling uncertainty conditional on the saved OOF predictions. They did not capture the full variability of retraining the models or choosing another fold partition. Metric differences were retained in their stated directions, so a negative Brier-score difference favored Fusion in the Fusion-minus-M1 comparison.',
  'chinese': '采用2000次患者聚类bootstrap抽样尝试概括不确定性，随机种子为20261030。每次有放回地抽取135个患者组，并保留每个抽中组的全部可用检查；同一组被多次抽中时，其检查也重复计入。所有模型使用完全相同的重采样检查索引，以计算配对指标差值。已保存的OOF概率保持固定，抽样过程中不重新拟合预测模型或其预处理。仅包含一种诊断类别的抽样不符合条件，且不追加抽样替代。针对ROC AUC、治疗相关改变AP、Brier评分及对数损失，以及这些指标的配对差值，取符合条件bootstrap值的第2.5和第97.5百分位数，得到95%百分位区间。这些区间表示条件于已保存OOF预测的患者抽样不确定性，不涵盖重新训练模型或选择其他分折的全部变异。指标差值保留既定相减方向，因此在Fusion减去M1的比较中，Brier评分差值为负有利于Fusion。'}
]

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def wc(s):return len(re.findall(r"[A-Za-z0-9]+(?:['/-][A-Za-z0-9]+)*",s))
def close(a,b):assert np.isclose(a,b,rtol=1e-12,atol=1e-12),(a,b)

def audit():
    old=json.loads((ROOT/'reports/ucsd_manuscript_step28/methods_section26.json').read_text());frozen=old['source_audit']['frozen_file_sha256'].copy()
    frozen.update(json.loads((ROOT/'reports/ucsd_manuscript_step28/artifact_verification.json').read_text())['sha256'])
    for p,h in frozen.items():assert sha(ROOT/p)==h,p
    sys.path.insert(0,str(ROOT/'src'))
    from sklearn.metrics import precision_recall_curve
    from rn_public.ucsd_fusion import EVALUATED,BOOT_METRICS,PAIRS,metric_vector,SEED
    pred=pd.read_csv(ROOT/'data/processed/ucsd_fusion_oof.csv');y=pred.tumor_label.to_numpy()
    metrics=json.loads((ROOT/'reports/ucsd_fusion/metrics.json').read_text());assert len(pred)==182 and pred.validation_group.nunique()==135
    primaries=[p for p in metrics['paired_comparisons'] if p['primary_comparison']]
    assert len(primaries)==1 and primaries[0]['comparison']=='Fusion_minus_M1' and primaries[0]['metric']=='roc_auc'
    main_checks={}
    for name in ['M0','M1','MRI','Fusion']:
        p=pred[name+'_oof_probability'].to_numpy();assert np.all((p>0)&(p<1))
        pairs=p[y==1][:,None]-p[y==0][None,:]
        auc=float(np.mean((pairs>0)+.5*(pairs==0)))
        def ap(label,score):
            precision,recall,_=precision_recall_curve(label,score)
            return float(-np.sum(np.diff(recall)*precision[:-1]))
        positive=p>=.5;tp=int(np.sum(positive&(y==1)));fp=int(np.sum(positive&(y==0)))
        tn=int(np.sum(~positive&(y==0)));fn=int(np.sum(~positive&(y==1)))
        values={'roc_auc':auc,'ap_tumor':ap(y,p),'ap_treatment_change':ap(1-y,1-p),
          'brier':float(np.mean((p-y)**2)),'log_loss':float(-np.mean(y*np.log(p)+(1-y)*np.log1p(-p))),
          'tn':tn,'fp':fp,'fn':fn,'tp':tp,'sensitivity_tumor':tp/(tp+fn),'specificity_treatment_change':tn/(tn+fp),
          'balanced_accuracy':.5*(tp/(tp+fn)+tn/(tn+fp)),'accuracy':(tp+tn)/len(y),
          'precision_tumor':None if tp+fp==0 else tp/(tp+fp),'precision_treatment_change':None if tn+fn==0 else tn/(tn+fn)}
        for key,value in values.items():
            saved=metrics['metrics'][name][key]
            assert saved is None if value is None else np.isclose(value,saved,rtol=1e-12,atol=1e-12)
        main_checks[name]=values
    close(primaries[0]['difference'],main_checks['Fusion']['roc_auc']-main_checks['M1']['roc_auc'])
    boot=pd.read_csv(ROOT/'reports/ucsd_fusion/paired_patient_bootstrap.csv');assert boot.shape==(2000,46)
    assert boot.draw.tolist()==list(range(1,2001)) and metrics['bootstrap_attempts']==metrics['bootstrap_valid']==2000
    assert metrics['bootstrap_one_class_skipped']==0 and metrics['bootstrap_models_refitted'] is False
    assert metrics['threshold']==.5 and SEED+1==20261030
    columns=[f'{name}__{m}' for name in EVALUATED for m in BOOT_METRICS]
    columns += [f'{a}_minus_{b}__{m}' for a,b in PAIRS for m in BOOT_METRICS]
    assert boot.columns.tolist()==['draw','resampled_scan_count',*columns]
    for a,b in PAIRS:
        for metric in BOOT_METRICS:
            np.testing.assert_allclose(boot[f'{a}_minus_{b}__{metric}'],boot[f'{a}__{metric}']-boot[f'{b}__{metric}'],rtol=1e-12,atol=1e-12)
    quant=boot[columns].quantile([.025,.975]);saved=metrics['patient_cluster_bootstrap_95ci']
    for key in columns:close(quant.loc[.025,key],saved[key]['lower']);close(quant.loc[.975,key],saved[key]['upper'])
    for p in metrics['paired_comparisons']:
        key=p['comparison']+'__'+p['metric'];close(p['lower'],saved[key]['lower']);close(p['upper'],saved[key]['upper'])
    clusters=[np.asarray(ix) for ix in pred.groupby('validation_group',sort=True).indices.values()]
    assert len(clusters)==135
    for cluster in clusters:
        assert pred.iloc[cluster].patient_id.nunique()==pred.iloc[cluster].brats_patient_id.nunique()==1
    rng=np.random.default_rng(20261030);selected_checks=[]
    for draw in range(1,2001):
        selections=rng.integers(0,len(clusters),len(clusters));indices=np.concatenate([clusters[j] for j in selections])
        assert len(indices)==int(boot.iloc[draw-1].resampled_scan_count) and np.unique(y[indices]).size==2
        if draw in [1,1000,2000]:
            for name in EVALUATED:
                actual=metric_vector(y[indices],pred[name+'_oof_probability'].to_numpy()[indices])
                for metric,value in zip(BOOT_METRICS,actual):close(value,boot.iloc[draw-1][name+'__'+metric])
            selected_checks.append({'draw':draw,'patient_group_draws':len(selections),'unique_groups':int(np.unique(selections).size),'examinations':len(indices)})
    extra=['reports/ucsd_fusion/paired_patient_bootstrap.csv','src/rn_public/ucsd_clinical_baseline.py']
    frozen.update({p:sha(ROOT/p) for p in extra})
    return {'frozen_file_sha256':frozen,'primary_comparison':'Fusion_minus_M1__roc_auc','main_metric_values_checked':60,
      'main_model_metric_reconstruction':main_checks,'bootstrap_attempts_checked':2000,'bootstrap_valid':2000,'single_class_skipped':0,
      'bootstrap_group_draws_per_attempt':135,'bootstrap_seed':20261030,'paired_difference_values_checked':40000,
      'percentile_endpoints_checked':88,'sampled_draw_metric_values_checked':72,'sampled_draws_checked':selected_checks,
      'resampled_examination_count_range':[int(boot.resampled_scan_count.min()),int(boot.resampled_scan_count.max())],
      'probabilities_fixed_during_resampling':True,'models_refitted':0,'new_hypothesis_tests':0,
      'documentation':['https://scikit-learn.org/1.8/modules/generated/sklearn.metrics.average_precision_score.html','https://scikit-learn.org/1.8/modules/generated/sklearn.metrics.brier_score_loss.html','https://scikit-learn.org/1.8/modules/generated/sklearn.metrics.log_loss.html']}

def equation():
    def model(name):return el('sSub',el('e',mr('AUC',True)),el('sub',mr(name,True)))
    return [mr('Δ'),mr('AUC',True),mr(' = '),model('Fusion'),mr(' − '),model('M1')]

def build():
    OUT.mkdir(parents=True,exist_ok=True);REPORT.mkdir(parents=True,exist_ok=True);a=audit()
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(.79)
    for name in ['Normal','Title','Heading 1']:
        st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        rf=st.element.get_or_add_rPr().rFonts;rf.set(qn('w:eastAsia'),'Noto Serif CJK SC')
        for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:rf.attrib.pop(qn('w:'+key),None)
    normal=doc.styles['Normal'];normal.font.size=Pt(12);normal.paragraph_format.line_spacing=1.;normal.paragraph_format.space_after=Pt(9)
    doc.styles['Title'].font.size=Pt(16);doc.styles['Title'].paragraph_format.space_after=Pt(12)
    doc.styles['Heading 1'].font.size=Pt(13);doc.styles['Heading 1'].paragraph_format.space_before=Pt(8);doc.styles['Heading 1'].paragraph_format.space_after=Pt(9)
    doc.core_properties.title='Materials and methods 2.7 Performance metrics and conditional uncertainty';doc.core_properties.author='Tianyu Feng'
    doc.add_paragraph('Materials and methods',style='Title');doc.add_paragraph('2.7 Performance metrics and conditional uncertainty',style='Heading 1')
    for item in PARAGRAPHS:
        en=doc.add_paragraph(item['english']);en.paragraph_format.page_break_before=item['id']=='M15'
        en.paragraph_format.keep_together=True;en.paragraph_format.keep_with_next=True
        zh=doc.add_paragraph(item['chinese']);zh.paragraph_format.keep_together=True
        for r in zh.runs:r.font.size=Pt(11)
        if item['id']=='M14':
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_together=True
            p.paragraph_format.space_before=Pt(6);p._p.append(el('oMath',equation()));p.add_run('    (8)')
    for e in [*doc.styles.element.xpath('.//w:pBdr'),*doc.element.xpath('.//w:pBdr')]:e.getparent().remove(e)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    path=OUT/'UCSD_Manuscript_Methods_2_7_EN_ZH.docx';doc.save(path)
    (OUT/'UCSD_Manuscript_Methods_2_7_EN.txt').write_text('\n\n'.join(p['english'] for p in PARAGRAPHS)+'\n')
    d={'stage':29,'date':'2026-10-29','subsection':'2.7 Performance metrics and conditional uncertainty','v2_paragraph_ids':['M14','M15'],'paragraphs':PARAGRAPHS,
      'english_word_count':sum(wc(p['english']) for p in PARAGRAPHS),'paragraph_word_counts':{p['id']:wc(p['english']) for p in PARAGRAPHS},
      'word_count_scope':'English narrative; excludes equation, headings and Chinese translation','v2_approximate_target_words':300,
      'equations':{'count':1,'format':'native editable OMML','numbers':[8],'meaning':'Prespecified primary Fusion-minus-M1 ROC AUC difference'},
      'citation_mapping':{},'new_manuscript_references':0,'cumulative_unique_cited_sources':4,'source_audit':a,'models_refitted':0,'new_hypothesis_tests':0,
      'next_section':'Calibration assessment (M16)','docx':str(path.relative_to(ROOT)),'docx_sha256':sha(path)}
    (REPORT/'methods_section27.json').write_text(json.dumps(d,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'docx':d['docx'],'english_word_count':d['english_word_count'],'paragraph_word_counts':d['paragraph_word_counts'],'prior_files':len(a['frozen_file_sha256'])},indent=2))

def verify():
    from pypdf import PdfReader
    d=json.loads((REPORT/'methods_section27.json').read_text())
    for p,h in d['source_audit']['frozen_file_sha256'].items():assert sha(ROOT/p)==h,p
    path=ROOT/d['docx'];assert sha(path)==d['docx_sha256'];texts=[p.text for p in Document(path).paragraphs]
    for p in PARAGRAPHS:
        assert p['english'] in texts and p['chinese'] in texts and texts.index(p['chinese'])==texts.index(p['english'])+1
        assert not re.search(r'[\u3400-\u9fff，。；：（）【】、]',p['english'])
    assert sum(wc(p['english']) for p in PARAGRAPHS)==d['english_word_count'] and '    (8)' in texts
    ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
    with zipfile.ZipFile(path) as z:
        t=etree.fromstring(z.read('word/document.xml'));assert len(t.xpath('//m:oMath',namespaces=ns))==1
        assert len(t.xpath('//m:sSub',namespaces=ns))==2 and not any(n.startswith('word/media/') for n in z.namelist())
    pages=len(PdfReader(OUT/'docx_render/UCSD_Manuscript_Methods_2_7_EN_ZH.pdf').pages);assert pages==2
    v={'success':True,'stage':29,'date':'2026-10-29','subsection':d['subsection'],'paragraph_ids':['M14','M15'],
      'english_word_count':d['english_word_count'],'paragraph_word_counts':d['paragraph_word_counts'],'matched_chinese_paragraphs':2,
      'word_pages':pages,'visual_review_pages':[1,2],'native_editable_equations':1,'equation_numbers':[8],'equation_images':0,
      'prior_files_unchanged':len(d['source_audit']['frozen_file_sha256']),'main_metric_values_checked':60,'bootstrap_attempts_checked':2000,
      'paired_difference_values_checked':40000,'percentile_endpoints_checked':88,'sampled_draw_metric_values_checked':72,
      'conditional_uncertainty_disclosed':True,'models_refitted':0,'new_hypothesis_tests':0,'new_manuscript_references':0,'next_section':d['next_section'],
      'sha256':{str(p.relative_to(ROOT)):sha(p) for p in OUT.iterdir() if p.is_file() and p.suffix in {'.docx','.txt'}}}
    (REPORT/'artifact_verification.json').write_text(json.dumps(v,ensure_ascii=False,indent=2)+'\n');print(json.dumps(v,indent=2))

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--verify',action='store_true');args=parser.parse_args()
    verify() if args.verify else build()
