"""Draft D2 using primary-source evidence with explicit reading limits."""
from pathlib import Path
import argparse, hashlib, json, re, zipfile
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'outputs/ucsd-manuscript-step42-20261029'
REPORT = ROOT/'reports/ucsd_manuscript_step42'
EVIDENCE = REPORT/'literature_evidence.json'
MAPPING = {'6': 'S06', '7': 'S12', '8': 'S19'}
PARAGRAPHS = [{
    'id': 'D2', 'citation_ids': ['S06', 'S12', 'S19'],
    'english': 'The observed increment has parallels in studies combining quantitative MRI with non-imaging information, although their predictor sets differ substantially from ours. In a single-center cohort of 222 patients with grade 3 or 4 glioma, Zhou et al. integrated a 17-feature radiomics score with relative cerebral blood volume, apparent diffusion coefficient, neutrophil-to-lymphocyte ratio, MGMT promoter methylation and temozolomide treatment (6). Their stratified 7:3 internal split yielded validation AUCs of 0.811 for the integrated model and 0.744 for the model excluding the radiomics score. That comparison therefore assessed radiomics beyond a baseline already containing advanced imaging and molecular information. McKenney et al. studied 74 patients with IDH-wildtype glioblastoma who underwent repeat surgery, using contrast-enhanced T1 and dynamic contrast-enhanced perfusion radiomics with MGMT status (7). Their repeated nested cross-validation showed higher accuracy for the combined classifier than for radiomics alone, with lesion labels determined by histopathology. These reports motivate evaluating complementary information but do not establish that our limited clinical inputs or eight MRI descriptors capture the same signal. Our outcome also encompasses broader treatment-related change, limiting alignment with their pseudoprogression endpoints. Related work by Reecher et al. used UCSD-PTGBM itself (8); its shared parent dataset precludes treating it as independent confirmation. Our study provides a paired assessment of fixed, low-dimensional clinical-MRI fusion on matched examinations. Differences in patient selection, reference standards, input availability and validation designs prevent direct ranking of our model against these reports by AUC.',
    'chinese': '本研究观察到的增量，与将定量MRI和非影像信息相结合的研究存在相似之处，但其预测变量组合与我们的研究有明显差异。在纳入222名3级或4级胶质瘤患者的单中心队列中，Zhou等将17项特征构成的影像组学评分与相对脑血容量、表观扩散系数、中性粒细胞与淋巴细胞比值、MGMT启动子甲基化及替莫唑胺治疗相整合（6）。其分层7∶3内部划分中，整合模型的验证AUC为0.811，不含影像组学评分的模型为0.744。因此，该比较评估的是影像组学在已经包含高级影像和分子信息的基线模型之上的价值。McKenney等研究了74名接受再次手术的IDH野生型胶质母细胞瘤患者，将增强T1及动态对比增强灌注影像组学与MGMT状态结合（7）。其重复嵌套交叉验证显示，组合分类器的准确率高于单独影像组学分类器，病灶标签由组织病理学确定。这些研究支持评估信息互补性的研究思路，但不能证明我们有限的临床输入或8项MRI描述特征捕获了相同信号。我们的结局还涵盖更广泛的治疗相关改变，因此与这些研究的假性进展终点并不完全一致。Reecher等的相关工作也使用了UCSD-PTGBM（8）；由于共享同一母数据集，不能将其视为独立确认。我们的研究在匹配检查上，对固定低维临床-MRI融合进行了配对评估。患者选择、参考标准、可用输入和验证设计的差异，使得无法依据AUC将我们的模型与这些研究直接排名比较。'
}]

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def wc(text):
    return len(re.findall(r"[A-Za-z0-9]+(?:['/-][A-Za-z0-9]+)*", text))

def audit():
    previous = json.loads((ROOT/'reports/ucsd_manuscript_step41/discussion_section41.json').read_text())
    frozen = previous['source_audit']['frozen_file_sha256'].copy()
    frozen.update(json.loads((ROOT/'reports/ucsd_manuscript_step41/artifact_verification.json').read_text())['sha256'])
    frozen[str(EVIDENCE.relative_to(ROOT))] = sha(EVIDENCE)
    for p, h in frozen.items(): assert sha(ROOT/p) == h, p
    evidence = json.loads(EVIDENCE.read_text()); studies = {r['reference_id']: r for r in evidence['studies']}
    assert set(studies) == set(MAPPING.values()) and evidence['cumulative_unique_cited_sources'] == 8
    library = json.loads((ROOT/'reports/ucsd_manuscript_step23/reference_library.json').read_text())
    refs = {r['id']: r for r in library['references']}
    selected = []
    for number, rid in MAPPING.items():
        ref = refs[rid]; record = studies[rid]
        assert ref['doi'] == record['doi'] and record['provisional_citation_number'] == int(number)
        assert ref['english_flag'] == ref['recent_5yr'] == 1
        selected.append({'number': int(number), 'reference_id': rid, 'citation': ref['citation'],
                         'doi': ref['doi'], 'eligibility_year': ref['eligibility_year'],
                         'read_level_this_stage': record['read_level_this_stage']})
    z = studies['S06']['facts']; m = studies['S12']['facts']; r = studies['S19']['facts']
    assert z['patients'] == 222 and z['radiomics_score_features'] == 17
    assert z['validation_integrated_auc'] == .811 and z['validation_baseline_without_radscore_auc'] == .744
    assert len(z['integrated_inputs']) == 6 and len(z['baseline_inputs']) == 5
    assert m['patients'] == 74 and m['reported_combined_vs_radiomics_accuracy_p_value'] == .03
    assert 'nested three-fold' in m['validation'] and r['parent_dataset'] == 'UCSD-PTGBM'
    assert studies['S19']['read_level_this_stage'] == 'Primary indexed abstract and bibliographic record only'
    main = json.loads((ROOT/'reports/ucsd_fusion/metrics.json').read_text())
    assert len(main['models']['MRI']) == 8 and len(main['models']['M1']) == 5
    assert main['same_cases_and_folds'] and not main['independent_external_validation']
    prior = []
    for pattern in ['methods_section*.json', 'results_section*.json', 'discussion_section*.json']:
        for p in ROOT.glob('reports/ucsd_manuscript_step*/'+pattern):
            if p.parent == REPORT: continue
            prior.append(json.loads(p.read_text()))
    prior_ids = {rid for s in prior for p in s['paragraphs'] for rid in p['citation_ids']}
    assert prior_ids == {'S01', 'S02', 'S03', 'S32', 'S40'} and not prior_ids.intersection(MAPPING.values())
    return {'frozen_file_sha256': frozen, 'reference_records': selected,
            'source_evidence_record': str(EVIDENCE.relative_to(ROOT)),
            'full_texts_checked_for_current_claims': ['S06', 'S12'],
            'abstract_only_limited_claims': {'S19': 'Same named parent dataset only; no detailed validation or exact-overlap claim'},
            'published_accuracy_contrast_not_misrepresented_as_AUC_test': True,
            'new_numeric_literature_values_checked': 5,
            'same_parent_dataset_is_independent_confirmation': False,
            'cross_paper_AUC_ranking_performed': False,
            'comparison_design_inferences_separated_from_reported_results': True,
            'previous_unique_cited_sources': sorted(prior_ids),
            'models_refitted': 0, 'new_hypothesis_tests': 0,
            'verification_scope': 'Validate source-to-claim mapping against manually reviewed primary evidence and frozen own-study records; no model fitting or new statistical comparisons'}

def build():
    a = audit(); OUT.mkdir(parents=True, exist_ok=True); REPORT.mkdir(parents=True, exist_ok=True)
    doc = Document(); sec = doc.sections[0]
    sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(.79)
    for name in ['Normal', 'Title', 'Heading 1']:
        st = doc.styles[name]; st.font.name = 'Times New Roman'; st.font.color.rgb = RGBColor(0, 0, 0)
        rf = st.element.get_or_add_rPr().rFonts; rf.set(qn('w:eastAsia'), 'Noto Serif CJK SC')
        for key in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']: rf.attrib.pop(qn('w:'+key), None)
    normal = doc.styles['Normal']; normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.; normal.paragraph_format.space_after = Pt(9)
    doc.styles['Title'].font.size = Pt(16); doc.styles['Title'].paragraph_format.space_after = Pt(12)
    heading = doc.styles['Heading 1']; heading.font.size = Pt(13)
    heading.paragraph_format.space_before = Pt(8); heading.paragraph_format.space_after = Pt(9)
    doc.core_properties.title = 'Discussion 4.2 Comparison with fusion studies'; doc.core_properties.author = 'Tianyu Feng'
    doc.add_paragraph('Discussion', style='Title'); doc.add_paragraph('4.2 Comparison with fusion studies', style='Heading 1')
    for item in PARAGRAPHS:
        en = doc.add_paragraph(item['english']); en.paragraph_format.keep_together = True
        en.paragraph_format.keep_with_next = True
        zh = doc.add_paragraph(item['chinese']); zh.paragraph_format.keep_together = True
        for run in zh.runs: run.font.size = Pt(11)
    title = doc.add_paragraph('References', style='Heading 1'); title.paragraph_format.page_break_before = True
    for ref in a['reference_records']:
        p = doc.add_paragraph(str(ref['number'])+'. '+ref['citation']); p.paragraph_format.keep_together = True
    for e in [*doc.styles.element.xpath('.//w:pBdr'), *doc.element.xpath('.//w:pBdr')]: e.getparent().remove(e)
    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    path = OUT/'UCSD_Manuscript_Discussion_4_2_EN_ZH.docx'; doc.save(path)
    (OUT/'UCSD_Manuscript_Discussion_4_2_EN.txt').write_text('\n\n'.join(p['english'] for p in PARAGRAPHS)+'\n')
    d = {'stage': 42, 'date': '2026-10-29', 'subsection': '4.2 Comparison with fusion studies', 'v2_paragraph_ids': ['D2'],
         'paragraphs': PARAGRAPHS, 'english_word_count': sum(wc(p['english']) for p in PARAGRAPHS),
         'paragraph_word_counts': {p['id']: wc(p['english']) for p in PARAGRAPHS},
         'word_count_scope': 'English narrative including in-text citation numbers; excludes headings, reference list and Chinese translation',
         'v2_approximate_target_words': 250, 'equations': {'count': 0, 'numbers': [], 'cumulative_previous_numbered_groups': 10},
         'citation_mapping': MAPPING, 'citation_numbering': 'Provisional first-appearance numbers 6-8; renumber globally when Introduction is assembled',
         'new_manuscript_references': 3, 'cumulative_unique_cited_sources': 8,
         'source_audit': a, 'models_refitted': 0, 'new_hypothesis_tests': 0,
         'next_section': 'Discussion 4.3 Conflicting evidence (D3)',
         'docx': str(path.relative_to(ROOT)), 'docx_sha256': sha(path)}
    (REPORT/'discussion_section42.json').write_text(json.dumps(d, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps({'docx': d['docx'], 'english_word_count': d['english_word_count'], 'frozen_files': len(a['frozen_file_sha256'])}))

def verify():
    from pypdf import PdfReader
    d = json.loads((REPORT/'discussion_section42.json').read_text()); a = audit()
    assert a == d['source_audit']
    path = ROOT/d['docx']; assert sha(path) == d['docx_sha256']
    doc = Document(path); texts = [p.text for p in doc.paragraphs]
    for p in PARAGRAPHS:
        assert p['english'] in texts and p['chinese'] in texts
        assert texts.index(p['chinese']) == texts.index(p['english'])+1
        assert not re.search(r'[\u3400-\u9fff，。；：（）【】、]', p['english'])
    assert sum(wc(p['english']) for p in PARAGRAPHS) == d['english_word_count']
    assert re.findall(r'\((\d+)\)', PARAGRAPHS[0]['english']) == list(MAPPING)
    for ref in a['reference_records']: assert str(ref['number'])+'. '+ref['citation'] in texts
    assert len(doc.tables) == 0 and len(doc.sections) == 1
    with zipfile.ZipFile(path) as z:
        xml = etree.fromstring(z.read('word/document.xml'))
        assert not xml.xpath('//m:oMath', namespaces={'m': 'http://schemas.openxmlformats.org/officeDocument/2006/math'})
        assert not [n for n in z.namelist() if n.startswith('word/media/')]
    pdf = PdfReader(OUT/'docx_render/UCSD_Manuscript_Discussion_4_2_EN_ZH.pdf'); pages = len(pdf.pages); assert pages == 2
    assert '4.2 Comparison' in pdf.pages[0].extract_text() and 'References' in pdf.pages[1].extract_text()
    v = {'success': True, 'stage': 42, 'date': '2026-10-29', 'subsection': d['subsection'], 'paragraph_ids': ['D2'],
         'english_word_count': d['english_word_count'], 'matched_chinese_paragraphs': 1, 'word_pages': pages,
         'visual_review_pages': [1, 2], 'native_editable_equations': 0, 'equation_images': 0,
         'frozen_files_unchanged': len(a['frozen_file_sha256']), 'prior_files_unchanged': len(a['frozen_file_sha256'])-1,
         'full_texts_checked_for_current_claims': ['S06', 'S12'], 'limited_abstract_claim_sources': ['S19'],
         'published_numeric_values_checked': 5, 'new_manuscript_references': 3,
         'provisional_citation_mapping': MAPPING, 'cumulative_unique_cited_sources': 8,
         'cross_paper_AUC_ranking_performed': False, 'models_refitted': 0, 'new_hypothesis_tests': 0,
         'next_section': d['next_section'],
         'sha256': {str(p.relative_to(ROOT)): sha(p) for p in OUT.iterdir() if p.is_file() and p.suffix in {'.docx', '.txt'}}}
    (REPORT/'artifact_verification.json').write_text(json.dumps(v, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps(v, indent=2))

if __name__ == '__main__':
    parser = argparse.ArgumentParser(); parser.add_argument('--verify', action='store_true'); args = parser.parse_args()
    verify() if args.verify else build()
