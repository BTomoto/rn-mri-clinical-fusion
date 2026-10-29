"""Draft R6 from frozen boundary-sensitivity predictions and paired intervals."""
from pathlib import Path
import argparse,hashlib,json,re,sys,zipfile
import numpy as np
import pandas as pd
from lxml import etree
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'outputs/ucsd-manuscript-step38-20261029'
REPORT=ROOT/'reports/ucsd_manuscript_step38'
PARAGRAPHS=[
 {'id':'R6','citation_ids':[],
  'english': 'All 182 examinations retained complete MRI features after the prespecified erosion, allowing comparison on the original case set and fold assignments. The eroded MRI model had an OOF AUC of 0.765, compared with 0.762 for the original MRI model (Supplementary Table S4; Figure S1). Eroded Fusion had an AUC of 0.788 (95% patient-bootstrap percentile interval, 0.706 to 0.867), compared with 0.782 for original Fusion. The paired eroded-minus-original Fusion AUC difference was 0.0062, with an interval of -0.0051 to 0.0183 that included zero. Eroded Fusion had a Brier score of 0.141 and treatment-related-change AP of 0.529, and correctly recognized 14 of 41 treatment-related-change examinations at the fixed threshold. The corresponding original Fusion values were 0.144, 0.505 and 12 of 41. These findings described sensitivity to one boundary perturbation under fixed processing and did not establish superiority of erosion or segmentation accuracy. The original ROI remained the primary analysis.',
  'chinese': '预设腐蚀后，全部182次检查仍具备完整MRI特征，因此可以在原有病例集合与分折划分上进行比较。腐蚀后MRI模型的OOF AUC为0.765，原MRI模型为0.762（Supplementary Table S4；Figure S1）。腐蚀后Fusion的AUC为0.788（95%患者bootstrap百分位区间0.706至0.867），原Fusion为0.782。腐蚀后减去原Fusion的配对AUC差值为0.0062，区间为−0.0051至0.0183，包含零。腐蚀后Fusion的Brier评分为0.141，治疗相关改变AP为0.529，在固定阈值下正确识别41次治疗相关改变检查中的14次；原Fusion的相应数值为0.144、0.505和12/41。上述结果描述固定处理规则下对单次边界扰动的敏感性，不能据此认定腐蚀更优或证明分割准确性。原ROI分析仍保留为主分析。'}
]
MODELS=['MRI','Fusion','MRI_eroded','Fusion_eroded']
MODEL_LABELS=['Original MRI','Original Fusion','Eroded MRI','Eroded Fusion']
FIGURE=ROOT/'outputs/ucsd-paper-materials-20261029/figures/FigureS1_Boundary_Sensitivity.png'
FIGURE_CAPTION='Figure S1. Boundary sensitivity in the same 182 examinations from 135 patients. Points show pooled OOF AUCs for the original and eroded MRI and Fusion models. Error bars are 95% patient-cluster bootstrap percentile intervals conditional on fixed OOF probabilities, without model refitting. The paired eroded-minus-original Fusion difference is reported in the text and Table S4 notes. Erosion used one iteration of the six-neighbor operation on the supplied voxel grid; this was not interpreted as a fixed millimeter distance. The original ROI analysis remained primary.'
TABLE_SPEC={'key':'TableS4','title':'Table S4. Original ROI and eroded-boundary sensitivity',
  'headers':['Model','ROC AUC (95% interval)','Brier score','AP for treatment change','Treatment change correct, n/N'],
  'notes':['All models used the same 182 examinations, patient groups, folds and fixed model settings. Eroded models replaced the eight MRI features while retaining the image-specific normalization parameters and full-ROI discretization boundaries.',
  'The paired eroded-minus-original Fusion AUC difference was 0.0062 (95% interval, -0.0051 to 0.0183). Intervals use 2000 patient-cluster bootstrap attempts with fixed OOF probabilities and no refitting; they do not include full retraining or partition variability.',
  'AP denotes average precision for treatment-related change, using one minus the tumor probability. Lower Brier scores indicate smaller probability error. Treatment-change correct counts use tumor probability below 0.5. The original ROI remained primary. Performance values are displayed to three decimals; the small paired AUC difference and its interval are shown to four decimals.']}

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def wc(s):return len(re.findall(r"[A-Za-z0-9]+(?:['/-][A-Za-z0-9]+)*",s))
def close(a,b):assert np.isclose(a,b,rtol=1e-12,atol=1e-12),(a,b)

def audit():
    old=json.loads((ROOT/'reports/ucsd_manuscript_step37/results_section35.json').read_text());frozen=old['source_audit']['frozen_file_sha256'].copy()
    frozen.update(json.loads((ROOT/'reports/ucsd_manuscript_step37/artifact_verification.json').read_text())['sha256'])
    for p,h in frozen.items():assert sha(ROOT/p)==h,p
    extra=['outputs/ucsd-paper-materials-20261029/tables/TableS4.csv',str(FIGURE.relative_to(ROOT)),str(FIGURE.with_suffix('.svg').relative_to(ROOT))]
    frozen.update({p:sha(ROOT/p) for p in extra})
    pred=pd.read_csv(ROOT/'data/processed/ucsd_fusion_oof.csv');y=pred.tumor_label.to_numpy()
    saved=json.loads((ROOT/'reports/ucsd_fusion/metrics.json').read_text());table=pd.read_csv(ROOT/'outputs/ucsd-paper-materials-20261029/tables/TableS4.csv')
    boot=pd.read_csv(ROOT/'reports/ucsd_fusion/paired_patient_bootstrap.csv');assert table.shape==(4,8) and len(pred)==182 and pred.validation_group.nunique()==135
    assert saved['same_cases_and_folds'] and saved['bootstrap_valid']==saved['bootstrap_attempts']==2000 and saved['bootstrap_models_refitted'] is False
    eroded=pd.read_csv(ROOT/'data/processed/ucsd_image_features_eroded.csv');assert len(eroded)==182 and set(eroded.ID)==set(pred.ID)
    assert np.isfinite(eroded.drop(columns='ID').to_numpy()).all()
    batch=json.loads((ROOT/'reports/ucsd_batch_features/batch_receipt.json').read_text());assert len(batch['cases'])==182
    assert batch['summary']['eroded_complete_timepoints']==182 and batch['summary']['target_failed_timepoints']==0
    for case in batch['cases']:
        assert 0<case['eroded_voxels']<case['roi_voxels'] and len(case['eroded'])==8
        assert all(d['erosion_error'] is None for d in case['diagnostics'])
    folds=json.loads((ROOT/'reports/ucsd_fusion/fold_audit.json').read_text());assert len(folds)==5
    for f in folds:
        assert set(f['test_IDs'])==set(pred.loc[pred.validation_fold.eq(f['fold']),'ID'])
        assert all(n in f['models'] for n in MODELS)
    def manual_metrics(p):
        diff=p[y==1][:,None]-p[y==0][None,:];auc=float(np.mean((diff>0)+.5*(diff==0)))
        labels=1-y;score=1-p;order=np.argsort(-score,kind='stable');labels=labels[order];score=score[order]
        ends=np.flatnonzero(np.r_[score[:-1]!=score[1:],True]);tp=np.cumsum(labels)[ends]
        ap=float(np.sum(np.diff(np.r_[0,tp])/labels.sum()*tp/(ends+1)))
        return dict(roc_auc=auc,ap_treatment_change=ap,brier=float(np.mean((p-y)**2)),log_loss=float(-np.mean(y*np.log(p)+(1-y)*np.log1p(-p))),tn=int(((p<.5)&(y==0)).sum()))
    records=[];rows=[];values={}
    for i,name in enumerate(MODELS):
        p=pred[name+'_oof_probability'].to_numpy();assert np.isfinite(p).all() and np.all((p>0)&(p<1));v=manual_metrics(p);values[name]=v
        for k,value in v.items():close(value,saved['metrics'][name][k])
        lower=float(boot[name+'__roc_auc'].quantile(.025));upper=float(boot[name+'__roc_auc'].quantile(.975))
        for key,value in [('lower',lower),('upper',upper)]:close(value,saved['patient_cluster_bootstrap_95ci'][name+'__roc_auc'][key])
        for key,value in [('AUC',v['roc_auc']),('CI_lower',lower),('CI_upper',upper),('Brier',v['brier']),('AP_TC',v['ap_treatment_change']),('TC_correct',v['tn']),('TC_total',41)]:close(value,table.iloc[i][key])
        records.append({'Model':MODEL_LABELS[i],'ROC_AUC':float(table.iloc[i].AUC),'AUC_interval_lower':float(table.iloc[i].CI_lower),
          'AUC_interval_upper':float(table.iloc[i].CI_upper),'Brier_score':float(table.iloc[i].Brier),'AP_treatment_change':float(table.iloc[i].AP_TC),
          'Treatment_change_correct':int(table.iloc[i].TC_correct),'Treatment_change_total':41})
        rows.append([MODEL_LABELS[i],f"{v['roc_auc']:.3f} [{lower:.3f}, {upper:.3f}]",f"{v['brier']:.3f}",f"{v['ap_treatment_change']:.3f}",f"{v['tn']}/41"])
    paired=[]
    for metric in ['roc_auc','ap_treatment_change','brier','log_loss']:
        key='Fusion_eroded_minus_Fusion__'+metric
        np.testing.assert_allclose(boot[key],boot['Fusion_eroded__'+metric]-boot['Fusion__'+metric],rtol=1e-12,atol=1e-12)
        row=next(r for r in saved['paired_comparisons'] if r['comparison']=='Fusion_eroded_minus_Fusion' and r['metric']==metric)
        close(values['Fusion_eroded'][metric]-values['Fusion'][metric],row['difference'])
        close(float(boot[key].quantile(.025)),row['lower']);close(float(boot[key].quantile(.975)),row['upper'])
        assert row['lower']<0<row['upper'] and row['primary_comparison'] is False;paired.append(row)
    a=paired[0];assert [f"{a[k]:.4f}" for k in ['difference','lower','upper']]==['0.0062','-0.0051','0.0183']
    receipt=json.loads((ROOT/'reports/ucsd_paper_materials/generation_receipt.json').read_text());assert receipt['output_sha256'][str(FIGURE.relative_to(ROOT))]==sha(FIGURE)
    svg=etree.parse(str(FIGURE.with_suffix('.svg')));texts=svg.xpath('//s:text/text()',namespaces={'s':'http://www.w3.org/2000/svg'})
    assert all(t in texts for t in ['MRI','Fusion','Eroded MRI','Eroded fusion']) and not re.search(r'[\u3400-\u9fff]',' '.join(texts))
    spec=dict(**TABLE_SPEC,rows=rows);assert not re.search(r'[\u3400-\u9fff]',' '.join(spec['headers']+sum(rows,[])+spec['notes']+[FIGURE_CAPTION]))
    return {'frozen_file_sha256':frozen,'translated_table':spec,'full_precision_table_records':records,
      'original_table_numeric_values_checked':28,'fixed_oof_metric_values_reconstructed':20,'model_auc_interval_endpoints_checked':8,
      'eroded_complete_cases_checked':182,'saved_fold_case_sets_checked':5,'paired_metric_differences_checked':4,
      'bootstrap_paired_difference_values_checked':8000,'paired_interval_endpoints_checked':8,'paired_eroded_minus_original_fusion_records':paired,
      'paired_intervals_include_zero':True,'figure_unchanged_from_saved_materials':True,
      'numeric_values_changed':0,'models_refitted':0,'features_reextracted':0,'new_hypothesis_tests':0,
      'verification_scope':'Check saved erosion completeness and common folds, reconstruct fixed-OOF metrics, verify original Table S4 and paired bootstrap values/endpoints, and attach unchanged Figure S1 without image processing or model reruns'}

def add_table(doc,spec):
    title=doc.add_paragraph(spec['title'],style='Heading 1');title.paragraph_format.page_break_before=True
    title.paragraph_format.keep_with_next=True
    table=doc.add_table(rows=1,cols=5);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    widths=[1.35,1.87,.85,1.23,1.62]
    for col,width in zip(table.columns,widths):col.width=Inches(width)
    for cell,width in zip(table.rows[0].cells,widths):cell.width=Inches(width)
    for cell,label in zip(table.rows[0].cells,spec['headers']):cell.text=label
    for values in spec['rows']:
        for cell,value,width in zip(table.add_row().cells,values,widths):cell.text=value;cell.width=Inches(width)
    props=table._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side)
        for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:e.set(qn('w:'+k),v)
        borders.append(e)
    props.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for side in ['top','left','bottom','right']:
        e=OxmlElement('w:'+side);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
    props.append(margins)
    for ri,row in enumerate(table.rows):
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if ri==0:trpr.append(OxmlElement('w:tblHeader'))
        for ci,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7EBEF');cell._tc.get_or_add_tcPr().append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:run.font.size=Pt(10.5);run.font.bold=ri==0
    for note in spec['notes']:
        p=doc.add_paragraph(note);p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(4)
        p.paragraph_format.keep_together=True
        for r in p.runs:r.font.size=Pt(10)

def build():
    OUT.mkdir(parents=True,exist_ok=True);REPORT.mkdir(parents=True,exist_ok=True);a=audit()
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(.79)
    for name in ['Normal','Title','Heading 1']:
        st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        rf=st.element.get_or_add_rPr().rFonts;rf.set(qn('w:eastAsia'),'Noto Serif CJK SC')
        for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:rf.attrib.pop(qn('w:'+key),None)
    normal=doc.styles['Normal'];normal.font.size=Pt(12);normal.paragraph_format.line_spacing=1.;normal.paragraph_format.space_after=Pt(9)
    doc.styles['Title'].font.size=Pt(16);doc.styles['Title'].paragraph_format.space_after=Pt(12)
    doc.styles['Heading 1'].font.size=Pt(13);doc.styles['Heading 1'].paragraph_format.space_before=Pt(8);doc.styles['Heading 1'].paragraph_format.space_after=Pt(9)
    doc.core_properties.title='Results 3.6 Boundary sensitivity';doc.core_properties.author='Tianyu Feng'
    doc.add_paragraph('Results',style='Title');doc.add_paragraph('3.6 Boundary sensitivity',style='Heading 1')
    for item in PARAGRAPHS:
        en=doc.add_paragraph(item['english']);en.paragraph_format.page_break_before=False
        en.paragraph_format.keep_together=True;en.paragraph_format.keep_with_next=True
        zh=doc.add_paragraph(item['chinese']);zh.paragraph_format.keep_together=True
        for r in zh.runs:r.font.size=Pt(11)
    add_table(doc,a['translated_table'])
    pd.DataFrame(a['full_precision_table_records']).to_csv(OUT/'TableS4_EN.csv',index=False)
    p=doc.add_paragraph('Figure S1',style='Heading 1');p.paragraph_format.page_break_before=True
    doc.add_picture(str(FIGURE),width=Inches(6.7))
    p=doc.add_paragraph(FIGURE_CAPTION);p.paragraph_format.keep_together=True
    for r in p.runs:r.font.size=Pt(10)
    for e in [*doc.styles.element.xpath('.//w:pBdr'),*doc.element.xpath('.//w:pBdr')]:e.getparent().remove(e)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    path=OUT/'UCSD_Manuscript_Results_3_6_EN_ZH.docx';doc.save(path)
    (OUT/'UCSD_Manuscript_Results_3_6_EN.txt').write_text('\n\n'.join(p['english'] for p in PARAGRAPHS)+'\n')
    d={'stage':38,'date':'2026-10-29','subsection':'3.6 Boundary sensitivity','v2_paragraph_ids':['R6'],'paragraphs':PARAGRAPHS,
      'english_word_count':sum(wc(p['english']) for p in PARAGRAPHS),'paragraph_word_counts':{p['id']:wc(p['english']) for p in PARAGRAPHS},
      'word_count_scope':'English narrative only; excludes headings, tables/notes, figure/caption and Chinese translation','v2_approximate_target_words':150,
      'equations':{'count':0,'numbers':[],'cumulative_previous_numbered_groups':10},
      'citation_mapping':{},'new_manuscript_references':0,'cumulative_unique_cited_sources':5,'editable_english_tables':['Table S4'],'table_numeric_values_changed':0,'included_figure':str(FIGURE.relative_to(ROOT)),'figure_sha256':sha(FIGURE),'figure_caption':FIGURE_CAPTION,'figure_reused_without_regeneration':True,'source_audit':a,'models_refitted':0,'new_hypothesis_tests':0,
      'next_section':'Results 3.7 Split stability (R7)','docx':str(path.relative_to(ROOT)),'docx_sha256':sha(path)}
    (REPORT/'results_section36.json').write_text(json.dumps(d,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'docx':d['docx'],'english_word_count':d['english_word_count'],'paragraph_word_counts':d['paragraph_word_counts'],'prior_files':len(a['frozen_file_sha256'])},indent=2))

def verify():
    from pypdf import PdfReader
    d=json.loads((REPORT/'results_section36.json').read_text())
    for p,h in d['source_audit']['frozen_file_sha256'].items():assert sha(ROOT/p)==h,p
    path=ROOT/d['docx'];assert sha(path)==d['docx_sha256'];texts=[p.text for p in Document(path).paragraphs]
    for p in PARAGRAPHS:
        assert p['english'] in texts and p['chinese'] in texts and texts.index(p['chinese'])==texts.index(p['english'])+1
        assert not re.search(r'[\u3400-\u9fff，。；：（）【】、]',p['english'])
    assert sum(wc(p['english']) for p in PARAGRAPHS)==d['english_word_count']
    doc=Document(path);assert len(doc.tables)==1
    spec=d['source_audit']['translated_table'];expected=[spec['headers'],*spec['rows']]
    assert [[c.text for c in row.cells] for row in doc.tables[0].rows]==expected
    assert all(not re.search(r'[\u3400-\u9fff]',c.text) for row in doc.tables[0].rows for c in row.cells)
    for note in spec['notes']:assert note in texts
    full=pd.read_csv(OUT/'TableS4_EN.csv');assert len(full)==4
    for (_,row),r in zip(full.iterrows(),d['source_audit']['full_precision_table_records']):
        assert row.Model==r['Model']
        for key,value in r.items():
            if key!='Model':close(row[key],value)
    assert len(doc.sections)==1 and FIGURE_CAPTION in texts
    ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
    with zipfile.ZipFile(path) as z:
        t=etree.fromstring(z.read('word/document.xml'));assert len(t.xpath('//m:oMath',namespaces=ns))==0
        media=[n for n in z.namelist() if n.startswith('word/media/')];assert len(media)==1
        assert hashlib.sha256(z.read(media[0])).hexdigest()==sha(FIGURE)
    pages=len(PdfReader(OUT/'docx_render/UCSD_Manuscript_Results_3_6_EN_ZH.pdf').pages);assert pages==3
    v={'success':True,'stage':38,'date':'2026-10-29','subsection':d['subsection'],'paragraph_ids':['R6'],
      'english_word_count':d['english_word_count'],'paragraph_word_counts':d['paragraph_word_counts'],'matched_chinese_paragraphs':1,
      'word_pages':pages,'visual_review_pages':[1,2,3],'native_editable_equations':0,'equation_numbers':[],'equation_images':0,
      'prior_files_unchanged':len(d['source_audit']['frozen_file_sha256']),'original_table_numeric_values_checked':28,'fixed_oof_metric_values_reconstructed':20,'model_auc_interval_endpoints_checked':8,'eroded_complete_cases_checked':182,'saved_fold_case_sets_checked':5,'paired_metric_differences_checked':4,'bootstrap_paired_difference_values_checked':8000,'paired_interval_endpoints_checked':8,'editable_english_tables':1,'table_numeric_values_changed':0,'paired_intervals_include_zero_disclosed':True,'unchanged_figure_embedded':True,
      'models_refitted':0,'new_hypothesis_tests':0,'new_manuscript_references':0,'cumulative_unique_cited_sources':5,'next_section':d['next_section'],
      'sha256':{str(p.relative_to(ROOT)):sha(p) for p in OUT.iterdir() if p.is_file() and p.suffix in {'.docx','.txt','.csv'}}}
    (REPORT/'artifact_verification.json').write_text(json.dumps(v,ensure_ascii=False,indent=2)+'\n');print(json.dumps(v,indent=2))

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--verify',action='store_true');args=parser.parse_args()
    verify() if args.verify else build()
