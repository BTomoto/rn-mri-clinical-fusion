"""Draft R5 with saved calibration diagnostics and fixed-threshold errors."""
from pathlib import Path
import argparse,hashlib,json,re,sys,zipfile
import numpy as np
import pandas as pd
from lxml import etree
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'outputs/ucsd-manuscript-step37-20261029'
REPORT=ROOT/'reports/ucsd_manuscript_step37'
PARAGRAPHS=[
 {'id':'R5','citation_ids':[],
  'english': 'Fusion had an offset calibration intercept of -0.121 and a calibration slope of 0.667 in the pooled OOF predictions (Figure 2C; Supplementary Table S3). These were descriptive calibration diagnostics and were not applied to modify the evaluated probabilities. At the fixed tumor-probability threshold of 0.5, Fusion correctly classified 130 of 141 tumor examinations (92.2%) and 12 of 41 treatment-related-change examinations (29.3%). Eleven tumor examinations were classified as treatment-related change, whereas 29 treatment-related-change examinations were classified as tumor. Overall accuracy was 142/182 (78.0%), compared with 141/182 (77.5%) for a reference rule assigning every examination to tumor. Balanced accuracy was 0.607. M1 correctly recognized two treatment-related-change examinations, while MRI and Fusion each recognized 12. The same MRI and Fusion models correctly classified 128 and 130 tumor examinations, respectively. Thus, the higher discrimination estimates coexisted with limited treatment-related-change recognition at the unchanged threshold, and overall accuracy remained close to the always-tumor reference.',
  'chinese': '在汇总的OOF预测中，Fusion固定offset的校准截距为−0.121，校准斜率为0.667（Figure 2C；Supplementary Table S3）。这些为描述性校准诊断，未用于修改被评价的概率。在固定的0.5肿瘤概率阈值下，Fusion正确分类141次肿瘤检查中的130次（92.2%），以及41次治疗相关改变检查中的12次（29.3%）。11次肿瘤检查被分为治疗相关改变，29次治疗相关改变检查被分为肿瘤。总体准确率为142/182（78.0%），将每次检查都分为肿瘤的参考规则准确率为141/182（77.5%）。平衡准确率为0.607。M1正确识别2次治疗相关改变检查，MRI和Fusion各正确识别12次；这两个MRI与Fusion模型分别正确分类128次和130次肿瘤检查。因此，较高的鉴别指标与固定阈值下有限的治疗相关改变识别能力同时存在，总体准确率仍接近全判肿瘤的参考规则。'}
]
MODELS=['M0','M1','MRI','Fusion']
TABLE_SPEC={'key':'TableS3','title':'Table S3. Fixed-threshold classification and calibration diagnostics',
  'headers':['Model','Treatment change correct, n/N','Tumor correct, n/N','Balanced accuracy','Calibration offset intercept','Calibration slope'],
  'notes':['All counts refer to the same 182 examinations from 135 patients. Tumor probability at least 0.5 is classified as tumor; a lower probability is classified as treatment-related change. Balanced accuracy is the mean of tumor sensitivity and treatment-change recognition.',
  'The calibration offset intercept is estimated with the predicted log odds as a fixed offset and slope set to one. The calibration slope is estimated jointly with a free intercept in a separate unpenalized logistic diagnostic regression. These two intercept definitions are distinct; the joint intercept is not shown here.',
  'Calibration uses the pooled OOF predictions without subsequent recalibration or threshold changes. Values are descriptive estimates without calibration confidence intervals. Metrics are displayed to three decimals; class counts retain their exact denominators.']}

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def wc(s):return len(re.findall(r"[A-Za-z0-9]+(?:['/-][A-Za-z0-9]+)*",s))
def close(a,b):assert np.isclose(a,b,rtol=1e-12,atol=1e-12),(a,b)

def audit():
    old=json.loads((ROOT/'reports/ucsd_manuscript_step36/results_section34.json').read_text());frozen=old['source_audit']['frozen_file_sha256'].copy()
    frozen.update(json.loads((ROOT/'reports/ucsd_manuscript_step36/artifact_verification.json').read_text())['sha256'])
    for p,h in frozen.items():assert sha(ROOT/p)==h,p
    extra=['outputs/ucsd-paper-materials-20261029/tables/TableS3.csv','reports/ucsd_fusion/calibration_bins.csv']
    frozen.update({p:sha(ROOT/p) for p in extra})
    from scipy.special import expit,logit
    pred=pd.read_csv(ROOT/'data/processed/ucsd_fusion_oof.csv');y=pred.tumor_label.to_numpy()
    saved=json.loads((ROOT/'reports/ucsd_fusion/metrics.json').read_text());table=pd.read_csv(ROOT/'outputs/ucsd-paper-materials-20261029/tables/TableS3.csv')
    bins=pd.read_csv(ROOT/'reports/ucsd_fusion/calibration_bins.csv');modelcsv=pd.read_csv(ROOT/'reports/ucsd_fusion/model_metrics.csv').set_index('model')
    assert len(pred)==182 and pred.validation_group.nunique()==135 and table.shape==(4,8) and saved['threshold']==.5
    records=[];rows=[];checks={};bin_rows=0
    for i,name in enumerate(MODELS):
        p=pred[name+'_oof_probability'].to_numpy();positive=p>=.5
        tp=int((positive&(y==1)).sum());tn=int((~positive&(y==0)).sum());fp=int((positive&(y==0)).sum());fn=int((~positive&(y==1)).sum())
        tumor_total=int((y==1).sum());tc_total=int((y==0).sum())
        values=dict(tp=tp,tn=tn,fp=fp,fn=fn,sensitivity_tumor=tp/tumor_total,specificity_treatment_change=tn/tc_total,
          balanced_accuracy=.5*(tp/tumor_total+tn/tc_total),accuracy=(tp+tn)/len(y))
        for k,v in values.items():close(v,saved['metrics'][name][k]);close(v,modelcsv.loc[name,k])
        c=saved['calibration'][name];assert c['error'] is None
        z=logit(np.clip(p,np.finfo(float).eps,1-np.finfo(float).eps));x=np.column_stack([np.ones(len(y)),z])
        off_score=float(abs(np.sum(expit(z+c['calibration_intercept_offset'])-y)))
        fit=expit(x@np.array([c['calibration_joint_intercept'],c['calibration_slope']]))
        joint_score=float(np.max(np.abs(x.T@(fit-y))));mineig=float(np.linalg.eigvalsh(x.T@((fit*(1-fit))[:,None]*x)).min())
        assert off_score<1e-8 and joint_score<1e-5 and mineig>=1e-8
        for k in ['calibration_intercept_offset','calibration_joint_intercept','calibration_slope']:close(c[k],modelcsv.loc[name,k])
        frame=pd.DataFrame(dict(y=y,p=p,bin=pd.qcut(pd.Series(p),5,duplicates='drop')));existing=bins[bins.model.eq(name)].reset_index(drop=True)
        for k,(_,g) in enumerate(frame.groupby('bin',observed=True),1):
            row=existing.iloc[k-1];assert row['bin']==k and row.n==len(g) and row.tumor_count==int(g.y.sum())
            close(row.mean_predicted,float(g.p.mean()));close(row.observed_tumor,float(g.y.mean()));bin_rows+=1
        assert len(existing)==5 and existing.n.sum()==182
        for k,v in [('TC_correct',tn),('TC_total',tc_total),('Tumor_correct',tp),('Tumor_total',tumor_total),
          ('Balanced_accuracy',values['balanced_accuracy']),('Calibration_offset',c['calibration_intercept_offset']),('Calibration_slope',c['calibration_slope'])]:close(v,table.iloc[i][k])
        records.append({'Model':name,'Treatment_change_correct':tn,'Treatment_change_total':tc_total,'Tumor_correct':tp,'Tumor_total':tumor_total,
          'Balanced_accuracy':float(table.iloc[i].Balanced_accuracy),'Calibration_offset_intercept':float(table.iloc[i].Calibration_offset),'Calibration_slope':float(table.iloc[i].Calibration_slope)})
        rows.append([name,f'{tn}/{tc_total}',f'{tp}/{tumor_total}',f"{values['balanced_accuracy']:.3f}",f"{c['calibration_intercept_offset']:.3f}",f"{c['calibration_slope']:.3f}"])
        checks[name]=dict(**values,offset_score_absolute_residual=off_score,joint_score_maximum_absolute_residual=joint_score,joint_hessian_minimum_eigenvalue=mineig)
    fusion=checks['Fusion'];assert [fusion[k] for k in ['tp','tn','fp','fn']]==[130,12,29,11]
    assert f"{100*fusion['sensitivity_tumor']:.1f}"=='92.2' and f"{100*fusion['specificity_treatment_change']:.1f}"=='29.3'
    always_correct=int((y==1).sum());always_accuracy=float(np.mean(y==1));assert always_correct==141
    assert f"{100*fusion['accuracy']:.1f}"=='78.0' and f'{100*always_accuracy:.1f}'=='77.5'
    assert fusion['tp']+fusion['tn']==142 and bin_rows==20
    spec=dict(**TABLE_SPEC,rows=rows);assert not re.search(r'[\u3400-\u9fff]',' '.join(spec['headers']+sum(rows,[])+spec['notes']))
    return {'frozen_file_sha256':frozen,'translated_table':spec,'full_precision_table_records':records,
      'original_table_numeric_values_checked':28,'classification_count_and_metric_values_reconstructed':32,
      'saved_calibration_checks':checks,'saved_calibration_variants_checked':4,'calibration_bin_rows_checked':20,'calibration_bin_numeric_fields_checked':100,
      'always_tumor_reference':{'rule':'assign tumor label to every examination; no probability-score AUC used','correct':always_correct,'total':182,'accuracy':always_accuracy},
      'fusion_reference_accuracy_checked':fusion['accuracy'],'calibration_diagnostics_refitted':0,'numeric_values_changed':0,'models_refitted':0,'features_reextracted':0,'new_hypothesis_tests':0,
      'verification_scope':'Reconstruct fixed-threshold counts and metrics from saved OOF probabilities, evaluate score residuals at saved calibration parameters, verify bin summaries and original Table S3, and compute the already-planned always-tumor accuracy reference without fitting'}

def add_table(doc,spec):
    title=doc.add_paragraph(spec['title'],style='Heading 1');title.paragraph_format.page_break_before=True
    title.paragraph_format.keep_with_next=True
    table=doc.add_table(rows=1,cols=6);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    widths=[.62,1.37,1.37,1.13,1.32,1.11]
    for col,width in zip(table.columns,widths):col.width=Inches(width)
    for cell,width in zip(table.rows[0].cells,widths):cell.width=Inches(width)
    for cell,label in zip(table.rows[0].cells,spec['headers']):cell.text=label
    for values in spec['rows']:
        for cell,value,width in zip(table.add_row().cells,values,widths):cell.text=value;cell.width=Inches(width)
    props=table._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side)
        for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:e.set(qn('w:'+k),v)
        borders.append(e)
    props.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for side in ['top','left','bottom','right']:
        e=OxmlElement('w:'+side);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
    props.append(margins)
    for ri,row in enumerate(table.rows):
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if ri==0:trpr.append(OxmlElement('w:tblHeader'))
        for ci,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7EBEF');cell._tc.get_or_add_tcPr().append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:run.font.size=Pt(10.5);run.font.bold=ri==0
    for note in spec['notes']:
        p=doc.add_paragraph(note);p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(4)
        p.paragraph_format.keep_together=True
        for r in p.runs:r.font.size=Pt(10)

def build():
    OUT.mkdir(parents=True,exist_ok=True);REPORT.mkdir(parents=True,exist_ok=True);a=audit()
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(.79)
    for name in ['Normal','Title','Heading 1']:
        st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        rf=st.element.get_or_add_rPr().rFonts;rf.set(qn('w:eastAsia'),'Noto Serif CJK SC')
        for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:rf.attrib.pop(qn('w:'+key),None)
    normal=doc.styles['Normal'];normal.font.size=Pt(12);normal.paragraph_format.line_spacing=1.;normal.paragraph_format.space_after=Pt(9)
    doc.styles['Title'].font.size=Pt(16);doc.styles['Title'].paragraph_format.space_after=Pt(12)
    doc.styles['Heading 1'].font.size=Pt(13);doc.styles['Heading 1'].paragraph_format.space_before=Pt(8);doc.styles['Heading 1'].paragraph_format.space_after=Pt(9)
    doc.core_properties.title='Results 3.5 Calibration and errors';doc.core_properties.author='Tianyu Feng'
    doc.add_paragraph('Results',style='Title');doc.add_paragraph('3.5 Calibration and errors',style='Heading 1')
    for item in PARAGRAPHS:
        en=doc.add_paragraph(item['english']);en.paragraph_format.page_break_before=False
        en.paragraph_format.keep_together=True;en.paragraph_format.keep_with_next=True
        zh=doc.add_paragraph(item['chinese']);zh.paragraph_format.keep_together=True
        for r in zh.runs:r.font.size=Pt(11)
    add_table(doc,a['translated_table'])
    pd.DataFrame(a['full_precision_table_records']).to_csv(OUT/'TableS3_EN.csv',index=False)
    for e in [*doc.styles.element.xpath('.//w:pBdr'),*doc.element.xpath('.//w:pBdr')]:e.getparent().remove(e)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    path=OUT/'UCSD_Manuscript_Results_3_5_EN_ZH.docx';doc.save(path)
    (OUT/'UCSD_Manuscript_Results_3_5_EN.txt').write_text('\n\n'.join(p['english'] for p in PARAGRAPHS)+'\n')
    d={'stage':37,'date':'2026-10-29','subsection':'3.5 Calibration and errors','v2_paragraph_ids':['R5'],'paragraphs':PARAGRAPHS,
      'english_word_count':sum(wc(p['english']) for p in PARAGRAPHS),'paragraph_word_counts':{p['id']:wc(p['english']) for p in PARAGRAPHS},
      'word_count_scope':'English narrative only; excludes headings, tables/notes and Chinese translation','v2_approximate_target_words':150,
      'equations':{'count':0,'numbers':[],'cumulative_previous_numbered_groups':10},
      'citation_mapping':{},'new_manuscript_references':0,'cumulative_unique_cited_sources':5,'editable_english_tables':['Table S3'],'table_numeric_values_changed':0,'source_audit':a,'models_refitted':0,'new_hypothesis_tests':0,
      'next_section':'Results 3.6 Boundary sensitivity (R6)','docx':str(path.relative_to(ROOT)),'docx_sha256':sha(path)}
    (REPORT/'results_section35.json').write_text(json.dumps(d,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'docx':d['docx'],'english_word_count':d['english_word_count'],'paragraph_word_counts':d['paragraph_word_counts'],'prior_files':len(a['frozen_file_sha256'])},indent=2))

def verify():
    from pypdf import PdfReader
    d=json.loads((REPORT/'results_section35.json').read_text())
    for p,h in d['source_audit']['frozen_file_sha256'].items():assert sha(ROOT/p)==h,p
    path=ROOT/d['docx'];assert sha(path)==d['docx_sha256'];texts=[p.text for p in Document(path).paragraphs]
    for p in PARAGRAPHS:
        assert p['english'] in texts and p['chinese'] in texts and texts.index(p['chinese'])==texts.index(p['english'])+1
        assert not re.search(r'[\u3400-\u9fff，。；：（）【】、]',p['english'])
    assert sum(wc(p['english']) for p in PARAGRAPHS)==d['english_word_count']
    doc=Document(path);assert len(doc.tables)==1
    spec=d['source_audit']['translated_table'];expected=[spec['headers'],*spec['rows']]
    assert [[c.text for c in row.cells] for row in doc.tables[0].rows]==expected
    assert all(not re.search(r'[\u3400-\u9fff]',c.text) for row in doc.tables[0].rows for c in row.cells)
    for note in spec['notes']:assert note in texts
    full=pd.read_csv(OUT/'TableS3_EN.csv');assert len(full)==4
    for (_,row),r in zip(full.iterrows(),d['source_audit']['full_precision_table_records']):
        assert row.Model==r['Model']
        for key,value in r.items():
            if key!='Model':close(row[key],value)
    assert len(doc.sections)==1
    ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
    with zipfile.ZipFile(path) as z:
        t=etree.fromstring(z.read('word/document.xml'));assert len(t.xpath('//m:oMath',namespaces=ns))==0
        assert not any(n.startswith('word/media/') for n in z.namelist())
    pages=len(PdfReader(OUT/'docx_render/UCSD_Manuscript_Results_3_5_EN_ZH.pdf').pages);assert pages==2
    v={'success':True,'stage':37,'date':'2026-10-29','subsection':d['subsection'],'paragraph_ids':['R5'],
      'english_word_count':d['english_word_count'],'paragraph_word_counts':d['paragraph_word_counts'],'matched_chinese_paragraphs':1,
      'word_pages':pages,'visual_review_pages':[1,2],'native_editable_equations':0,'equation_numbers':[],'equation_images':0,
      'prior_files_unchanged':len(d['source_audit']['frozen_file_sha256']),'original_table_numeric_values_checked':28,'classification_count_and_metric_values_reconstructed':32,'saved_calibration_variants_checked':4,'calibration_bin_rows_checked':20,'editable_english_tables':1,'table_numeric_values_changed':0,'always_tumor_reference_disclosed':True,'calibration_diagnostics_refitted':0,
      'models_refitted':0,'new_hypothesis_tests':0,'new_manuscript_references':0,'cumulative_unique_cited_sources':5,'next_section':d['next_section'],
      'sha256':{str(p.relative_to(ROOT)):sha(p) for p in OUT.iterdir() if p.is_file() and p.suffix in {'.docx','.txt','.csv'}}}
    (REPORT/'artifact_verification.json').write_text(json.dumps(v,ensure_ascii=False,indent=2)+'\n');print(json.dumps(v,indent=2))

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--verify',action='store_true');args=parser.parse_args()
    verify() if args.verify else build()
