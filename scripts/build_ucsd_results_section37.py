"""Draft R7 from frozen repeated-split predictions and descriptive distributions."""
from pathlib import Path
import argparse,hashlib,json,re,sys,zipfile
import numpy as np
import pandas as pd
from lxml import etree
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION,WD_ORIENT

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'outputs/ucsd-manuscript-step39-20261029'
REPORT=ROOT/'reports/ucsd_manuscript_step39'
PARAGRAPHS=[
 {'id':'R7','citation_ids':[],
  'english': 'All 20 prespecified repetitions completed the patient-grouped five-fold validation without failed fold-model fits. Across repetitions, median AUCs were 0.556 for M0, 0.667 for M1, 0.775 for MRI and 0.784 for Fusion (Supplementary Table S5A; Figure 3A). Fusion had first and third AUC quartiles of 0.775 and 0.795, with an observed range of 0.759 to 0.809. The median within-repetition Fusion-minus-M1 AUC difference was 0.125, with a range of 0.103 to 0.150; all 20 differences were positive (Supplementary Table S5B; Figure 3B). AP, Brier-score and logarithmic-loss differences also favored Fusion in all 20 repetitions. These sign counts described the observed partitions and did not replace the primary paired bootstrap intervals. Because the same patients were reused, the repetitions were not independent validation cohorts. Medians, quartiles and ranges summarized sensitivity to partitioning and were not confidence intervals. The first repetition reproduced the original partition and predictions, and its Fusion AUC of 0.782 remained the primary estimate rather than the maximum observed across repetitions.',
  'chinese': '20套预定重复均完成患者分组五折验证，没有折内模型拟合失败。跨重复的AUC中位数分别为M0的0.556、M1的0.667、MRI的0.775和Fusion的0.784（Supplementary Table S5A；Figure 3A）。Fusion AUC的第一和第三四分位数为0.775及0.795，实际范围为0.759至0.809。每套重复内部Fusion减去M1的AUC差值中位数为0.125，范围为0.103至0.150，20个差值均为正（Supplementary Table S5B；Figure 3B）。AP、Brier评分和对数损失差值也在全部20套重复中有利于Fusion。这些方向计数描述所观察的划分，不替代主要配对bootstrap区间。由于重复使用同一批患者，20套重复不是独立验证队列。中位数、四分位数及范围概括对划分的敏感性，不是置信区间。第一套重复复现原划分及预测，其Fusion AUC为0.782，仍作为主估计，不以全部重复中的最高值替代。'}
]
MODELS=['M0','M1','MRI','Fusion']
METRICS=['roc_auc','ap_treatment_change','brier','log_loss']
METRIC_LABELS=['ROC AUC','Treatment-related-change AP','Brier score','Logarithmic loss']
FIGURE=ROOT/'outputs/ucsd-paper-materials-20261029/figures/Figure3_Repeat_Stability.png'
FIGURE_CAPTION='Figure 3. Sensitivity to 20 prespecified patient-grouped five-fold partitions. A, Pooled OOF AUC for each of the four fixed models within each repetition. B, Fusion-minus-M1 AUC differences calculated within the corresponding repetition; the dashed line marks zero. All repetitions reused the same 182 examinations from 135 patients. The repetitions are not independent validation cohorts, and their distribution does not define a confidence interval. Repetition 1 reproduces the primary partition. The highest AUC was not used to select a replacement analysis.'
TABLE_SPECS=[
 {'key':'TableS5A','title':'Table S5A. Performance across 20 fixed patient-grouped partitions',
  'headers':['Model','Median ROC AUC [Q1, Q3]','ROC AUC range','Median Brier score'],
  'notes':['All 20 repetitions were complete. Each repetition evaluated the same 182 examinations from 135 patients, using patient-grouped five-fold validation and the same fixed settings for M0, M1, MRI and Fusion.',
  'Q1 and Q3 are the first and third quartiles across the 20 repetition-level estimates. Ranges are observed minima and maxima, not 95% confidence intervals. Repetitions reuse the same patients and are not independent studies. The primary Fusion AUC remains 0.782 from repetition 1.']},
 {'key':'TableS5B','title':'Table S5B. Fusion-minus-M1 differences across the 20 partitions',
  'headers':['Metric','Median difference [Q1, Q3]','Difference range','Favorable repetitions, n/N'],
  'notes':['Differences are calculated as Fusion minus M1 within each repetition, then summarized across repetitions. The median paired difference is not obtained by subtracting the two models’ separate medians.',
  'Positive AUC or treatment-change AP differences and negative Brier-score or logarithmic-loss differences favor Fusion. AP denotes average precision for treatment-related change. Q1 and Q3 are quartiles, and ranges are observed minima and maxima.',
  'Sign counts, including 20/20, are descriptive summaries of reused patients across fixed partitions. They are not independent replications or hypothesis tests and do not replace the primary paired patient-bootstrap intervals. Values are displayed to three decimals.']}
]

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def wc(s):return len(re.findall(r"[A-Za-z0-9]+(?:['/-][A-Za-z0-9]+)*",s))
def close(a,b):assert np.isclose(a,b,rtol=1e-12,atol=1e-12),(a,b)

def audit():
    old=json.loads((ROOT/'reports/ucsd_manuscript_step38/results_section36.json').read_text());frozen=old['source_audit']['frozen_file_sha256'].copy()
    frozen.update(json.loads((ROOT/'reports/ucsd_manuscript_step38/artifact_verification.json').read_text())['sha256'])
    for p,h in frozen.items():assert sha(ROOT/p)==h,p
    extra=['outputs/ucsd-paper-materials-20261029/tables/TableS5A.csv','outputs/ucsd-paper-materials-20261029/tables/TableS5B.csv',
      str(FIGURE.relative_to(ROOT)),str(FIGURE.with_suffix('.svg').relative_to(ROOT))]
    frozen.update({p:sha(ROOT/p) for p in extra})
    directory=ROOT/'reports/ucsd_stability';saved=json.loads((directory/'metrics.json').read_text());lock=json.loads((directory/'analysis_lock.json').read_text())
    status=pd.read_csv(directory/'repeat_status.csv');manifest=pd.read_csv(directory/'repeated_fold_manifest.csv');pred=pd.read_csv(ROOT/'data/processed/ucsd_stability_repeated_oof.csv')
    metrics=pd.read_csv(directory/'repeated_metrics.csv');deltas=pd.read_csv(directory/'repeated_paired_differences.csv')
    assert saved['repeats_complete']==20 and saved['repeats_failed']==saved['model_fold_failure_records']==0
    assert status['repeat'].tolist()==list(range(1,21)) and status.seed.tolist()==list(range(20261029,20261049))==lock['seeds']
    assert status.complete.all() and status.failed_fold_model_records.eq(0).all() and json.loads((directory/'failures.json').read_text())['records']==[]
    assert not any(saved[k] for k in ['repeats_are_independent_samples','repeated_ranges_are_confidence_intervals','probability_ensemble_created','prior_primary_analysis_replaced'])
    assert len(pred)==len(manifest)==3640 and len(metrics)==len(deltas)==80
    main=pd.read_csv(ROOT/'data/processed/ucsd_fusion_oof.csv').set_index('ID');all_values={};fold_count=0;fit_records=0
    def manual_values(y,p):
        diff=p[y==1][:,None]-p[y==0][None,:];auc=float(np.mean((diff>0)+.5*(diff==0)))
        labels=1-y;score=1-p;order=np.argsort(-score,kind='stable');labels=labels[order];score=score[order]
        ends=np.flatnonzero(np.r_[score[:-1]!=score[1:],True]);tp=np.cumsum(labels)[ends]
        ap=float(np.sum(np.diff(np.r_[0,tp])/labels.sum()*tp/(ends+1)))
        return dict(roc_auc=auc,ap_treatment_change=ap,brier=float(np.mean((p-y)**2)),log_loss=float(-np.mean(y*np.log(p)+(1-y)*np.log1p(-p))))
    for repetition,seed in enumerate(lock['seeds'],1):
        p=pred[pred['repeat'].eq(repetition)];mf=manifest[manifest['repeat'].eq(repetition)]
        assert len(p)==182 and p.ID.nunique()==182 and set(p.ID)==set(main.index) and p.validation_group.nunique()==135
        pd.testing.assert_frame_equal(p[mf.columns].reset_index(drop=True),mf.reset_index(drop=True))
        assert p.seed.eq(seed).all();audit=json.loads((directory/f'repeated_fold_audits/repeat_{repetition:02d}.json').read_text())
        assert audit['complete'] and audit['errors']==[] and len(audit['folds'])==5
        seen=[]
        for f in audit['folds']:
            tr=set(f['train_IDs']);te=set(f['test_IDs']);assert not tr&te and tr|te==set(p.ID)
            assert te==set(p.loc[p.validation_fold.eq(f['fold']),'ID']) and f['valid'] and not f['errors']
            train=p[p.ID.isin(tr)];test=p[p.ID.isin(te)]
            for key in ['patient_id','brats_patient_id','validation_group']:assert not set(train[key])&set(test[key])
            assert train.tumor_label.nunique()==test.tumor_label.nunique()==2
            assert set(f['models'])==set(MODELS) and all(v['success'] for v in f['models'].values());seen.extend(te);fold_count+=1;fit_records+=4
        assert len(seen)==182 and set(seen)==set(p.ID)
        y=p.tumor_label.to_numpy()
        for name in MODELS:
            probability=p[name+'_oof_probability'].to_numpy();assert np.isfinite(probability).all() and np.all((probability>0)&(probability<1))
            value=manual_values(y,probability);all_values[(repetition,name)]=value
            row=metrics[metrics['repeat'].eq(repetition)&metrics.model.eq(name)].iloc[0]
            for key,v in value.items():close(v,row[key])
            if repetition==1:
                np.testing.assert_allclose(probability,main.loc[p.ID,name+'_oof_probability'],rtol=1e-12,atol=1e-12)
                assert p.validation_fold.tolist()==main.loc[p.ID,'validation_fold'].tolist()
        for metric in METRICS:
            row=deltas[deltas['repeat'].eq(repetition)&deltas.metric.eq(metric)].iloc[0]
            close(row.difference,all_values[(repetition,'Fusion')][metric]-all_values[(repetition,'M1')][metric])
    original_a=pd.read_csv(ROOT/'outputs/ucsd-paper-materials-20261029/tables/TableS5A.csv')
    original_b=pd.read_csv(ROOT/'outputs/ucsd-paper-materials-20261029/tables/TableS5B.csv')
    assert original_a.shape==(4,7) and original_b.shape==(4,8)
    record_a=[];rows_a=[];record_b=[];rows_b=[]
    for i,name in enumerate(MODELS):
        values=np.array([all_values[(r,name)]['roc_auc'] for r in range(1,21)]);br=np.array([all_values[(r,name)]['brier'] for r in range(1,21)])
        q1,med,q3=np.quantile(values,[.25,.5,.75],method='linear');lo=float(values.min());hi=float(values.max());bm=float(np.median(br))
        for key,v in [('AUC_median',med),('AUC_q25',q1),('AUC_q75',q3),('AUC_min',lo),('AUC_max',hi),('Brier_median',bm)]:close(v,original_a.iloc[i][key])
        record_a.append({'Model':name,**{k:float(original_a.iloc[i][k]) for k in ['AUC_median','AUC_q25','AUC_q75','AUC_min','AUC_max','Brier_median']}})
        rows_a.append([name,f'{med:.3f} [{q1:.3f}, {q3:.3f}]',f'[{lo:.3f}, {hi:.3f}]',f'{bm:.3f}'])
    distribution=pd.read_csv(directory/'repeated_delta_distribution.csv').set_index('metric')
    for i,metric in enumerate(METRICS):
        values=np.array([all_values[(r,'Fusion')][metric]-all_values[(r,'M1')][metric] for r in range(1,21)])
        q1,med,q3=np.quantile(values,[.25,.5,.75],method='linear');lo=float(values.min());hi=float(values.max())
        favorable=int(np.sum(values>0 if metric in ['roc_auc','ap_treatment_change'] else values<0))
        for key,v in [('Median',med),('Q25',q1),('Q75',q3),('Minimum',lo),('Maximum',hi),('Favourable_repeats',favorable),('Total_repeats',20)]:close(v,original_b.iloc[i][key])
        for key,v in [('median',med),('q25',q1),('q75',q3),('minimum',lo),('maximum',hi)]:close(v,distribution.loc[metric,key])
        assert favorable==20 and distribution.loc[metric,'interpretation']=='descriptive_split_distribution_not_confidence_interval'
        record_b.append({'Metric':METRIC_LABELS[i],**{k:float(original_b.iloc[i][k]) for k in ['Median','Q25','Q75','Minimum','Maximum']},'Favorable_repetitions':favorable,'Total_repetitions':20})
        rows_b.append([METRIC_LABELS[i],f'{med:+.3f} [{q1:.3f}, {q3:.3f}]',f'[{lo:.3f}, {hi:.3f}]',f'{favorable}/20'])
    receipt=json.loads((ROOT/'reports/ucsd_paper_materials/generation_receipt.json').read_text());assert receipt['output_sha256'][str(FIGURE.relative_to(ROOT))]==sha(FIGURE)
    svg=etree.parse(str(FIGURE.with_suffix('.svg')));text_values=svg.xpath('//s:text/text()',namespaces={'s':'http://www.w3.org/2000/svg'})
    assert all(name in text_values for name in MODELS) and not re.search(r'[\u3400-\u9fff]',' '.join(text_values))
    specs=[dict(**TABLE_SPECS[0],rows=rows_a),dict(**TABLE_SPECS[1],rows=rows_b)]
    for spec in specs:assert not re.search(r'[\u3400-\u9fff]',' '.join(spec['headers']+sum(spec['rows'],[])+spec['notes']))
    return {'frozen_file_sha256':frozen,'translated_tables':specs,'full_precision_table_records':{'TableS5A':record_a,'TableS5B':record_b},
      'repetitions_checked':20,'repeated_prediction_rows_checked':3640,'folds_checked':fold_count,'saved_successful_fold_model_records_checked':fit_records,
      'fixed_oof_metric_values_reconstructed':320,'within_repetition_metric_differences_checked':80,'table_s5a_numeric_values_checked':24,'table_s5b_numeric_values_checked':28,
      'first_repetition_main_oof_values_checked':728,'figure_unchanged_from_saved_materials':True,
      'descriptive_distributions_not_intervals':True,'numeric_values_changed':0,'models_refitted':0,'new_hypothesis_tests':0,
      'verification_scope':'Inspect all 20 saved status and fold records, reconstruct four metrics from fixed OOF probabilities, verify within-repeat differences and 52 original table numbers, and retain the unchanged figure without refitting or selecting a partition'}

def add_table(doc,spec):
    title=doc.add_paragraph(spec['title'],style='Heading 1');title.paragraph_format.page_break_before=True
    title.paragraph_format.keep_with_next=True
    table=doc.add_table(rows=1,cols=4);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    widths=[2.10,1.84,1.84,1.14]
    for col,width in zip(table.columns,widths):col.width=Inches(width)
    for cell,width in zip(table.rows[0].cells,widths):cell.width=Inches(width)
    for cell,label in zip(table.rows[0].cells,spec['headers']):cell.text=label
    for values in spec['rows']:
        for cell,value,width in zip(table.add_row().cells,values,widths):cell.text=value;cell.width=Inches(width)
    props=table._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side)
        for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:e.set(qn('w:'+k),v)
        borders.append(e)
    props.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for side in ['top','left','bottom','right']:
        e=OxmlElement('w:'+side);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
    props.append(margins)
    for ri,row in enumerate(table.rows):
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if ri==0:trpr.append(OxmlElement('w:tblHeader'))
        for ci,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7EBEF');cell._tc.get_or_add_tcPr().append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:run.font.size=Pt(10.5);run.font.bold=ri==0
    for note in spec['notes']:
        p=doc.add_paragraph(note);p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(4)
        p.paragraph_format.keep_together=True
        for r in p.runs:r.font.size=Pt(10)

def build():
    OUT.mkdir(parents=True,exist_ok=True);REPORT.mkdir(parents=True,exist_ok=True);a=audit()
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(.79)
    for name in ['Normal','Title','Heading 1']:
        st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        rf=st.element.get_or_add_rPr().rFonts;rf.set(qn('w:eastAsia'),'Noto Serif CJK SC')
        for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:rf.attrib.pop(qn('w:'+key),None)
    normal=doc.styles['Normal'];normal.font.size=Pt(12);normal.paragraph_format.line_spacing=1.;normal.paragraph_format.space_after=Pt(9)
    doc.styles['Title'].font.size=Pt(16);doc.styles['Title'].paragraph_format.space_after=Pt(12)
    doc.styles['Heading 1'].font.size=Pt(13);doc.styles['Heading 1'].paragraph_format.space_before=Pt(8);doc.styles['Heading 1'].paragraph_format.space_after=Pt(9)
    doc.core_properties.title='Results 3.7 Split stability';doc.core_properties.author='Tianyu Feng'
    doc.add_paragraph('Results',style='Title');doc.add_paragraph('3.7 Split stability',style='Heading 1')
    for item in PARAGRAPHS:
        en=doc.add_paragraph(item['english']);en.paragraph_format.page_break_before=False
        en.paragraph_format.keep_together=True;en.paragraph_format.keep_with_next=True
        zh=doc.add_paragraph(item['chinese']);zh.paragraph_format.keep_together=True
        for r in zh.runs:r.font.size=Pt(11)
    for spec in a['translated_tables']:
        add_table(doc,spec)
        pd.DataFrame(a['full_precision_table_records'][spec['key']]).to_csv(OUT/(spec['key']+'_EN.csv'),index=False)
    landscape=doc.add_section(WD_SECTION.NEW_PAGE);landscape.orientation=WD_ORIENT.LANDSCAPE
    landscape.page_width=Inches(11);landscape.page_height=Inches(8.5)
    landscape.top_margin=landscape.bottom_margin=landscape.left_margin=landscape.right_margin=Inches(.79)
    doc.add_paragraph('Figure 3',style='Heading 1');doc.add_picture(str(FIGURE),width=Inches(9.3))
    p=doc.add_paragraph(FIGURE_CAPTION);p.paragraph_format.keep_together=True
    for r in p.runs:r.font.size=Pt(10)
    for e in [*doc.styles.element.xpath('.//w:pBdr'),*doc.element.xpath('.//w:pBdr')]:e.getparent().remove(e)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    path=OUT/'UCSD_Manuscript_Results_3_7_EN_ZH.docx';doc.save(path)
    (OUT/'UCSD_Manuscript_Results_3_7_EN.txt').write_text('\n\n'.join(p['english'] for p in PARAGRAPHS)+'\n')
    d={'stage':39,'date':'2026-10-29','subsection':'3.7 Split stability','v2_paragraph_ids':['R7'],'paragraphs':PARAGRAPHS,
      'english_word_count':sum(wc(p['english']) for p in PARAGRAPHS),'paragraph_word_counts':{p['id']:wc(p['english']) for p in PARAGRAPHS},
      'word_count_scope':'English narrative only; excludes headings, tables/notes, figure/caption and Chinese translation','v2_approximate_target_words':150,
      'equations':{'count':0,'numbers':[],'cumulative_previous_numbered_groups':10},
      'citation_mapping':{},'new_manuscript_references':0,'cumulative_unique_cited_sources':5,'editable_english_tables':['Table S5A','Table S5B'],'table_numeric_values_changed':0,'included_figure':str(FIGURE.relative_to(ROOT)),'figure_sha256':sha(FIGURE),'figure_caption':FIGURE_CAPTION,'figure_reused_without_regeneration':True,'source_audit':a,'models_refitted':0,'new_hypothesis_tests':0,
      'next_section':'Results 3.8 Single-scan sensitivity (R8)','docx':str(path.relative_to(ROOT)),'docx_sha256':sha(path)}
    (REPORT/'results_section37.json').write_text(json.dumps(d,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'docx':d['docx'],'english_word_count':d['english_word_count'],'paragraph_word_counts':d['paragraph_word_counts'],'prior_files':len(a['frozen_file_sha256'])},indent=2))

def verify():
    from pypdf import PdfReader
    d=json.loads((REPORT/'results_section37.json').read_text())
    for p,h in d['source_audit']['frozen_file_sha256'].items():assert sha(ROOT/p)==h,p
    path=ROOT/d['docx'];assert sha(path)==d['docx_sha256'];texts=[p.text for p in Document(path).paragraphs]
    for p in PARAGRAPHS:
        assert p['english'] in texts and p['chinese'] in texts and texts.index(p['chinese'])==texts.index(p['english'])+1
        assert not re.search(r'[\u3400-\u9fff，。；：（）【】、]',p['english'])
    assert sum(wc(p['english']) for p in PARAGRAPHS)==d['english_word_count']
    doc=Document(path);assert len(doc.tables)==2
    for table,spec in zip(doc.tables,d['source_audit']['translated_tables']):
        assert [[c.text for c in row.cells] for row in table.rows]==[spec['headers'],*spec['rows']]
        assert all(not re.search(r'[\u3400-\u9fff]',c.text) for row in table.rows for c in row.cells)
        for note in spec['notes']:assert note in texts
        full=pd.read_csv(OUT/(spec['key']+'_EN.csv'));assert len(full)==4
        for (_,row),r in zip(full.iterrows(),d['source_audit']['full_precision_table_records'][spec['key']]):
            for key,value in r.items():
                if isinstance(value,str):assert row[key]==value
                else:close(row[key],value)
    assert len(doc.sections)==2 and doc.sections[1].orientation==WD_ORIENT.LANDSCAPE and FIGURE_CAPTION in texts
    ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
    with zipfile.ZipFile(path) as z:
        t=etree.fromstring(z.read('word/document.xml'));assert len(t.xpath('//m:oMath',namespaces=ns))==0
        media=[n for n in z.namelist() if n.startswith('word/media/')];assert len(media)==1
        assert hashlib.sha256(z.read(media[0])).hexdigest()==sha(FIGURE)
    pages=len(PdfReader(OUT/'docx_render/UCSD_Manuscript_Results_3_7_EN_ZH.pdf').pages);assert pages==4
    v={'success':True,'stage':39,'date':'2026-10-29','subsection':d['subsection'],'paragraph_ids':['R7'],
      'english_word_count':d['english_word_count'],'paragraph_word_counts':d['paragraph_word_counts'],'matched_chinese_paragraphs':1,
      'word_pages':pages,'visual_review_pages':[1,2,3,4],'native_editable_equations':0,'equation_numbers':[],'equation_images':0,
      'prior_files_unchanged':len(d['source_audit']['frozen_file_sha256']),'repetitions_checked':20,'repeated_prediction_rows_checked':3640,'folds_checked':100,'saved_successful_fold_model_records_checked':400,
      'fixed_oof_metric_values_reconstructed':320,'within_repetition_metric_differences_checked':80,'table_s5a_numeric_values_checked':24,'table_s5b_numeric_values_checked':28,
      'first_repetition_main_oof_values_checked':728,'editable_english_tables':2,'table_numeric_values_changed':0,'descriptive_distributions_not_intervals':True,'unchanged_figure_embedded':True,'figure_page_landscape':True,
      'models_refitted':0,'new_hypothesis_tests':0,'new_manuscript_references':0,'cumulative_unique_cited_sources':5,'next_section':d['next_section'],
      'sha256':{str(p.relative_to(ROOT)):sha(p) for p in OUT.iterdir() if p.is_file() and p.suffix in {'.docx','.txt','.csv'}}}
    (REPORT/'artifact_verification.json').write_text(json.dumps(v,ensure_ascii=False,indent=2)+'\n');print(json.dumps(v,indent=2))

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--verify',action='store_true');args=parser.parse_args()
    verify() if args.verify else build()
