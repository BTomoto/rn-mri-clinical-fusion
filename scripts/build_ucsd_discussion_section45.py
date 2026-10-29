"""Draft D5 using frozen feature definitions and checked primary sources."""
from pathlib import Path
import argparse, ast, hashlib, json, re, zipfile
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'outputs/ucsd-manuscript-step45-20261030'
REPORT = ROOT/'reports/ucsd_manuscript_step45'
EVIDENCE = REPORT/'literature_evidence.json'
MAPPING = {'4': 'S32', '12': 'S39'}
PARAGRAPHS = [{
    'id': 'D5', 'citation_ids': ['S32', 'S39'],
    'english': 'The eight MRI inputs summarize different statistical aspects of the supplied annotation region. Mean intensity describes its average percentile-scaled signal, and standard deviation describes the spread of those values. Histogram entropy summarizes uncertainty in the discretized gray-level distribution without retaining voxel positions. GLCM contrast weights squared gray-level differences between neighboring voxels and therefore incorporates local spatial variation (4). These descriptors capture aspects of regional image appearance, with their interpretation depending on the supplied images and extraction rules. The union of all positive annotation labels does not isolate a histologically verified viable-tumor compartment. Duenweg et al. showed that normalization choices affected radiomic comparability across acquisition settings, including in brain imaging (12). Their findings do not establish an optimal transformation for our derived images. Here, percentile scaling and ROI-based discretization define the numerical context of the features; unknown spatial units further limit interpretation of texture at a physical distance. The fitted coefficients describe conditional associations within a penalized model, with their interpretation depending on feature scaling and the other included inputs. They do not establish independent risk factors or causal effects. Accordingly, the observed incremental discrimination supports a collective predictive signal within this dataset, without identifying a specific cellular mechanism.',
    'chinese': '8项MRI输入概括了现成标注区域的不同统计特征。平均强度描述该区域经百分位缩放后的平均信号，标准差描述这些数值的离散程度。直方图熵概括离散化灰度分布的不确定性，不保留体素的位置。GLCM对比度对相邻体素间灰度差的平方进行加权，因此纳入了局部空间变化（4）。这些描述指标反映区域影像外观的某些方面，其解释取决于所提供的图像和提取规则。全部正值标注标签的并集，并不单独对应经组织学验证的存活肿瘤区室。Duenweg等发现，归一化选择会影响不同采集条件间的影像组学可比性，其中包括脑部影像（12）。其结果并不能确定我们这些派生图像的最优变换。在本研究中，百分位缩放和基于ROI的离散化界定了特征的数值含义；空间单位未知进一步限制了在物理距离上解释纹理的能力。拟合系数描述正则化模型内的条件关联，其解释取决于特征缩放和其他纳入的输入，并不能确立独立危险因素或因果效应。因此，所观察到的增量鉴别能力支持该数据集内存在整体预测信号，但不能据此确定某种具体的细胞机制。'
}]

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def wc(text):
    return len(re.findall(r"[A-Za-z0-9]+(?:['/-][A-Za-z0-9]+)*", text))

def audit():
    previous = json.loads((ROOT/'reports/ucsd_manuscript_step44/discussion_section44.json').read_text())
    frozen = previous['source_audit']['frozen_file_sha256'].copy()
    frozen.update(json.loads((ROOT/'reports/ucsd_manuscript_step44/artifact_verification.json').read_text())['sha256'])
    extra = ['reports/ucsd_manuscript_step26/methods_section24.json', 'reports/ucsd_manuscript_step27/methods_section25.json']
    frozen.update({p: sha(ROOT/p) for p in extra if p not in frozen})
    frozen[str(EVIDENCE.relative_to(ROOT))] = sha(EVIDENCE)
    for p, h in frozen.items(): assert sha(ROOT/p) == h, p
    evidence = json.loads(EVIDENCE.read_text()); records = {r['reference_id']: r for r in evidence['studies']}
    library = {r['id']: r for r in json.loads((ROOT/'reports/ucsd_manuscript_step23/reference_library.json').read_text())['references']}
    feature_section = json.loads((ROOT/extra[1]).read_text())
    reference = feature_section['reference_used']; assert reference['pages_or_article'] == '328-338' and ':328-338.' in reference['citation']
    refs = {'S32': reference, 'S39': library['S39']}; selected = []
    for number, rid in MAPPING.items():
        ref = refs[rid]; r = records[rid]
        assert ref['doi'] == r['doi'] and r['provisional_citation_number'] == int(number)
        selected.append({'number': int(number), 'reference_id': rid, 'citation': ref['citation'], 'doi': ref['doi'],
                         'eligibility_year': ref['eligibility_year'], 'read_level_this_stage': r['read_level_this_stage']})
    assert library['S39']['english_flag'] == library['S39']['recent_5yr'] == library['S39']['recent_3yr'] == 1
    assert library['S39']['eligibility_year'] == records['S39']['facts']['publication_year'] == 2025
    code = (ROOT/'src/rn_public/ucsd_batch_features.py').read_text(); tree = ast.parse(code)
    seq = ast.get_source_segment(code, next(n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name == 'sequence_features'))
    checks = ['mean=float(values.mean())', 'std=float(values.std(ddof=0))',
              'entropy=float(-np.sum(p*np.log2(p)))', 'glcm_contrast=float(np.sum(glcm*difference**2))']
    assert all(s in seq for s in checks)
    assert 'p=counts/counts.sum();p=p[p>0]' in seq and 'quantized,lo,hi=fixed_bins(v)' in seq
    assert 'bins[roi]=quantized' in seq and 'difference=np.arange(32)[:,None]-np.arange(32)[None,:]' in seq
    glcm = ast.get_source_segment(code, next(n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name == 'symmetric_merged_glcm'))
    assert 'matrix+=counts+counts.T' in glcm and 'return matrix/total,total,used' in glcm
    roi = (ROOT/'src/rn_public/ucsd_roi_pilot.py').read_text(); rtree = ast.parse(roi)
    union = ast.get_source_segment(roi, next(n for n in rtree.body if isinstance(n, ast.FunctionDef) and n.name == 'region_union'))
    bins = ast.get_source_segment(roi, next(n for n in rtree.body if isinstance(n, ast.FunctionDef) and n.name == 'fixed_bins'))
    assert 'mask = a > 0' in union and 'return mask' in union
    assert 'lo = float(v.min())' in bins and 'hi = float(v.max())' in bins
    affine = (ROOT/'src/rn_public/ucsd_intensity_followup.py').read_text(); atree = ast.parse(affine)
    transform = ast.get_source_segment(affine, next(n for n in atree.body if isinstance(n, ast.FunctionDef) and n.name == 'affine_values'))
    assert "return (a-params['p01'])/(params['p99']-params['p01'])" in transform and 'clip' not in transform
    saved = feature_section['source_audit']; assert saved['saved_feature_values_checked'] == 1456 and len(saved['offsets']) == 13
    expected = [f'{s}_{f}' for s in ['FLAIR', 'T1post'] for f in ['mean', 'std', 'entropy', 'glcm_contrast']]
    main = json.loads((ROOT/'reports/ucsd_fusion/metrics.json').read_text()); assert main['models']['MRI'] == expected == saved['MRI_feature_names']
    assert not saved['full_IBSI_benchmark_validation_claimed'] and saved['study_pixels_redecoded'] == 0
    image = json.loads((ROOT/extra[0]).read_text())['source_audit']
    assert image['saved_QC_file_records_checked'] == 546 and image['spatial_unit'] == 'unknown'
    assert image['stored_dtype'] == 'uint8' and not image['annotation_blinding_verified']
    assert not main['independent_external_validation'] and main['C'] == 1
    primary = next(r for r in main['paired_comparisons'] if r['comparison'] == 'Fusion_minus_M1' and r['metric'] == 'roc_auc')
    assert primary['difference'] > 0 and primary['lower'] > 0
    prior_ids = set()
    for pattern in ['methods_section*.json', 'results_section*.json', 'discussion_section*.json']:
        for p in ROOT.glob('reports/ucsd_manuscript_step*/'+pattern):
            if int(p.parent.name.rsplit('step', 1)[1]) < 45:
                prior_ids.update(r for item in json.loads(p.read_text())['paragraphs'] for r in item['citation_ids'])
    assert len(prior_ids) == 11 and 'S32' in prior_ids and 'S39' not in prior_ids
    return {'frozen_file_sha256': frozen, 'reference_records': selected,
            'source_evidence_record': str(EVIDENCE.relative_to(ROOT)), 'new_full_texts_checked_for_current_claims': ['S39'],
            'official_feature_specification_reviewed': 'IBSI statistical features, histogram entropy TLU2 and GLCM contrast ACUI',
            'feature_definition_checks': 4, 'MRI_feature_names_checked': 8, 'GLCM_offsets_in_existing_definition_record': 13,
            'saved_image_QC_record_count': 546, 'unknown_spatial_units_retained': True,
            'annotation_region_is_union_of_positive_labels': True, 'full_IBSI_benchmark_validation_claimed': False,
            'independent_risk_factor_or_causal_effect_claimed': False, 'specific_cellular_mechanism_identified': False,
            'normalization_study_claims_our_transform_optimal': False, 'all_features_claimed_sensitive_to_every_affine_transform': False,
            'S32_complete_pagination_reused': '328-338', 'previous_unique_cited_sources': sorted(prior_ids),
            'models_refitted': 0, 'features_reextracted': 0, 'new_hypothesis_tests': 0,
            'verification_scope': 'Check frozen definitions, prior QC and own-study comparison plus manually reviewed external evidence; no image decoding, feature extraction, model fits or new association analyses'}

def build():
    a = audit(); OUT.mkdir(parents=True, exist_ok=True); REPORT.mkdir(parents=True, exist_ok=True)
    doc = Document(); sec = doc.sections[0]; sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(.79)
    for name in ['Normal', 'Title', 'Heading 1']:
        st = doc.styles[name]; st.font.name = 'Times New Roman'; st.font.color.rgb = RGBColor(0, 0, 0)
        rf = st.element.get_or_add_rPr().rFonts; rf.set(qn('w:eastAsia'), 'Noto Serif CJK SC')
        for key in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']: rf.attrib.pop(qn('w:'+key), None)
    normal = doc.styles['Normal']; normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.; normal.paragraph_format.space_after = Pt(9)
    doc.styles['Title'].font.size = Pt(16); doc.styles['Title'].paragraph_format.space_after = Pt(12)
    heading = doc.styles['Heading 1']; heading.font.size = Pt(13)
    heading.paragraph_format.space_before = Pt(8); heading.paragraph_format.space_after = Pt(9)
    doc.core_properties.title = 'Discussion 4.5 Interpretation of image features'; doc.core_properties.author = 'Tianyu Feng'
    doc.add_paragraph('Discussion', style='Title'); doc.add_paragraph('4.5 Interpretation of image features', style='Heading 1')
    for item in PARAGRAPHS:
        en = doc.add_paragraph(item['english']); en.paragraph_format.keep_together = True; en.paragraph_format.keep_with_next = True
        zh = doc.add_paragraph(item['chinese']); zh.paragraph_format.keep_together = True
        for run in zh.runs: run.font.size = Pt(11)
    title = doc.add_paragraph('References', style='Heading 1'); title.paragraph_format.page_break_before = True
    for ref in a['reference_records']:
        p = doc.add_paragraph(str(ref['number'])+'. '+ref['citation']); p.paragraph_format.keep_together = True
    for e in [*doc.styles.element.xpath('.//w:pBdr'), *doc.element.xpath('.//w:pBdr')]: e.getparent().remove(e)
    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    path = OUT/'UCSD_Manuscript_Discussion_4_5_EN_ZH.docx'; doc.save(path)
    (OUT/'UCSD_Manuscript_Discussion_4_5_EN.txt').write_text('\n\n'.join(p['english'] for p in PARAGRAPHS)+'\n')
    d = {'stage': 45, 'date': '2026-10-30', 'subsection': '4.5 Interpretation of image features', 'v2_paragraph_ids': ['D5'],
         'paragraphs': PARAGRAPHS, 'english_word_count': sum(wc(p['english']) for p in PARAGRAPHS),
         'paragraph_word_counts': {p['id']: wc(p['english']) for p in PARAGRAPHS},
         'word_count_scope': 'English narrative including in-text citation numbers; excludes headings, reference list and Chinese translation',
         'v2_approximate_target_words': 200, 'equations': {'count': 0, 'numbers': [], 'cumulative_previous_numbered_groups': 10},
         'citation_mapping': MAPPING, 'citation_numbering': 'Reuse number 4 and add first-appearance number 12; renumber globally when Introduction is assembled',
         'new_manuscript_references': 1, 'cumulative_unique_cited_sources': 12,
         'source_audit': a, 'models_refitted': 0, 'features_reextracted': 0, 'new_hypothesis_tests': 0,
         'next_section': 'Discussion 4.6 Boundary and split robustness (D6)',
         'docx': str(path.relative_to(ROOT)), 'docx_sha256': sha(path)}
    (REPORT/'discussion_section45.json').write_text(json.dumps(d, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps({'docx': d['docx'], 'english_word_count': d['english_word_count'], 'frozen_files': len(a['frozen_file_sha256'])}))

def verify():
    from pypdf import PdfReader
    d = json.loads((REPORT/'discussion_section45.json').read_text()); a = audit(); assert a == d['source_audit']
    path = ROOT/d['docx']; assert sha(path) == d['docx_sha256']; doc = Document(path); texts = [p.text for p in doc.paragraphs]
    for p in PARAGRAPHS:
        assert p['english'] in texts and p['chinese'] in texts and texts.index(p['chinese']) == texts.index(p['english'])+1
        assert not re.search(r'[\u3400-\u9fff，。；：（）【】、]', p['english'])
    assert sum(wc(p['english']) for p in PARAGRAPHS) == d['english_word_count']
    assert re.findall(r'\((\d+)\)', PARAGRAPHS[0]['english']) == list(MAPPING)
    for ref in a['reference_records']: assert str(ref['number'])+'. '+ref['citation'] in texts
    assert not doc.tables and len(doc.sections) == 1
    with zipfile.ZipFile(path) as z:
        xml = etree.fromstring(z.read('word/document.xml'))
        assert not xml.xpath('//m:oMath', namespaces={'m': 'http://schemas.openxmlformats.org/officeDocument/2006/math'})
        assert not [n for n in z.namelist() if n.startswith('word/media/')]
    pdf = PdfReader(OUT/'docx_render/UCSD_Manuscript_Discussion_4_5_EN_ZH.pdf'); pages = len(pdf.pages); assert pages == 2
    assert '4.5 Interpretation' in pdf.pages[0].extract_text() and 'References' in pdf.pages[1].extract_text()
    v = {'success': True, 'stage': 45, 'date': '2026-10-30', 'subsection': d['subsection'], 'paragraph_ids': ['D5'],
         'english_word_count': d['english_word_count'], 'matched_chinese_paragraphs': 1, 'word_pages': pages,
         'visual_review_pages': [1, 2], 'native_editable_equations': 0, 'equation_images': 0,
         'frozen_files_unchanged': len(a['frozen_file_sha256']), 'prior_files_unchanged': len(a['frozen_file_sha256'])-1,
         'new_full_texts_checked_for_current_claims': ['S39'], 'official_IBSI_definitions_checked': True,
         'feature_definition_checks': 4, 'MRI_feature_names_checked': 8, 'new_manuscript_references': 1,
         'provisional_citation_mapping': MAPPING, 'cumulative_unique_cited_sources': 12,
         'biological_or_causal_attribution_claimed': False, 'unknown_spatial_units_retained': True,
         'models_refitted': 0, 'features_reextracted': 0, 'new_hypothesis_tests': 0, 'next_section': d['next_section'],
         'sha256': {str(p.relative_to(ROOT)): sha(p) for p in OUT.iterdir() if p.is_file() and p.suffix in {'.docx', '.txt'}}}
    (REPORT/'artifact_verification.json').write_text(json.dumps(v, ensure_ascii=False, indent=2)+'\n'); print(json.dumps(v, indent=2))

if __name__ == '__main__':
    parser = argparse.ArgumentParser(); parser.add_argument('--verify', action='store_true'); args = parser.parse_args()
    verify() if args.verify else build()
