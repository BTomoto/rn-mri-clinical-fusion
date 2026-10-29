"""Draft D7 from frozen single-scan records and primary validation guidance."""
from pathlib import Path
import argparse, hashlib, json, re, zipfile
import pandas as pd
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'outputs/ucsd-manuscript-step47-20261030'
REPORT = ROOT/'reports/ucsd_manuscript_step47'
EVIDENCE = REPORT/'literature_evidence.json'
MAPPING = {'15': 'S25', '16': 'S26'}
PARAGRAPHS = [{
    'id': 'D7', 'citation_ids': ['S25', 'S26'],
    'english': "The single-scan analysis leaves Fusion's incremental discrimination uncertain under the selected sampling scheme. Its conditional paired AUC-difference interval ranged from -0.015 to 0.252, accommodating a small disadvantage as well as a larger gain. Selection retained the same 135 patients but reduced treatment-related-change examinations from 41 to 22 and changed the label distribution, training samples and fold assignments. These concurrent changes prevent attributing the difference from the primary analysis solely to repeated scans. The smallest available numeric session also did not establish the first clinical visit. At the fixed threshold, Fusion correctly recognized only 2/22 treatment-related-change examinations, despite identifying most tumor examinations. This class-specific limitation remains relevant even when ranking performance is summarized by AUC. Equalizing contributions to one examination per patient does not create a cohort independent of the primary analysis or verify the source diagnoses. Guidance for external validation emphasizes a dataset representative of the target population and sample-size planning for precise performance estimates (15,16). Accordingly, the selected analysis is best interpreted as an internal sensitivity check. Confirmation of the increment requires independent data with adequate representation of both outcomes and evaluation of discrimination, calibration and class-specific performance.",
    'chinese': '单扫描分析中，Fusion在所选抽样方案下的增量鉴别仍不确定。其条件性配对AUC差值区间为−0.015至0.252，既容许小幅劣势，也容许较大的增益。选择后仍为同一批135名患者，但治疗相关改变检查由41次减少至22次，标签分布、训练样本及分折分配也发生了变化。这些同时发生的变化，使我们不能将相较主分析的差异仅归因于重复扫描。最小可用数字会话编号也不能确定首次临床就诊。在固定阈值下，Fusion仅正确识别2/22次治疗相关改变检查，尽管识别出了大多数肿瘤检查。即使采用AUC概括排序能力，这一类别特异性的不足仍然重要。将贡献统一为每位患者一次检查，并不会形成独立于主分析的队列，也不能核实源诊断。外部验证指南强调，数据集应代表目标人群，并应为精确估计性能而规划样本量（15,16）。因此，所选分析更适合解释为内部敏感性检查。确认增量需要独立数据，充分覆盖两种结局，并评价鉴别能力、校准及各类别的表现。'
}]

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def wc(text):
    return len(re.findall(r"[A-Za-z0-9]+(?:['/-][A-Za-z0-9]+)*", text))

def audit():
    previous = json.loads((ROOT/'reports/ucsd_manuscript_step46/discussion_section46.json').read_text())
    frozen = previous['source_audit']['frozen_file_sha256'].copy()
    frozen.update(json.loads((ROOT/'reports/ucsd_manuscript_step46/artifact_verification.json').read_text())['sha256'])
    extra = 'reports/ucsd_manuscript_step40/results_section38.json'
    if extra not in frozen: frozen[extra] = sha(ROOT/extra)
    frozen[str(EVIDENCE.relative_to(ROOT))] = sha(EVIDENCE)
    for p, h in frozen.items(): assert sha(ROOT/p) == h, p
    evidence = json.loads(EVIDENCE.read_text()); records = {r['reference_id']: r for r in evidence['studies']}
    refs = {r['id']: r for r in json.loads((ROOT/'reports/ucsd_manuscript_step23/reference_library.json').read_text())['references']}
    selected = []
    for number, rid in MAPPING.items():
        ref = refs[rid]; r = records[rid]
        assert ref['doi'] == r['doi'] and r['provisional_citation_number'] == int(number)
        assert ref['english_flag'] == ref['recent_5yr'] == ref['recent_3yr'] == 1 and ref['eligibility_year'] == 2024
        selected.append({'number': int(number), 'reference_id': rid, 'citation': ref['citation'], 'doi': ref['doi'],
                         'eligibility_year': 2024, 'read_level_this_stage': r['read_level_this_stage']})
    main = json.loads((ROOT/'reports/ucsd_fusion/metrics.json').read_text())
    stability = json.loads((ROOT/'reports/ucsd_stability/metrics.json').read_text()); single = stability['single_scan']
    ci = single['bootstrap_95ci']['Fusion_minus_M1__roc_auc']
    assert f"{ci['lower']:.3f}" == '-0.015' and f"{ci['upper']:.3f}" == '0.252' and ci['lower'] < 0 < ci['upper']
    full = pd.read_csv(ROOT/'data/processed/ucsd_clinical_image_features.csv')
    prior = pd.read_csv(ROOT/'data/processed/ucsd_fusion_oof.csv').set_index('ID')
    pred = pd.read_csv(ROOT/'data/processed/ucsd_stability_single_scan_oof.csv').set_index('ID')
    ledger = pd.read_csv(ROOT/'reports/ucsd_stability/single_scan_selection.csv')
    assert len(full) == len(prior) == 182 and len(pred) == 135
    assert pred.patient_id.nunique() == full.patient_id.nunique() == 135 and set(pred.patient_id) == set(full.patient_id)
    identifiers = full[['ID', 'patient_id']].copy(); identifiers['session'] = identifiers.ID.str.rsplit('_', n=1).str[-1].astype(int)
    chosen = set(identifiers.sort_values(['patient_id', 'session', 'ID']).drop_duplicates('patient_id').ID)
    assert chosen == set(pred.index) == set(ledger.loc[ledger.selected, 'ID'])
    assert pred.tumor_label.equals(prior.loc[pred.index, 'tumor_label'])
    assert int(prior.tumor_label.eq(0).sum()) == main['treatment_change_scans'] == 41
    assert int(pred.tumor_label.eq(0).sum()) == single['treatment_change_scans'] == 22
    assert int(pred.tumor_label.eq(1).sum()) == 113 and pred.index.is_unique
    fold_cross = pd.crosstab(prior.loc[pred.index, 'validation_fold'], pred.validation_fold)
    assert fold_cross.gt(0).sum(axis=1).gt(1).any()  # Different partitions, not merely different fold numbers.
    negative = pred.Fusion_oof_probability.lt(.5)
    tn = int((negative & pred.tumor_label.eq(0)).sum()); tp = int((~negative & pred.tumor_label.eq(1)).sum())
    assert tn == single['metrics']['Fusion']['tn'] == 2 and tp == single['metrics']['Fusion']['tp'] == 109
    assert main['threshold'] == .5 and not stability['threshold_optimization'] and not stability['external_validation']
    assert single['complete'] and stability['bootstrap_models_refitted'] == 0
    lock = json.loads((ROOT/'reports/ucsd_stability/analysis_lock.json').read_text())
    assert lock['single_scan_rule'] == 'minimum_numeric_session_within_current_imaged_subset_then_ID'
    old = json.loads((ROOT/extra).read_text())['source_audit']
    assert old['selected_patients'] == old['selected_examinations_checked'] == 135 and old['models_refitted'] == 0
    prior_ids = set()
    for pattern in ['methods_section*.json', 'results_section*.json', 'discussion_section*.json']:
        for p in ROOT.glob('reports/ucsd_manuscript_step*/'+pattern):
            if int(p.parent.name.rsplit('step', 1)[1]) < 47:
                prior_ids.update(r for item in json.loads(p.read_text())['paragraphs'] for r in item['citation_ids'])
    assert len(prior_ids) == 14 and not prior_ids.intersection(MAPPING.values())
    return {'frozen_file_sha256': frozen, 'reference_records': selected, 'source_evidence_record': str(EVIDENCE.relative_to(ROOT)),
            'new_full_texts_checked_for_current_claims': ['S25', 'S26'], 'conditional_paired_AUC_interval_checked': ci,
            'same_primary_and_selected_patient_set_verified': True, 'selected_examinations_and_patients': 135,
            'treatment_change_counts_checked': {'primary': 41, 'single_scan': 22},
            'per_examination_source_labels_unchanged': True, 'identifier_only_selection_reconstructed': True,
            'selected_partitions_differ_beyond_fold_numbering': True,
            'fixed_threshold_recognition_checked': {'treatment_change': [2, 22], 'tumor': [109, 113]},
            'absence_of_increment_claimed_proven': False, 'performance_change_attributed_only_to_repeated_scans': False,
            'new_external_cohort_or_label_validation_claimed': False, 'retrospective_power_or_sample_size_calculated': False,
            'previous_unique_cited_sources': sorted(prior_ids), 'models_refitted': 0, 'new_hypothesis_tests': 0,
            'verification_scope': 'Inspect saved interval, identifier-only selection, unchanged labels, same-patient set, fold membership and fixed-threshold counts, plus manually reviewed guidance; no fitting, resampling or threshold search'}

def build():
    a = audit(); OUT.mkdir(parents=True, exist_ok=True); REPORT.mkdir(parents=True, exist_ok=True)
    doc = Document(); sec = doc.sections[0]; sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(.79)
    for name in ['Normal', 'Title', 'Heading 1']:
        st = doc.styles[name]; st.font.name = 'Times New Roman'; st.font.color.rgb = RGBColor(0, 0, 0)
        rf = st.element.get_or_add_rPr().rFonts; rf.set(qn('w:eastAsia'), 'Noto Serif CJK SC')
        for key in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']: rf.attrib.pop(qn('w:'+key), None)
    normal = doc.styles['Normal']; normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.; normal.paragraph_format.space_after = Pt(9)
    doc.styles['Title'].font.size = Pt(16); doc.styles['Title'].paragraph_format.space_after = Pt(12)
    heading = doc.styles['Heading 1']; heading.font.size = Pt(13)
    heading.paragraph_format.space_before = Pt(8); heading.paragraph_format.space_after = Pt(9)
    doc.core_properties.title = 'Discussion 4.7 Meaning of the single-scan result'; doc.core_properties.author = 'Tianyu Feng'
    doc.add_paragraph('Discussion', style='Title'); doc.add_paragraph('4.7 Meaning of the single-scan result', style='Heading 1')
    for item in PARAGRAPHS:
        en = doc.add_paragraph(item['english']); en.paragraph_format.keep_together = True; en.paragraph_format.keep_with_next = True
        zh = doc.add_paragraph(item['chinese']); zh.paragraph_format.keep_together = True
        for run in zh.runs: run.font.size = Pt(11)
    title = doc.add_paragraph('References', style='Heading 1'); title.paragraph_format.page_break_before = True
    for ref in a['reference_records']:
        p = doc.add_paragraph(str(ref['number'])+'. '+ref['citation']); p.paragraph_format.keep_together = True
    for e in [*doc.styles.element.xpath('.//w:pBdr'), *doc.element.xpath('.//w:pBdr')]: e.getparent().remove(e)
    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    path = OUT/'UCSD_Manuscript_Discussion_4_7_EN_ZH.docx'; doc.save(path)
    (OUT/'UCSD_Manuscript_Discussion_4_7_EN.txt').write_text('\n\n'.join(p['english'] for p in PARAGRAPHS)+'\n')
    d = {'stage': 47, 'date': '2026-10-30', 'subsection': '4.7 Meaning of the single-scan result', 'v2_paragraph_ids': ['D7'],
         'paragraphs': PARAGRAPHS, 'english_word_count': sum(wc(p['english']) for p in PARAGRAPHS),
         'paragraph_word_counts': {p['id']: wc(p['english']) for p in PARAGRAPHS},
         'word_count_scope': 'English narrative including in-text citation numbers; excludes headings, reference list and Chinese translation',
         'v2_approximate_target_words': 200, 'equations': {'count': 0, 'numbers': [], 'cumulative_previous_numbered_groups': 10},
         'citation_mapping': MAPPING, 'citation_numbering': 'First-appearance numbers 15-16; renumber globally when Introduction is assembled',
         'new_manuscript_references': 2, 'cumulative_unique_cited_sources': 16,
         'source_audit': a, 'models_refitted': 0, 'new_hypothesis_tests': 0,
         'next_section': 'Discussion 4.8 Reference-standard limitations (D8)',
         'docx': str(path.relative_to(ROOT)), 'docx_sha256': sha(path)}
    (REPORT/'discussion_section47.json').write_text(json.dumps(d, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps({'docx': d['docx'], 'english_word_count': d['english_word_count'], 'frozen_files': len(a['frozen_file_sha256'])}))

def verify():
    from pypdf import PdfReader
    d = json.loads((REPORT/'discussion_section47.json').read_text()); a = audit(); assert a == d['source_audit']
    path = ROOT/d['docx']; assert sha(path) == d['docx_sha256']; doc = Document(path); texts = [p.text for p in doc.paragraphs]
    for p in PARAGRAPHS:
        assert p['english'] in texts and p['chinese'] in texts and texts.index(p['chinese']) == texts.index(p['english'])+1
        assert not re.search(r'[\u3400-\u9fff，。；：（）【】、]', p['english'])
    assert sum(wc(p['english']) for p in PARAGRAPHS) == d['english_word_count']
    assert re.findall(r'\((\d+),(\d+)\)', PARAGRAPHS[0]['english']) == [('15', '16')]
    for ref in a['reference_records']: assert str(ref['number'])+'. '+ref['citation'] in texts
    assert not doc.tables and len(doc.sections) == 1
    with zipfile.ZipFile(path) as z:
        xml = etree.fromstring(z.read('word/document.xml'))
        assert not xml.xpath('//m:oMath', namespaces={'m': 'http://schemas.openxmlformats.org/officeDocument/2006/math'})
        assert not [n for n in z.namelist() if n.startswith('word/media/')]
    pdf = PdfReader(OUT/'docx_render/UCSD_Manuscript_Discussion_4_7_EN_ZH.pdf'); pages = len(pdf.pages); assert pages == 2
    assert '4.7 Meaning' in pdf.pages[0].extract_text() and 'References' in pdf.pages[1].extract_text()
    v = {'success': True, 'stage': 47, 'date': '2026-10-30', 'subsection': d['subsection'], 'paragraph_ids': ['D7'],
         'english_word_count': d['english_word_count'], 'matched_chinese_paragraphs': 1, 'word_pages': pages,
         'visual_review_pages': [1, 2], 'native_editable_equations': 0, 'equation_images': 0,
         'frozen_files_unchanged': len(a['frozen_file_sha256']), 'prior_files_unchanged': len(a['frozen_file_sha256'])-1,
         'new_full_texts_checked_for_current_claims': ['S25', 'S26'], 'same_patient_set_verified': True,
         'selected_examinations_and_patients': 135, 'identifier_only_selection_reconstructed': True,
         'per_examination_source_labels_unchanged': True, 'selected_partitions_differ_beyond_fold_numbering': True,
         'fixed_threshold_recognition_checked': {'treatment_change': [2, 22], 'tumor': [109, 113]},
         'new_manuscript_references': 2, 'provisional_citation_mapping': MAPPING, 'cumulative_unique_cited_sources': 16,
         'new_external_validation_claimed': False, 'retrospective_power_or_sample_size_calculated': False,
         'models_refitted': 0, 'new_hypothesis_tests': 0, 'next_section': d['next_section'],
         'sha256': {str(p.relative_to(ROOT)): sha(p) for p in OUT.iterdir() if p.is_file() and p.suffix in {'.docx', '.txt'}}}
    (REPORT/'artifact_verification.json').write_text(json.dumps(v, ensure_ascii=False, indent=2)+'\n'); print(json.dumps(v, indent=2))

if __name__ == '__main__':
    parser = argparse.ArgumentParser(); parser.add_argument('--verify', action='store_true'); args = parser.parse_args()
    verify() if args.verify else build()
