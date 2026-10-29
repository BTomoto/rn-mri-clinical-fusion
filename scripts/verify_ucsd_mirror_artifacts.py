"""Reconcile real image bytes, notebook output, Excel data and preserved clinical files."""
from pathlib import Path
import base64
import hashlib
import json
import subprocess

import nbformat
from openpyxl import load_workbook
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT/'reports/ucsd_mri_mirror'
OUT = ROOT/'outputs/ucsd-mri-mirror-20261029'
REFERENCE = '76a6d7f045690820ef4847ab03f812054b0d7cd6'
PRESERVE = [
    'data/raw/reference/ucsd_ptgbm_clinical_v3.xlsx',
    'data/raw/reference/ucsd_ptgbm_negative_cases.xlsx',
    'data/processed/ucsd_clinical_revalidated_cohort.csv',
    'data/processed/ucsd_clinical_revalidated_oof.csv',
    'reports/ucsd_clinical_revalidation/metrics.json',
    'outputs/ucsd-clinical-revalidation-20261029/UCSD_Clinical_Revalidation.xlsx',
    'src/rn_public/ucsd_clinical_revalidation.py',
    'src/rn_public/ucsd_mri_intake.py',
    'notebooks/13_ucsd_clinical_revalidation.ipynb',
    'notebooks/14_ucsd_mri_intake_and_qc.ipynb',
]


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def main():
    qc = json.loads((BASE/'sample_qc_receipt.json').read_text())
    s = qc['summary']
    assert s['code_sha256'] == digest(ROOT/'src/rn_public/ucsd_mri_mirror.py')
    assert s['plan_sha256'] == digest(ROOT/'reports/ucsd_mri_mirror_sample_plan.md')
    assert len(qc['files']) == 12 and s['successful_download_files'] == s['decoded_image_files'] == 12
    assert s['technical_pass_image_files'] == 0 and s['relative_native_grid_matches'] == 12
    image_hashes = {}
    for r in qc['files']:
        path = ROOT/r['relative_path']
        image_hashes[r['relative_path']] = digest(path)
        assert image_hashes[r['relative_path']] == r['sha256'] == r['mirror_sha256']
        assert path.stat().st_size == r['bytes']
        assert r['gzip_crc_verified'] and r['decoded'] and r['nonfinite_voxels'] == 0
        assert r['spatial_units'] == 'unknown' and not r['technical_pass']
        assert r['relative_grid_to_FLAIR']['physical_units_confirmed'] is False
    assert len(set(image_hashes.values())) == 12
    data = json.loads((ROOT/'reports/ucsd_mirror_workbook_input.json').read_text())
    path = OUT/'UCSD_MRI_Mirror_Samples.xlsx'
    wb = load_workbook(path, data_only=False)
    assert wb.sheetnames == ['Summary', 'Files', 'Coverage']
    for tab, key in [('Files', 'files'), ('Coverage', 'coverage')]:
        matrix = [data[key]['headers'], *data[key]['rows']]
        sh = wb[tab]
        assert (sh.max_row, sh.max_column) == (len(matrix), len(matrix[0]))
        for i, row in enumerate(matrix, 1):
            for j, value in enumerate(row, 1):
                actual = sh.cell(i, j).value
                assert actual == (None if value == '' else value), (tab, i, j, actual, value)
    cached = load_workbook(path, data_only=True)['Summary']
    expected = {'C5': 2, 'C6': 12, 'C7': 14251565, 'C8': 12, 'C9': 12,
                'C10': 12, 'C11': 0, 'C12': 12, 'C16': 182, 'C17': 141,
                'C18': 41, 'C19': 56, 'C20': 135}
    for address, value in expected.items():
        assert cached[address].value == value, (address, cached[address].value, value)
    for sh in wb:
        for row in sh:
            for cell in row:
                assert cell.data_type != 'e', (sh.title, cell.coordinate, cell.value)
    nb = nbformat.read(ROOT/'notebooks/15_ucsd_mri_mirror_samples.ipynb', as_version=4)
    nbformat.validate(nb)
    cells = [c for c in nb.cells if c.cell_type == 'code']
    kernel = json.loads((ROOT/'reports/kernel_execution_15.json').read_text())
    assert kernel['success'] and kernel['notebooks'][0]['errors'] == 0
    code_hash = hashlib.sha256('\n'.join(c.source for c in cells).encode()).hexdigest()
    assert kernel['notebooks'][0]['code_sha256'] == code_hash
    assert [c.execution_count for c in cells] == [1, 2, 3, 4, 5]
    images = [o.data['image/png'] for c in cells for o in c.outputs
              if o.output_type in ['display_data', 'execute_result'] and 'image/png' in o.data]
    names = ['actual_FLAIR_planes.png', 'actual_multisequence_axial.png', 'actual_segmentation_overlay.png']
    assert len(images) == 3
    for encoded, name in zip(images, names):
        assert base64.b64decode(encoded) == (OUT/name).read_bytes()
    tests = json.loads((BASE/'software_test_receipt.json').read_text())
    assert tests['success'] and tests['tests_run'] == 61 and tests['errors'] == tests['failures'] == 0
    preserved = {}
    for item in PRESERVE:
        prior = subprocess.run(['git', 'show', f'{REFERENCE}:{item}'], cwd=ROOT, check=True, capture_output=True).stdout
        assert prior == (ROOT/item).read_bytes(), item
        preserved[item] = hashlib.sha256(prior).hexdigest()
    result = dict(success=True, actual_NIfTI_files_verified=12, actual_NIfTI_bytes=14251565,
                  source_sample_sha256=image_hashes, excel_file_rows_verified=12,
                  excel_coverage_rows_verified=238, workbook_sha256=digest(path),
                  notebook_code_sha256=code_hash, actual_notebook_images_verified=3,
                  MRI_model_fitted=False, prediction_ROI_approved=False,
                  physical_scale_QC_pass=False, preserved_reference_commit=REFERENCE,
                  preserved_artifacts_sha256=preserved)
    (BASE/'artifact_verification.json').write_text(json.dumps(result, indent=2)+'\n')
    print(json.dumps({k: v for k, v in result.items() if not isinstance(v, dict)}, indent=2))


if __name__ == '__main__':
    main()
