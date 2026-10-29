"""Author the first manuscript subsection in English with matched Chinese text."""
from pathlib import Path
import hashlib,json,re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
REPORT=ROOT/'reports/ucsd_manuscript_step23'
OUT=ROOT/'outputs/ucsd-manuscript-step23-20261029'

PARAGRAPHS=[
 {'id':'M1','citation_ids':['S01','S02','S03'],
  'english': 'We performed a retrospective secondary analysis to evaluate whether quantitative magnetic resonance imaging (MRI) features added discriminatory information to clinical variables for distinguishing residual or recurrent tumor from treatment-related change on post-treatment glioma examinations. The target was the diagnostic status assigned to each MRI examination. Clinical, MRI, and combined models were evaluated in a common cohort using cross-validation with patients kept separate between training and validation folds. The examination was the unit of prediction, and all examinations used for modeling from a given patient were assigned to the same fold. The analysis used publicly available, de-identified records and images, without participant recruitment, patient contact, additional imaging, or specimen collection. The dataset authors reported institutional review board approval and a waiver of informed consent for the original data collection (1). Reporting was organized with reference to TRIPOD+AI guidance for prediction models and the Checklist for Artificial Intelligence in Medical Imaging (CLAIM) (2, 3).',
  'chinese': '本研究对公开资料进行回顾性二次分析，评估定量磁共振成像（MRI）特征能否在临床变量基础上，增加对胶质瘤治疗后残余或复发肿瘤与治疗相关改变的鉴别信息。目标为每次MRI检查所对应的诊断状态。临床模型、MRI模型与融合模型在同一队列中接受交叉验证，训练折与验证折之间保持患者隔离。预测单位为单次检查，同一患者用于建模的全部检查均分配至同一折。分析使用公开、去标识的记录与影像，没有新增受试者招募、患者接触、影像检查或标本采集。数据集作者报告，原始资料收集已获机构伦理审查批准并豁免知情同意（1）。论文报告参照预测模型TRIPOD+AI指南及医学影像人工智能报告清单CLAIM组织（2，3）。'},
 {'id':'M2','citation_ids':['S01'],
  'english': 'The source collection was the University of California San Diego Post-Treatment Glioblastoma (UCSD-PTGBM) dataset, distributed through The Cancer Imaging Archive under a Creative Commons Attribution 4.0 license (1). We used the public version 3 clinical workbook and the accompanying list of examinations with treatment-related changes. Imaging inputs were obtained from the publicly accessible data-nih/tcia release mirror on GitHub and comprised derived fluid-attenuated inversion recovery (FLAIR) volumes, postcontrast T1-weighted volumes, and existing BraTS annotation masks in Neuroimaging Informatics Technology Initiative format. Source clinical identifiers, BraTS identifiers, mirror filenames, and the correspondence between them were retained in the linkage records. The mirror files were treated as a specific derived input release. Their observed intensity representation and header metadata determined the processing rules used in this analysis. Download addresses, file sizes, and SHA256 checksums were retained to identify the exact files used and support restoration of the analysis inputs.',
  'chinese': '资料来源为加州大学圣迭戈分校治疗后胶质母细胞瘤数据集（UCSD-PTGBM），该数据集通过癌症影像档案库发布，采用知识共享署名4.0许可（1）。本研究使用公开的第3版临床工作簿及其配套的治疗相关改变检查名单。影像输入来自GitHub上公开可访问的data-nih/tcia发布镜像，包括派生的液体衰减反转恢复（FLAIR）图像、增强T1加权图像及现成的BraTS标注掩膜，文件格式为NIfTI。连接记录保留临床源标识、BraTS标识、镜像文件名及其对应关系。镜像文件作为特定的派生输入版本处理，本分析的影像处理规则依据其实际观察到的强度表示和头信息确定。下载地址、文件大小及SHA256校验值均予保留，以准确识别所用文件并支持恢复分析输入。'},
 {'id':'M3','citation_ids':['S01'],
  'english': 'The published source collection contained 243 examinations from 178 patients (1). An identifier audit identified discordant relationships between the clinical and BraTS patient identifiers in five examinations from two source patients. These examinations were quarantined before clinical-model revalidation, leaving 238 examinations from 176 patients in the clinical cohort. Of these, 182 examinations from 135 patients had corresponding FLAIR images, postcontrast T1-weighted images, and annotation masks in the mirror and yielded the complete eight-feature MRI matrix. The remaining 56 clinical examinations were retained in the coverage records but did not enter the imaging-model comparison because the required mirror inputs were unavailable. The final imaging cohort comprised 141 examinations labeled as residual or recurrent tumor and 41 labeled as treatment-related change. Cohort assembly is summarized in Figure 1, and clinical characteristics by imaging coverage are provided in Supplementary Table S1. All four primary models were evaluated on this same imaging cohort.',
  'chinese': '已发表的源数据集包含178名患者的243次检查（1）。标识核查发现，2名源患者的5次检查存在临床与BraTS患者标识对应关系冲突。这些检查在临床模型重验证之前即被隔离，临床队列因此保留176名患者的238次检查。其中，135名患者的182次检查在镜像中具有对应的FLAIR图像、增强T1加权图像及标注掩膜，并形成完整的8项MRI特征矩阵。其余56次临床检查保留在覆盖记录中，但因缺少所需镜像输入，未进入影像模型比较。最终影像队列包括141次标记为残余或复发肿瘤的检查，以及41次标记为治疗相关改变的检查。队列构建流程见Figure 1，按影像覆盖情况汇总的临床特征见Supplementary Table S1。四个主要模型均在同一影像队列中接受评价。'},
]


def wc(s):return len(re.findall(r"[A-Za-z0-9]+(?:['/-][A-Za-z0-9]+)*",s))


def build():
    REPORT.mkdir(parents=True,exist_ok=True);OUT.mkdir(parents=True,exist_ok=True)
    refs=json.loads((REPORT/'reference_library.json').read_text())['references'];used=refs[:3]
    plan=json.loads((ROOT/'reports/ucsd_writing_requirements_v2/requirements_and_outline.json').read_text())
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Inches(.79);sec.bottom_margin=Inches(.79);sec.left_margin=Inches(.79);sec.right_margin=Inches(.79)
    for name in ['Normal','Title','Heading 1','Heading 2']:
        st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        rp=st.element.get_or_add_rPr();rp.rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC')
        for attribute in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
            rp.rFonts.attrib.pop(qn('w:'+attribute),None)
    st=doc.styles['Normal'];st.font.size=Pt(12);st.paragraph_format.line_spacing=1.;st.paragraph_format.space_after=Pt(9)
    doc.styles['Title'].font.size=Pt(16);doc.styles['Title'].paragraph_format.space_after=Pt(12)
    doc.styles['Heading 1'].font.size=Pt(13);doc.styles['Heading 1'].paragraph_format.space_before=Pt(8);doc.styles['Heading 1'].paragraph_format.space_after=Pt(9)
    doc.core_properties.title='Materials and methods 2.1 Study design and data sources';doc.core_properties.author='Tianyu Feng'
    doc.add_paragraph('Materials and methods',style='Title')
    doc.add_paragraph('2.1 Study design and data sources',style='Heading 1')
    for r in PARAGRAPHS:
        p=doc.add_paragraph(r['english']);p.paragraph_format.keep_together=True
        q=doc.add_paragraph(r['chinese']);q.paragraph_format.keep_together=True
        for run in q.runs:run.font.size=Pt(11)
        assert not re.search(r'[\u3400-\u9fff，。；：（）【】、]',r['english'])
    p=doc.add_paragraph('References',style='Heading 1');p.paragraph_format.page_break_before=True
    for i,r in enumerate(used,1):
        p=doc.add_paragraph(f"{i}. {r['citation']}")
        for run in p.runs:run.font.size=Pt(11)
    # Bilingual reading copy: an existing English figure is supplied at the end.
    p=doc.add_paragraph('Figure 1',style='Heading 1');p.paragraph_format.page_break_before=True
    fig=ROOT/'outputs/ucsd-paper-materials-20261029/figures/Figure1_Cohort_Flow.png'
    doc.add_picture(str(fig),width=Inches(6.7))
    p=doc.add_paragraph('Figure 1. Cohort assembly and analysis populations. Counts refer to examinations and source patients. After identifier quarantine, the imaging cohort included 182 examinations from 135 patients. The one-scan-per-patient sensitivity analysis selected the smallest available session number within this imaging cohort. It does not necessarily represent the first examination in the clinical course.')
    for run in p.runs:run.font.size=Pt(10)
    for e in [*doc.styles.element.xpath('.//w:pBdr'),*doc.element.xpath('.//w:pBdr')]:e.getparent().remove(e)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    path=OUT/'UCSD_Manuscript_Methods_2_1_EN_ZH.docx';doc.save(path)
    content={'stage':23,'subsection':'2.1 Study design and data sources','v2_paragraph_ids':['M1','M2','M3'],
      'paragraphs':PARAGRAPHS,'english_word_count':sum(wc(p['english']) for p in PARAGRAPHS),
      'paragraph_word_counts':{p['id']:wc(p['english']) for p in PARAGRAPHS},
      'citation_mapping':{str(i):r['id'] for i,r in enumerate(used,1)},
      'citation_numbering':'Current subsection numbering only; regenerate in first-occurrence order when Introduction is integrated.',
      'included_figure':str(fig.relative_to(ROOT)),'figure_sha256':hashlib.sha256(fig.read_bytes()).hexdigest(),
      'v2_outline_sha256':hashlib.sha256((ROOT/'reports/ucsd_writing_requirements_v2/requirements_and_outline.json').read_bytes()).hexdigest(),
      'docx':str(path.relative_to(ROOT)),'docx_sha256':hashlib.sha256(path.read_bytes()).hexdigest(),
      'models_refitted':0,'next_section':'2.2 Reference standard (M4)'}
    (REPORT/'methods_section21.json').write_text(json.dumps(content,ensure_ascii=False,indent=2)+'\n')
    (OUT/'UCSD_Manuscript_Methods_2_1_EN.txt').write_text('\n\n'.join(p['english'] for p in PARAGRAPHS)+'\n')
    print(json.dumps({k:content[k] for k in ['docx','english_word_count','paragraph_word_counts','citation_mapping','next_section']},indent=2))


if __name__=='__main__':build()
