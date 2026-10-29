"""Read-only reconciliation of the exported audit workbook and executed Notebook12."""
from pathlib import Path
import hashlib,json
import nbformat
import numpy as np
import pandas as pd
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
path=ROOT/'outputs/ucsd-feature-audit-20261029/UCSD_Clinical_MRI_Feasibility.xlsx'
w=load_workbook(path,read_only=True,data_only=True)
assert w.sheetnames==['Summary','Fields','Clinical_raw','Candidates','Issues','MRI']
sources={'Fields':'reports/ucsd_feature_audit/clinical_field_audit.csv',
         'Clinical_raw':'data/processed/ucsd_clinical_all_fields_raw.csv',
         'Candidates':'data/processed/ucsd_clinical_feature_candidates.csv',
         'Issues':'reports/ucsd_feature_audit/case_issues.csv'}
counts={}
for name,source in sources.items():
    d=pd.read_csv(ROOT/source)
    rows=list(w[name].iter_rows(max_row=len(d)+1,max_col=len(d.columns),values_only=True))
    x=pd.DataFrame(rows[1:],columns=rows[0]);assert x.columns.tolist()==d.columns.tolist()
    for c in d:
        if pd.api.types.is_numeric_dtype(d[c]):
            np.testing.assert_allclose(pd.to_numeric(x[c]),d[c],atol=1e-9,rtol=1e-9,equal_nan=True)
        else:assert x[c].fillna('').tolist()==d[c].fillna('').tolist(),(name,c)
    counts[name]=len(x)
s=w['Summary']
assert [s[p].value for p in ['C5','C6','C7','C8','C9','C10']]==[243,178,36,5,7,7]
assert [s[p].value for p in ['C15','C16','C17','C18','C19','C20']]==[149,168,112,114,197,117]
errors={'#REF!','#DIV/0!','#VALUE!','#NAME?','#NUM!','#NULL!','#SPILL!','#CALC!'}
for name,cols,rows in [('Summary',7,28),('Fields',19,37),('Clinical_raw',36,244),('Candidates',40,244),('Issues',7,34),('MRI',7,24)]:
    assert not any(isinstance(v,str) and v in errors for r in w[name].iter_rows(max_col=cols,max_row=rows,values_only=True) for v in r)
w.close()
w=load_workbook(path,read_only=True,data_only=False);assert w['Summary']['C19'].value.startswith('=COUNT(');w.close()
receipt=json.loads((ROOT/'reports/kernel_execution_12.json').read_text())
nb=nbformat.read(ROOT/'notebooks/12_ucsd_clinical_and_mri_feasibility.ipynb',as_version=4)
cells=[c for c in nb.cells if c.cell_type=='code']
assert len(cells)==5 and all(c.execution_count is not None for c in cells)
assert receipt['success'] and receipt['notebooks'][0]['errors']==0
assert receipt['notebooks'][0]['code_sha256']==hashlib.sha256('\n'.join(c.source for c in cells).encode()).hexdigest()
inv=json.loads((ROOT/'reports/ucsd_feature_audit/mri_inventory_checks.json').read_text())
assert inv['sample_case_file_entries']==inv['sample_entries_symbolic_links']==46
assert inv['sample_entries_known_byte_size']==0 and inv['source_NIfTI_bytes_downloaded']==0
report=dict(success=True,rows_verified=counts,excel_formula_errors=0,notebook_code_cells=5,notebook_errors=0,
            workbook_sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
            clinical_audit_code_sha256=hashlib.sha256((ROOT/'src/rn_public/ucsd_feature_audit.py').read_bytes()).hexdigest(),
            MRI_download_or_pixel_QC_claimed=False,
            note='Saved-file checks; formulas recalculated by artifact-tool, not native Excel')
(ROOT/'reports/ucsd_feature_audit/artifact_verification.json').write_text(json.dumps(report,indent=2));print(json.dumps(report,indent=2))
