"""Draft M9-M11 with native editable Word equations and frozen-feature checks."""
from pathlib import Path
from copy import deepcopy
import argparse,csv,hashlib,json,re,sys,zipfile
import numpy as np
from lxml import etree
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'outputs/ucsd-manuscript-step27-20261029'
REPORT=ROOT/'reports/ucsd_manuscript_step27'
PARAGRAPHS=[
 {'id':'M9','citation_ids':[],
  'english': 'Before feature extraction, intensity processing was performed separately for each FLAIR and postcontrast T1-weighted image. We calculated the first and 99th percentiles from finite voxel values strictly greater than zero, using linear percentile interpolation. This support region was determined from the image rather than from the annotation mask. Equation 1 defines the scaled intensity u, with x denoting a supplied voxel value and the percentile symbols denoting the corresponding image-specific reference values. The transformation was applied to all ROI voxels, including voxels whose original intensity was zero. Values below zero or above one after scaling were retained. The unclipped transformation was adopted after a technical pilot showed that percentile clipping merged some higher ROI intensities; it was an engineering amendment made after observing processing diagnostics, without using classification performance to choose the method. Image-specific parameters were calculated without pooling information across patients or examinations. Degenerate percentile ranges were treated as processing failures.',
  'chinese': '在特征提取之前，分别对每份FLAIR和增强T1加权图像进行强度处理。采用线性分位数插值，从有限且严格大于零的体素值计算第1和第99百分位数。该支持区域由图像本身确定，不由标注掩膜确定。Equation 1定义缩放后强度u，其中x表示所提供的体素值，分位数符号表示对应图像的参考值。变换应用于全部ROI体素，包括原始强度为零的体素。缩放后低于零或高于一的数值均保留。技术试验发现百分位裁切会合并部分较高ROI强度后，采用了不裁切变换；这是观察处理诊断后作出的工程修订，未依据分类性能选择方法。各图像独立计算参数，不跨患者或检查汇总信息。分位数范围退化时按处理失败记录。'},
 {'id':'M10','citation_ids':['S32'],
  'english': 'Four features were calculated from each sequence, yielding eight MRI inputs: mean, population standard deviation, Shannon entropy and three-dimensional gray-level co-occurrence matrix contrast. Mean and standard deviation used the scaled ROI intensities directly, with the voxel count N as the denominator in Equation 2. Entropy and texture calculations used fixed-bin-number discretization following the Image Biomarker Standardization Initiative (IBSI) framework (4). For each sequence and examination, a and b denoted the minimum and maximum scaled intensities within the full ROI. Equation 3 assigned these values to 32 bins; the minimum entered bin 1 and the maximum entered bin 32. In Equation 4, each bin probability is the fraction of ROI voxels assigned to that bin, and bins with zero probability do not contribute. Entropy was expressed in bits using base-2 logarithms. Constant ROI intensities were treated as processing failures. The feature set and bin count were fixed before batch extraction, without outcome-based feature selection.',
  'chinese': '每个序列计算四项特征，得到八个MRI输入：均值、总体标准差、Shannon熵和三维灰度共生矩阵对比度。均值和标准差直接由缩放后的ROI强度计算，Equation 2中的分母为体素数N。熵与纹理计算采用基于影像生物标志物标准化倡议（IBSI）框架的固定箱数离散化（4）。对于每个序列和每次检查，a和b分别表示完整ROI内缩放强度的最小值和最大值。Equation 3将这些数值分入32箱，最小值进入第1箱，最大值进入第32箱。Equation 4中的每个箱概率为分配至该箱的ROI体素比例，概率为零的箱不参与求和。熵使用以2为底的对数，以bit表示。ROI强度恒定时按处理失败记录。特征集合与箱数在批量提取前固定，未依据结局筛选特征。'},
 {'id':'M11','citation_ids':['S32'],
  'english': 'The gray-level co-occurrence matrix (GLCM) was constructed from neighboring voxel pairs contained entirely within the ROI. We used 13 undirected offsets representing the 26 immediate neighbors at a Chebyshev grid distance of one supplied voxel. This distance referred to the file grid and was not interpreted as one millimeter. For each offset in the set D, the directional count matrix C recorded pairs with bin indices i and j. Equation 5 defines the pooled symmetric counts S and the normalized matrix P. Contrast was calculated from this single matrix as shown in Equation 6 (4). We combined counts before calculating the feature, without averaging separate directional contrast values or applying distance weights. Each sequence produced one contrast value. A region with no valid neighboring pairs was recorded as a processing failure. We checked the implementation with small arrays having known results; this did not establish complete IBSI benchmark validation.',
  'chinese': '灰度共生矩阵（GLCM）由两端均位于ROI内的相邻体素对构建。采用13个无向偏移，代表所提供体素网格中Chebyshev距离为1的26个直接邻居。该距离指文件网格上的距离，不解释为1毫米。对于集合D中的每个偏移，方向计数矩阵C记录箱索引为i与j的体素对。Equation 5定义合并后的对称计数S及归一化矩阵P。对比度由这一矩阵按Equation 6计算（4）。我们先合并计数再计算特征，不对各方向分别计算的对比度求平均，也不使用距离权重。每个序列产生一个对比度值。没有有效相邻体素对的区域按处理失败记录。实现以具有已知结果的小数组进行核查；这不等于完成了全部IBSI基准验证。'}
]

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def wc(s):return len(re.findall(r"[A-Za-z0-9]+(?:['/-][A-Za-z0-9]+)*",s))
def rows(p):
    with p.open() as f:return list(csv.DictReader(f))

# Native OMML helpers. All fractions, subscripts, sums, delimiters and roots remain editable.
def el(tag,*children):
    n=OxmlElement('m:'+tag)
    for child in children:
        for c in child if isinstance(child,list) else [child]:n.append(deepcopy(c))
    return n
def mr(value,plain=False):
    n=el('r')
    if plain:
        pr=el('rPr');normal=el('nor');normal.set(qn('m:val'),'1');pr.append(normal);n.append(pr)
    pr=OxmlElement('w:rPr');font=OxmlElement('w:rFonts');font.set(qn('w:ascii'),'Cambria Math');font.set(qn('w:hAnsi'),'Cambria Math');pr.append(font)
    size=OxmlElement('w:sz');size.set(qn('w:val'),'24');pr.append(size);n.append(pr)
    t=el('t');t.set(qn('xml:space'),'preserve');t.text=value;n.append(t);return n
def sub(base,index):return el('sSub',el('e',base),el('sub',mr(index)))
def sup(base,index):return el('sSup',el('e',base),el('sup',mr(index)))
def frac(top,bottom):return el('f',el('num',top),el('den',bottom))
def bracket(content,left='(',right=')'):
    pr=el('dPr')
    for k,v in [('begChr',left),('endChr',right),('grow','1')]:n=el(k);n.set(qn('m:val'),v);pr.append(n)
    return el('d',pr,el('e',content))
def summation(low,high,body):
    pr=el('naryPr')
    for k,v in [('chr','∑'),('limLoc','undOvr'),('grow','1')]+([] if high else [('supHide','1')]):n=el(k);n.set(qn('m:val'),v);pr.append(n)
    return el('nary',pr,el('sub',mr(low)),el('sup',mr(high)),el('e',body))
def root(content):
    pr=el('radPr');h=el('degHide');h.set(qn('m:val'),'1');pr.append(h)
    return el('rad',pr,el('deg'),el('e',content))
def function(name,argument):return el('func',el('fName',name),el('e',argument))
def equations():
    uv=lambda:sub(mr('u'),'v')
    sij=lambda indices:sub(mr('S'),indices)
    cij=lambda indices:[sub(mr('C'),indices),bracket(mr('d'))]
    return {
      1:[mr('u = '),frac([mr('x − '),sub(mr('p'),'1')],[sub(mr('p'),'99'),mr(' − '),sub(mr('p'),'1')])],
      2:[mr('μ = '),frac(mr('1'),mr('N')),summation('v=1','N',uv()),mr(' ,     σ = '),root([frac(mr('1'),mr('N')),summation('v=1','N',sup(bracket([uv(),mr(' − μ')]),'2'))])],
      3:[mr('q'),bracket(mr('u')),mr(' = '),function(mr('min',True),bracket([mr('32, 1 + '),bracket([mr('32'),frac(mr('u − a'),mr('b − a'))],'⌊','⌋')]))],
      4:[mr('H = −'),summation('k=1','32',[sub(mr('p'),'k'),mr(' '),function(sub(mr('log',True),'2'),bracket(sub(mr('p'),'k')))])],
      5:[sij('ij'),mr(' = '),summation('d∈D','',bracket([*cij('ij'),mr(' + '),*cij('ji')],'[',']')),mr(' ,    '),sub(mr('P'),'ij'),mr(' = '),frac(sij('ij'),summation('r=1','32',summation('s=1','32',sij('rs'))))],
      6:[mr('Contrast',True),mr(' = '),summation('i=1','32',summation('j=1','32',[sup(bracket(mr('i − j')),'2'),sub(mr('P'),'ij')]))]
    }

def audit():
    old=json.loads((ROOT/'reports/ucsd_manuscript_step26/methods_section24.json').read_text());frozen=old['source_audit']['frozen_file_sha256'].copy()
    frozen.update(json.loads((ROOT/'reports/ucsd_manuscript_step26/artifact_verification.json').read_text())['sha256'])
    for p,h in frozen.items():assert sha(ROOT/p)==h,p
    choice=json.loads((ROOT/'reports/ucsd_roi_pilot/normalization_choice.json').read_text())
    assert choice['added_after_observing_clipping'] and not choice['model_performance_used'] and not choice['uses_clinical_class']
    assert choice['original_clipping_outputs_preserved'] and not choice['independent_optimality_validated']
    assert sha(ROOT/'reports/ucsd_roi_pilot_intensity_amendment.md')==choice['amendment_sha256']
    assert sha(ROOT/'src/rn_public/ucsd_intensity_followup.py')==choice['code_sha256']
    sys.path.insert(0,str(ROOT/'src'))
    from rn_public.ucsd_batch_features import OFFSETS,FEATURES,sequence_features
    from rn_public.ucsd_intensity_followup import affine_values
    from rn_public.ucsd_roi_pilot import fixed_bins
    assert len(OFFSETS)==13 and all(max(abs(k) for k in d)==1 for d in OFFSETS)
    assert len(set(OFFSETS)|{tuple(-k for k in d) for d in OFFSETS})==26
    a=np.asarray([0,20,80,100]);v=affine_values(a,{'p01':20,'p99':80});np.testing.assert_allclose(v,[-1/3,0,1,4/3])
    bins,lo,hi=fixed_bins(np.asarray([-.5,0,.5,1,1.5]));assert bins.tolist()==[1,9,17,25,32]
    formula_bins=np.minimum(32,1+np.floor(32*(np.asarray([-.5,0,.5,1,1.5])-lo)/(hi-lo))).astype(int)
    np.testing.assert_array_equal(bins,formula_bins)
    checker=(np.indices((2,2,2)).sum(axis=0)%2+1).astype(np.uint8)
    values,_,diag=sequence_features(checker,np.ones_like(checker,dtype=bool),np.zeros_like(checker,dtype=bool))
    expected={'mean':.5,'std':.5,'entropy':1.,'glcm_contrast':961*16/28}
    for k,val in expected.items():assert np.isclose(values[k],val,rtol=1e-12,atol=1e-12)
    assert diag['neighbour_pairs_symmetric']==56 and diag['directions_with_pairs']==13
    batch=json.loads((ROOT/'reports/ucsd_batch_features/batch_receipt.json').read_text())
    cases={r['ID']:r for r in batch['cases']};data=rows(ROOT/'data/processed/ucsd_clinical_image_features.csv')
    for r in data:
        for feature in FEATURES:assert np.isclose(float(r[feature]),cases[r['ID']]['primary'][feature],rtol=1e-12,atol=1e-12)
    diagnostics=rows(ROOT/'reports/ucsd_batch_features/normalization_and_texture_diagnostics.csv');assert len(diagnostics)==364
    for r in diagnostics:
        assert float(r['p99'])>float(r['p01']) and int(r['support_voxels'])>0
        assert 2<=int(r['occupied_bins'])<=32 and int(r['directions_with_pairs'])==13
        assert np.isclose(float(r['bin_minimum']),float(r['normalized_minimum'])) and np.isclose(float(r['bin_maximum']),float(r['normalized_maximum']))
    files=['reports/ucsd_roi_pilot/normalization_choice.json','reports/ucsd_roi_pilot_intensity_amendment.md','reports/ucsd_batch_features/feature_dictionary.csv','reports/ucsd_batch_features/normalization_and_texture_diagnostics.csv','src/rn_public/ucsd_intensity_followup.py']
    frozen.update({p:sha(ROOT/p) for p in files})
    return {'frozen_file_sha256':frozen,'saved_feature_values_checked':len(data)*len(FEATURES),'sequence_diagnostic_records_checked':364,
      'MRI_feature_names':list(FEATURES),'post_observation_engineering_amendment_disclosed':True,'offsets':[list(d) for d in OFFSETS],
      'hand_array_verification':{'scaled_tail_values':v.tolist(),'bin_indices':bins.tolist(),'checkerboard_features':values,'independent_expected':expected,'symmetric_pair_count':56},
      'full_IBSI_benchmark_validation_claimed':False,'study_pixels_redecoded':0,'models_refitted':0}

def reference():
    r=deepcopy(next(r for r in json.loads((ROOT/'reports/ucsd_manuscript_step23/reference_library.json').read_text())['references'] if r['id']=='S32'))
    r['pages_or_article']='328-338';r['citation']=r['citation'].replace(':328.',':328-338.')
    r['current_review_level']='Official IBSI processing and feature specification sections reviewed; primary publication metadata rechecked'
    r['metadata_correction']='Complete pagination 328-338 checked against PubMed; prior candidate pool retained unchanged'
    r['specification_sources']=['https://ibsi.readthedocs.io/en/latest/02_Image_processing.html#fixed-bin-number','https://ibsi.readthedocs.io/en/latest/03_Image_features.html#grey-level-co-occurrence-matrix']
    return r

def build():
    OUT.mkdir(parents=True,exist_ok=True);REPORT.mkdir(parents=True,exist_ok=True);a=audit();ref=reference()
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(.79)
    for name in ['Normal','Title','Heading 1']:
        st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        rf=st.element.get_or_add_rPr().rFonts;rf.set(qn('w:eastAsia'),'Noto Serif CJK SC')
        for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:rf.attrib.pop(qn('w:'+key),None)
    normal=doc.styles['Normal'];normal.font.size=Pt(12);normal.paragraph_format.line_spacing=1.;normal.paragraph_format.space_after=Pt(9)
    doc.styles['Title'].font.size=Pt(16);doc.styles['Title'].paragraph_format.space_after=Pt(12)
    doc.styles['Heading 1'].font.size=Pt(13);doc.styles['Heading 1'].paragraph_format.space_before=Pt(8);doc.styles['Heading 1'].paragraph_format.space_after=Pt(9)
    doc.core_properties.title='Materials and methods 2.5 Intensity normalization and feature definitions';doc.core_properties.author='Tianyu Feng'
    doc.add_paragraph('Materials and methods',style='Title');doc.add_paragraph('2.5 Intensity normalization and feature definitions',style='Heading 1')
    forms=equations();groups={'M9':[1],'M10':[2,3,4],'M11':[5,6]}
    for item in PARAGRAPHS:
        en=doc.add_paragraph(item['english']);en.paragraph_format.page_break_before=item['id']!='M9'
        en.paragraph_format.keep_together=True;en.paragraph_format.keep_with_next=True
        zh=doc.add_paragraph(item['chinese']);zh.paragraph_format.keep_together=True
        for r in zh.runs:r.font.size=Pt(11)
        for index in groups[item['id']]:
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_together=True
            p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(7)
            p._p.append(el('oMath',forms[index]));p.add_run(f'    ({index})')
    doc.add_paragraph('References',style='Heading 1');p=doc.add_paragraph('4. '+ref['citation']);p.paragraph_format.keep_together=True
    for r in p.runs:r.font.size=Pt(11)
    for e in [*doc.styles.element.xpath('.//w:pBdr'),*doc.element.xpath('.//w:pBdr')]:e.getparent().remove(e)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    path=OUT/'UCSD_Manuscript_Methods_2_5_EN_ZH.docx';doc.save(path)
    text='\n\n'.join(p['english'] for p in PARAGRAPHS)+'\n'
    (OUT/'UCSD_Manuscript_Methods_2_5_EN.txt').write_text(text)
    d={'stage':27,'date':'2026-10-29','subsection':'2.5 Intensity normalization and feature definitions','v2_paragraph_ids':['M9','M10','M11'],'paragraphs':PARAGRAPHS,
      'english_word_count':sum(wc(p['english']) for p in PARAGRAPHS),'paragraph_word_counts':{p['id']:wc(p['english']) for p in PARAGRAPHS},
      'word_count_scope':'English narrative including in-text citations; excludes equations, headings, reference entry and Chinese translation',
      'v2_approximate_target_words':450,'equations':{'count':6,'format':'native editable OMML','numbers':[1,2,3,4,5,6],'meaning':{'1':'Unclipped per-image scaling','2':'Mean and population SD','3':'32-bin discretization within ROI range','4':'Base-2 histogram entropy','5':'Symmetric direction-count pooling and normalization','6':'GLCM contrast'}},
      'citation_mapping':{'4':'S32'},'cumulative_provisional_citation_mapping':{'1':'S01','2':'S02','3':'S03','4':'S32'},'reference_used':ref,
      'citation_numbering':'New source 4 follows existing 1-3; regenerate globally when Introduction is assembled',
      'source_audit':a,'models_refitted':0,'new_hypothesis_tests':0,'next_section':'Fixed logistic models and patient-grouped training (M12-M13)',
      'docx':str(path.relative_to(ROOT)),'docx_sha256':sha(path)}
    (REPORT/'methods_section25.json').write_text(json.dumps(d,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({k:d[k] for k in ['docx','english_word_count','paragraph_word_counts']},indent=2))

def verify():
    from pypdf import PdfReader
    d=json.loads((REPORT/'methods_section25.json').read_text())
    for p,h in d['source_audit']['frozen_file_sha256'].items():assert sha(ROOT/p)==h,p
    path=ROOT/d['docx'];assert sha(path)==d['docx_sha256']
    doc=Document(path);texts=[p.text for p in doc.paragraphs]
    for p in PARAGRAPHS:
        assert p['english'] in texts and p['chinese'] in texts and texts.index(p['chinese'])==texts.index(p['english'])+1
        assert not re.search(r'[\u3400-\u9fff，。；：（）【】、]',p['english'])
    assert sum(wc(p['english']) for p in PARAGRAPHS)==d['english_word_count']
    ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
    with zipfile.ZipFile(path) as z:
        tree=etree.fromstring(z.read('word/document.xml'));assert len(tree.xpath('//m:oMath',namespaces=ns))==6
        assert len(tree.xpath('//m:f',namespaces=ns))==5 and len(tree.xpath('//m:rad',namespaces=ns))==1
        assert len(tree.xpath('//m:nary',namespaces=ns))==8
        assert not any(n.startswith('word/media/') for n in z.namelist())
    assert all(f'    ({n})' in texts for n in range(1,7))
    assert any(s.startswith('4. Zwanenburg A') and ':328-338.' in s for s in texts)
    pages=len(PdfReader(OUT/'docx_render/UCSD_Manuscript_Methods_2_5_EN_ZH.pdf').pages)
    v={'success':True,'stage':27,'date':'2026-10-29','subsection':d['subsection'],'paragraph_ids':['M9','M10','M11'],
      'english_word_count':d['english_word_count'],'paragraph_word_counts':d['paragraph_word_counts'],'matched_chinese_paragraphs':3,
      'word_pages':pages,'visual_review_pages':list(range(1,pages+1)),'native_editable_equations':6,'equation_images':0,'equation_structure_verified':True,
      'prior_files_unchanged':len(d['source_audit']['frozen_file_sha256']),'saved_feature_values_checked':1456,'sequence_diagnostic_records_checked':364,
      'formula_hand_array_checks_passed':True,'models_refitted':0,'study_pixels_redecoded':0,'new_hypothesis_tests':0,
      'current_section_citations':1,'cumulative_unique_cited_sources':4,'citation_mapping':{'4':'S32'},'next_section':d['next_section'],
      'sha256':{str(p.relative_to(ROOT)):sha(p) for p in OUT.iterdir() if p.is_file() and p.suffix in {'.docx','.txt'}}}
    (REPORT/'artifact_verification.json').write_text(json.dumps(v,ensure_ascii=False,indent=2)+'\n');print(json.dumps(v,indent=2))

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--verify',action='store_true');args=parser.parse_args()
    verify() if args.verify else build()
