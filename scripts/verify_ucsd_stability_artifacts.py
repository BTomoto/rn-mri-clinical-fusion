"""Independent numerical and artifact audit of all fixed Stage19 repetitions."""
from pathlib import Path
import base64,hashlib,json,math,subprocess,sys
import nbformat,numpy as np,pandas as pd
from scipy.special import expit
from scipy.stats import rankdata
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'src'))
from rn_public.ucsd_stability import freeze_plan,MODELS,KEYS


def verify_prediction_parameters(d,pred,records):
 count=0
 for f in records:
  train=d[d.ID.isin(f['train_IDs'])];test=d[d.ID.isin(f['test_IDs'])]
  for key in ['patient_id','brats_patient_id','validation_group']:assert set(train[key]).isdisjoint(test[key])
  assert len(train)+len(test)==len(d)
  for name,r in f['models'].items():
   assert r['success'];p=r['parameters'];columns={}
   for block,part in p['continuous'].items():
    cols=part['columns'];a=train[cols].to_numpy().copy();b=test[cols].to_numpy().copy()
    if block=='interval':
     median=float(np.nanmedian(a));assert math.isclose(median,part['impute_median'],abs_tol=1e-12)
     a=np.where(np.isnan(a),median,a);b=np.where(np.isnan(b),median,b)
    mean=a.mean(axis=0);scale=a.std(axis=0,ddof=0);scale=np.where(scale==0,1,scale)
    np.testing.assert_allclose(mean,part['mean'],atol=1e-12,rtol=1e-12);np.testing.assert_allclose(scale,part['scale'],atol=1e-12,rtol=1e-12)
    b=(b-mean)/scale;columns.update({block+'__'+key:b[:,j] for j,key in enumerate(cols)})
   for key in p['binary']:columns['binary__'+key]=test[key].to_numpy()
   x=np.column_stack([columns[key] for key in p['transformed_names']]);prob=expit(x@np.asarray(p['coefficients'])+p['intercept'])
   actual=pred.set_index('ID').loc[test.ID,name+'_oof_probability'].to_numpy()
   np.testing.assert_allclose(actual,prob,atol=1e-12,rtol=1e-12);count+=len(test)
 return count


def verify_metrics(pred,metrics):
 y=pred.tumor_label.to_numpy();n1=int(y.sum());n0=int((1-y).sum())
 for name in MODELS:
  p=pred[name+'_oof_probability'].to_numpy();r=metrics[metrics.model.eq(name)].iloc[0]
  assert math.isclose((rankdata(p)[y==1].sum()-n1*(n1+1)/2)/(n1*n0),r.roc_auc,abs_tol=1e-12)
  assert math.isclose(np.mean((p-y)**2),r.brier,abs_tol=1e-12)
  assert int(((y==0)&(p<.5)).sum())==r.tn and int(((y==1)&(p>=.5)).sum())==r.tp


def main():
 out=ROOT/'reports/ucsd_stability';s=json.loads((out/'metrics.json').read_text());d,single,m,plans,sm,sp,lock=freeze_plan(ROOT)
 pred=pd.read_csv(ROOT/'data/processed/ucsd_stability_repeated_oof.csv');metrics=pd.read_csv(out/'repeated_metrics.csv')
 assert len(pred)==3640 and not pred.duplicated(['repeat','ID']).any() and len(metrics)==80
 reconstructed=0
 for p in plans:
  k=p['repeat'];pr=pred[pred['repeat'].eq(k)];mr=metrics[metrics['repeat'].eq(k)]
  assert len(pr)==182 and pr.ID.tolist()==d.ID.tolist()
  expected=m[m['repeat'].eq(k)];assert pr.validation_fold.tolist()==expected.validation_fold.tolist()
  for key in ['patient_id','brats_patient_id','validation_group']:assert pr.groupby(key).validation_fold.nunique().eq(1).all()
  audit=json.loads((out/f'repeated_fold_audits/repeat_{k:02d}.json').read_text());assert audit['complete'] and not audit['errors']
  reconstructed+=verify_prediction_parameters(d,pr,audit['folds']);verify_metrics(pr,mr)
 delta=pd.read_csv(out/'repeated_paired_differences.csv')
 for row in delta.itertuples():
  mm=metrics[metrics['repeat'].eq(row.repeat)].set_index('model')
  assert math.isclose(row.difference,mm.loc['Fusion',row.metric]-mm.loc['M1',row.metric],abs_tol=1e-12)
 distribution=pd.read_csv(out/'repeated_metric_distribution.csv')
 for row in distribution.itertuples():
  values=metrics[metrics.model.eq(row.model)][row.metric]
  np.testing.assert_allclose([row.median,row.q25,row.q75,row.minimum,row.maximum],[values.median(),values.quantile(.25),values.quantile(.75),values.min(),values.max()],atol=1e-12)
 assert distribution.interpretation.eq('descriptive_split_distribution_not_confidence_interval').all()
 single_pred=pd.read_csv(ROOT/'data/processed/ucsd_stability_single_scan_oof.csv');single_metrics=pd.read_csv(out/'single_scan_metrics.csv')
 assert len(single_pred)==135 and single_pred.patient_id.is_unique and single_pred.brats_patient_id.is_unique
 assert set(single_pred.ID)==set(single.ID) and int(single_pred.tumor_label.sum())==113
 audit=json.loads((out/'single_scan_fold_audit.json').read_text());assert audit['complete']
 reconstructed+=verify_prediction_parameters(single,single_pred,audit['folds']);verify_metrics(single_pred,single_metrics)
 draws=pd.read_csv(out/'single_scan_patient_bootstrap.csv');assert len(draws)==2000 and draws.draw.is_unique
 for key,ci in s['single_scan']['bootstrap_95ci'].items():np.testing.assert_allclose(draws[key].quantile([.025,.975]),[ci['lower'],ci['upper']],atol=1e-12)
 for metric in KEYS:np.testing.assert_allclose(draws['Fusion_minus_M1__'+metric],draws['Fusion__'+metric]-draws['M1__'+metric],atol=1e-12)
 assert json.loads((out/'failures.json').read_text())['records']==[]
 nb=nbformat.read(ROOT/'notebooks/19_ucsd_model_stability.ipynb',as_version=4);code=[c for c in nb.cells if c.cell_type=='code'];kr=json.loads((ROOT/'reports/kernel_execution_19.json').read_text())
 assert kr['success'] and len(code)==5 and [c.execution_count for c in code]==[1,2,3,4,5]
 assert hashlib.sha256('\n'.join(c.source for c in code).encode()).hexdigest()==kr['notebooks'][0]['code_sha256']
 assert not [o for c in code for o in c.outputs if o.output_type=='error']
 images=[base64.b64decode(o.data['image/png']) for c in code for o in c.outputs if o.output_type in ['display_data','execute_result'] and 'image/png' in o.data];assert len(images)==2
 for data,name in zip(images,['repeated_validation.png','single_scan_sensitivity.png']):assert data==(ROOT/'figures/ucsd_stability'/name).read_bytes()
 payload=json.loads((ROOT/'reports/ucsd_stability_workbook_input.json').read_text());wb=load_workbook(ROOT/'outputs/ucsd-stability-20261029/UCSD_Model_Stability.xlsx',data_only=True,read_only=True);rows_checked=0
 for name,data in payload['datasets'].items():
  rows=list(wb[name].iter_rows(min_row=1,max_row=len(data['rows'])+1,max_col=len(data['headers']),values_only=True));assert list(rows[0])==data['headers']
  for actual,want in zip(rows[1:],data['rows']):
   for a,b in zip(actual,want):
    if isinstance(b,bool):b=int(b)
    if b is None or b=='':assert a is None or a==''
    elif isinstance(b,(int,float)):assert isinstance(a,(int,float)) and math.isclose(a,b,abs_tol=1e-12,rel_tol=1e-12),(name,a,b)
    else:assert a==b,(name,a,b)
   rows_checked+=1
 for i,name in enumerate(MODELS,5):assert math.isclose(wb['Summary'][f'C{i}'].value,metrics.loc[metrics.model.eq(name),'roc_auc'].median(),abs_tol=1e-12)
 for i,name in enumerate(MODELS,13):assert math.isclose(wb['Summary'][f'C{i}'].value,s['single_scan']['metrics'][name]['roc_auc'],abs_tol=1e-12)
 for cell,value in {'C30':182,'C31':135,'C32':113,'C33':22,'C34':20,'F20':20,'F21':20}.items():assert wb['Summary'][cell].value==value
 assert not [(sh.title,cell.coordinate) for sh in wb for row in sh for cell in row if cell.data_type=='e']
 changed=subprocess.check_output(['git','diff','950460995d71a271b234186c1b12a0c691c46dc0','--name-only'],cwd=ROOT,text=True).splitlines()
 for path in changed:
  existed=subprocess.run(['git','cat-file','-e','950460995d71a271b234186c1b12a0c691c46dc0:'+path],cwd=ROOT,capture_output=True).returncode==0
  assert not existed or path in ['README.md','.gitignore','reports/project_status.json'],path
 result=dict(success=True,repeated_OOF_rows=3640,single_scan_OOF_rows=135,repeats_checked=20,
  model_probabilities_independently_reconstructed=reconstructed,patient_overlap=0,BraTS_patient_overlap=0,
  bootstrap_draws_checked=2000,workbook_data_rows_checked=rows_checked,workbook_error_cells=0,
  notebook_code_cells=5,notebook_errors=0,notebook_embedded_figures_byte_identical=2,
  prior_stage_artifacts_preserved=True,external_Excel_application_tested=False)
 (out/'artifact_verification.json').write_text(json.dumps(result,indent=2)+'\n');print(json.dumps(result,indent=2))


if __name__=='__main__':main()
