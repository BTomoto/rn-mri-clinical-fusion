"""Author V2 M5-M6 after checking saved clinical inputs and fold preprocessing."""
from pathlib import Path
import argparse,csv,hashlib,json,re
from collections import Counter
import numpy as np
from openpyxl import load_workbook
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'outputs/ucsd-manuscript-step25-20261029'
REPORT=ROOT/'reports/ucsd_manuscript_step25'
INITIAL='Days from Acquisition to Date of initial surgery, treatment or diagnosis'
PARAGRAPHS=[
 {'id':'M5','citation_ids':['S01'],
  'english': 'We used four clinical variables from the public clinical workbook: age, sex at birth, the interval from the initial clinical event to the examination, and other treatment before the scan (1). Age was entered as a continuous variable in years at the time of scanning. Sex was encoded as 1 for male and 0 for female. The other-treatment field was encoded as 1 for Yes and 0 for No, retaining its broad source definition without assigning a specific treatment modality or drug. An additional binary indicator identified a missing initial-event interval, with 1 denoting missing and 0 denoting observed. These encodings provided two inputs for the age-and-sex model and five clinical inputs for both the expanded clinical model and the fusion model. The same variable definitions were used throughout the comparisons. Follow-up, survival outcomes and diagnostic category fields were excluded from the clinical inputs. Clinical input definitions are summarized in Supplementary Table S2.',
  'chinese': '我们使用公开临床工作簿中的四项临床变量：年龄、出生性别、初始临床事件至检查的间隔，以及扫描前其他治疗（1）。年龄采用扫描时的岁数，作为连续变量输入。性别编码为男性1、女性0。其他治疗字段编码为Yes对应1、No对应0，保留其宽泛的源定义，不将其指定为某种治疗方式或药物。另设一个二分类指示变量识别初始事件间隔缺失，缺失为1、已观测为0。这些编码为年龄与性别模型提供两个输入，并为扩展临床模型和融合模型各提供五个临床输入。各项比较始终使用相同的变量定义。随访、生存结局及诊断类别字段不纳入临床输入。临床输入定义汇总于Supplementary Table S2。'},
 {'id':'M6','citation_ids':[],
  'english': 'The source time field expressed the initial-event date as a signed day offset relative to MRI acquisition. We accepted only strictly negative numeric offsets, reversed their signs to obtain positive elapsed days, and applied the natural logarithm of one plus the interval. The reference event could be initial surgery, another initial treatment, or diagnosis in an untreated patient, so the resulting variable did not represent a uniform interval since radiotherapy or surgery. Nonnumeric entries, blank values and nonnegative offsets were treated as missing; a zero offset did not establish that the event preceded the examination. In the 182-examination imaging cohort, 36 intervals were missing under this rule. Within each training fold, missing transformed intervals were replaced with the median observed transformed interval, followed by standardization. Age was also standardized within the training fold. Learned medians and scaling parameters were applied unchanged to the corresponding validation fold. Binary inputs were passed through without scaling. Age, sex and the other-treatment indicator were complete in the imaging cohort.',
  'chinese': '源时间字段以相对于MRI采集日期的有符号天数偏移表示初始事件日期。我们仅接受严格为负的数值型偏移，将其取负号得到正的经过天数，再对间隔加1后取自然对数。参考事件可能是首次手术、其他初始治疗，或未接受治疗患者的诊断，因此所得变量不表示统一的放疗后或手术后间隔。非数值条目、空白及非负偏移均按缺失处理；零偏移不能确认事件发生在检查之前。按照这一规则，182次检查的影像队列中有36条间隔缺失。在每个训练折内，使用该训练折已观测到的变换后间隔中位数填补缺失，随后进行标准化。年龄也在训练折内标准化。所得中位数和缩放参数不作变更，直接用于对应验证折。二分类输入保留原编码，不进行缩放。影像队列的年龄、性别及其他治疗指示变量均无缺失。'}
]

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def wc(s):return len(re.findall(r"[A-Za-z0-9]+(?:['/-][A-Za-z0-9]+)*",s))
def rows(p):
    with p.open() as f:return list(csv.DictReader(f))
def numeric(v):
    try:return float(v)
    except (ValueError,TypeError):return np.nan
def same(a,b):assert np.isclose(a,b,rtol=1e-12,atol=1e-12),(a,b)

def audit():
    old=json.loads((ROOT/'reports/ucsd_manuscript_step24/methods_section22.json').read_text())
    frozen=old['source_audit']['frozen_file_sha256'].copy()
    frozen.update(json.loads((ROOT/'reports/ucsd_manuscript_step24/artifact_verification.json').read_text())['sha256'])
    for p,h in frozen.items():assert sha(ROOT/p)==h,p
    sources=json.loads((ROOT/'data/metadata/public_label_audit_20261029/downloaded_sources.json').read_text())
    for s in sources:assert sha(ROOT/s['local_path'])==s['sha256']
    inputs=['src/rn_public/ucsd_clinical_revalidation.py','src/rn_public/ucsd_fusion.py','reports/ucsd_fusion/fold_audit.json','reports/ucsd_fusion/cohort_and_fold_manifest.csv','reports/ucsd_feature_audit/clinical_field_audit.csv','reports/ucsd_clinical_revalidation_protocol.md']
    frozen.update({p:sha(ROOT/p) for p in inputs})
    wb=load_workbook(ROOT/'data/raw/reference/ucsd_ptgbm_clinical_v3.xlsx',read_only=True,data_only=True)
    raw=list(wb['Cinical info'].iter_rows(max_col=36,values_only=True));header=[str(v).strip() for v in raw[0]]
    lookup={r['ID']:r for r in [dict(zip(header,a)) for a in raw[1:]]}
    dictionary={str(r[0]).strip():r[1] for r in wb['Data Dictionary'].iter_rows(min_row=1,max_row=80,max_col=2,values_only=True) if r[0] is not None};wb.close()
    assert dictionary["Patient's Age"]=='At time of scan (in years)'
    assert INITIAL in dictionary and 'diagnosis' in dictionary[INITIAL]
    d=rows(ROOT/'data/processed/ucsd_clinical_image_features.csv');assert len(d)==182
    missing=[];origin=Counter()
    for r in d:
        src=lookup[r['ID']];same(float(r['age_years']),float(src["Patient's Age"]))
        assert int(r['sex_male'])=={'Male':1,'Female':0}[src['Sex at birth']]
        assert int(r['other_treatment_prior'])=={'Yes':1,'No':0}[src['Other treatment prior to scan']]
        offset=numeric(src[INITIAL]);eligible=np.isfinite(offset) and offset<0
        assert int(r['initial_event_missing'])==int(not eligible)
        if eligible:
            same(float(r['days_since_initial_event']),-offset);same(float(r['log_initial_event_days']),np.log1p(-offset))
        else:
            assert r['days_since_initial_event']==r['log_initial_event_days']==''
            missing.append(r['ID']);origin['nonnumeric_or_blank' if not np.isfinite(offset) else 'zero' if offset==0 else 'positive']+=1
    assert len(missing)==36 and dict(origin)=={'nonnumeric_or_blank':35,'zero':1}
    indexed={r['ID']:r for r in d};folds=json.loads((ROOT/'reports/ucsd_fusion/fold_audit.json').read_text())
    manifest=rows(ROOT/'reports/ucsd_fusion/cohort_and_fold_manifest.csv');foldmap={r['ID']:int(r['validation_fold']) for r in manifest}
    checks=[]
    for f in folds:
        assert set(f['train_IDs']).isdisjoint(f['test_IDs'])
        assert set(f['train_IDs'])|set(f['test_IDs'])==set(indexed)
        assert all(foldmap[i]!=f['fold'] for i in f['train_IDs']) and all(foldmap[i]==f['fold'] for i in f['test_IDs'])
        train=[indexed[i] for i in f['train_IDs']]
        age=np.asarray([float(r['age_years']) for r in train]);log=np.asarray([numeric(r['log_initial_event_days']) for r in train]);median=float(np.nanmedian(log));filled=np.where(np.isfinite(log),log,median)
        for model in ['M0','M1','Fusion']:
            p=f['models'][model]['parameters'];pa=p['continuous']['age'];same(pa['mean'][0],age.mean());same(pa['scale'][0],age.std(ddof=0))
            assert p['binary']==(['sex_male'] if model=='M0' else ['sex_male','other_treatment_prior','initial_event_missing'])
            if model!='M0':
                pi=p['continuous']['interval'];same(pi['impute_median'],median);same(pi['mean'][0],filled.mean());same(pi['scale'][0],filled.std(ddof=0))
                checks.append({'fold':f['fold'],'model':model,'train_examinations':len(train),'observed_intervals':int(np.isfinite(log).sum()),'saved_median':pi['impute_median'],'reconstructed_median':median,'saved_mean':pi['mean'][0],'saved_sd':pi['scale'][0]})
        assert set(indexed[i]['patient_id'] for i in f['train_IDs']).isdisjoint(indexed[i]['patient_id'] for i in f['test_IDs'])
    assert len(folds)==5 and len(checks)==10
    return {'frozen_file_sha256':frozen,'source_dictionary':{k:dictionary[k] for k in ["Patient's Age",'Sex at birth',INITIAL,'Other treatment prior to scan']},
      'clinical_records_matched_to_raw_workbook':182,'missing_intervals':36,'missing_origin':dict(origin),'missing_interval_IDs':missing,
      'core_clinical_variables_missing':0,'training_age_scalers_checked':15,'training_interval_pipelines_checked':10,'fold_preprocessing_checks':checks,
      'no_model_fitting_performed':True}

def build():
    OUT.mkdir(parents=True,exist_ok=True);REPORT.mkdir(parents=True,exist_ok=True);a=audit()
    source=json.loads((ROOT/'reports/ucsd_manuscript_step23/reference_library.json').read_text())['references'][0]
    doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(.79)
    for name in ['Normal','Title','Heading 1']:
        st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        rf=st.element.get_or_add_rPr().rFonts;rf.set(qn('w:eastAsia'),'Noto Serif CJK SC')
        for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:rf.attrib.pop(qn('w:'+key),None)
    normal=doc.styles['Normal'];normal.font.size=Pt(12);normal.paragraph_format.line_spacing=1.;normal.paragraph_format.space_after=Pt(10)
    doc.styles['Title'].font.size=Pt(16);doc.styles['Title'].paragraph_format.space_after=Pt(12)
    doc.styles['Heading 1'].font.size=Pt(13);doc.styles['Heading 1'].paragraph_format.space_before=Pt(8);doc.styles['Heading 1'].paragraph_format.space_after=Pt(9)
    doc.core_properties.title='Materials and methods 2.3 Clinical predictors and missing data';doc.core_properties.author='Tianyu Feng'
    doc.add_paragraph('Materials and methods',style='Title');doc.add_paragraph('2.3 Clinical predictors and missing data',style='Heading 1')
    for p in PARAGRAPHS:
        en=doc.add_paragraph(p['english']);en.paragraph_format.keep_together=True;en.paragraph_format.keep_with_next=True
        en.paragraph_format.page_break_before=p['id']=='M6'
        z=doc.add_paragraph(p['chinese'])
        z.paragraph_format.keep_together=True
        for r in z.runs:r.font.size=Pt(11)
    doc.add_paragraph('References',style='Heading 1');z=doc.add_paragraph('1. '+source['citation'])
    z.paragraph_format.keep_together=True
    for r in z.runs:r.font.size=Pt(11)
    for e in [*doc.styles.element.xpath('.//w:pBdr'),*doc.element.xpath('.//w:pBdr')]:e.getparent().remove(e)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    p=OUT/'UCSD_Manuscript_Methods_2_3_EN_ZH.docx';doc.save(p)
    (OUT/'UCSD_Manuscript_Methods_2_3_EN.txt').write_text('\n\n'.join(r['english'] for r in PARAGRAPHS)+'\n')
    content={'stage':25,'date':'2026-10-29','subsection':'2.3 Clinical predictors and missing data','v2_paragraph_ids':['M5','M6'],'paragraphs':PARAGRAPHS,
      'english_word_count':sum(wc(r['english']) for r in PARAGRAPHS),'paragraph_word_counts':{r['id']:wc(r['english']) for r in PARAGRAPHS},
      'v2_approximate_target_words':300,'citation_mapping':{'1':'S01'},'citation_numbering':'Preserves S01 as 1; regenerate globally when Introduction is assembled',
      'source_audit':a,'implementation_documentation':[{'url':'https://scikit-learn.org/1.8/modules/generated/sklearn.impute.SimpleImputer.html','checked_date':'2026-10-29'},{'url':'https://scikit-learn.org/1.8/common_pitfalls.html#data-leakage','checked_date':'2026-10-29'}],
      'models_refitted':0,'new_hypothesis_tests':0,'next_section':'MRI inputs and annotation region (M7-M8)',
      'docx':str(p.relative_to(ROOT)),'docx_sha256':sha(p)}
    (REPORT/'methods_section23.json').write_text(json.dumps(content,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({k:content[k] for k in ['docx','english_word_count','paragraph_word_counts']},indent=2))

def verify():
    from pypdf import PdfReader
    d=json.loads((REPORT/'methods_section23.json').read_text())
    for p,h in d['source_audit']['frozen_file_sha256'].items():assert sha(ROOT/p)==h,p
    assert sha(ROOT/d['docx'])==d['docx_sha256']
    paragraphs=[p.text for p in Document(ROOT/d['docx']).paragraphs]
    for i,r in enumerate(PARAGRAPHS):
        assert paragraphs[2+2*i:4+2*i]==[r['english'],r['chinese']]
        assert not re.search(r'[\u3400-\u9fff，。；：（）【】、]',r['english'])
    assert sum(wc(r['english']) for r in PARAGRAPHS)==d['english_word_count']
    assert paragraphs[6]=='References' and paragraphs[7].startswith('1. Gagnon L')
    pages=len(PdfReader(OUT/'docx_render/UCSD_Manuscript_Methods_2_3_EN_ZH.pdf').pages)
    assert pages==2
    v={'success':True,'stage':25,'date':'2026-10-29','subsection':d['subsection'],'paragraph_ids':['M5','M6'],
      'english_word_count':d['english_word_count'],'paragraph_word_counts':d['paragraph_word_counts'],'matched_chinese_paragraphs':2,
      'word_pages':pages,'visual_review_pages':list(range(1,pages+1)),'current_section_citations':1,'existing_citation_number_1_preserved':True,
      'prior_files_unchanged':len(d['source_audit']['frozen_file_sha256']),'clinical_records_checked':182,'missing_intervals_verified':36,
      'training_age_scalers_checked':15,'training_interval_pipelines_checked':10,'all_saved_preprocessing_parameters_match':True,
      'models_refitted':0,'new_hypothesis_tests':0,'next_section':d['next_section'],
      'sha256':{str(p.relative_to(ROOT)):sha(p) for p in OUT.iterdir() if p.is_file() and p.suffix in {'.docx','.txt'}}}
    (REPORT/'artifact_verification.json').write_text(json.dumps(v,ensure_ascii=False,indent=2)+'\n');print(json.dumps(v,indent=2))

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--verify',action='store_true');args=parser.parse_args()
    verify() if args.verify else build()
