"""Draft D1 from frozen results, without fitting or changing scientific outputs."""
from pathlib import Path
import argparse, hashlib, json, re, zipfile
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'outputs/ucsd-manuscript-step41-20261029'
REPORT = ROOT/'reports/ucsd_manuscript_step41'
PARAGRAPHS = [{
    'id': 'D1', 'citation_ids': [],
    'english': 'In this secondary analysis of public data, adding eight fixed MRI descriptors to five clinical inputs improved internal discrimination between residual or recurrent tumor and treatment-related change after treatment for high-grade glioma. Fusion achieved an OOF AUC of 0.782, compared with 0.672 for M1, with a paired difference of 0.110 (95% patient-bootstrap percentile interval, 0.021 to 0.214). This comparison used the same examinations and patient-grouped folds, limiting differences attributable to case selection or partition assignment. The AUC increment remained positive across all 20 repeated partitions, although these repetitions reused the same cohort and provided descriptive evidence of split stability. Incremental discrimination was less certain when one available examination per patient was retained: the AUC-difference interval extended from -0.015 to 0.252. These intervals were conditional on fixed OOF predictions and did not capture full retraining variability. Moreover, Fusion recognized only 12/41 treatment-related-change examinations in the primary analysis and 2/22 in the single-scan analysis at the fixed threshold. Its primary calibration slope of 0.667, together with intervals crossing zero for the Brier-score and logarithmic-loss differences, also limited claims about probability accuracy. The findings therefore support an incremental MRI signal within this dataset, while leaving its clinical usefulness and performance in independent cohorts unresolved.',
    'chinese': '在这项基于公开数据的二次分析中，将8项固定MRI描述特征加入5项临床输入，提高了对高级别胶质瘤治疗后残余或复发肿瘤与治疗相关改变的内部鉴别能力。Fusion的OOF AUC为0.782，M1为0.672，配对差值为0.110（95%患者bootstrap百分位区间为0.021至0.214）。该比较使用相同检查和相同的患者分组折，从而减少了病例选择或分折分配差异的影响。全部20套重复分折的AUC增量均为正，但这些重复使用了同一队列，提供的是分折稳定性的描述性证据。在每位患者仅保留一次可用检查时，增量鉴别更不确定，AUC差值区间为−0.015至0.252。这些区间以固定OOF预测为条件，未涵盖完整的重新训练变异。此外，在固定阈值下，Fusion在主分析中仅正确识别12/41次治疗相关改变检查，在单扫描分析中仅正确识别2/22次。其主分析校准斜率为0.667，且Brier评分和对数损失差值的区间均跨零，这也限制了对概率准确性作出结论。因此，研究结果支持该数据集内存在MRI增量信号，但其临床实用性和独立队列中的表现仍有待确定。'
}]

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def wc(text):
    return len(re.findall(r"[A-Za-z0-9]+(?:['/-][A-Za-z0-9]+)*", text))

def audit():
    previous = json.loads((ROOT/'reports/ucsd_manuscript_step40/results_section38.json').read_text())
    frozen = previous['source_audit']['frozen_file_sha256'].copy()
    frozen.update(json.loads((ROOT/'reports/ucsd_manuscript_step40/artifact_verification.json').read_text())['sha256'])
    extra = ['reports/ucsd_fusion/metrics.json', 'reports/ucsd_stability/metrics.json',
             'reports/ucsd_writing_requirements_v2/requirements_and_outline.json']
    for p in extra:
        if p not in frozen: frozen[p] = sha(ROOT/p)
    for p, h in frozen.items(): assert sha(ROOT/p) == h, p
    main = json.loads((ROOT/'reports/ucsd_fusion/metrics.json').read_text())
    stability = json.loads((ROOT/'reports/ucsd_stability/metrics.json').read_text())
    pair = {r['metric']: r for r in main['paired_comparisons'] if r['comparison'] == 'Fusion_minus_M1'}
    single = stability['single_scan']
    repeated = next(r for r in stability['repeated_primary_difference_distribution'] if r['metric'] == 'roc_auc')
    assert len(main['models']['MRI']) == 8 and len(main['models']['M1']) == 5 and len(main['models']['Fusion']) == 13
    assert main['same_cases_and_folds'] and main['scans'] == 182 and main['source_patients'] == 135
    numeric_claims = {
        'primary_Fusion_AUC': [main['metrics']['Fusion']['roc_auc'], '0.782'],
        'primary_M1_AUC': [main['metrics']['M1']['roc_auc'], '0.672'],
        'primary_paired_AUC_difference': [pair['roc_auc']['difference'], '0.110'],
        'primary_paired_AUC_lower': [pair['roc_auc']['lower'], '0.021'],
        'primary_paired_AUC_upper': [pair['roc_auc']['upper'], '0.214'],
        'single_paired_AUC_lower': [single['bootstrap_95ci']['Fusion_minus_M1__roc_auc']['lower'], '-0.015'],
        'single_paired_AUC_upper': [single['bootstrap_95ci']['Fusion_minus_M1__roc_auc']['upper'], '0.252'],
        'primary_Fusion_calibration_slope': [main['calibration']['Fusion']['calibration_slope'], '0.667']}
    for value, displayed in numeric_claims.values(): assert f'{value:.3f}' == displayed
    assert repeated['positive_count'] == repeated['n_complete'] == stability['repeats_complete'] == 20
    assert repeated['minimum'] > 0 and not stability['repeats_are_independent_samples']
    assert not stability['repeated_ranges_are_confidence_intervals']
    for metric in ['brier', 'log_loss']: assert pair[metric]['lower'] < 0 < pair[metric]['upper']
    assert main['threshold'] == .5 and main['metrics']['Fusion']['tn'] == 12 and main['treatment_change_scans'] == 41
    assert single['metrics']['Fusion']['tn'] == 2 and single['treatment_change_scans'] == 22
    assert main['bootstrap_models_refitted'] == stability['bootstrap_models_refitted'] == 0
    assert not main['independent_external_validation'] and not stability['external_validation']
    return {'frozen_file_sha256': frozen, 'numeric_claim_mapping': numeric_claims,
            'model_input_counts_checked': {'MRI': 8, 'M1': 5, 'Fusion': 13},
            'primary_treatment_change_recognition': [12, 41], 'single_scan_treatment_change_recognition': [2, 22],
            'positive_repeated_AUC_increments': 20, 'repetitions_are_independent_samples': False,
            'primary_Brier_and_log_loss_intervals_include_zero': True,
            'bootstrap_intervals_conditional_on_fixed_predictions': True,
            'external_validation_performed': False, 'models_refitted': 0, 'new_hypothesis_tests': 0,
            'interpretation': 'Within-dataset incremental discrimination; no established clinical usefulness or external performance',
            'verification_scope': 'Verify frozen files and saved numerical claims; no feature extraction, training, diagnostic refit or resampling'}

def build():
    a = audit()
    OUT.mkdir(parents=True, exist_ok=True); REPORT.mkdir(parents=True, exist_ok=True)
    doc = Document(); sec = doc.sections[0]
    sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(.79)
    for name in ['Normal', 'Title', 'Heading 1']:
        style = doc.styles[name]; style.font.name = 'Times New Roman'; style.font.color.rgb = RGBColor(0, 0, 0)
        rf = style.element.get_or_add_rPr().rFonts; rf.set(qn('w:eastAsia'), 'Noto Serif CJK SC')
        for key in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']: rf.attrib.pop(qn('w:'+key), None)
    normal = doc.styles['Normal']; normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.; normal.paragraph_format.space_after = Pt(9)
    doc.styles['Title'].font.size = Pt(16); doc.styles['Title'].paragraph_format.space_after = Pt(12)
    heading = doc.styles['Heading 1']; heading.font.size = Pt(13)
    heading.paragraph_format.space_before = Pt(8); heading.paragraph_format.space_after = Pt(9)
    doc.core_properties.title = 'Discussion 4.1 Principal findings'; doc.core_properties.author = 'Tianyu Feng'
    doc.add_paragraph('Discussion', style='Title'); doc.add_paragraph('4.1 Principal findings', style='Heading 1')
    for item in PARAGRAPHS:
        en = doc.add_paragraph(item['english']); en.paragraph_format.keep_together = True
        en.paragraph_format.keep_with_next = True
        zh = doc.add_paragraph(item['chinese']); zh.paragraph_format.keep_together = True
        for run in zh.runs: run.font.size = Pt(11)
    for e in [*doc.styles.element.xpath('.//w:pBdr'), *doc.element.xpath('.//w:pBdr')]: e.getparent().remove(e)
    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    path = OUT/'UCSD_Manuscript_Discussion_4_1_EN_ZH.docx'; doc.save(path)
    (OUT/'UCSD_Manuscript_Discussion_4_1_EN.txt').write_text('\n\n'.join(p['english'] for p in PARAGRAPHS)+'\n')
    d = {'stage': 41, 'date': '2026-10-29', 'subsection': '4.1 Principal findings', 'v2_paragraph_ids': ['D1'],
         'paragraphs': PARAGRAPHS, 'english_word_count': sum(wc(p['english']) for p in PARAGRAPHS),
         'paragraph_word_counts': {p['id']: wc(p['english']) for p in PARAGRAPHS},
         'word_count_scope': 'English narrative only; excludes headings and Chinese translation', 'v2_approximate_target_words': 200,
         'equations': {'count': 0, 'numbers': [], 'cumulative_previous_numbered_groups': 10},
         'citation_mapping': {}, 'new_manuscript_references': 0, 'cumulative_unique_cited_sources': 5,
         'source_audit': a, 'models_refitted': 0, 'new_hypothesis_tests': 0,
         'next_section': 'Discussion 4.2 Comparison with fusion studies (D2)',
         'docx': str(path.relative_to(ROOT)), 'docx_sha256': sha(path)}
    (REPORT/'discussion_section41.json').write_text(json.dumps(d, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps({'docx': d['docx'], 'english_word_count': d['english_word_count'], 'prior_files': len(a['frozen_file_sha256'])}))

def verify():
    from pypdf import PdfReader
    d = json.loads((REPORT/'discussion_section41.json').read_text()); a = audit()
    assert a == d['source_audit']
    path = ROOT/d['docx']; assert sha(path) == d['docx_sha256']
    doc = Document(path); texts = [p.text for p in doc.paragraphs]
    for p in PARAGRAPHS:
        assert p['english'] in texts and p['chinese'] in texts
        assert texts.index(p['chinese']) == texts.index(p['english'])+1
        assert not re.search(r'[\u3400-\u9fff，。；：（）【】、]', p['english'])
    assert sum(wc(p['english']) for p in PARAGRAPHS) == d['english_word_count']
    assert len(doc.tables) == 0 and len(doc.sections) == 1
    with zipfile.ZipFile(path) as z:
        xml = etree.fromstring(z.read('word/document.xml'))
        assert not xml.xpath('//m:oMath', namespaces={'m': 'http://schemas.openxmlformats.org/officeDocument/2006/math'})
        assert not [n for n in z.namelist() if n.startswith('word/media/')]
    pages = len(PdfReader(OUT/'docx_render/UCSD_Manuscript_Discussion_4_1_EN_ZH.pdf').pages); assert pages == 1
    v = {'success': True, 'stage': 41, 'date': '2026-10-29', 'subsection': d['subsection'], 'paragraph_ids': ['D1'],
         'english_word_count': d['english_word_count'], 'matched_chinese_paragraphs': 1, 'word_pages': pages,
         'visual_review_pages': [1], 'native_editable_equations': 0, 'equation_images': 0,
         'prior_files_unchanged': len(a['frozen_file_sha256']), 'saved_numeric_claim_values_checked': 8,
         'model_input_counts_checked': 3, 'recognition_counts_checked': 4,
         'positive_repeated_AUC_increments_checked': 20, 'conditional_interval_limit_disclosed': True,
         'clinical_usefulness_not_established': True, 'new_manuscript_references': 0,
         'cumulative_unique_cited_sources': 5, 'models_refitted': 0, 'new_hypothesis_tests': 0,
         'next_section': d['next_section'],
         'sha256': {str(p.relative_to(ROOT)): sha(p) for p in OUT.iterdir() if p.is_file() and p.suffix in {'.docx', '.txt'}}}
    (REPORT/'artifact_verification.json').write_text(json.dumps(v, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps(v, indent=2))

if __name__ == '__main__':
    parser = argparse.ArgumentParser(); parser.add_argument('--verify', action='store_true'); args = parser.parse_args()
    verify() if args.verify else build()
