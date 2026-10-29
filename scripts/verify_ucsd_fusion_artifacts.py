"""Independently reconcile Stage18 folds, preprocessing, probabilities and artifacts."""
from pathlib import Path
import base64,hashlib,json,math,subprocess,sys
import nbformat,numpy as np,pandas as pd
from scipy.special import expit
from scipy.stats import rankdata
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'src'))
from rn_public.ucsd_fusion import prepare,freeze_plan,EVALUATED,BOOT_METRICS,FEATURES


def main():
 out=ROOT/'reports/ucsd_fusion';s=json.loads((out/'metrics.json').read_text());d,e,m,lock=freeze_plan(ROOT)
 pred=pd.read_csv(ROOT/'data/processed/ucsd_fusion_oof.csv');folds=json.loads((out/'fold_audit.json').read_text())
 assert len(pred)==182 and pred.ID.tolist()==d.ID.tolist() and pred.validation_fold.tolist()==m.validation_fold.tolist()
 assert not pred.filter(like='_oof_probability').isna().any().any()
 for key in ['patient_id','brats_patient_id','validation_group']:assert pred.groupby(key).validation_fold.nunique().eq(1).all()
 checked=0
 for f in folds:
  tr=d[d.ID.isin(f['train_IDs'])];te=d[d.ID.isin(f['test_IDs'])]
  assert set(tr.ID).isdisjoint(te.ID) and len(tr)+len(te)==182
  for namespace in ['patient_id','brats_patient_id','validation_group']:assert set(tr[namespace]).isdisjoint(te[namespace])
  for name,record in f['models'].items():
   params=record['parameters'];source=d.copy()
   if name.endswith('_eroded'):source.loc[:,list(e.columns[1:])]=e.iloc[:,1:].to_numpy()
   train=source[source.ID.isin(tr.ID)];test=source[source.ID.isin(te.ID)]
   scaled={}
   for block,q in params['continuous'].items():
    columns=q['columns'];values=train[columns].to_numpy().copy()
    if block=='interval':
     median=float(np.nanmedian(values));assert math.isclose(median,q['impute_median'],abs_tol=1e-12)
     values=np.where(np.isnan(values),median,values)
    means=values.mean(axis=0);sd=values.std(axis=0,ddof=0);sd=np.where(sd==0,1,sd)
    np.testing.assert_allclose(means,q['mean'],rtol=1e-12,atol=1e-12);np.testing.assert_allclose(sd,q['scale'],rtol=1e-12,atol=1e-12)
    vals=test[columns].to_numpy()
    if block=='interval':vals=np.where(np.isnan(vals),q['impute_median'],vals)
    vals=(vals-np.asarray(q['mean']))/np.asarray(q['scale'])
    scaled.update({block+'__'+column:vals[:,i] for i,column in enumerate(columns)})
   for column in params['binary']:scaled['binary__'+column]=test[column].to_numpy()
   assert set(x.split('__',1)[1] for x in params['transformed_names'])==set(FEATURES[name.removesuffix('_eroded')])
   x=np.column_stack([scaled[column] for column in params['transformed_names']]);prob=expit(x@np.asarray(params['coefficients'])+params['intercept'])
   actual=pred.set_index('ID').loc[test.ID,name+'_oof_probability'].to_numpy()
   np.testing.assert_allclose(actual,prob,rtol=1e-12,atol=1e-12);checked+=len(test)
  np.testing.assert_allclose(pred.loc[pred.validation_fold.eq(f['fold']),'prior_oof_probability'],tr.tumor_label.mean())
 y=pred.tumor_label.to_numpy();n1=int(y.sum());n0=int((1-y).sum())
 for name in EVALUATED:
  p=pred[name+'_oof_probability'].to_numpy();metrics=s['metrics'][name]
  auc=(rankdata(p)[y==1].sum()-n1*(n1+1)/2)/(n1*n0)
  assert math.isclose(auc,metrics['roc_auc'],abs_tol=1e-12)
  assert math.isclose(np.mean((p-y)**2),metrics['brier'],abs_tol=1e-12)
  for key,value in dict(tn=int(((y==0)&(p<.5)).sum()),fp=int(((y==0)&(p>=.5)).sum()),fn=int(((y==1)&(p<.5)).sum()),tp=int(((y==1)&(p>=.5)).sum())).items():assert metrics[key]==value
 draws=pd.read_csv(out/'paired_patient_bootstrap.csv');assert len(draws)==2000 and draws.draw.is_unique
 for key,ci in s['patient_cluster_bootstrap_95ci'].items():np.testing.assert_allclose(draws[key].quantile([.025,.975]),[ci['lower'],ci['upper']],atol=1e-12)
 for metric in BOOT_METRICS:np.testing.assert_allclose(draws['Fusion_minus_M1__'+metric],draws['Fusion__'+metric]-draws['M1__'+metric],atol=1e-12)
 calibration=pd.read_csv(out/'calibration_bins.csv');assert calibration.groupby('model').n.sum().eq(182).all()
 nb=nbformat.read(ROOT/'notebooks/18_ucsd_clinical_mri_fusion.ipynb',as_version=4);code=[c for c in nb.cells if c.cell_type=='code']
 kr=json.loads((ROOT/'reports/kernel_execution_18.json').read_text());assert kr['success'] and len(code)==5
 assert hashlib.sha256('\n'.join(c.source for c in code).encode()).hexdigest()==kr['notebooks'][0]['code_sha256']
 assert [c.execution_count for c in code]==[1,2,3,4,5] and not [o for c in code for o in c.outputs if o.output_type=='error']
 images=[base64.b64decode(o.data['image/png']) for c in code for o in c.outputs if o.output_type in ['display_data','execute_result'] and 'image/png' in o.data];assert len(images)==2
 for data,name in zip(images,['model_comparison.png','paired_and_boundary_comparison.png']):assert data==(ROOT/'figures/ucsd_fusion'/name).read_bytes()
 payload=json.loads((ROOT/'reports/ucsd_fusion_workbook_input.json').read_text());wb=load_workbook(ROOT/'outputs/ucsd-fusion-20261029/UCSD_Clinical_MRI_Fusion.xlsx',data_only=True,read_only=True)
 rows_checked=0
 for name,dataset in payload['datasets'].items():
  actual=list(wb[name].iter_rows(min_row=1,max_row=len(dataset['rows'])+1,max_col=len(dataset['headers']),values_only=True));assert list(actual[0])==dataset['headers']
  for row,want in zip(actual[1:],dataset['rows']):
   for a,b in zip(row,want):
    if isinstance(b,bool):b=int(b)
    if b is None or b=='':assert a is None or a==''
    elif isinstance(b,(int,float)):assert isinstance(a,(int,float)) and math.isclose(a,b,rel_tol=1e-12,abs_tol=1e-12),(name,a,b)
    else:assert a==b,(name,a,b)
   rows_checked+=1
 for i,name in enumerate(EVALUATED,5):
  assert math.isclose(wb['Summary'][f'C{i}'].value,s['metrics'][name]['roc_auc'],abs_tol=1e-12)
  assert wb['Summary'][f'G{i}'].value==s['metrics'][name]['tn']
 assert wb['Summary']['C20'].value==182 and wb['Summary']['C22'].value==41 and wb['Summary']['C23'].value==36
 assert not [(sh.title,cell.coordinate,cell.value) for sh in wb for row in sh for cell in row if cell.data_type=='e']
 changed=subprocess.check_output(['git','diff','b2401beba321b65c8356df16aa0c9d44c2c9c31e','--name-only'],cwd=ROOT,text=True).splitlines()
 preserved=True
 for path in changed:
  existed=subprocess.run(['git','cat-file','-e','b2401beba321b65c8356df16aa0c9d44c2c9c31e:'+path],cwd=ROOT,capture_output=True).returncode==0
  if existed and path not in ['README.md','.gitignore','reports/project_status.json']:preserved=False
 assert preserved
 result=dict(success=True,source_input_hashes_verified=4,scans=182,patients=135,folds=5,patient_overlap=0,BraTS_patient_overlap=0,
  OOF_model_probabilities_independently_reconstructed=checked,bootstrap_draws_checked=2000,
  workbook_data_rows_checked=rows_checked,workbook_error_cells=0,notebook_code_cells=5,notebook_errors=0,
  notebook_embedded_figures_byte_identical=2,prior_stage_artifacts_preserved=True,external_Excel_application_tested=False)
 (out/'artifact_verification.json').write_text(json.dumps(result,indent=2)+'\n');print(json.dumps(result,indent=2))


if __name__=='__main__':main()
