"""Draft C1 from frozen main and sensitivity results without modifying the analysis."""
from pathlib import Path
import argparse, hashlib, json, re, zipfile
import pandas as pd
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'outputs/ucsd-manuscript-step53-20261030'
REPORT = ROOT/'reports/ucsd_manuscript_step53'
EVIDENCE = REPORT/'literature_evidence.json'
MAPPING = {}
PARAGRAPHS = [{'id': 'C1', 'citation_ids': [], 'english': 'This secondary analysis of public UCSD-PTGBM data supports an incremental MRI signal for distinguishing residual or recurrent tumor from treatment-related change after treatment for high-grade glioma. In 182 examinations from 135 patients, the fixed clinical-MRI fusion model achieved an out-of-fold AUC of 0.782, exceeding the expanded clinical model by 0.110, with a conditional paired 95% percentile interval of 0.021 to 0.214. The increment remained positive across 20 repeated patient-grouped partitions, while its presence and magnitude were uncertain in the single-scan analysis. Limited recognition of treatment-related change at the fixed threshold and unresolved improvements in probability error constrain clinical interpretation. The reproducible workflow provides a computational basis for evaluating low-dimensional fusion, subject to inherited diagnoses and the supplied derived images. Independent diagnostic evaluation and assessment of clinically justified decisions remain necessary to determine whether the observed increment generalizes and can support care.', 'chinese': '本项基于公开UCSD-PTGBM数据的二次分析支持存在MRI增量信号，用于鉴别高级别胶质瘤治疗后的残余或复发肿瘤与治疗相关改变。在135名患者的182次检查中，固定的临床–MRI融合模型获得0.782的折外AUC，较扩展临床模型提高0.110，条件性配对95%百分位区间为0.021至0.214。在20套重复患者分组分折中，增量均保持为正；但在单扫描分析中，增量是否存在及其大小仍不确定。固定阈值下对治疗相关改变的识别有限，且概率误差改善仍未确定，限制了临床解释。可复现的工作流程为低维融合评价提供了计算基础，其结论受沿用诊断及所用派生影像的约束。仍需独立诊断评价，并对具有临床依据的决策进行评估，才能判断所观察到的增量能否推广并支持诊疗。'}]

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def wc(text):
    return len(re.findall(r"[A-Za-z0-9]+(?:['/-][A-Za-z0-9]+)*", text))

def audit():
    previous=json.loads((ROOT/'reports/ucsd_manuscript_step52/discussion_section52.json').read_text())
    frozen=previous['source_audit']['frozen_file_sha256'].copy()
    frozen.update(json.loads((ROOT/'reports/ucsd_manuscript_step52/artifact_verification.json').read_text())['sha256'])
    for rel in ['reports/ucsd_manuscript_step52/discussion_section52.json','reports/ucsd_manuscript_step32/methods_section210.json']:
        if rel not in frozen:frozen[rel]=sha(ROOT/rel)
    frozen[str(EVIDENCE.relative_to(ROOT))]=sha(EVIDENCE)
    for p,h in frozen.items():assert sha(ROOT/p)==h,p
    m=json.loads((ROOT/'reports/ucsd_fusion/metrics.json').read_text());sm=json.loads((ROOT/'reports/ucsd_stability/metrics.json').read_text())
    assert m['scans']==182 and m['source_patients']==135 and len(m['models']['M1'])==5 and len(m['models']['Fusion'])==13
    assert f"{m['metrics']['Fusion']['roc_auc']:.3f}"=='0.782'
    paired={p['metric']:p for p in m['paired_comparisons'] if p['comparison']=='Fusion_minus_M1'}
    auc=paired['roc_auc'];assert [f"{auc[k]:.3f}" for k in ['difference','lower','upper']]==['0.110','0.021','0.214']
    for key in ['brier','log_loss']:assert paired[key]['lower']<0<paired[key]['upper']
    assert m['threshold']==.5 and m['metrics']['Fusion']['tn']==12 and m['treatment_change_scans']==41
    assert not m['bootstrap_models_refitted'] and not m['independent_external_validation']
    assert not m['source_annotation_blinding_confirmed'] and not m['physical_scale_confirmed']
    repeated=pd.read_csv(ROOT/'reports/ucsd_stability/repeated_paired_differences.csv')
    rows=repeated.loc[repeated.comparison.eq('Fusion_minus_M1') & repeated.metric.eq('roc_auc')]
    assert len(rows)==20 and rows.difference.gt(0).all() and sm['repeats_complete']==20 and sm['repeats_failed']==0
    assert not sm['repeats_are_independent_samples'] and not sm['external_validation']
    ci=sm['single_scan']['bootstrap_95ci']['Fusion_minus_M1__roc_auc'];assert ci['lower']<0<ci['upper']
    prior_ids=set()
    for pattern in ['methods_section*.json','results_section*.json','discussion_section*.json']:
        for p in ROOT.glob('reports/ucsd_manuscript_step*/'+pattern):
            if int(p.parent.name.rsplit('step',1)[1])<53:prior_ids.update(r for item in json.loads(p.read_text())['paragraphs'] for r in item['citation_ids'])
    assert len(prior_ids)==17
    return {'frozen_file_sha256':frozen,'reference_records':[],'source_evidence_record':str(EVIDENCE.relative_to(ROOT)),
            'new_full_texts_checked_for_current_claims':[],'primary_examinations':182,'primary_patients':135,
            'primary_fusion_AUC':m['metrics']['Fusion']['roc_auc'],'paired_AUC_difference_checked':auc,
            'positive_repeated_AUC_increments_checked':20,'single_scan_conditional_AUC_interval_checked':ci,
            'probability_error_difference_intervals_cross_zero':True,'fixed_threshold_treatment_change_recognition':[12,41],
            'inherited_label_and_derived_image_limits_preserved':True,'no_external_validation_or_clinical_utility_claimed':True,
            'no_new_citations_figures_or_tables':True,'declarations_and_acknowledgments_deferred_by_user':True,
            'remaining_user_requested_writing':['Abstract','Introduction'],'previous_unique_cited_sources':sorted(prior_ids),
            'models_refitted':0,'new_hypothesis_tests':0,'new_bootstrap_draws':0,
            'verification_scope':'Read-only main and single-scan interval checks, saved repeated increment directions and model scope; conclude from existing results without new fitting, calculations or declarations'}

def build():
    a = audit(); OUT.mkdir(parents=True, exist_ok=True); REPORT.mkdir(parents=True, exist_ok=True)
    doc = Document(); sec = doc.sections[0]; sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(.79)
    for name in ['Normal', 'Title', 'Heading 1']:
        st = doc.styles[name]; st.font.name = 'Times New Roman'; st.font.color.rgb = RGBColor(0, 0, 0)
        rf = st.element.get_or_add_rPr().rFonts; rf.set(qn('w:eastAsia'), 'Noto Serif CJK SC')
        for key in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']: rf.attrib.pop(qn('w:'+key), None)
    normal = doc.styles['Normal']; normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.; normal.paragraph_format.space_after = Pt(9)
    doc.styles['Title'].font.size = Pt(16); doc.styles['Title'].paragraph_format.space_after = Pt(12)
    heading = doc.styles['Heading 1']; heading.font.size = Pt(13)
    heading.paragraph_format.space_before = Pt(8); heading.paragraph_format.space_after = Pt(9)
    doc.core_properties.title = '5 Conclusion'; doc.core_properties.author = 'Tianyu Feng'
    doc.add_paragraph('5 Conclusion', style='Title')
    for item in PARAGRAPHS:
        en = doc.add_paragraph(item['english']); en.paragraph_format.keep_together = True; en.paragraph_format.keep_with_next = True
        zh = doc.add_paragraph(item['chinese']); zh.paragraph_format.keep_together = True
        for run in zh.runs: run.font.size = Pt(11)
    if a['reference_records']:
        title = doc.add_paragraph('References', style='Heading 1'); title.paragraph_format.page_break_before = False
    for ref in a['reference_records']:
        p = doc.add_paragraph(str(ref['number'])+'. '+ref['citation']); p.paragraph_format.keep_together = True
    for e in [*doc.styles.element.xpath('.//w:pBdr'), *doc.element.xpath('.//w:pBdr')]: e.getparent().remove(e)
    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    path = OUT/'UCSD_Manuscript_Conclusion_EN_ZH.docx'; doc.save(path)
    (OUT/'UCSD_Manuscript_Conclusion_EN.txt').write_text('\n\n'.join(p['english'] for p in PARAGRAPHS)+'\n')
    d = {'stage': 53, 'date': '2026-10-30', 'subsection': '5 Conclusion', 'v2_paragraph_ids': ['C1'],
         'paragraphs': PARAGRAPHS, 'english_word_count': sum(wc(p['english']) for p in PARAGRAPHS),
         'paragraph_word_counts': {p['id']: wc(p['english']) for p in PARAGRAPHS},
         'word_count_scope': 'English narrative including in-text citation numbers; excludes headings, reference list and Chinese translation',
         'v2_approximate_target_words': 150, 'equations': {'count': 0, 'numbers': [], 'cumulative_previous_numbered_groups': 10},
         'citation_mapping': MAPPING, 'citation_numbering': 'No citations added; existing manuscript numbering unchanged',
         'new_manuscript_references': 0, 'cumulative_unique_cited_sources': 17,
         'source_audit': a, 'models_refitted': 0, 'new_hypothesis_tests': 0,
         'next_section': 'Abstract (A1-A4); Introduction to follow; declarations deferred by user',
         'docx': str(path.relative_to(ROOT)), 'docx_sha256': sha(path)}
    (REPORT/'conclusion_section53.json').write_text(json.dumps(d, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps({'docx': d['docx'], 'english_word_count': d['english_word_count'], 'frozen_files': len(a['frozen_file_sha256'])}))

def verify():
    from pypdf import PdfReader
    d = json.loads((REPORT/'conclusion_section53.json').read_text()); a = audit(); assert a == d['source_audit']
    path = ROOT/d['docx']; assert sha(path) == d['docx_sha256']; doc = Document(path); texts = [p.text for p in doc.paragraphs]
    for p in PARAGRAPHS:
        assert p['english'] in texts and p['chinese'] in texts and texts.index(p['chinese']) == texts.index(p['english'])+1
        assert not re.search(r'[\u3400-\u9fff，。；：（）【】、]', p['english'])
    assert sum(wc(p['english']) for p in PARAGRAPHS) == d['english_word_count']
    assert re.findall(r'\((\d+)\)', PARAGRAPHS[0]['english']) == []
    for ref in a['reference_records']: assert str(ref['number'])+'. '+ref['citation'] in texts
    assert not doc.tables and len(doc.sections) == 1
    with zipfile.ZipFile(path) as z:
        xml = etree.fromstring(z.read('word/document.xml'))
        assert not xml.xpath('//m:oMath', namespaces={'m': 'http://schemas.openxmlformats.org/officeDocument/2006/math'})
        assert not [n for n in z.namelist() if n.startswith('word/media/')]
    pdf = PdfReader(OUT/'docx_render/UCSD_Manuscript_Conclusion_EN_ZH.pdf'); pages = len(pdf.pages); assert pages == 1
    assert '5 Conclusion' in pdf.pages[0].extract_text() and 'References' not in texts
    v = {'success': True, 'stage': 53, 'date': '2026-10-30', 'subsection': d['subsection'], 'paragraph_ids': ['C1'],
         'english_word_count': d['english_word_count'], 'matched_chinese_paragraphs': 1, 'word_pages': pages,
         'visual_review_pages': [1], 'native_editable_equations': 0, 'equation_images': 0,
         'frozen_files_unchanged': len(a['frozen_file_sha256']), 'prior_files_unchanged': len(a['frozen_file_sha256'])-1,
         'new_full_texts_checked_for_current_claims': [], 'primary_AUC_and_paired_interval_checked': True,
         'positive_repeated_AUC_increments_checked': 20, 'single_scan_uncertainty_preserved': True,
         'probability_error_and_minority_recognition_limits_preserved': True,
         'no_external_validation_or_clinical_utility_claimed': True, 'new_bootstrap_draws': 0,
         'conclusion_v2_paragraph_complete': True, 'declarations_and_acknowledgments_deferred_by_user': True,
         'remaining_user_requested_writing': ['Abstract','Introduction'],
         'new_manuscript_references': 0, 'provisional_citation_mapping': MAPPING, 'cumulative_unique_cited_sources': 17,
         'models_refitted': 0, 'new_hypothesis_tests': 0, 'next_section': d['next_section'],
         'sha256': {str(p.relative_to(ROOT)): sha(p) for p in OUT.iterdir() if p.is_file() and p.suffix in {'.docx', '.txt'}}}
    (REPORT/'artifact_verification.json').write_text(json.dumps(v, ensure_ascii=False, indent=2)+'\n'); print(json.dumps(v, indent=2))

if __name__ == '__main__':
    parser = argparse.ArgumentParser(); parser.add_argument('--verify', action='store_true'); args = parser.parse_args()
    verify() if args.verify else build()
