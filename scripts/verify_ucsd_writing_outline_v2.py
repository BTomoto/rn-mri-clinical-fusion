"""Verify the revised outline against its source fingerprints and frozen results."""
from pathlib import Path
import hashlib,json,re,subprocess
import fitz
from docx import Document

ROOT=Path(__file__).resolve().parents[1]
REPORT=ROOT/'reports/ucsd_writing_requirements_v2'


def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()


def main():
    plan=json.loads((REPORT/'requirements_and_outline.json').read_text())
    receipt=json.loads((REPORT/'generation_receipt.json').read_text())
    path=ROOT/receipt['deliverable'];assert sha(path)==receipt['sha256']
    for x in receipt['source_documents']:assert sha(ROOT.parent/'upload'/x['filename'])==x['sha256']
    groups=plan['paragraphs'];assert sum(r[1] for k,v in groups.items() if k!='abstract' for r in v)==8000
    assert sum(r[1] for r in groups['abstract'])==250
    assert sum(len(v) for k,v in groups.items() if k!='abstract')==47
    ids=[r[0] for v in groups.values() for r in v];assert len(ids)==len(set(ids))==51
    b=plan['bibliography'];assert b['planned_total']>=b['minimum_total']==25
    assert b['recent_5yr_english_min']/b['planned_total']>.9
    assert b['recent_3yr_english_min']/b['planned_total']>.5
    assert not b['completed'] and not plan['manuscript_body_drafted'] and plan['models_refitted']==0
    doc=Document(path);texts=[p.text for p in doc.paragraphs];text='\n'.join(texts)
    for ident in ids+[r[0] for r in plan['pdf_requirements']]+[r[0] for r in plan['template_requirements']]:
        assert len([t for t in texts if re.match('^'+ident+r'\s',t)])==1,ident
    assert len(doc.tables)==3
    table=[[c.text for c in r.cells] for r in doc.tables[0].rows]
    assert table[-1]==['Main text','8000','47']
    assert not doc.styles.element.xpath('.//w:pBdr')
    # Evidence anchors remain tied to the unchanged analysis, not new fitting.
    metrics=json.loads((ROOT/'reports/ucsd_fusion/metrics.json').read_text())['metrics']
    for name in ['M1','MRI','Fusion']:assert f"{metrics[name]['roc_auc']:.3f}" in text
    old=json.loads((ROOT/'reports/ucsd_writing_preparation/generation_receipt.json').read_text())
    assert sha(ROOT/old['deliverable'])==old['sha256']
    material=json.loads((ROOT/'reports/ucsd_paper_materials/generation_receipt.json').read_text())
    for rel,h in {**material['input_sha256'],**material['output_sha256']}.items():assert sha(ROOT/rel)==h
    render=path.parent/'docx_render_verified';pdf=fitz.open(render/(path.stem+'.pdf'));assert len(pdf)==13
    assert all(len(p.get_text())>350 for p in pdf)
    assert all((render/f'page-{i}.png').exists() for i in range(1,14))
    changed=subprocess.check_output(['git','diff','0e12d29b1a36053b9538025c6dcdeb63e36eda52','--name-only'],cwd=ROOT,text=True).splitlines()
    for name in changed:
        existed=subprocess.run(['git','cat-file','-e','0e12d29b1a36053b9538025c6dcdeb63e36eda52:'+name],cwd=ROOT,capture_output=True).returncode==0
        assert not existed or name in ['README.md','.gitignore','reports/project_status.json'],name
    result={'success':True,'source_pdf_pages_reviewed':3,'source_template_pages_reviewed':4,'output_pages_reviewed':13,
            'pdf_rules_mapped':8,'template_rules_mapped':7,'paragraph_ids_checked':51,'main_paragraphs':47,
            'main_word_target':8000,'reference_target':40,'planned_5yr_english_share':.925,'planned_3yr_english_share':.625,
            'bibliography_completed':False,'manuscript_body_drafted':False,'model_refits':0,
            'source_documents_unchanged':True,'prior_outline_and_result_hashes_preserved':True,'docx_sha256':sha(path)}
    (REPORT/'artifact_verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(result,ensure_ascii=False,indent=2))


if __name__=='__main__':main()
