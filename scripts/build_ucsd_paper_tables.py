"""Editable Word tables only; manuscript narrative is intentionally not drafted."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]


def shown_table(t):
 h=t['headers'];r=t['rows'];key=t['key']
 if key in ['Table2','TableS4','TableS6']:
  return ['模型','AUC（95%区间）','Brier','治疗改变AP','治疗改变识别'],[[x[0],f'{x[1]:.4f}\n({x[2]:.4f}, {x[3]:.4f})',f'{x[4]:.4f}',f'{x[5]:.4f}',f'{int(x[6])}/{int(x[7])}'] for x in r],[1.05,1.9,1.,1.15,1.2]
 if key in ['Table3','TableS7']:
  return ['指标','差值','95%区间'],[[x[0],f'{x[1]:.4f}',f'({x[2]:.4f}, {x[3]:.4f})'] for x in r],[2.4,1.2,2.7]
 if key=='TableS3':
  return ['模型','治疗改变\n正确识别','肿瘤\n正确识别','平衡\n准确率','校准\n截距','校准\n斜率'],[[x[0],f'{x[1]}/{x[2]}',f'{x[3]}/{x[4]}',f'{100*x[5]:.1f}%',f'{x[6]:.4f}',f'{x[7]:.4f}'] for x in r],[1.05,1.15,1.15,1.,1.,.95]
 if key=='TableS5A':
  return ['模型','AUC中位数\n[Q1, Q3]','AUC范围','Brier\n中位数'],[[x[0],f'{x[1]:.4f}\n[{x[2]:.4f}, {x[3]:.4f}]',f'{x[4]:.4f}–{x[5]:.4f}',f'{x[6]:.4f}'] for x in r],[1.2,2.1,1.8,1.2]
 if key=='TableS5B':
  return ['指标','差值中位数\n[Q1, Q3]','范围','有利方向\n重复数'],[[x[0],f'{x[1]:.4f}\n[{x[2]:.4f}, {x[3]:.4f}]',f'{x[4]:.4f}–{x[5]:.4f}',f'{int(x[6])}/{int(x[7])}'] for x in r],[1.55,2.,1.75,1.]
 if key=='TableS2':return h,r,[1.65,1.1,1.15,2.4]
 return h,r,[2.35,1.4,1.4,1.15]


def build(root=ROOT):
 root=Path(root);payload=json.loads((root/'reports/ucsd_paper_materials/table_package.json').read_text())
 doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
 sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.75);sec.right_margin=Inches(.75)
 for name in ['Normal','Title','Heading 1','Heading 2']:
  style=doc.styles[name];style.font.name='Nimbus Roman';style.font.color.rgb=RGBColor(0,0,0)
  style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC')
 doc.styles['Normal'].font.size=Pt(10.5);doc.styles['Normal'].paragraph_format.space_after=Pt(5)
 doc.styles['Title'].font.size=Pt(17);doc.styles['Title'].paragraph_format.space_after=Pt(10)
 doc.styles['Heading 1'].font.size=Pt(12);doc.styles['Heading 1'].paragraph_format.space_before=Pt(8);doc.styles['Heading 1'].paragraph_format.space_after=Pt(6)
 doc.add_paragraph('胶质瘤治疗后鉴别研究数据表',style='Title')
 doc.add_paragraph('主影像队列包含182次扫描、135名源患者。模型性能来自患者分组的折外预测，表中保留主分析及全部预设敏感性结果。')
 specs={t['key']:t for t in payload['tables']}
 groups=[['Table1'],['Table2','Table3'],['TableS1'],['TableS2'],['TableS3','TableS4'],['TableS5A','TableS5B'],['TableS6','TableS7']]
 for page,keys in enumerate(groups):
  if page:doc.add_page_break()
  for key in keys:
   t=specs[key];doc.add_paragraph(t['title'],style='Heading 1')
   headers,rows,widths=shown_table(t);table=doc.add_table(rows=1,cols=len(headers));table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
   total=sum(widths);widths=[w*7/total for w in widths]
   for column,w in zip(table.columns,widths):column.width=Inches(w)
   for j,text in enumerate(headers):table.rows[0].cells[j].text=text
   for row in rows:
    cells=table.add_row().cells
    for j,value in enumerate(row):cells[j].text=str(value)
   for i,row in enumerate(table.rows):
    trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
    if i==0:
     repeat=OxmlElement('w:tblHeader');repeat.set(qn('w:val'),'true');trpr.append(repeat)
    for j,cell in enumerate(row.cells):
     cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
     tcpr=cell._tc.get_or_add_tcPr();margins=OxmlElement('w:tcMar')
     for side in ['top','left','bottom','right']:
      e=OxmlElement('w:'+side);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margins.append(e)
     tcpr.append(margins);borders=OxmlElement('w:tcBorders')
     for side in ['top','left','bottom','right']:
      e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
     tcpr.append(borders)
     if i==0:
      fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'EDF0F2');tcpr.append(fill)
     for p in cell.paragraphs:
      p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 or key=='TableS2' else WD_ALIGN_PARAGRAPH.CENTER
      p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.08
      for run in p.runs:run.font.size=Pt(9.5 if key=='TableS2' else 10);run.bold=i==0
   for note in t['notes']:
    p=doc.add_paragraph('注：'+note);p.paragraph_format.space_before=Pt(4)
    for run in p.runs:run.font.size=Pt(9)
  if page==2:
   doc.add_paragraph('标签与数据口径',style='Heading 1')
   doc.add_paragraph('治疗改变来自公开阴性病例表；肿瘤标签沿用已冻结的公开总体分组补集规则，并非全部逐例病理确证。治疗改变包含多种治疗后改变，不等同于仅放射性脑坏死。')
   doc.add_paragraph('普通BraTS全部正标签并集称为源标注区域。源标注盲法未确认，图像为8位派生版本、空间单位unknown，没有外部验证。')
  if page==6:
   doc.add_paragraph('数据来源',style='Heading 1')
   doc.add_paragraph('UCSD-PTGBM，数据DOI 10.7937/fwv2-dt74；描述论文DOI 10.1038/s41597-025-06499-z。公开镜像：https://github.com/data-nih/tcia/releases。数据许可CC BY 4.0，具体文件清单及校验值保留于项目。')
 # The bundled base template carries a Title border. Remove it explicitly.
 for element in [*doc.styles.element.xpath('.//w:pBdr'),*doc.element.xpath('.//w:pBdr')]:element.getparent().remove(element)
 path=root/'outputs/ucsd-paper-materials-20261029/UCSD_Paper_Tables.docx';path.parent.mkdir(parents=True,exist_ok=True);doc.save(path)
 return path


if __name__=='__main__':print(build())
