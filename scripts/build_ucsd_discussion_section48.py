"""Draft D8 from frozen source labels and distinct source annotation procedures."""
from pathlib import Path
import argparse, hashlib, json, re, zipfile
import pandas as pd
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'outputs/ucsd-manuscript-step48-20261030'
REPORT = ROOT/'reports/ucsd_manuscript_step48'
EVIDENCE = REPORT/'literature_evidence.json'
MAPPING = {'1': 'S01', '3': 'S03'}
PARAGRAPHS = [{'id': 'D8', 'citation_ids': ['S01', 'S03'], 'english': 'The principal limitation is that discrimination was evaluated against inherited examination labels rather than independently readjudicated diagnoses. We assigned tumor status to the complement of the 51-case negative list, following the published binary partition. This reproduces the source grouping but cannot resolve individual diagnostic uncertainty. The dataset authors reported pathological confirmation for 45 of 192 tumor examinations in the full collection; this aggregate cannot establish the pathology-confirmed fraction of our imaging subset (1). Other source classifications incorporated clinical and imaging assessment. The target may therefore partly reflect the imaging appearances from which predictors were derived, creating possible dependence between the reference standard and the index data. Its direction and magnitude were not measured. Treatment-related change also combined pseudoprogression, radiation necrosis and nonspecific changes. Disagreement between workbook and published subtype counts remained unresolved despite agreement in the binary total, limiting subtype-specific interpretation. Annotation introduces a separate uncertainty. The source paper describes automated BraTS segmentation followed by radiologist correction and approval, whereas its cellular-tumor masks used clinical, pathological and longitudinal imaging information (1). We used the BraTS-region union, not those cellular-tumor masks; the information available during correction of the supplied BraTS masks could not be confirmed. Unverified blinding remains a potential bias concern rather than evidence that leakage occurred. Reporting guidance calls for documenting reference-standard methods and annotation sources (3). Patient grouping prevents cross-fold patient overlap but cannot repair these upstream uncertainties, which warrant independent diagnostic adjudication and documented annotation procedures in subsequent evaluation.', 'chinese': '本研究的主要局限在于，鉴别能力是对照沿用的检查标签进行评价，而非对照独立重新裁定的诊断。我们依据已发表的二分类划分，将51例阴性清单之外的检查赋为肿瘤类。这能够重现源分组，但无法解决个体诊断的不确定性。数据集作者报告，完整数据集中192次肿瘤检查有45次获得病理证实；这一汇总不能确定本研究影像子集中病理证实检查的比例（1）。其他源分类纳入了临床与影像评估。因此，目标标签可能部分反映了用于提取预测变量的影像表现，使参考标准与受评价的影像数据之间存在潜在依赖，其方向和程度尚未测量。治疗相关改变还合并了假性进展、放射性坏死及非特异性改变。尽管二分类总数一致，工作簿与已发表亚型数量的差异仍未解决，限制了亚型特异性的解释。标注还带来另一项不确定性。源论文描述了BraTS自动分割后由放射科医师修正与审核的流程，而其细胞肿瘤掩膜使用了临床、病理及纵向影像信息（1）。本研究使用的是BraTS区域并集，而非这些细胞肿瘤掩膜；无法确认所用BraTS掩膜在修正过程中可获得哪些信息。盲法未经确认属于潜在偏倚问题，并非已经发生信息泄漏的证据。报告指南要求记录参考标准的获取方法及标注来源（3）。患者分组能够防止患者跨折重叠，但不能弥补这些上游不确定性，因此后续评价需要独立诊断裁定及有记录的标注流程。'}]

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def wc(text):
    return len(re.findall(r"[A-Za-z0-9]+(?:['/-][A-Za-z0-9]+)*", text))

def audit():
    from openpyxl import load_workbook
    from collections import Counter
    previous = json.loads((ROOT/'reports/ucsd_manuscript_step47/discussion_section47.json').read_text())
    frozen = previous['source_audit']['frozen_file_sha256'].copy()
    frozen.update(json.loads((ROOT/'reports/ucsd_manuscript_step47/artifact_verification.json').read_text())['sha256'])
    for rel in ['reports/ucsd_manuscript_step24/methods_section22.json', 'reports/ucsd_manuscript_step26/methods_section24.json', 'data/raw/reference/ucsd_ptgbm_negative_cases.xlsx', 'data/processed/ucsd_clinical_baseline_cohort.csv', 'data/processed/ucsd_ptgbm_subtype_count_comparison.csv']:
        if rel not in frozen: frozen[rel] = sha(ROOT/rel)
    frozen[str(EVIDENCE.relative_to(ROOT))] = sha(EVIDENCE)
    for p,h in frozen.items(): assert sha(ROOT/p) == h,p
    evidence = json.loads(EVIDENCE.read_text()); records = {r['reference_id']:r for r in evidence['studies']}
    refs = {r['id']:r for r in json.loads((ROOT/'reports/ucsd_manuscript_step23/reference_library.json').read_text())['references']}
    selected=[]
    for number,rid in MAPPING.items():
        ref=refs[rid]; r=records[rid]
        assert ref['doi']==r['doi'] and r['provisional_citation_number']==int(number)
        selected.append({'number':int(number),'reference_id':rid,'citation':ref['citation'],'doi':ref['doi'],'read_level_this_stage':r['read_level_this_stage']})
    wb=load_workbook(ROOT/'data/raw/reference/ucsd_ptgbm_negative_cases.xlsx',read_only=True,data_only=True)
    rows=list(wb['all_negative_cases_TCIA'].iter_rows(values_only=True)); wb.close()
    negatives=[dict(zip(rows[0],r)) for r in rows[1:]]; ids={r['ID'] for r in negatives}
    assert len(ids)==len(negatives)==51
    subtypes=dict(Counter(r['Dx'] for r in negatives))
    assert subtypes=={'NS':29,'PsP (clinically confirmed)':9,'Radiation necrosis':7,'PsP (bx proven)':6}
    counts={}
    for name,filename in [('source','ucsd_clinical_baseline_cohort.csv'),('clinical','ucsd_clinical_revalidated_cohort.csv'),('imaging','ucsd_clinical_image_features.csv')]:
        df=pd.read_csv(ROOT/'data/processed'/filename)
        assert df.ID.is_unique and df.tumor_label.eq((~df.ID.isin(ids)).astype(int)).all()
        counts[name]={'examinations':len(df),'patients':int(df.patient_id.nunique()),'tumor':int(df.tumor_label.sum()),'treatment_change':int(df.tumor_label.eq(0).sum())}
    old=json.loads((ROOT/'reports/ucsd_manuscript_step24/methods_section22.json').read_text())['source_audit']
    assert counts==old['label_counts'] and subtypes==old['subtype_counts_in_frozen_workbook']
    assert old['pathology_confirmed_tumor_examinations_in_published_full_collection']==45
    assert records['S01']['facts']['published_subtypes']==old['subtype_counts_in_published_full_collection']
    folds=json.loads((ROOT/'reports/ucsd_fusion/fold_audit.json').read_text())
    assert len(folds)==5
    for f in folds:
        for key,value in f.items():
            if 'overlap' in key: assert value==0,(key,value)
    prior_ids=set()
    for pattern in ['methods_section*.json','results_section*.json','discussion_section*.json']:
        for p in ROOT.glob('reports/ucsd_manuscript_step*/'+pattern):
            if int(p.parent.name.rsplit('step',1)[1])<48: prior_ids.update(r for item in json.loads(p.read_text())['paragraphs'] for r in item['citation_ids'])
    assert len(prior_ids)==16 and set(MAPPING.values()).issubset(prior_ids)
    return {'frozen_file_sha256':frozen,'reference_records':selected,'source_evidence_record':str(EVIDENCE.relative_to(ROOT)),
            'new_full_texts_checked_for_current_claims':[],'previously_reviewed_full_text_revisited':['S01'],'official_full_checklist_revisited':['S03'],
            'label_counts':counts,'label_rows_checked':sum(v['examinations'] for v in counts.values()),'all_labels_reconstructed_and_matched':True,
            'subtype_counts_in_frozen_workbook':subtypes,'subtype_counts_in_published_full_collection':old['subtype_counts_in_published_full_collection'],
            'subtype_discrepancy_resolved':False,'pathology_confirmed_tumor_examinations_in_published_full_collection':45,
            'pathology_count_extrapolated_to_imaging_subset':False,'distinct_BraTS_and_cellular_annotation_procedures_preserved':True,
            'supplied_BraTS_annotation_blinding_confirmed':False,'information_leakage_claimed_proven':False,'independent_diagnostic_readjudication_performed':False,
            'five_saved_folds_checked_for_no_patient_overlap':True,'previous_unique_cited_sources':sorted(prior_ids),'models_refitted':0,'new_hypothesis_tests':0,
            'verification_scope':'Read-only reconstruction of 663 existing labels and subtype counts plus saved fold-overlap checks and manual reading of source annotation procedures; no diagnostic readjudication or fitting'}

def build():
    a = audit(); OUT.mkdir(parents=True, exist_ok=True); REPORT.mkdir(parents=True, exist_ok=True)
    doc = Document(); sec = doc.sections[0]; sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(.79)
    for name in ['Normal', 'Title', 'Heading 1']:
        st = doc.styles[name]; st.font.name = 'Times New Roman'; st.font.color.rgb = RGBColor(0, 0, 0)
        rf = st.element.get_or_add_rPr().rFonts; rf.set(qn('w:eastAsia'), 'Noto Serif CJK SC')
        for key in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']: rf.attrib.pop(qn('w:'+key), None)
    normal = doc.styles['Normal']; normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.; normal.paragraph_format.space_after = Pt(9)
    doc.styles['Title'].font.size = Pt(16); doc.styles['Title'].paragraph_format.space_after = Pt(12)
    heading = doc.styles['Heading 1']; heading.font.size = Pt(13)
    heading.paragraph_format.space_before = Pt(8); heading.paragraph_format.space_after = Pt(9)
    doc.core_properties.title = 'Discussion 4.8 Reference-standard limitations'; doc.core_properties.author = 'Tianyu Feng'
    doc.add_paragraph('Discussion', style='Title'); doc.add_paragraph('4.8 Reference-standard limitations', style='Heading 1')
    for item in PARAGRAPHS:
        en = doc.add_paragraph(item['english']); en.paragraph_format.keep_together = True; en.paragraph_format.keep_with_next = True
        zh = doc.add_paragraph(item['chinese']); zh.paragraph_format.keep_together = True
        for run in zh.runs: run.font.size = Pt(11)
    title = doc.add_paragraph('References', style='Heading 1'); title.paragraph_format.page_break_before = True
    for ref in a['reference_records']:
        p = doc.add_paragraph(str(ref['number'])+'. '+ref['citation']); p.paragraph_format.keep_together = True
    for e in [*doc.styles.element.xpath('.//w:pBdr'), *doc.element.xpath('.//w:pBdr')]: e.getparent().remove(e)
    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    path = OUT/'UCSD_Manuscript_Discussion_4_8_EN_ZH.docx'; doc.save(path)
    (OUT/'UCSD_Manuscript_Discussion_4_8_EN.txt').write_text('\n\n'.join(p['english'] for p in PARAGRAPHS)+'\n')
    d = {'stage': 48, 'date': '2026-10-30', 'subsection': '4.8 Reference-standard limitations', 'v2_paragraph_ids': ['D8'],
         'paragraphs': PARAGRAPHS, 'english_word_count': sum(wc(p['english']) for p in PARAGRAPHS),
         'paragraph_word_counts': {p['id']: wc(p['english']) for p in PARAGRAPHS},
         'word_count_scope': 'English narrative including in-text citation numbers; excludes headings, reference list and Chinese translation',
         'v2_approximate_target_words': 250, 'equations': {'count': 0, 'numbers': [], 'cumulative_previous_numbered_groups': 10},
         'citation_mapping': MAPPING, 'citation_numbering': 'Existing numbers 1 and 3 preserved; renumber globally when Introduction is assembled',
         'new_manuscript_references': 0, 'cumulative_unique_cited_sources': 16,
         'source_audit': a, 'models_refitted': 0, 'new_hypothesis_tests': 0,
         'next_section': 'Discussion 4.9 Image-processing limitations (D9)',
         'docx': str(path.relative_to(ROOT)), 'docx_sha256': sha(path)}
    (REPORT/'discussion_section48.json').write_text(json.dumps(d, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps({'docx': d['docx'], 'english_word_count': d['english_word_count'], 'frozen_files': len(a['frozen_file_sha256'])}))

def verify():
    from pypdf import PdfReader
    d = json.loads((REPORT/'discussion_section48.json').read_text()); a = audit(); assert a == d['source_audit']
    path = ROOT/d['docx']; assert sha(path) == d['docx_sha256']; doc = Document(path); texts = [p.text for p in doc.paragraphs]
    for p in PARAGRAPHS:
        assert p['english'] in texts and p['chinese'] in texts and texts.index(p['chinese']) == texts.index(p['english'])+1
        assert not re.search(r'[\u3400-\u9fff，。；：（）【】、]', p['english'])
    assert sum(wc(p['english']) for p in PARAGRAPHS) == d['english_word_count']
    assert re.findall(r'\((\d+)\)', PARAGRAPHS[0]['english']) == ['1','1','3']
    for ref in a['reference_records']: assert str(ref['number'])+'. '+ref['citation'] in texts
    assert not doc.tables and len(doc.sections) == 1
    with zipfile.ZipFile(path) as z:
        xml = etree.fromstring(z.read('word/document.xml'))
        assert not xml.xpath('//m:oMath', namespaces={'m': 'http://schemas.openxmlformats.org/officeDocument/2006/math'})
        assert not [n for n in z.namelist() if n.startswith('word/media/')]
    pdf = PdfReader(OUT/'docx_render/UCSD_Manuscript_Discussion_4_8_EN_ZH.pdf'); pages = len(pdf.pages); assert pages == 2
    assert '4.8 Reference' in pdf.pages[0].extract_text() and 'References' in pdf.pages[1].extract_text()
    v = {'success': True, 'stage': 48, 'date': '2026-10-30', 'subsection': d['subsection'], 'paragraph_ids': ['D8'],
         'english_word_count': d['english_word_count'], 'matched_chinese_paragraphs': 1, 'word_pages': pages,
         'visual_review_pages': [1, 2], 'native_editable_equations': 0, 'equation_images': 0,
         'frozen_files_unchanged': len(a['frozen_file_sha256']), 'prior_files_unchanged': len(a['frozen_file_sha256'])-1,
         'new_full_texts_checked_for_current_claims': [], 'label_rows_checked': a['label_rows_checked'],
         'all_labels_reconstructed_and_matched': True, 'subtype_discrepancy_resolved': False,
         'pathology_count_extrapolated_to_imaging_subset': False, 'distinct_BraTS_and_cellular_annotation_procedures_preserved': True,
         'supplied_BraTS_annotation_blinding_confirmed': False, 'information_leakage_claimed_proven': False,
         'independent_diagnostic_readjudication_performed': False, 'five_saved_folds_checked_for_no_patient_overlap': True,
         'new_manuscript_references': 0, 'provisional_citation_mapping': MAPPING, 'cumulative_unique_cited_sources': 16,
         'models_refitted': 0, 'new_hypothesis_tests': 0, 'next_section': d['next_section'],
         'sha256': {str(p.relative_to(ROOT)): sha(p) for p in OUT.iterdir() if p.is_file() and p.suffix in {'.docx', '.txt'}}}
    (REPORT/'artifact_verification.json').write_text(json.dumps(v, ensure_ascii=False, indent=2)+'\n'); print(json.dumps(v, indent=2))

if __name__ == '__main__':
    parser = argparse.ArgumentParser(); parser.add_argument('--verify', action='store_true'); args = parser.parse_args()
    verify() if args.verify else build()
