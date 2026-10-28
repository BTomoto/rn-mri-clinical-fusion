"""Reconcile author-hosted clinical versions and audit follow-up availability.

Availability of a scan is deliberately not an event-free outcome. No model fits.
"""
from pathlib import Path
import datetime as dt
import json
import math
import numbers
from collections import Counter
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

ORIGIN=dt.datetime(1900,1,1)
SCHEMAS={
 'figshare':dict(rn='AG',srs='W',death='BC',death_occurred='BB',death_cause='BD',surgery='AH',surgery_day='AI',
                 tnm='K',stage='L',mri_start='BE',wbrt=['M','Q'],srs_groups=['U','Y','AC']),
 'molab':dict(rn='AR',srs='AH',death='BU',death_occurred='BT',death_cause='BV',surgery='AS',surgery_day='AT',
             tnm='V',stage='W',mri_start='BW',wbrt=['X','AB'],srs_groups=['AF','AJ','AN'])}


def clean(x):
    if x is None:return None
    if isinstance(x,str):return x.strip() or None
    if isinstance(x,dt.datetime):return x.isoformat()
    return x


def number(x):
    if isinstance(x,numbers.Real) and not isinstance(x,bool) and math.isfinite(x):return float(x)
    return None


def read_clinical(path,schema_name):
    schema=SCHEMAS[schema_name]
    wb=load_workbook(path,data_only=True,read_only=True)
    s=wb.active
    # Fail on layout changes rather than applying an old positional parser.
    if not (s[f'{schema["rn"]}14'].value=='NECROSIS' and
            s[f'{schema["rn"]}15'].value=='Time' and
            s[f'{schema["srs"]}15'].value=='Begin' and
            s[f'{schema["death"]}14'].value=='Time'):
        raise ValueError('Clinical header layout changed; review the schema before parsing')
    records=[];treatments=[];scans=[]
    for i,row in enumerate(s.iter_rows(values_only=True),1):
        if i<16:continue
        try:pid=str(int(row[0]));lid=int(row[4])
        except (TypeError,ValueError):continue
        key=f'{pid}_L{lid}'
        def raw(c):return clean(row[column_index_from_string(c)-1])
        def num(c):return number(row[column_index_from_string(c)-1])
        r=dict(patient_id=pid,patient_id_source=str(row[0]),lesion_id=lid,lesion_key=key,
            source_row_excel=i,source_version=schema_name,
            age_at_first_mri_years=raw('B'),sex_code=raw('C'),gpa_availability_raw=raw('D'),
            primary_tumor_raw=raw('F'),primary_subtype_raw=raw('G'),
            egfr_raw=raw('H'),alk_raw=raw('I'),ros1_raw=raw('J'),
            tn_stage_raw=raw(schema['tnm']),tumor_stage_raw=raw(schema['stage']),
            rn_time_day=num(schema['rn']),rn_raw=raw(schema['rn']),
            rn_source_cell=f'Hoja1!{schema["rn"]}{i}',
            first_srs_start_day=num(schema['srs']),death_day=num(schema['death']),
            death_recorded=raw(schema['death_occurred']),death_cause_raw=raw(schema['death_cause']),
            surgery_type_raw=raw(schema['surgery']),surgery_day_raw=raw(schema['surgery_day']),
            rn_status='recorded_event_time' if num(schema['rn']) is not None else 'unknown',
            verified_non_event=0,outcome_adjudicated=0,final_eligible=0)
        if schema_name=='molab':
            for name,col in [('kras','K'),('ttf1','L'),('synaptophysin','M'),('pdl1','N'),('brca1','O'),
                             ('er','P'),('pr','Q'),('her2','R'),('braf','S'),('nras','T'),('ckit','U')]:
                r[name+'_raw']=raw(col)
            r['comments_raw']=raw('BS')
        for typ,starts in [('WBRT',schema['wbrt']),('SRS',schema['srs_groups'])]:
            for n,start in enumerate(starts,1):
                j=column_index_from_string(start)-1
                vals=[clean(x) for x in row[j:j+4]]
                for field,value in zip(['dose_per_fraction_gy','fractions','start_day','end_day'],vals):
                    r[f'{typ.lower()}{n}_{field}']=value
                if any(x is not None for x in vals):
                    treatments.append(dict(lesion_key=key,patient_id=pid,lesion_id=lid,treatment_type=typ,treatment_number=n,
                        dose_per_fraction_gy=vals[0],fractions=vals[1],start_day=vals[2],end_day=vals[3]))
        for j,value in enumerate(row[column_index_from_string(schema['mri_start'])-1:]):
            if isinstance(value,dt.datetime):
                scans.append(dict(lesion_key=key,patient_id=pid,source_scan_number=j+1,
                    date_code=value.strftime('%Y%m%d'),study_day=(value-ORIGIN).days,
                    source_version=schema_name,source_row_excel=i))
        records.append(r)
    wb.close()
    df=pd.DataFrame(records)
    if df.lesion_key.duplicated().any():raise ValueError('Duplicate patient-lesion key')
    return df,pd.DataFrame(treatments),pd.DataFrame(scans)


def compare_versions(old,new):
    j=old.merge(new,on='lesion_key',how='outer',suffixes=('_old','_current'),indicator=True,validate='one_to_one')
    fields=[c for c in old if c in new and c not in ['lesion_key','patient_id_source','source_version','rn_source_cell','rn_status']]
    changes=[]
    for r in j.to_dict('records'):
        for field in fields:
            a,b=clean(r.get(field+'_old')),clean(r.get(field+'_current'))
            a=None if isinstance(a,float) and pd.isna(a) else a
            b=None if isinstance(b,float) and pd.isna(b) else b
            if a==b:continue
            changes.append(dict(lesion_key=r['lesion_key'],field=field,figshare_value=a,molab_current_value=b,
                old_source_row=r.get('source_row_excel_old'),current_source_row=r.get('source_row_excel_current'),
                action='retain_both_versions_use_author_hosted_current_in_new_audit'))
    return pd.DataFrame(changes,columns=['lesion_key','field','figshare_value','molab_current_value',
        'old_source_row','current_source_row','action'])


def followup_availability(current,scans,studies,masks):
    masks=masks.copy()
    masks['study_day']=masks.study_date_code.astype(str).map(lambda x:(dt.datetime.strptime(x,'%Y%m%d')-ORIGIN).days)
    rows=[];horizons=[]
    for r in current.to_dict('records'):
        key,pid,srs,rn,death=r['lesion_key'],r['patient_id'],r['first_srs_start_day'],r['rn_time_day'],r['death_day']
        clinical=scans[scans.lesion_key==key].study_day.drop_duplicates().sort_values()
        directory=studies[studies.patient_id==pid].study_day.drop_duplicates().sort_values()
        lesion_masks=masks[masks.lesion_key==key]
        if pd.isna(rn):timing='unknown'
        elif pd.isna(srs):timing='recorded_rn_no_index_srs'
        elif rn<=srs:timing='rn_at_or_before_index_srs'
        elif pd.notna(death) and rn>death:timing='recorded_rn_after_death'
        else:timing='recorded_rn_after_index_srs'
        valid_index=pd.notna(srs) and (pd.isna(death) or srs<=death)
        # Strictly before SRS. These remain image-availability candidates.
        before=directory[(directory>=srs-30)&(directory<srs)] if valid_index and timing not in ['rn_at_or_before_index_srs','recorded_rn_after_death'] else directory.iloc[:0]
        base_masks=lesion_masks[(lesion_masks.study_day>=srs-30)&(lesion_masks.study_day<srs)] if len(before) else lesion_masks.iloc[:0]
        actual=base_masks.sort_values('study_day').iloc[-1] if len(base_masks) else None
        q=dict(lesion_key=key,patient_id=pid,lesion_id=r['lesion_id'],rn_time_day=rn,rn_status=r['rn_status'],
            first_srs_start_day=srs,death_day=death,rn_timing_class=timing,
            source_mri_dates=len(clinical),directory_mri_dates=len(directory),
            source_last_mri_day=int(clinical.max()) if len(clinical) else None,
            source_mri_dates_after_death=int((clinical>death).sum()) if pd.notna(death) else None,
            source_post_srs_mri_dates=int(((clinical>srs)&(clinical<=death)).sum()) if valid_index else None,
            directory_post_srs_mri_dates=int(((directory>srs)&(directory<=death)).sum()) if valid_index else None,
            baseline_directory_candidate=int(len(before)>0),
            baseline_directory_day=int(before.max()) if len(before) else None,
            baseline_mask_available=int(actual is not None),
            baseline_mask_date_code=str(actual.study_date_code) if actual is not None else None,
            baseline_mask_file=str(actual.mask_file) if actual is not None else None,
            baseline_mask_day=int(actual.study_day) if actual is not None else None,
            verified_event_free_followup_day=None,verified_non_event=0,outcome_adjudicated=0,final_eligible=0)
        rows.append(q)
        if not valid_index:continue
        for h in [180,365]:
            horizons.append(dict(lesion_key=key,patient_id=pid,horizon_days=h,rn_timing_class=timing,
                rn_time_day=rn,first_srs_start_day=srs,death_day=death,
                recorded_rn_time_within_horizon=int(timing=='recorded_rn_after_index_srs' and rn<=srs+h),
                death_recorded_by_horizon=int(pd.notna(death) and death<=srs+h),
                source_scan_on_or_after_horizon_before_death=int(((clinical>=srs+h)&(clinical<=death)).any()),
                directory_scan_on_or_after_horizon_before_death=int(((directory>=srs+h)&(directory<=death)).any()),
                lesion_mask_on_or_after_horizon_before_death=int(((lesion_masks.study_day>=srs+h)&(lesion_masks.study_day<=death)).any()),
                verified_non_event=0,verified_event_free_followup_day=None,
                interpretation='scan_availability_only_no_event_free_status_inferred'))
    return pd.DataFrame(rows),pd.DataFrame(horizons)


def run(root):
    root=Path(root)
    old,_,_=read_clinical(root/'data/raw/openbtai/OpenBTAI_METS_ClinicalData.xlsx','figshare')
    current,treatments,scans=read_clinical(root/'data/raw/openbtai_molab_current/OpenBTAI_METS_ClinicalData_Nov2023.xlsx','molab')
    sept,_,_=read_clinical(root/'data/raw/openbtai_molab_current/OpenBTAI_METS_ClinicalData_Sept2023.xlsx','molab')
    changes=compare_versions(old,current)
    sept_changes=compare_versions(sept,current)
    studies=pd.read_csv(root/'data/processed/image_studies.csv',dtype={'patient_id':str})
    masks=pd.read_csv(root/'data/processed/mask_lesion_audit.csv',dtype={'patient_id':str,'study_date_code':str})
    available,horizons=followup_availability(current,scans,studies,masks)
    prior=Counter()
    for line in (root/'data/metadata/downloaded_imaging.jsonl').read_text().splitlines():
        entry=json.loads(line);prior[(entry['patient_id'],entry['study_date_code'])]+=1
    available['baseline_study_date_code']=available.baseline_directory_day.map(
        lambda x:(ORIGIN+dt.timedelta(days=int(x))).strftime('%Y%m%d') if pd.notna(x) else None)
    available['previous_archive_dicom_files']=available.apply(
        lambda r:prior[(r.patient_id,r.baseline_study_date_code)],axis=1)
    new_studies=available[(available.baseline_directory_candidate==1)&(available.previous_archive_dicom_files==0)][['patient_id','baseline_study_date_code']].drop_duplicates()
    radio=pd.read_csv(root/'data/processed/radiomics_existing_total.csv',dtype={'Patient':str,'Timepoint':str})
    radio['lesion_key']=radio.Patient+'_L'+radio.Lesion.astype(int).astype(str)
    sel=available[available.baseline_mask_available==1]
    features=sel.merge(radio,left_on=['lesion_key','baseline_mask_date_code','baseline_mask_file'],
        right_on=['lesion_key','Timepoint','Mask'],validate='one_to_one')
    features['use_for_model']=0
    features['clinical_source_version']='molab_Nov2023_author_hosted'
    features['radiomics_source_version']='figshare_22277638_v1'
    horizons['baseline_mask_candidate']=horizons.lesion_key.isin(features.lesion_key).astype(int)
    frames={'clinical_current':current,'treatments_current':treatments,'clinical_mri_dates_current':scans,
            'clinical_version_changes':changes,'clinical_sept_nov_changes':sept_changes,
            'followup_availability_current':available,'horizon_availability_current':horizons,
            'baseline_features_current':features}
    for name,df in frames.items():df.to_csv(root/f'data/processed/{name}.csv',index=False)
    horizon_columns=['recorded_rn_time_within_horizon','death_recorded_by_horizon',
        'source_scan_on_or_after_horizon_before_death','directory_scan_on_or_after_horizon_before_death',
        'lesion_mask_on_or_after_horizon_before_death']
    def horizon_summary(df):
        g=df.groupby('horizon_days')[horizon_columns].sum()
        g['lesion_records']=df.groupby('horizon_days').size()
        return g.reset_index().to_dict('records')
    summary=dict(audit_date='2026-10-28',current_clinical_source='author-hosted file named Nov2023',
        patients=int(current.patient_id.nunique()),lesions=len(current),
        original_rn_times=int(old.rn_time_day.notna().sum()),current_rn_times=int(current.rn_time_day.notna().sum()),
        current_rn_patients=int(current.loc[current.rn_time_day.notna(),'patient_id'].nunique()),
        current_unknown_rn=int(current.rn_time_day.isna().sum()),
        added_rn_times=int(((changes.field=='rn_time_day')&changes.figshare_value.isna()&changes.molab_current_value.notna()).sum()),
        changed_existing_rn_times=int(((changes.field=='rn_time_day')&changes.figshare_value.notna()&changes.molab_current_value.notna()).sum()),
        original_srs_records=int(old.first_srs_start_day.notna().sum()),current_srs_records=int(current.first_srs_start_day.notna().sum()),
        current_rn_timing_classes=available.rn_timing_class.value_counts().to_dict(),
        baseline_directory_candidates=int(available.baseline_directory_candidate.sum()),
        baseline_mask_candidates=len(features),baseline_mask_patients=int(features.patient_id.nunique()),
        baseline_mask_recorded_rn_times=int(features.rn_time_day.notna().sum()),
        new_baseline_studies_outside_previous_archive=len(new_studies),
        verified_non_events=0,published39_vs_current35='unresolved_4_lesion_difference',
        modeling_status='blocked_no_confirmed_non_events_or_lesion_level_reference_standard',models_fitted=False,
        sept_nov_shared_field_differences=len(sept_changes),
        horizon_availability_all_valid_srs=horizon_summary(horizons),
        horizon_availability_baseline_mask_cohort=horizon_summary(horizons[horizons.lesion_key.isin(features.lesion_key)]))
    (root/'reports/endpoint_followup_summary.json').write_text(json.dumps(summary,indent=2))
    return summary
