"""Outcome/timing and existing mask audit. Never infer non-events from blanks."""
from pathlib import Path
import datetime as dt
import hashlib
import json
import zipfile
import numpy as np
import pandas as pd
from openpyxl import load_workbook

ORIGIN = dt.datetime(1900, 1, 1)


def outcome_class(rn, srs):
    if pd.isna(rn):
        return 'unknown'
    if pd.isna(srs):
        return 'recorded_rn_no_index_srs'
    return 'recorded_rn_after_index_srs' if rn > srs else 'rn_at_or_before_index_srs'


def audit_outcomes(root):
    root = Path(root)
    lesions = pd.read_csv(root/'data/processed/lesions.csv', dtype={'patient_id':str})
    baselines = pd.read_csv(root/'data/processed/baseline_matches.csv',dtype={'patient_id':str})
    morph = pd.read_csv(root/'data/processed/morphology_existing.csv',dtype={'patient_id':str})
    studies = pd.read_csv(root/'data/processed/image_studies.csv',dtype={'patient_id':str})
    mask_path=root/'data/processed/mask_lesion_audit.csv'
    mask_table=pd.read_csv(mask_path,dtype={'patient_id':str,'study_date_code':str}) if mask_path.exists() else pd.DataFrame()
    if len(mask_table):
        mask_table['study_day']=mask_table.study_date_code.map(lambda x:(dt.datetime.strptime(x,'%Y%m%d')-ORIGIN).days)
    wb = load_workbook(root/'data/raw/openbtai/OpenBTAI_METS_ClinicalData.xlsx', data_only=True, read_only=True)
    source_rows = {i: row for i,row in enumerate(wb.active.iter_rows(values_only=True),1)
                   if i in set(lesions.source_row_excel)}
    baseline_map = baselines.set_index('lesion_key').to_dict('index')
    records=[]
    for r in lesions.to_dict('records'):
        key, pid, srs, rn = r['lesion_key'],r['patient_id'],r['first_srs_start_day'],r['rn_time_day']
        row = source_rows[r['source_row_excel']]
        dates = [x for x in row[56:] if isinstance(x,dt.datetime)]
        days = sorted({(x-ORIGIN).days for x in dates})
        matching = baseline_map.get(key,{})
        m = morph[morph.lesion_key == key]
        strict_m = m[(m.study_day >= srs-30)&(m.study_day < srs)] if pd.notna(srs) else m.iloc[:0]
        inclusive_m = m[(m.study_day >= srs-30)&(m.study_day <= srs)] if pd.notna(srs) else m.iloc[:0]
        actual_masks=mask_table[(mask_table.lesion_key==key)&(mask_table.study_day>=srs-30)&(mask_table.study_day<srs)] if len(mask_table) and pd.notna(srs) else mask_table.iloc[:0]
        patient_studies = studies[studies.patient_id == pid]
        out = {k:r[k] for k in ['lesion_key','patient_id','lesion_id','source_row_excel','first_srs_start_day','first_srs_end_day','rn_time_day','death_day']}
        out.update(rn_source_cell=f'Hoja1!AG{r["source_row_excel"]}',
            rn_evidence_class=outcome_class(rn,srs),
            rn_from_index_srs_days=rn-srs if pd.notna(rn) and pd.notna(srs) else None,
            rn_after_death=int(pd.notna(rn) and pd.notna(r['death_day']) and rn>r['death_day']),
            index_srs_after_death=int(pd.notna(srs) and srs>r['death_day']),
            source_followup_dates=len(days),
            source_last_mri_day=max(days) if days else None,
            source_mri_dates_after_death=sum(x>r['death_day'] for x in days),
            directory_last_mri_day=int(patient_studies.study_day.max()) if len(patient_studies) else None,
            baseline_directory_candidate=int(bool(matching)),
            baseline_date_code=matching.get('baseline_date_code'),
            baseline_day=matching.get('baseline_day'),
            baseline_same_day_as_srs=matching.get('baseline_same_day_as_srs',0),
            baseline_strictly_pre_srs=int(bool(matching) and matching['baseline_same_day_as_srs']==0),
            existing_morphology_rows=len(m),
            baseline_morphology_inclusive_rows=len(inclusive_m),
            baseline_morphology_strict_rows=len(strict_m),
            baseline_actual_mask_rows=len(actual_masks),
            baseline_actual_mask_available=int(len(actual_masks)>0),
            baseline_actual_mask_date_code=actual_masks.sort_values('study_day').iloc[-1].study_date_code if len(actual_masks) else None,
            early_mri_day=matching.get('early_day'),
            landmark90_directory_candidate=int(matching.get('dynamic_day90_time_candidate',0)),
            # These are not silently filled by the existence of later scans or death.
            verified_event_free_followup_day=None,
            verified_non_event=0,
            outcome_adjudicated=0,final_eligible=0)
        if len(inclusive_m):
            first = inclusive_m.sort_values('study_day').iloc[-1]
            out['baseline_mask_date_code'] = str(int(first.study_date_code))
        else:
            out['baseline_mask_date_code'] = None
        reasons=[]
        if out['rn_evidence_class']=='unknown':reasons.append('need_explicit_outcome_and_surveillance')
        else:reasons.append('need_lesion_level_reference_standard')
        if out['rn_evidence_class']=='rn_at_or_before_index_srs':reasons.append('resolve_treatment_chronology')
        if out['index_srs_after_death'] or out['rn_after_death']:reasons.append('resolve_event_after_death')
        if pd.isna(srs):reasons.append('no_numeric_index_srs')
        elif not matching:reasons.append('no_directory_baseline_in_30d')
        elif matching['baseline_same_day_as_srs']:reasons.append('same_day_order_unknown')
        if not len(actual_masks):reasons.append('baseline_target_segmentation_not_found_in_masks')
        reasons.append('target_visibility_and_full_image_qc_pending')
        out['remaining_requirements']=';'.join(reasons)
        records.append(out)
    wb.close()
    result=pd.DataFrame(records)
    result.to_csv(root/'data/processed/outcome_timing_audit.csv',index=False)
    counts=dict(clinical_lesions=len(result),patients=int(result.patient_id.nunique()),
        recorded_rn_times=int(result.rn_time_day.notna().sum()),unknown_rn=int(result.rn_time_day.isna().sum()),
        recorded_rn_after_index_srs=int((result.rn_evidence_class=='recorded_rn_after_index_srs').sum()),
        rn_at_or_before_index_srs=int((result.rn_evidence_class=='rn_at_or_before_index_srs').sum()),
        recorded_rn_no_index_srs=int((result.rn_evidence_class=='recorded_rn_no_index_srs').sum()),
        rn_after_death=int(result.rn_after_death.sum()),index_srs_after_death=int(result.index_srs_after_death.sum()),
        baseline_directory_candidates=int(result.baseline_directory_candidate.sum()),
        strictly_pre_srs_directory_candidates=int(result.baseline_strictly_pre_srs.sum()),
        same_day_directory_candidates=int(result.baseline_same_day_as_srs.sum()),
        baseline_existing_morphology_lesions=int((result.baseline_morphology_inclusive_rows>0).sum()),
        baseline_existing_morphology_rn_times=int(((result.baseline_morphology_inclusive_rows>0)&result.rn_time_day.notna()).sum()),
        baseline_actual_mask_lesions=int(result.baseline_actual_mask_available.sum()),
        baseline_actual_mask_patients=int(result.loc[result.baseline_actual_mask_available==1,'patient_id'].nunique()),
        baseline_actual_mask_rn_times=int(((result.baseline_actual_mask_available==1)&result.rn_time_day.notna()).sum()),
        strict_baseline_existing_morphology_lesions=int((result.baseline_morphology_strict_rows>0).sum()),
        landmark90_directory_candidates=int(result.landmark90_directory_candidate.sum()),
        verified_non_events=0,final_eligible=0)
    if len(mask_table):
        radio=pd.read_csv(root/'data/processed/radiomics_existing_total.csv',dtype={'Patient':str,'Timepoint':str})
        radio['lesion_key']=radio.Patient+'_L'+radio.Lesion.astype(int).astype(str)
        selected=result[result.baseline_actual_mask_available==1][['lesion_key','baseline_actual_mask_date_code','rn_evidence_class','rn_time_day','first_srs_start_day','death_day']]
        features=selected.merge(radio,left_on=['lesion_key','baseline_actual_mask_date_code'],right_on=['lesion_key','Timepoint'],validate='one_to_one')
        features['baseline_total_volume_cm3']=features.original_shape_VoxelVolume/1000
        features['candidate_only']=1
        features['use_for_model']=0
        features.to_csv(root/'data/processed/baseline_existing_features.csv',index=False)
    return result,counts


def mask_labels_to_lesion(label):
    """Release encoding inferred from volume reconciliation, not RN status."""
    label=int(label)
    if label <= 0 or label%10 not in [1,2]:
        raise ValueError(f'Unexpected mask label {label}')
    return label//10,label%10


def audit_masks(root):
    import nibabel as nib
    root=Path(root)
    morph=pd.read_csv(root/'data/processed/morphology_existing.csv',dtype={'patient_id':str,'study_date_code':str})
    radio=pd.read_csv(root/'data/processed/radiomics_existing_total.csv',dtype={'Patient':str,'Timepoint':str})
    images,components=[],[]
    with zipfile.ZipFile(root/'data/raw/openbtai/NIFTIs_Masks.zip') as z:
        for info in sorted(z.infolist(),key=lambda x:x.filename):
            if not info.filename.endswith('.nii') or '__MACOSX' in info.filename:continue
            name=Path(info.filename).name
            pid,date,*_=name.split('_');pid=str(int(pid))
            b=z.read(info)
            rec=dict(mask_file=name,patient_id=pid,study_date_code=date,sha256=hashlib.sha256(b).hexdigest(),
                zip_crc32=info.CRC,read_status='not_read',full_image_qc_complete=0)
            try:
                im=nib.Nifti1Image.from_bytes(b)
                arr=np.asanyarray(im.dataobj)
                # Background dominates a brain volume; count the sparse labels only.
                foreground=arr[arr!=0]
                labels,counts=np.unique(foreground,return_counts=True)
                voxel=float(abs(np.linalg.det(im.affine[:3,:3])))
                rec.update(read_status='readable',shape='x'.join(map(str,arr.shape)),
                    voxel_mm3=voxel,finite_affine=int(np.isfinite(im.affine).all()),
                    nonzero_labels=','.join(str(int(v)) for v in labels if v),
                    positive_voxels=int(np.count_nonzero(arr)),
                    spacing_mm=','.join(str(float(v)) for v in im.header.get_zooms()[:3]))
                per={}
                for label,count in zip(labels,counts):
                    if not label:continue
                    lid,comp=mask_labels_to_lesion(label)
                    per.setdefault(lid,{1:0,2:0})[comp]=int(count)
                rec['inferred_lesions']=len(per)
                for lid,c in per.items():
                    total=(c[1]+c[2])*voxel
                    m=morph[(morph.patient_id==pid)&(morph.study_date_code==date)&(morph.lesion_id==lid)]
                    rr=radio[(radio.Patient==pid)&(radio.Timepoint==date)&(radio.Lesion==lid)&(radio.Mask==name)]
                    # Exact file/lesion/date join; never choose an arbitrary duplicate.
                    mv=float(m.iloc[0].TOTALVOLUME) if len(m)==1 else None
                    rv=float(rr.iloc[0].original_shape_VoxelVolume) if len(rr)==1 else None
                    components.append(dict(patient_id=pid,study_date_code=date,lesion_id=lid,lesion_key=f'{pid}_L{lid}',
                        mask_file=name,component1_voxels=c[1],component2_voxels=c[2],voxel_mm3=voxel,
                        mask_total_mm3=total,mask_total_cm3=total/1000,
                        morphology_matches=len(m),morphology_total_cm3=mv,
                        morphology_abs_difference_cm3=abs(total/1000-mv) if mv is not None else None,
                        radiomics_matches=len(rr),radiomics_voxel_mm3=rv,
                        radiomics_abs_difference_mm3=abs(total-rv) if rv is not None else None,
                        radiomics_relative_difference=abs(total-rv)/abs(rv) if rv else None,
                        morphology_relative_difference=abs(total/1000-mv)/abs(mv) if mv else None,
                        encoding_status='inferred_10x_lesion_plus_component_verify_with_source',
                        rn_status_from_mask='not_encoded',anatomical_qc_complete=0))
            except Exception as exc:
                rec['read_status']=f'{type(exc).__name__}: {exc}'
            images.append(rec)
    files=pd.DataFrame(images);lesions=pd.DataFrame(components)
    files.to_csv(root/'data/processed/mask_file_audit.csv',index=False)
    lesions.to_csv(root/'data/processed/mask_lesion_audit.csv',index=False)
    counts=dict(mask_files=len(files),readable_mask_files=int((files.read_status=='readable').sum()),
        mask_lesion_time_records=len(lesions),mask_unique_lesions=int(lesions.lesion_key.nunique()),
        morphology_unique_join_records=int((lesions.morphology_matches==1).sum()),
        morphology_volume_agrees_1e_5_cm3=int((lesions.morphology_abs_difference_cm3<=1e-5).sum()),
        radiomics_unique_join_records=int((lesions.radiomics_matches==1).sum()),
        radiomics_volume_agrees_0_01_mm3=int((lesions.radiomics_abs_difference_mm3<=0.01).sum()),
        radiomics_volume_agrees_relative_0_1_percent=int((lesions.radiomics_relative_difference<=0.001).sum()),
        morphology_volume_agrees_relative_0_1_percent=int((lesions.morphology_relative_difference<=0.001).sum()),
        morphology_year_code_2000_or_later=int((morph.study_date_code.str[:4].astype(int)>=2000).sum()))
    return files,lesions,counts


def run_audit(root):
    root=Path(root)
    files,mask_lesions,mask_counts=audit_masks(root)
    outcomes,counts=audit_outcomes(root)
    summary=dict(audit_date='2026-10-28',outcomes=counts,masks=mask_counts,
        modeling_status='not_ready_unresolved_outcome_labels_and_baseline_segmentation',
        source_rn_blanks_treated_as_negative=False,
        published39_vs_release32='unresolved_no_lesion_level_crosswalk_found',
        dicom_scope='two_archive_order_samples_per_selected_directory_when_available',
        models_fitted=False)
    p=root/'data/processed/dicom_sample_audit.csv'
    if p.exists():
        ds=pd.read_csv(p)
        summary['dicom_samples']=dict(directories=len(ds),sampled_files=int(ds.sampled_files.sum()),
            readable_headers=int((ds.sample_read_status=='readable_headers').sum()),
            directories_with_mixed_sample_series_uid=int((ds.sample_same_series_uid!=1).sum()),
            directories_with_differing_sample_geometry=int((ds.sample_same_geometry!=1).sum()),
            directories_matching_original_sample_sha256=int(ds.sample_matches_original_download_sha256.sum()),
            full_series_qc_complete=False)
    (root/'reports/readiness_audit.json').write_text(json.dumps(summary,indent=2))
    return summary
