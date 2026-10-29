"""Prespecified age/sex vs four-variable clinical comparison after dual-ID quarantine."""
from pathlib import Path
import importlib.metadata
import json
import warnings
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from sklearn.compose import ColumnTransformer
from sklearn.exceptions import ConvergenceWarning
from sklearn.impute import SimpleImputer
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import roc_auc_score, average_precision_score, brier_score_loss
from sklearn.model_selection import StratifiedGroupKFold
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import StandardScaler
from .ucsd_clinical_baseline import make_cohort, sha256, scores, SOURCE_HASHES
from .ucsd_feature_audit import conservative_groups, INITIAL

SEED=20261029
EXCLUDED_IDS=['UCSD-PTGBM-0036_01','UCSD-PTGBM-0036_02','UCSD-PTGBM-0038_01','UCSD-PTGBM-0038_02','UCSD-PTGBM-0038_03']
FEATURES={'M0':['age_years','sex_male'],
          'M1':['age_years','log_initial_event_days','sex_male','other_treatment_prior','initial_event_missing']}


def prepare_cohort(root):
    root=Path(root);prior=make_cohort(root)
    w=load_workbook(root/'data/raw/reference/ucsd_ptgbm_clinical_v3.xlsx',read_only=True,data_only=True)
    rows=list(w['Cinical info'].iter_rows(values_only=True));w.close()
    raw=pd.DataFrame([r[:36] for r in rows[1:]],columns=[str(v).strip() for v in rows[0][:36]])
    d=prior[['ID','patient_id','tumor_label','label_source','age_years','sex_male','primary_diagnosis_raw']].merge(
        raw[['ID','BraTS ID',INITIAL,'Other treatment prior to scan']],on='ID',validate='one_to_one')
    d['brats_patient_id']=d['BraTS ID'].str.rsplit('-',n=1).str[0]
    d['validation_group']=conservative_groups(d)
    cp=d.groupby('patient_id').brats_patient_id.nunique();bp=d.groupby('brats_patient_id').patient_id.nunique()
    ambiguous=d.loc[d.patient_id.isin(cp[cp.gt(1)].index)|d.brats_patient_id.isin(bp[bp.gt(1)].index),'validation_group']
    flag=d.validation_group.isin(ambiguous)
    if sorted(d.loc[flag,'ID'])!=EXCLUDED_IDS:raise ValueError('Identity conflict set changed; review protocol')
    excluded=d.loc[flag,['ID','patient_id','BraTS ID','brats_patient_id','tumor_label']].copy()
    excluded['reason']='dual_identifier_mapping_conflict; not outcome/performance based'
    d=d.loc[~flag].copy().sort_values('ID').reset_index(drop=True)
    offset=pd.to_numeric(d[INITIAL],errors='coerce')
    d['initial_event_offset_raw']=d[INITIAL]
    d['days_since_initial_event']=-offset.where(offset.lt(0))
    d['initial_event_missing']=d.days_since_initial_event.isna().astype(int)
    d['log_initial_event_days']=np.log1p(d.days_since_initial_event)
    d['other_treatment_prior']=d['Other treatment prior to scan'].map({'Yes':1,'No':0})
    if d[['age_years','sex_male','other_treatment_prior']].isna().any().any():raise ValueError('Unexpected missing core covariate')
    if (len(d),d.patient_id.nunique(),d.tumor_label.sum())!=(238,176,187):raise ValueError('Unexpected post-quarantine cohort')
    if d.groupby('patient_id').brats_patient_id.nunique().ne(1).any() or d.groupby('brats_patient_id').patient_id.nunique().ne(1).any():raise ValueError('Remaining identity ambiguity')
    cols=['ID','patient_id','BraTS ID','brats_patient_id','validation_group','tumor_label','label_source','age_years','sex_male',
          'other_treatment_prior','initial_event_offset_raw','days_since_initial_event','log_initial_event_days','initial_event_missing','primary_diagnosis_raw']
    return d[cols],excluded


def make_model(name):
    if name not in FEATURES:raise ValueError(name)
    terms=[('age',StandardScaler(),['age_years'])]
    if name=='M1':
        terms.append(('initial_event',Pipeline([('imputer',SimpleImputer(strategy='median')),('scale',StandardScaler())]),['log_initial_event_days']))
    binary=['sex_male'] if name=='M0' else ['sex_male','other_treatment_prior','initial_event_missing']
    terms.append(('binary','passthrough',binary))
    return Pipeline([('preprocessing',ColumnTransformer(terms,remainder='drop')),
                     ('logistic',LogisticRegression(C=1,solver='lbfgs',max_iter=2000,random_state=SEED))])


def parameter_record(model,name):
    pre=model.named_steps['preprocessing'];lr=model.named_steps['logistic']
    r=dict(feature_order=FEATURES[name],intercept=float(lr.intercept_[0]),coefficients=lr.coef_[0].tolist(),
           age_mean=float(pre.named_transformers_['age'].mean_[0]),age_sd=float(pre.named_transformers_['age'].scale_[0]))
    if name=='M1':
        it=pre.named_transformers_['initial_event']
        r.update(log_interval_impute_median=float(it.named_steps['imputer'].statistics_[0]),
                 log_interval_mean_after_imputation=float(it.named_steps['scale'].mean_[0]),
                 log_interval_sd_after_imputation=float(it.named_steps['scale'].scale_[0]))
    return r


def grouped_oof(d):
    d=d.copy().reset_index(drop=True);y=d.tumor_label.to_numpy()
    p={k:np.full(len(d),np.nan) for k in ['M0','M1','prior']};foldid=np.zeros(len(d),dtype=int);coverage=np.zeros(len(d),dtype=int)
    records=[];cv=StratifiedGroupKFold(n_splits=5,shuffle=True,random_state=SEED)
    for k,(tr,te) in enumerate(cv.split(d,y,d.validation_group),1):
        train,test=d.iloc[tr],d.iloc[te]
        for namespace in ['patient_id','brats_patient_id','validation_group']:
            if set(train[namespace])&set(test[namespace]):raise ValueError('Patient leakage: '+namespace)
        if train.tumor_label.nunique()!=2 or test.tumor_label.nunique()!=2:raise ValueError('Fold lacks a class')
        record=dict(fold=k,train_IDs=train.ID.tolist(),test_IDs=test.ID.tolist(),
                    train_patients=sorted(train.patient_id.unique()),test_patients=sorted(test.patient_id.unique()),
                    train_brats_patients=sorted(train.brats_patient_id.unique()),test_brats_patients=sorted(test.brats_patient_id.unique()),
                    n_train=len(tr),n_test=len(te),n_test_patients=int(test.patient_id.nunique()),
                    test_tumor=int(test.tumor_label.sum()),test_treatment_change=int(test.tumor_label.eq(0).sum()),
                    train_log_interval_observed=int(train.log_initial_event_days.notna().sum()),models={})
        for name in ['M0','M1']:
            model=make_model(name)
            with warnings.catch_warnings():
                warnings.simplefilter('error',ConvergenceWarning);model.fit(train,y[tr])
            p[name][te]=model.predict_proba(test)[:,1]
            record['models'][name]=dict(parameters=parameter_record(model,name),metrics=scores(y[te],p[name][te]))
        p['prior'][te]=y[tr].mean();record['train_prevalence']=float(y[tr].mean())
        foldid[te]=k;coverage[te]+=1;records.append(record)
    if not np.all(coverage==1):raise ValueError('OOF coverage error')
    d['validation_fold']=foldid
    for name in p:
        if not np.isfinite(p[name]).all():raise ValueError('Invalid probabilities')
        d[name+'_oof_probability']=p[name]
        d[name+'_prediction_0_5']=(p[name]>=.5).astype(int)
    return d,records


def paired_intervals(d,n=2000):
    rng=np.random.default_rng(SEED+1)
    clusters=[np.asarray(v) for v in d.groupby('validation_group',sort=True).indices.values()]
    y=d.tumor_label.to_numpy();p={k:d[k+'_oof_probability'].to_numpy() for k in ['M0','M1']}
    keys=['roc_auc','ap_tumor','ap_treatment_change','brier'];draws=[]
    for _ in range(n):
        ix=np.concatenate([clusters[k] for k in rng.integers(0,len(clusters),len(clusters))]);yy=y[ix]
        if len(np.unique(yy))<2:continue
        vals=[]
        for name in p:
            pp=p[name][ix];vals.extend([roc_auc_score(yy,pp),average_precision_score(yy,pp),average_precision_score(1-yy,1-pp),brier_score_loss(yy,pp)])
        vals.extend([vals[4]-vals[0],vals[7]-vals[3]]);draws.append(vals)
    if len(draws)<.95*n:raise ValueError('Too many one-class bootstrap draws')
    qs=np.quantile(np.asarray(draws),[.025,.975],axis=0)
    intervals={name:{key:dict(lower=float(qs[0,j*4+i]),upper=float(qs[1,j*4+i])) for i,key in enumerate(keys)} for j,name in enumerate(['M0','M1'])}
    intervals['M1_minus_M0']={key:dict(lower=float(qs[0,8+i]),upper=float(qs[1,8+i])) for i,key in enumerate(['roc_auc','brier'])}
    delta=pd.DataFrame(np.asarray(draws)[:,8:],columns=['delta_auc_M1_minus_M0','delta_brier_M1_minus_M0'])
    return intervals,delta


def make_plots(d,output):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from sklearn.metrics import roc_curve,precision_recall_curve
    from sklearn.calibration import calibration_curve
    y=d.tumor_label.to_numpy();fig,axes=plt.subplots(1,3,figsize=(13,4),layout='constrained')
    calibration=[]
    for name,color,label in [('M0','#72818C','Age + sex'),('M1','#235F87','Expanded clinical')]:
        p=d[name+'_oof_probability'].to_numpy();fpr,tpr,_=roc_curve(y,p);prec,rec,_=precision_recall_curve(1-y,1-p)
        axes[0].plot(fpr,tpr,color=color,label=f'{label}: {roc_auc_score(y,p):.3f}')
        axes[1].plot(rec,prec,color=color,label=f'{label}: AP={average_precision_score(1-y,1-p):.3f}')
        observed,predicted=calibration_curve(y,p,n_bins=5,strategy='quantile')
        axes[2].plot(predicted,observed,'o-',color=color,label=label)
        calibration.extend([dict(model=name,mean_predicted=float(a),observed_proportion=float(b)) for a,b in zip(predicted,observed)])
    axes[0].plot([0,1],[0,1],'--',color='gray');axes[1].axhline((1-y).mean(),linestyle='--',color='gray',label='Treatment-change prevalence');axes[2].plot([0,1],[0,1],'--',color='gray')
    for ax in axes:ax.set(xlim=(0,1),ylim=(0,1));ax.spines[['top','right']].set_visible(False);ax.legend(fontsize=8,loc='lower right' if ax is not axes[1] else 'upper right')
    axes[0].set(title='Patient-grouped OOF ROC',xlabel='False positive rate',ylabel='Tumor sensitivity')
    axes[1].set(title='Treatment-change precision-recall',xlabel='Treatment-change recall',ylabel='Treatment-change precision')
    axes[2].set(title='Descriptive OOF calibration',xlabel='Predicted tumor probability',ylabel='Observed tumor proportion')
    fig.savefig(output/'clinical_revalidation_comparison.png',dpi=180);fig.savefig(output/'clinical_revalidation_comparison.pdf');plt.close(fig)
    return pd.DataFrame(calibration)


def run(root):
    root=Path(root);r=root/'reports/ucsd_clinical_revalidation';r.mkdir(parents=True,exist_ok=True)
    out=root/'outputs/ucsd-clinical-revalidation-20261029';out.mkdir(parents=True,exist_ok=True)
    cohort,excluded=prepare_cohort(root)
    cohort.to_csv(root/'data/processed/ucsd_clinical_revalidated_cohort.csv',index=False)
    excluded.to_csv(r/'excluded_identity_conflicts.csv',index=False)
    pred,folds=grouped_oof(cohort)
    pred.to_csv(root/'data/processed/ucsd_clinical_revalidated_oof.csv',index=False,float_format='%.12g')
    metrics={name:scores(pred.tumor_label,pred[name+'_oof_probability']) for name in ['M0','M1','prior']}
    intervals,draws=paired_intervals(pred);draws.to_csv(r/'paired_bootstrap_differences.csv',index=False,float_format='%.12g')
    complete=cohort[cohort.initial_event_missing.eq(0)].copy()
    secondary,secondary_folds=grouped_oof(complete)
    secondary.to_csv(r/'complete_interval_sensitivity_oof.csv',index=False,float_format='%.12g')
    sec_metrics={name:scores(secondary.tumor_label,secondary[name+'_oof_probability']) for name in ['M0','M1','prior']}
    final={name:parameter_record(make_model(name).fit(cohort,cohort.tumor_label),name) for name in ['M0','M1']}
    summary=dict(step=13,date='2026-10-29',timepoints=len(cohort),source_patients=int(cohort.patient_id.nunique()),
                 tumor_timepoints=int(cohort.tumor_label.sum()),treatment_change_timepoints=int(cohort.tumor_label.eq(0).sum()),
                 tumor_prevalence=float(cohort.tumor_label.mean()),excluded_timepoints=len(excluded),
                 input_features={'M0':['age_years','sex_male'],'M1':['age_years','sex_male','days_since_initial_event','other_treatment_prior']},
                 transformed_feature_order=FEATURES,missing_interval_timepoints=int(cohort.initial_event_missing.sum()),
                 metrics=metrics,patient_cluster_bootstrap_95ci=intervals,
                 paired_differences_M1_minus_M0={k:metrics['M1'][k]-metrics['M0'][k] for k in ['roc_auc','brier']},
                 bootstrap_valid=len(draws),bootstrap_models_refitted=False,
                 bootstrap_limit='Conditional on fixed OOF predictions; does not include all training/split variability',
                 secondary_complete_interval=dict(timepoints=len(complete),source_patients=int(complete.patient_id.nunique()),
                     tumor_timepoints=int(complete.tumor_label.sum()),treatment_change_timepoints=int(complete.tumor_label.eq(0).sum()),metrics=sec_metrics),
                 seed=SEED,n_splits=5,threshold=.5,class_weight=None,C=1,hyperparameter_search=False,
                 same_cases_and_folds_for_models=True,clinical_ID_patient_overlap=0,BraTS_ID_patient_overlap=0,
                 MRI_used=False,molecular_features_used=False,external_validation=False,final_full_cohort_parameters=final,
                 source_sha256=SOURCE_HASHES,protocol_sha256=sha256(root/'reports/ucsd_clinical_revalidation_protocol.md'),
                 analysis_code_sha256=sha256(__file__),
                 versions={k:importlib.metadata.version(k) for k in ['scikit-learn','pandas','numpy','scipy','matplotlib']},
                 next_step_requires_user_approval=True)
    (r/'metrics.json').write_text(json.dumps(summary,indent=2))
    (r/'fold_audit.json').write_text(json.dumps(folds,indent=2))
    (r/'complete_interval_fold_audit.json').write_text(json.dumps(secondary_folds,indent=2))
    fm=[]
    for f in folds:
        for name in ['M0','M1']:
            fm.append(dict(fold=f['fold'],model=name,n_train=f['n_train'],n_test=f['n_test'],n_test_patients=f['n_test_patients'],
                           test_tumor=f['test_tumor'],test_treatment_change=f['test_treatment_change'],**f['models'][name]['metrics']))
    pd.DataFrame(fm).to_csv(r/'fold_metrics.csv',index=False)
    make_plots(pred,out).to_csv(r/'calibration_bins.csv',index=False)
    return summary
