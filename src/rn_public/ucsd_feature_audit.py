"""Clinical availability, temporal validity and dual-ID audit; does not fit models."""
from pathlib import Path
import json
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from .ucsd_clinical_baseline import make_cohort, sha256

INITIAL = 'Days from Acquisition to Date of initial surgery, treatment or diagnosis'
LAST_SURGERY = 'Days from Acquisition to Date of last surgery prior to scan if different'
FIRST_RT = 'Days from Acquistion to Date of first radiation'
LAST_RT = 'Days from Acquistion to Date of last radiation prior to scan'
FIRST_CHEMO = 'Days from Acquisition to Date of 1st chemo start'
LAST_BEV = 'Days from Acquisition to Date of last Avastin treatment prior to scan'
FUTURE_FIELDS = ['Days from Acquisition to Date of last follow-up',
                 'Days from Acquisition to Date of death', 'Overall survival', 'Progression Free survival']
MISSING_STRINGS = {'', 'unknown', 'unknown or not reported', 'not reported', 'nan', 'none'}


def observed(s):
    return s.notna() & ~s.astype(str).str.strip().str.lower().isin(MISSING_STRINGS)


def conservative_groups(d):
    """Union both public identifier namespaces; no claim to adjudicate true identities."""
    parent = {}
    def find(x):
        parent.setdefault(x, x)
        if parent[x] != x:
            parent[x] = find(parent[x])
        return parent[x]
    for _, r in d.iterrows():
        a, b = find('clinical:'+r.patient_id), find('brats:'+r.brats_patient_id)
        parent[b] = a
    names = {}
    for p in sorted(d.patient_id.unique()):
        names.setdefault(find('clinical:'+p), []).append(p)
    return d.patient_id.map(lambda p: 'dualid:'+min(names[find('clinical:'+p)]))


def role(field):
    if field in ['ID', 'BraTS ID']:
        return 'identifier_only', '仅用于连接/患者分组，不能作预测变量'
    if field in FUTURE_FIELDS:
        return 'exclude_future_outcome', '扫描后随访/结局，禁止进入当前状态鉴别模型'
    if field in ["Patient's Age", 'Sex at birth', 'Other treatment prior to scan']:
        return 'candidate', '扫描时人口学/字段明确指定扫描前；其他治疗为宽泛二分类，异质性须说明'
    if field == INITIAL:
        return 'candidate_with_missingness', '有效非正偏移可取负号；混合初始手术/治疗/诊断事件，不能命名为距手术时间'
    if field in ['MGMT', 'IDH', '1p19q', 'ATRX']:
        return 'conditional_molecular', '有缺失，缺检测时间戳；患者内冲突须标记，先作为条件/敏感性候选，不跨扫描补值'
    if field in ['Surgery', 'Number of surgeries', 'Surgery extend', 'Radiation', 'Number of radiation courses',
                 '1st Chemo type', 'Avastin® (bevacizumab)', LAST_SURGERY, FIRST_RT, LAST_RT, FIRST_CHEMO, LAST_BEV,
                 'Days from Acquisition to Other treatment dates']:
        return 'conditional_treatment', '需核实扫描前可知性、条件性空值和日期一致性；未来/同日先后不明不作为已发生治疗'
    if field in ['Magnetic Field Strength', 'Manufacturer', "Manufacturer's Model Name"]:
        return 'acquisition_audit_only', '常量或采集批次信息，先作质控/偏倚审查，不作为默认临床输入'
    if field in ['Primary Diagnosis', 'WHO 2021 Diagnosis', 'Non WHO 2021 Diagnosis', 'Grade']:
        return 'cohort_definition_audit', '原诊断与分级保留；避免把回顾性分类或小样本分级自动加入预测'
    if field in ['Race', 'Ethnicity']:
        return 'not_primary_candidate', '未知和小类别较多，主模型暂不使用'
    raise ValueError('Unclassified clinical source field: '+field)


def run(root):
    root = Path(root)
    out = root/'reports/ucsd_feature_audit'; out.mkdir(parents=True, exist_ok=True)
    # This validates both original XLSX byte fingerprints and the prior label join.
    prior = make_cohort(root)
    w = load_workbook(root/'data/raw/reference/ucsd_ptgbm_clinical_v3.xlsx', read_only=True, data_only=True)
    rows = list(w['Cinical info'].iter_rows(values_only=True))
    dictionary = list(w['Data Dictionary'].iter_rows(min_row=2, max_row=37, values_only=True)); w.close()
    fields = [str(v).strip() for v in rows[0][:36]]
    raw = pd.DataFrame([r[:36] for r in rows[1:]], columns=fields)
    if len(fields) != 36 or len(dictionary) != 36:
        raise ValueError('Clinical field layout changed')
    raw.to_csv(root/'data/processed/ucsd_clinical_all_fields_raw.csv', index=False)
    d = raw.merge(prior[['ID', 'patient_id', 'tumor_label', 'label_source']], on='ID', validate='one_to_one')
    d['brats_patient_id'] = d['BraTS ID'].str.rsplit('-', n=1).str[0]
    d['validation_group_conservative'] = conservative_groups(d)
    issues = []
    def issue(ix, field, kind, explanation, action):
        for i in d.index[ix]:
            issues.append(dict(ID=d.at[i,'ID'], patient_id=d.at[i,'patient_id'], field=field,
                               raw_value=str(d.at[i,field]) if field in d else '', kind=kind,
                               explanation=explanation, action=action))
    cp = d.groupby('patient_id').brats_patient_id.nunique()
    bp = d.groupby('brats_patient_id').patient_id.nunique()
    bad_components = d.loc[d.patient_id.isin(cp[cp.gt(1)].index) | d.brats_patient_id.isin(bp[bp.gt(1)].index), 'validation_group_conservative']
    d['id_mapping_ambiguous'] = d.validation_group_conservative.isin(bad_components).astype(int)
    issue(d.id_mapping_ambiguous.eq(1), 'BraTS ID', 'cross_namespace_identity_conflict',
          '临床患者与BraTS患者不是一一对应；保守合并分组仅防泄漏，不解决真实身份/时间点连接',
          '暂停这5条MRI连接；未来可先排除两个源患者作主分析，或经证据核实后更正')
    d['candidate_age_years'] = pd.to_numeric(d["Patient's Age"])
    d['candidate_sex_male'] = d['Sex at birth'].map({'Male':1, 'Female':0})
    d['candidate_other_treatment_prior'] = d['Other treatment prior to scan'].map({'Yes':1, 'No':0})
    time_rows = []
    for field in [f for f in fields if f.startswith('Days')]:
        z = pd.to_numeric(d[field], errors='coerce')
        time_rows.append(dict(field=field, numeric_timepoints=int(z.notna().sum()), before_scan=int(z.lt(0).sum()),
                              same_day=int(z.eq(0).sum()), after_scan=int(z.gt(0).sum()),
                              minimum=None if z.isna().all() else float(z.min()), maximum=None if z.isna().all() else float(z.max())))
        if field not in FUTURE_FIELDS:
            issue(z.gt(0),field,'future_treatment_offset','正偏移表示事件晚于扫描，与扫描前预测输入不相容；未翻转符号或改写日期', '排除该事件信息；保留原值及冲突标记')
            issue(z.eq(0),field,'same_day_order_unknown','日期相同，无法确认事件在扫描前后', '严格扫描前分析中不认定已发生，可另作敏感性规则')
    initial = pd.to_numeric(d[INITIAL], errors='coerce')
    d['candidate_days_since_initial_event'] = -initial.where(initial.lt(0))
    d['initial_event_same_day'] = initial.eq(0).astype(int)
    first = pd.to_numeric(d[FIRST_RT], errors='coerce')
    last = pd.to_numeric(d[LAST_RT], errors='coerce')
    surgery = pd.to_numeric(d[LAST_SURGERY], errors='coerce')
    chemo = pd.to_numeric(d[FIRST_CHEMO], errors='coerce')
    bev = pd.to_numeric(d[LAST_BEV], errors='coerce')
    rt_order = first.notna() & last.notna() & first.gt(last)
    issue(rt_order, FIRST_RT, 'radiation_order_conflict', '首次放疗偏移晚于最后放疗偏移，不能自行交换/更正', '本条放疗间隔整体隔离')
    d['radiation_temporal_conflict'] = (first.gt(0)|last.ge(0)|rt_order).astype(int)
    good_last = last.lt(0) & ~d.radiation_temporal_conflict.astype(bool) & d.Radiation.eq('Yes')
    d['conditional_days_since_last_radiation'] = -last.where(good_last)
    d['radiation_evidence_state'] = np.select([
        d.radiation_temporal_conflict.eq(1), good_last,
        d.Radiation.eq('No') & first.isna() & last.isna()],
        ['temporal_conflict_or_same_day', 'dated_before_scan', 'source_No_without_dates'], default='source_Yes_timing_unverified')
    d['conditional_days_since_last_other_surgery'] = -surgery.where(surgery.lt(0))
    d['conditional_days_since_first_chemo'] = -chemo.where(chemo.lt(0))
    d['bevacizumab_evidence_state'] = np.select([
        d['Avastin® (bevacizumab)'].eq('Yes') & bev.lt(0),
        d['Avastin® (bevacizumab)'].eq('No'),
        d['Avastin® (bevacizumab)'].eq('Yes') & bev.ge(0)],
        ['dated_before_scan','source_No','future_or_same_day'], default='unknown_or_undated')
    d['conditional_bevacizumab_prior_strict'] = d.bevacizumab_evidence_state.map({'dated_before_scan':1, 'source_No':0})
    issue(d['Avastin® (bevacizumab)'].eq('N'), 'Avastin® (bevacizumab)', 'nonstandard_category', 'N不是字典定义的Yes/No，本步不自动等同No', '保留原文，清洗候选设未知')
    for field in ['MGMT', 'IDH', '1p19q', 'ATRX']:
        known = observed(d[field])
        norm = d[field].where(known).astype('string').str.strip()
        conflict = norm.groupby(d.patient_id).nunique().gt(1)
        flag = d.patient_id.isin(conflict[conflict].index)
        d[field.lower()+'_within_patient_conflict'] = flag.astype(int)
        d['conditional_'+field.lower()] = norm.mask(flag)
        issue(flag,field,'within_patient_molecular_conflict', '同一源患者不同扫描的已记录结果不一致；无标本日期说明', '不跨时间点补值；该患者此分子变量隔离为未知，原值保留')
    issue(d.Grade.lt(4),'Grade','lower_grade_in_source_collection', '源集合包含Grade 2/3，影响纯GBM/WHO分类的纳入定义', '本步保留源集合，后续预先指定组织学敏感性分析')
    issue(d['Primary Diagnosis'].str.contains(r'Grade\s*5',na=False),'Primary Diagnosis','invalid_diagnosis_grade_text', '源诊断文本含Grade 5；未改写成Grade 4', '保留异常标记，不将此文本作预测输入')
    issue(d.Surgery.eq('No') & d['Surgery extend'].isin(['GTR','STR']),'Surgery extend','surgery_extent_flag_conflict', 'Surgery=No却记录切除范围GTR/STR，字段含义/时点需澄清', '切除范围暂不作为主模型输入')
    a = []
    for j, field in enumerate(fields):
        known = observed(d[field]); r, note = role(field)
        vals = d.loc[known,field].astype(str).str.strip().value_counts()
        a.append(dict(field=field,dictionary_field=str(dictionary[j][0]).strip(),definition=dictionary[j][1],
                      role=r,observed_timepoints=int(known.sum()),missing_or_unknown=int((~known).sum()),
                      observed_fraction=float(known.mean()),observed_patients=int(d.loc[known,'patient_id'].nunique()),
                      observed_tumor_timepoints=int((known&d.tumor_label.eq(1)).sum()),
                      observed_change_timepoints=int((known&d.tumor_label.eq(0)).sum()),
                      unique_nonmissing_values=int(len(vals)),example_counts=json.dumps(vals.head(5).to_dict(),ensure_ascii=False),note=note))
    audit=pd.DataFrame(a);audit.to_csv(out/'clinical_field_audit.csv',index=False)
    pd.DataFrame(time_rows).to_csv(out/'temporal_field_audit.csv',index=False)
    issue_table=pd.DataFrame(issues);issue_table.to_csv(out/'case_issues.csv',index=False)
    base=pd.read_csv(root/'data/processed/ucsd_clinical_oof_predictions.csv')[['ID','validation_fold']]
    d=d.merge(base,on='ID',validate='one_to_one')
    cross = d.groupby('brats_patient_id').validation_fold.nunique()
    leakage=d[d.brats_patient_id.isin(cross[cross.gt(1)].index)]
    leakage[['ID','patient_id','BraTS ID','brats_patient_id','validation_fold']].to_csv(out/'prior_validation_identity_overlap.csv',index=False)
    identity_cols=['ID','patient_id','BraTS ID','brats_patient_id','validation_group_conservative','id_mapping_ambiguous','validation_fold']
    d[identity_cols].to_csv(out/'identity_mapping.csv',index=False)
    prepared_cols=identity_cols+['tumor_label','label_source']+[c for c in d if c.startswith(('candidate_','conditional_'))]+[
        'initial_event_same_day','radiation_temporal_conflict','radiation_evidence_state','bevacizumab_evidence_state',
        'mgmt_within_patient_conflict','idh_within_patient_conflict','1p19q_within_patient_conflict','atrx_within_patient_conflict']
    prepared=d[prepared_cols];prepared.to_csv(root/'data/processed/ucsd_clinical_feature_candidates.csv',index=False)
    combos={}
    for name,cols in {'age_sex':['candidate_age_years','candidate_sex_male'],
                      'age_sex_other_prior':['candidate_age_years','candidate_sex_male','candidate_other_treatment_prior'],
                      'plus_initial_event':['candidate_age_years','candidate_sex_male','candidate_other_treatment_prior','candidate_days_since_initial_event'],
                      'plus_MGMT_IDH_raw':['MGMT','IDH'],
                      'strict_last_radiation':['conditional_days_since_last_radiation']}.items():
        mask=d[cols].notna().all(axis=1)
        combos[name]=dict(timepoints=int(mask.sum()),patients=int(d.loc[mask,'patient_id'].nunique()),
                          tumor=int(d.loc[mask,'tumor_label'].sum()),treatment_change=int(d.loc[mask,'tumor_label'].eq(0).sum()))
    summary=dict(step=12,date='2026-10-29',clinical_fields=36,timepoints=243,source_patients=178,
                 source_sha256=sha256(root/'data/raw/reference/ucsd_ptgbm_clinical_v3.xlsx'),
                 clinical_patient_count=178,brats_patient_count=int(d.brats_patient_id.nunique()),
                 conservative_connected_groups=int(d.validation_group_conservative.nunique()),
                 conservative_groups_are_not_adjudicated_patient_count=True,
                 ambiguous_mapping_timepoints=int(d.id_mapping_ambiguous.sum()),
                 ambiguous_mapping_patients=sorted(d.loc[d.id_mapping_ambiguous.eq(1),'patient_id'].unique()),
                 prior_validation_shared_brats_patients=int(cross.gt(1).sum()),
                 prior_baseline_metrics_status='provisional_due_to_cross_namespace_identity_conflict; not_refitted_in_step12',
                 role_counts=audit.role.value_counts().to_dict(),molecular_observed={c:int(observed(d[c]).sum()) for c in ['MGMT','IDH','1p19q','ATRX']},
                 issue_counts=issue_table.kind.value_counts().to_dict(),complete_case_examples=combos,
                 issue_rows=len(issue_table),issues_not_necessarily_distinct_timepoints=True,
                 future_fields_excluded=FUTURE_FIELDS,models_fitted_this_step=False,
                 original_OpenBTAI_model_ready=False,next_step_requires_user_approval=True)
    (out/'clinical_summary.json').write_text(json.dumps(summary,indent=2))
    return summary
