"""Fixed, patient-grouped UCSD-PTGBM age/sex diagnostic baseline; no tuning."""
from pathlib import Path
import hashlib
import importlib.metadata
import json
import warnings

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from sklearn.compose import ColumnTransformer
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import (average_precision_score, brier_score_loss,
                             confusion_matrix, roc_auc_score)
from sklearn.model_selection import StratifiedGroupKFold
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import StandardScaler
from sklearn.exceptions import ConvergenceWarning

SEED = 20261029
FEATURES = ['age_years', 'sex_male']
SOURCE_HASHES = {
    'ucsd_ptgbm_clinical_v3.xlsx': 'bbf40bbdd289266924869052806de4d7e8fc9ee73286cdc0cc066672ef5bb51d',
    'ucsd_ptgbm_negative_cases.xlsx': '6d99db8ac1feac569e50c18ed214a1a761858fb8bfa614a29ec5f036946b96e4',
}


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def make_cohort(root):
    """Rebuild labels from pinned original bytes; do not reuse historical enable flags."""
    root = Path(root)
    for filename, expected in SOURCE_HASHES.items():
        if sha256(root / 'data/raw/reference' / filename) != expected:
            raise ValueError('Source fingerprint changed: ' + filename)
    w = load_workbook(root / 'data/raw/reference/ucsd_ptgbm_clinical_v3.xlsx', read_only=True, data_only=True)
    rows = list(w['Cinical info'].iter_rows(values_only=True)); w.close()
    c = pd.DataFrame([r[:36] for r in rows[1:]], columns=[str(v).strip() for v in rows[0][:36]])
    c['source_row_clinical'] = np.arange(2, len(c) + 2)
    w = load_workbook(root / 'data/raw/reference/ucsd_ptgbm_negative_cases.xlsx', read_only=True, data_only=True)
    rows = list(w['all_negative_cases_TCIA'].iter_rows(values_only=True)); w.close()
    neg = pd.DataFrame(rows[1:], columns=rows[0])
    neg['source_row_negative_list'] = np.arange(2, len(neg) + 2)
    if len(c) != 243 or len(neg) != 51 or c.ID.duplicated().any() or neg.ID.duplicated().any():
        raise ValueError('Unexpected totals or duplicate source IDs')
    if not c.ID.str.fullmatch(r'UCSD-PTGBM-\d{4}_\d{2}').all() or not neg.ID.isin(c.ID).all():
        raise ValueError('Unmatched/invalid source ID')
    allowed = {'NS', 'PsP (clinically confirmed)', 'PsP (bx proven)', 'Radiation necrosis'}
    if not neg.Dx.isin(allowed).all():
        raise ValueError('Unreviewed negative subtype')
    d = c.merge(neg.rename(columns={'Dx': 'negative_list_dx'}), on='ID', how='left', validate='one_to_one')
    d['patient_id'] = d.ID.str.rsplit('_', n=1).str[0]
    d['tumor_label'] = d.negative_list_dx.isna().astype(int)
    d['label_source'] = np.where(d.tumor_label.eq(0), 'explicit_negative_case_list',
                                'derived_from_published_243_equals_192_plus_51_partition')
    d['age_years'] = pd.to_numeric(d["Patient's Age"], errors='raise')
    d['sex_at_birth_raw'] = d['Sex at birth']
    d['sex_male'] = d.sex_at_birth_raw.map({'Female': 0, 'Male': 1})
    if not d.age_years.between(18, 100).all() or d[FEATURES].isna().any().any():
        raise ValueError('Invalid/missing age or sex; review rather than silently exclude')
    if d.patient_id.nunique() != 178 or int(d.tumor_label.sum()) != 192:
        raise ValueError('Unexpected patient/label partition')
    if d.groupby('patient_id').sex_male.nunique().gt(1).any():
        raise ValueError('Within-patient sex inconsistency')
    d['primary_diagnosis_raw'] = d['Primary Diagnosis']
    d['diagnosis_text_anomaly'] = d.primary_diagnosis_raw.str.contains(r'Grade\s*5', case=False, na=False).astype(int)
    d['analysis_question_approved'] = 1
    d['included_in_clinical_baseline'] = 1
    d['eligibility_reason'] = 'source_HGG_collection; complete_age_sex; broad_partition_label'
    # Independent cross-check against the frozen prior feasibility output; keep its historical flags intact.
    audit = pd.read_csv(root / 'data/processed/ucsd_ptgbm_label_feasibility.csv').set_index('ID')
    if not d.set_index('ID').tumor_label.sort_index().equals(audit.broad_tumor_label_candidate.sort_index().rename('tumor_label')):
        raise ValueError('Prior label audit differs')
    cols = ['ID', 'patient_id', 'tumor_label', 'label_source', 'negative_list_dx', 'age_years',
            'sex_at_birth_raw', 'sex_male', 'primary_diagnosis_raw', 'diagnosis_text_anomaly',
            'source_row_clinical', 'source_row_negative_list', 'analysis_question_approved',
            'included_in_clinical_baseline', 'eligibility_reason']
    return d[cols].sort_values('ID').reset_index(drop=True)


def make_model():
    preprocessing = ColumnTransformer([('age', StandardScaler(), ['age_years']),
                                       ('sex', 'passthrough', ['sex_male'])], remainder='drop')
    return Pipeline([('preprocessing', preprocessing),
                     ('logistic', LogisticRegression(C=1.0, solver='lbfgs', max_iter=2000, random_state=SEED))])


def scores(y, p):
    y, p = np.asarray(y), np.asarray(p)
    tn, fp, fn, tp = confusion_matrix(y, p >= .5, labels=[0, 1]).ravel()
    sens, spec = tp / (tp + fn), tn / (tn + fp)
    return dict(roc_auc=float(roc_auc_score(y, p)), ap_tumor=float(average_precision_score(y, p)),
                ap_treatment_change=float(average_precision_score(1-y, 1-p)),
                brier=float(brier_score_loss(y, p)), sensitivity_tumor=float(sens),
                specificity_treatment_change=float(spec), balanced_accuracy=float((sens+spec)/2),
                accuracy=float((tp+tn)/len(y)), tn=int(tn), fp=int(fp), fn=int(fn), tp=int(tp))


def grouped_predictions(d):
    d = d.copy()
    x, y = d[FEATURES], d.tumor_label.to_numpy()
    group = d.patient_id.to_numpy()
    folds = StratifiedGroupKFold(n_splits=5, shuffle=True, random_state=SEED)
    p, prior = np.full(len(d), np.nan), np.full(len(d), np.nan)
    fold_id, coverage = np.zeros(len(d), dtype=int), np.zeros(len(d), dtype=int)
    details = []
    for k, (tr, te) in enumerate(folds.split(x, y, group), 1):
        train_patients, test_patients = sorted(set(group[tr])), sorted(set(group[te]))
        if set(train_patients) & set(test_patients):
            raise ValueError('Patient leakage')
        if len(np.unique(y[tr])) != 2 or len(np.unique(y[te])) != 2:
            raise ValueError('Fold missing a class')
        model = make_model()
        with warnings.catch_warnings():
            warnings.simplefilter('error', ConvergenceWarning)
            model.fit(x.iloc[tr], y[tr])
        p[te] = model.predict_proba(x.iloc[te])[:, 1]
        prior[te] = y[tr].mean()
        fold_id[te] = k; coverage[te] += 1
        scaler = model.named_steps['preprocessing'].named_transformers_['age']
        details.append(dict(fold=k, train_patients=train_patients, test_patients=test_patients,
                            n_train_patients=len(train_patients), n_test_patients=len(test_patients),
                            n_train_timepoints=len(tr), n_test_timepoints=len(te),
                            test_tumor=int(y[te].sum()), test_treatment_change=int((1-y[te]).sum()),
                            train_prevalence=float(y[tr].mean()), train_age_mean=float(scaler.mean_[0]),
                            train_age_sd=float(scaler.scale_[0]),
                            logistic=scores(y[te], p[te]), prevalence_only=scores(y[te], prior[te])))
    if not np.all(coverage == 1) or not np.isfinite(p).all():
        raise ValueError('Invalid OOF coverage')
    d['validation_fold'] = fold_id
    d['oof_tumor_probability'] = p
    d['oof_prevalence_only_probability'] = prior
    d['oof_prediction_at_0_5'] = (p >= .5).astype(int)
    return d, details


def cluster_intervals(d, n_bootstrap=2000):
    """Resample patients with all their scans; probabilities remain fixed OOF values."""
    rng = np.random.default_rng(SEED + 1)
    clusters = [np.asarray(ix) for ix in d.groupby('patient_id', sort=True).indices.values()]
    y, p, q = (d[c].to_numpy() for c in ['tumor_label', 'oof_tumor_probability', 'oof_prevalence_only_probability'])
    draws = []
    for _ in range(n_bootstrap):
        ix = np.concatenate([clusters[i] for i in rng.integers(0, len(clusters), len(clusters))])
        if np.unique(y[ix]).size < 2:
            continue
        b = brier_score_loss(y[ix], p[ix])
        draws.append([roc_auc_score(y[ix], p[ix]), average_precision_score(y[ix], p[ix]),
                      average_precision_score(1-y[ix], 1-p[ix]), b,
                      b-brier_score_loss(y[ix], q[ix])])
    keys = ['roc_auc', 'ap_tumor', 'ap_treatment_change', 'brier', 'brier_difference_logistic_minus_prior']
    if len(draws) < n_bootstrap * .95:
        raise ValueError('Too many invalid bootstrap samples')
    interval = np.quantile(np.asarray(draws), [.025, .975], axis=0)
    return {k: dict(lower=float(interval[0, i]), upper=float(interval[1, i])) for i, k in enumerate(keys)}, len(draws)


def write_plots(d, outdir):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from sklearn.metrics import roc_curve, precision_recall_curve
    from sklearn.calibration import calibration_curve
    plt.rcParams.update({'font.size': 10, 'axes.spines.top': False, 'axes.spines.right': False})
    y = d.tumor_label.to_numpy(); p = d.oof_tumor_probability.to_numpy()
    fig, axes = plt.subplots(1, 3, figsize=(13, 4), constrained_layout=True)
    fpr, tpr, _ = roc_curve(y, p)
    axes[0].plot(fpr, tpr, color='#275d80', label=f'Age + sex AUC = {roc_auc_score(y,p):.3f}')
    axes[0].plot([0, 1], [0, 1], '--', color='gray'); axes[0].legend(loc='lower right')
    axes[0].set(xlabel='False positive rate (treatment change)', ylabel='Sensitivity (tumor)', title='Patient-grouped OOF ROC', xlim=(0,1), ylim=(0,1))
    precision, recall, _ = precision_recall_curve(y, p)
    axes[1].plot(recall, precision, color='#275d80', label=f'Tumor AP = {average_precision_score(y,p):.3f}')
    axes[1].axhline(y.mean(), linestyle='--', color='gray', label=f'Tumor prevalence = {y.mean():.3f}')
    axes[1].set(xlabel='Recall (tumor)', ylabel='Precision (tumor)', title='OOF precision-recall', xlim=(0,1), ylim=(0,1)); axes[1].legend(loc='lower left')
    observed, predicted = calibration_curve(y, p, n_bins=5, strategy='quantile')
    axes[2].plot(predicted, observed, 'o-', color='#275d80', label='5 quantile bins')
    axes[2].plot([0,1], [0,1], '--', color='gray')
    axes[2].set(xlabel='Mean predicted tumor probability', ylabel='Observed tumor proportion', title='Descriptive OOF calibration', xlim=(0,1), ylim=(0,1)); axes[2].legend(loc='lower right')
    fig.savefig(outdir / 'clinical_baseline_validation.png', dpi=180)
    fig.savefig(outdir / 'clinical_baseline_validation.pdf')
    plt.close(fig)
    return pd.DataFrame({'mean_predicted_tumor_probability': predicted, 'observed_tumor_proportion': observed})


def run(root):
    root = Path(root)
    outdir = root / 'outputs/ucsd-clinical-baseline-20261029'; outdir.mkdir(parents=True, exist_ok=True)
    reportdir = root / 'reports/ucsd_clinical_baseline'; reportdir.mkdir(parents=True, exist_ok=True)
    cohort = make_cohort(root)
    cohort.to_csv(root / 'data/processed/ucsd_clinical_baseline_cohort.csv', index=False)
    d, folds = grouped_predictions(cohort)
    d.to_csv(root / 'data/processed/ucsd_clinical_oof_predictions.csv', index=False, float_format='%.12g')
    main = scores(d.tumor_label, d.oof_tumor_probability)
    prior = scores(d.tumor_label, d.oof_prevalence_only_probability)
    intervals, valid_bootstraps = cluster_intervals(d)
    final = make_model().fit(d[FEATURES], d.tumor_label)
    sc = final.named_steps['preprocessing'].named_transformers_['age']
    lr = final.named_steps['logistic']; coef = lr.coef_[0]
    parameters = dict(features=FEATURES, age_mean=float(sc.mean_[0]), age_sd=float(sc.scale_[0]),
                      intercept=float(lr.intercept_[0]), coefficients_standardized_age_and_male=coef.tolist(),
                      age_odds_ratio_per_10_years=float(np.exp(coef[0]*10/sc.scale_[0])),
                      male_vs_female_odds_ratio=float(np.exp(coef[1])),
                      note='Full-cohort descriptive fit, not OOF performance or independent risk-factor inference')
    summary = dict(question='HGG post-treatment residual/recurrent tumor vs treatment-related change',
                   user_approved_scope_change=True, original_OpenBTAI_model_ready=False,
                   source_sha256=SOURCE_HASHES, protocol_sha256=sha256(root/'reports/ucsd_clinical_baseline_protocol.md'),
                   analysis_code_sha256=sha256(__file__), seed=SEED, folds=5, predictors=FEATURES,
                   timepoints=len(d), patients=int(d.patient_id.nunique()), tumor_timepoints=int(d.tumor_label.sum()),
                   treatment_change_timepoints=int((1-d.tumor_label).sum()),
                   mixed_class_patients=int(d.groupby('patient_id').tumor_label.nunique().gt(1).sum()),
                   diagnosis_text_anomalies=int(d.diagnosis_text_anomaly.sum()), excluded_timepoints=0,
                   sex_counts_timepoints=d.sex_at_birth_raw.value_counts().to_dict(),
                   age_min=float(d.age_years.min()), age_max=float(d.age_years.max()),
                   tumor_prevalence=float(d.tumor_label.mean()), logistic_oof=main, prevalence_only_oof=prior,
                   brier_difference_logistic_minus_prior=main['brier']-prior['brier'],
                   patient_cluster_bootstrap_95ci=intervals, bootstrap_requested=2000, bootstrap_valid=valid_bootstraps,
                   bootstrap_refits_models=False, bootstrap_limit='Conditional on fixed OOF predictions; excludes full training/split variability',
                   grouped_patient_overlap=0, each_timepoint_predicted_once=True, threshold=0.5,
                   model_selection_performed=False, mri_features_used=False, external_validation=False,
                   final_fit_parameters=parameters,
                   versions={k: importlib.metadata.version(k) for k in ['scikit-learn','numpy','pandas','scipy','matplotlib']},
                   next_step_requires_user_approval=True)
    (reportdir/'metrics.json').write_text(json.dumps(summary, indent=2))
    (reportdir/'fold_audit.json').write_text(json.dumps(folds, indent=2))
    pd.DataFrame([{k: v for k, v in f.items() if not isinstance(v, (dict, list))} |
                  {'roc_auc': f['logistic']['roc_auc'], 'brier': f['logistic']['brier']} for f in folds]).to_csv(reportdir/'fold_metrics.csv', index=False)
    write_plots(d, outdir).to_csv(reportdir/'calibration_bins.csv', index=False)
    (reportdir/'final_fit_parameters.json').write_text(json.dumps(parameters, indent=2))
    return summary
